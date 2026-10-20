"""Observation-only workbook audits built on the pure-Python formula engine.

Every function takes an evaluated workbook (``Evaluation``) and returns plain,
JSON-safe observations that cite exact cells. Nothing here repairs a workbook,
asserts that a flagged cell is wrong, or supplies company content: heuristics say
"candidate" and name their evidence. An empty finding list means the heuristic saw
nothing, not that the model is correct.
"""
from __future__ import annotations

import difflib
import math
import re
from collections import Counter, defaultdict

from openpyxl.utils.cell import get_column_letter

from .formula_engine import (Evaluation, XlErr, a1, formula_refs, tokenize, walk, _is_num, _cell,
                             WORKBOOK_ERROR_CODES)

MAX_EXAMPLES = 12


def _range_text(sheet, r1, c1, r2, c2):
    a = a1(r1, c1)
    return f'{sheet}!{a}' if (r1, c1) == (r2, c2) else f'{sheet}!{a}:{a1(r2, c2)}'


def _val(ev, sheet, r, c):
    return ev.cell_value(sheet, r, c)


def row_label(ev, sd, r, max_col=4):
    for c in range(1, max_col + 1):
        v = ev.cell_value(sd.name, r, c)
        if isinstance(v, str) and v.strip():
            return v.strip()
    return ''


def _formula_keys(ev):
    for key in sorted(ev.values):
        yield key, ev.model.sheets[key[0]]


# ------------------------------------------------------------------ (a) hidden
def hidden_inventory(ev: Evaluation, *, max_examples=MAX_EXAMPLES):
    model = ev.model
    sheets, rows, cols = [], [], []
    for sd in model.sheets.values():
        nonblank = len(sd.cells) + len(sd.formulas)
        if sd.state != 'visible':
            sheets.append({'sheet': sd.name, 'state': sd.state, 'formula_cells': len(sd.formulas), 'literal_cells': len(sd.cells)})
        if sd.hidden_rows:
            n = sum(1 for (r, _c) in list(sd.cells) + list(sd.formulas) if r in sd.hidden_rows)
            rows.append({'sheet': sd.name, 'rows': sorted(sd.hidden_rows)[:200], 'row_count': len(sd.hidden_rows), 'nonblank_cells': n})
        if sd.hidden_cols:
            n = sum(1 for (_r, c) in list(sd.cells) + list(sd.formulas) if c in sd.hidden_cols)
            cols.append({'sheet': sd.name, 'columns': [get_column_letter(c) for c in sorted(sd.hidden_cols)][:200],
                         'column_count': len(sd.hidden_cols), 'nonblank_cells': n})
        del nonblank
    readers = defaultdict(lambda: {'readers': 0, 'examples': []})
    span_hidden = []
    for (sk, r, c), sd in _formula_keys(ev):
        visible_reader = sd.state == 'visible' and r not in sd.hidden_rows and c not in sd.hidden_cols
        for tsheet, r1, c1, r2, c2 in formula_refs(model, sd.name, r, c):
            tsd = model.sheet(tsheet)
            reason = None
            if tsd.state != 'visible':
                reason = f'hidden sheet {tsd.name}'
            else:
                hr = [x for x in tsd.hidden_rows if r1 <= x <= r2]
                hc = [x for x in tsd.hidden_cols if c1 <= x <= c2]
                if hr or hc:
                    area = (r2 - r1 + 1) * (c2 - c1 + 1)
                    if area > 1:
                        span_hidden.append({'cell': f'{sd.name}!{a1(r, c)}', 'range': _range_text(tsd.name, r1, c1, r2, c2),
                                            'hidden_rows_inside': hr[:10], 'hidden_columns_inside': [get_column_letter(x) for x in hc[:10]]})
                    reason = f'hidden rows/columns in {tsd.name}'
            if reason and visible_reader:
                slot = readers[reason]
                slot['readers'] += 1
                if len(slot['examples']) < max_examples:
                    slot['examples'].append(f'{sd.name}!{a1(r, c)} reads {_range_text(tsd.name, r1, c1, r2, c2)}')
    return {'hidden_sheets': sheets, 'hidden_rows': rows, 'hidden_columns': cols,
            'hidden_defined_names': list(model.hidden_names),
            'visible_cells_reading_hidden': [{'reason': k, **v} for k, v in sorted(readers.items())],
            'ranges_spanning_hidden_rows_or_columns': span_hidden[:50],
            'span_hidden_total': len(span_hidden)}


# ------------------------------------------------- (b) range-width consistency
def _sig(ev, sd, r, c):
    refs = [x for x in formula_refs(ev.model, sd.name, r, c) if x[0].lower() != sd.name.lower()]
    if not refs:
        return None
    targets = {x[0] for x in refs}
    if len(targets) != 1:
        return None
    start = min(x[2] for x in refs)
    end = max(x[4] for x in refs)
    return {'target': next(iter(targets)), 'start': start, 'width': end - start + 1,
            'rows': (min(x[1] for x in refs), max(x[3] for x in refs))}


def range_width_audit(ev: Evaluation, *, min_modal_rows=2, modal_share=0.6, max_examples=MAX_EXAMPLES):
    """Rows whose adjacent period columns pull a different stride/width of another sheet.

    A row is described by (stride, width): how far its source start column advances
    between adjacent cells, and how many source columns each cell reads. Rows that
    disagree with the sheet's dominant pattern are reported with their cells.
    """
    model, findings = ev.model, []
    for sd in model.sheets.values():
        row_cells = defaultdict(list)
        for (r, c) in sorted(sd.formulas):
            v = ev.cell_value(sd.name, r, c)
            if isinstance(v, str):                      # label/header pulls are not period data
                continue
            s = _sig(ev, sd, r, c)
            if s:
                row_cells[r].append((c, s))
        summaries = {}
        for r, items in row_cells.items():
            dominant = Counter(x[1]['target'] for x in items).most_common(1)[0][0]
            items = [x for x in items if x[1]['target'] == dominant]
            if len(items) < 2:
                continue
            adjacent = [(items[i], items[i + 1]) for i in range(len(items) - 1)
                        if items[i + 1][0] - items[i][0] == 1 and items[i][1]['target'] == items[i + 1][1]['target']]
            if not adjacent:
                continue
            strides = Counter(b[1]['start'] - a[1]['start'] for a, b in adjacent)
            widths = Counter(x[1]['width'] for x in items)
            summaries[r] = {'stride': strides.most_common(1)[0][0], 'stride_counts': dict(strides),
                            'width': widths.most_common(1)[0][0], 'width_counts': dict(widths),
                            'target': items[0][1]['target'], 'cells': [c for c, _ in items], 'items': items}
        if not summaries:
            continue
        # Intra-row inconsistency: adjacent cells advance / read different widths.
        for r, s in summaries.items():
            if len(s['stride_counts']) > 1 or len(s['width_counts']) > 1:
                findings.append({'kind': 'inconsistent_within_row', 'severity': 'high', 'sheet': sd.name, 'row': r,
                                 'label': row_label(ev, sd, r), 'target_sheet': s['target'],
                                 'stride_counts': s['stride_counts'], 'width_counts': s['width_counts'],
                                 'cells': [f'{sd.name}!{a1(r, c)}' for c in s['cells']][:30],
                                 'formulas': [sd.formulas[(r, c)] for c in s['cells']][:6]})
        # Stride 0 means every cell reads the same source column (a fixed parameter pull), not a period roll.
        summaries = {r: s for r, s in summaries.items() if s['stride'] != 0 or len(s['stride_counts']) > 1}
        if not summaries:
            continue
        pattern = Counter((s['stride'], s['width']) for s in summaries.values())
        (m_stride, m_width), m_rows = pattern.most_common(1)[0]
        if m_rows < min_modal_rows or m_rows / len(summaries) < modal_share:
            continue
        siblings = [r for r, s in summaries.items() if (s['stride'], s['width']) == (m_stride, m_width)]
        for r, s in sorted(summaries.items()):
            if (s['stride'], s['width']) == (m_stride, m_width) or len(s['stride_counts']) > 1 or len(s['width_counts']) > 1:
                continue
            stride_diff = s['stride'] != m_stride and max(s['stride'], m_stride) > 1
            example = siblings[0]
            ex_cells = summaries[example]['items'][:2]
            severity = 'high' if stride_diff else 'info'
            if not stride_diff and s['width'] == m_width:
                continue
            if stride_diff:
                text = (f'Adjacent period columns step {s["stride"]} source column(s) and read {s["width"]} column(s) of '
                        f'{s["target"]}, while {len(siblings)} sibling rows on this sheet step {m_stride} and read {m_width}.')
            else:
                text = (f'Cells read {s["width"]} source column(s) per period; sibling rows read {m_width}. May be a deliberate '
                        'point-in-time pull or an omitted period range.')
            findings.append({'kind': 'row_pattern_differs_from_sheet', 'severity': severity, 'sheet': sd.name, 'row': r,
                             'label': row_label(ev, sd, r), 'target_sheet': s['target'],
                             'row_stride': s['stride'], 'row_width': s['width'], 'modal_stride': m_stride, 'modal_width': m_width,
                             'modal_row_count': len(siblings),
                             'cells': [f'{sd.name}!{a1(r, c)}' for c in s['cells']][:30],
                             'formulas': [sd.formulas[(r, c)] for c in s['cells']][:6],
                             'sibling_example': {'row': example, 'label': row_label(ev, sd, example),
                                                 'formulas': [sd.formulas[(example, c)] for c, _ in ex_cells]},
                             'observation': text})
    findings.sort(key=lambda f: (f['severity'] != 'high', f['sheet'], f['row']))
    return {'findings': findings[:200], 'finding_count': len(findings),
            'high_severity_count': sum(1 for f in findings if f['severity'] == 'high'),
            'method': 'per-row source stride/width vs sheet-modal pattern; observation only'}


# ------------------------------------------------------------ (c) unit / scale
_CURRENCY = re.compile(r"(INR|Rs\.?|₹|USD|US\$|\$|EUR|€|GBP|£|Rupees?)", re.I)
_SCALE = re.compile(r"(?<![A-Za-z])(mn|mm|million|millions|m|cr|crore|crores|lakh|lakhs|lac|lacs|k|'000|000s|thousand|thousands|bn|billion|billions)(?![A-Za-z])", re.I)
_SCALE_VALUE = {'mn': 1e6, 'mm': 1e6, 'million': 1e6, 'millions': 1e6, 'm': 1e6, 'cr': 1e7, 'crore': 1e7, 'crores': 1e7,
                'lakh': 1e5, 'lakhs': 1e5, 'lac': 1e5, 'lacs': 1e5, 'k': 1e3, "'000": 1e3, '000s': 1e3, 'thousand': 1e3,
                'thousands': 1e3, 'bn': 1e9, 'billion': 1e9, 'billions': 1e9}
_CUR_NORM = {'rs': 'INR', 'rs.': 'INR', '₹': 'INR', 'rupee': 'INR', 'rupees': 'INR', 'us$': 'USD', '$': 'USD',
             '€': 'EUR', '£': 'GBP'}
SCALE_CONSTANTS = {1e3: 'thousand', 1e5: 'lakh', 1e6: 'million', 1e7: 'crore', 1e9: 'billion'}


def parse_unit(text):
    """Return {'currency','scale','token'} when short label text declares a unit, else None."""
    if not isinstance(text, str) or len(text) > 60:
        return None
    cur = _CURRENCY.search(text)
    sc = _SCALE.search(text)
    if cur is None and sc is None:
        return None
    word = sc.group(1).lower() if sc else None
    if sc and cur is None:
        bare = text.strip().lower()
        if len(word) == 1 or not (bare.startswith(('in ', '(')) or bare == word):
            return None
    if cur and sc and len(word) == 1 and abs(sc.start() - cur.end()) > 3:
        sc, word = None, None
    currency = None
    if cur:
        low = cur.group(1).lower()
        currency = _CUR_NORM.get(low, low.upper())
    scale = _SCALE_VALUE[word] if word else 1.0
    token = (cur.group(0) if cur else '') + (' ' + sc.group(0) if sc else '')
    return {'currency': currency, 'scale': scale, 'token': token.strip()}


def _unit_key(u):
    return None if u is None else (u['currency'], u['scale'])


def _sheet_unit(ev, sd):
    for r in range(1, min(sd.max_row, 4) + 1):
        for c in range(1, min(sd.max_col, 6) + 1):
            u = parse_unit(ev.cell_value(sd.name, r, c))
            if u:
                return u, f'{sd.name}!{a1(r, c)}'
    return None, None


def _row_unit(ev, sd, r, sheet_unit):
    for c in range(1, 5):
        v = ev.cell_value(sd.name, r, c)
        if isinstance(v, str) and v.strip():
            u = parse_unit(v)
            if u:
                if u['scale'] == 1.0 and sheet_unit and sheet_unit['scale'] != 1.0 and u['currency'] == sheet_unit['currency']:
                    return sheet_unit, f'{sd.name}!{a1(r, c)}'
                return u, f'{sd.name}!{a1(r, c)}'
            break
    return sheet_unit, None


def _scale_constants(node):
    found = []
    for n in walk(node):
        if n[0] == 'num' and n[1] in SCALE_CONSTANTS:
            found.append(n[1])
        elif n[0] == 'bin' and n[1] == '^' and n[2][0] == 'num' and n[2][1] == 10 and n[3][0] == 'num' and n[3][1] in (3, 5, 6, 7, 9):
            found.append(10.0 ** n[3][1])
    return found


def unit_scale_audit(ev: Evaluation, *, max_examples=MAX_EXAMPLES):
    model = ev.model
    sheet_units, declared = {}, []
    for sd in model.sheets.values():
        u, where = _sheet_unit(ev, sd)
        sheet_units[sd.name] = u
        if u:
            declared.append({'sheet': sd.name, 'cell': where, 'unit': u['token'], 'currency': u['currency'], 'scale': u['scale']})
    row_units = {}
    mixed_in_sheet = []
    for sd in model.sheets.values():
        seen = {}
        for r in range(1, sd.max_row + 1):
            u, where = _row_unit(ev, sd, r, sheet_units[sd.name])
            row_units[(sd.name, r)] = u
            if where and u:
                seen.setdefault(_unit_key(u), []).append(where)
        keys = set(seen)
        if sheet_units[sd.name]:
            keys.add(_unit_key(sheet_units[sd.name]))
        if len(keys) > 1:
            mixed_in_sheet.append({'sheet': sd.name, 'units': [{'unit': k, 'cells': v[:5]} for k, v in seen.items()]})
    mismatches, conversions, mixed_formulas, constants = [], [], [], []
    for (sk, r, c), sd in _formula_keys(ev):
        node = model.ast(sd.name, r, c)
        own = row_units.get((sd.name, r))
        consts = _scale_constants(node)
        has_name = any(n[0] in ('name', 'struct') for n in walk(node))
        if consts:
            constants.append({'cell': f'{sd.name}!{a1(r, c)}', 'constants': sorted(set(consts)), 'formula': sd.formulas[(r, c)][:140],
                              'own_unit': own['token'] if own else None})
        prec_units = {}
        for tsheet, r1, c1, r2, c2 in formula_refs(model, sd.name, r, c):
            tsd = model.sheet(tsheet)
            for rr in range(r1, min(r2, r1 + 300) + 1):
                u = row_units.get((tsd.name, rr))
                if u:
                    prec_units.setdefault(_unit_key(u), (u, f'{tsd.name}!{a1(rr, c1)}'))
        if own and prec_units:
            differing = {k: v for k, v in prec_units.items() if k != _unit_key(own)}
            if differing:
                converted = bool(consts) or has_name
                record = {'cell': f'{sd.name}!{a1(r, c)}', 'cell_unit': own['token'],
                          'precedent_units': [{'unit': v[0]['token'], 'example': v[1]} for v in differing.values()][:4],
                          'formula': sd.formulas[(r, c)][:140]}
                (conversions if converted else mismatches).append(record)
        if len(prec_units) > 1:
            fn_names = {n[1] for n in walk(node) if n[0] == 'func'}
            ops = {n[1] for n in walk(node) if n[0] == 'bin'}
            if (fn_names & {'SUM', 'AVERAGE', 'MIN', 'MAX'}) or ({'+', '-'} & ops):
                mixed_formulas.append({'cell': f'{sd.name}!{a1(r, c)}', 'units': [v[0]['token'] for v in prec_units.values()],
                                       'formula': sd.formulas[(r, c)][:140]})
    return {'sheet_units': declared, 'units_mixed_within_sheet': mixed_in_sheet,
            'cross_sheet_unit_mismatch_without_conversion': mismatches[:100], 'mismatch_count': len(mismatches),
            'cross_sheet_unit_change_with_conversion_constant_or_name': conversions[:max_examples],
            'conversion_count': len(conversions),
            'additive_formulas_mixing_units': mixed_formulas[:100], 'mixed_formula_count': len(mixed_formulas),
            'scale_constants_in_formulas': constants[:100], 'scale_constant_cell_count': len(constants),
            'limits': 'unit labels are parsed from short header/label text; cells without a label are unknown, not clean'}


# --------------------------------------------------- (d) driver growth double count
def _growth_signature(ev, sd, r, c):
    """For a self-referencing growth cell (X = X[prev col] * g...), return its growth-series key."""
    node = ev.model.ast(sd.name, r, c)
    own_prev = None
    others = []
    for n in walk(node):
        if n[0] != 'ref':
            continue
        sheet = (n[1] or sd.name).lower()
        if sheet == sd.name.lower() and n[2] == r and n[4] == r and n[3] == n[5] and n[3] < c and own_prev is None:
            own_prev = n
        elif (n[2], n[3]) == (n[4], n[5]):
            others.append((sheet, n[2], n[3] - c))
    ops = {n[1] for n in walk(node) if n[0] == 'bin'}
    if own_prev is None or '*' not in ops or not others:
        return None
    return tuple(sorted(set(others)))


_PRICE = re.compile(r'price|rate|arpu|arppu|ticket|per |yield|premium|fee|cost per|take|aov|average|asp|ltv|cac', re.I)
_COUNT = re.compile(r'user|customer|count|number|no\.|volume|policies|orders|subscriber|units|merchant|partner|gwp|lives|vehicles|nos', re.I)


def driver_growth_double_count(ev: Evaluation, *, max_examples=MAX_EXAMPLES):
    model = ev.model
    by_row = defaultdict(list)
    for (sk, r, c), sd in _formula_keys(ev):
        sig = _growth_signature(ev, sd, r, c)
        if sig:
            by_row[(sd.name, r)].append((c, sig))
    groups = defaultdict(list)
    for (sheet, r), items in by_row.items():
        if len(items) < 3:
            continue
        common = Counter(sig for _, sig in items).most_common(1)[0]
        if common[1] / len(items) >= 0.6:
            groups[common[0]].append((sheet, r))
    findings = []
    for sig, rows in groups.items():
        if len(rows) < 2:
            continue
        rowset = {(s.lower(), r): (s, r) for s, r in rows}
        combos = []
        for (sk, r, c), sd in _formula_keys(ev):
            node = model.ast(sd.name, r, c)
            hits = set()
            for n in _product_leaves(node):
                if n[0] == 'ref' and (n[2], n[3]) == (n[4], n[5]):
                    key = ((n[1] or sd.name).lower(), n[2])
                    if key in rowset:
                        hits.add(key)
            if len(hits) >= 2:
                combos.append({'cell': f'{sd.name}!{a1(r, c)}', 'driver_rows': [rowset[h] for h in sorted(hits)],
                               'formula': sd.formulas[(r, c)][:140]})
        if not combos:
            continue
        labels = {(s, r): row_label(ev, model.sheet(s), r) for s, r in rows}
        price_like = [k for k, v in labels.items() if _PRICE.search(v)]
        count_like = [k for k, v in labels.items() if _COUNT.search(v)]
        findings.append({'growth_series': [{'sheet': s, 'row': r, 'col_offset': o} for s, r, o in sig],
                         'driver_rows': [{'sheet': s, 'row': r, 'label': labels[(s, r)]} for s, r in rows],
                         'combined_by_multiplication_in': combos[:max_examples], 'combination_count': len(combos),
                         'label_hint': 'price_like_and_count_like' if price_like and count_like and set(price_like) != set(count_like) else 'labels_not_conclusive',
                         'confidence': 'candidate',
                         'observation': 'Two or more driver rows are each rolled forward by the same growth series and then multiplied together; '
                                        'growth compounds in both factors. Candidate only: confirm the intent from the model owner.'})
    return {'findings': findings, 'finding_count': len(findings),
            'method': 'rows with X[t]=X[t-1]*f(series) grouped by shared series, then co-multiplication downstream'}


def _product_leaves(node):
    stack, out = [node], []
    while stack:
        n = stack.pop()
        if n[0] == 'bin' and n[1] == '*':
            stack += [n[2], n[3]]
        elif n[0] == 'bin' and n[1] == '/':
            stack.append(n[2])
        elif n[0] in ('paren',):
            stack.append(n[1])
        elif n[0] == 'un':
            stack.append(n[2])
        elif n[0] == 'func' and n[1] in ('SUM', 'ROUND') and n[2]:
            stack.append(n[2][0])
        else:
            out.append(n)
    return out


# ------------------------------------------------- (e) balance sheet: plug / checks
_BS_NAME = re.compile(r'balance|\bbs\b|financial position', re.I)
_LABEL_PLUG = re.compile(r'plug|balancing|suspense|unreconciled|difference|reconcil', re.I)
_ASSETLIKE = re.compile(r'cash|bank|receivable|inventor|fixed asset|\bppe\b|investment|total asset|current asset|prepaid|intangible', re.I)


def _same_col_refs(ev, sd, r, c):
    out = []
    for tsheet, r1, c1, r2, c2 in formula_refs(ev.model, sd.name, r, c):
        if tsheet.lower() == sd.name.lower() and c1 <= c <= c2:
            out += [(rr, c) for rr in range(r1, r2 + 1)]
    return out


def _column_closure(ev, sd, r, c, exclude=None, depth=6):
    seen, frontier = set(), [(r, c)]
    for _ in range(depth):
        nxt = []
        for cell in frontier:
            if cell not in sd.formulas:
                continue
            for ref in _same_col_refs(ev, sd, *cell):
                if ref != exclude and ref not in seen:
                    seen.add(ref)
                    nxt.append(ref)
        frontier = nxt
    return seen


def balance_sheet_audit(ev: Evaluation, *, sheet=None, tolerance=1e-6, max_examples=MAX_EXAMPLES):
    model = ev.model
    candidates = [model.sheet(sheet)] if sheet else [s for s in model.sheets.values() if _BS_NAME.search(s.name)
                                                       or any(re.search(r'total assets', str(v), re.I) for v in s.cells.values())]
    candidates = [s for s in candidates if s is not None]
    report = {'sheets': [], 'method': 'structural plug detection, label-based totals check, zero-sum rows, sign pattern'}
    for sd in candidates:
        plugs, label_plugs, totals, zero_rows, signs = [], [], [], [], []
        formula_rows = defaultdict(list)
        for (r, c) in sorted(sd.formulas):
            formula_rows[r].append(c)
        # (i) structural plug: F = -(others) while a total T reads F and the same others.
        for (r, c) in sorted(sd.formulas):
            node = model.ast(sd.name, r, c)
            if not any(n[0] == 'un' and n[1] == '-' for n in walk(node)) and not any(n[0] == 'bin' and n[1] == '-' for n in walk(node)):
                continue
            prec = set(_same_col_refs(ev, sd, r, c)) - {(r, c)}
            if len(prec) < 2:
                continue
            for (r2, c2) in sorted(sd.formulas):
                if c2 != c or r2 == r:
                    continue
                reach = _column_closure(ev, sd, r2, c2, exclude=(r, c))
                if (r, c) not in _column_closure(ev, sd, r2, c2) or len(prec & reach) < 2:
                    continue
                val = ev.cell_value(sd.name, r2, c2)
                plugs.append({'plug_cell': f'{sd.name}!{a1(r, c)}', 'plug_label': row_label(ev, sd, r),
                              'plug_formula': sd.formulas[(r, c)][:140], 'balancing_total_cell': f'{sd.name}!{a1(r2, c2)}',
                              'balancing_total_formula': sd.formulas[(r2, c2)][:140],
                              'balancing_total_value': val if not isinstance(val, XlErr) else val.code,
                              'observation': 'Cell is defined from other lines that a total also reads, so that total is balanced by construction.'})
                break
        # (ii) label-based plugs
        for r in range(1, sd.max_row + 1):
            lab = row_label(ev, sd, r)
            if lab and _LABEL_PLUG.search(lab):
                label_plugs.append({'row': r, 'label': lab, 'cells': [f'{sd.name}!{a1(r, c)}' for c in formula_rows.get(r, [])][:10]})
        # (iii) cash defined by subtracting from totals
        cash_plugs = []
        for r in range(1, sd.max_row + 1):
            lab = row_label(ev, sd, r)
            if lab and re.search(r'cash|bank', lab, re.I):
                for c in formula_rows.get(r, []):
                    node = model.ast(sd.name, r, c)
                    refs = _same_col_refs(ev, sd, r, c)
                    if refs and any(n[0] == 'bin' and n[1] == '-' for n in walk(node)) and any(
                            re.search(r'total|liabilit|equity', row_label(ev, sd, rr), re.I) for rr, _ in refs):
                        cash_plugs.append({'cell': f'{sd.name}!{a1(r, c)}', 'formula': sd.formulas[(r, c)][:140]})
        # label-based totals check
        a_rows = [r for r in range(1, sd.max_row + 1) if re.search(r'^total assets?\b', row_label(ev, sd, r), re.I)]
        l_rows = [r for r in range(1, sd.max_row + 1) if re.search(r'^total (equity (and|&) liabilit|liabilit(y|ies) (and|&) (shareholders|equity))', row_label(ev, sd, r), re.I)]
        if a_rows and l_rows:
            for c in range(1, sd.max_col + 1):
                va, vl = ev.cell_value(sd.name, a_rows[0], c), ev.cell_value(sd.name, l_rows[0], c)
                if _is_num(va) and _is_num(vl):
                    diff = va - vl
                    totals.append({'column': get_column_letter(c), 'assets': va, 'liabilities_and_equity': vl, 'difference': diff,
                                   'balanced': abs(diff) <= tolerance * max(1.0, abs(va))})
        # (iv) zero-sum rows: formulas summing subtotals evaluating ~0 in every formula column
        for r, cols in formula_rows.items():
            if len(cols) < 2:
                continue
            vals = [ev.cell_value(sd.name, r, c) for c in cols]
            vals = [v for v in vals if _is_num(v)]
            if vals and all(abs(v) <= max(tolerance, 1e-9) * 1e3 for v in vals):
                refs = sum(len(_same_col_refs(ev, sd, r, c)) for c in cols[:1])
                if refs >= 2:
                    zero_rows.append({'row': r, 'label': row_label(ev, sd, r), 'cells': [f'{sd.name}!{a1(r, c)}' for c in cols][:8],
                                      'formula': sd.formulas[(r, cols[0])][:140]})
        # sign pattern
        for r in range(1, sd.max_row + 1):
            lab = row_label(ev, sd, r)
            nums = [ev.cell_value(sd.name, r, c) for c in range(1, sd.max_col + 1) if (r, c) in sd.formulas or (r, c) in sd.cells]
            nums = [v for v in nums if _is_num(v) and not isinstance(v, bool) and v != 0]
            if len(nums) < 2:
                continue
            pos, neg = sum(1 for v in nums if v > 0), sum(1 for v in nums if v < 0)
            if pos and neg:
                signs.append({'row': r, 'label': lab, 'positive_cells': pos, 'negative_cells': neg, 'kind': 'sign_flips_across_periods'})
            elif neg and lab and _ASSETLIKE.search(lab):
                signs.append({'row': r, 'label': lab, 'positive_cells': pos, 'negative_cells': neg, 'kind': 'asset_like_row_negative'})
        report['sheets'].append({'sheet': sd.name, 'state': sd.state, 'plug_candidates': plugs[:max_examples], 'plug_candidate_count': len(plugs),
                                 'label_plug_rows': label_plugs, 'cash_defined_by_subtraction': cash_plugs[:max_examples],
                                 'assets_vs_liabilities_equity': totals[:30], 'zero_sum_rows': zero_rows[:10], 'sign_observations': signs[:40]})
    return report


# --------------------------------------------------- (f) error propagation map
def error_propagation_map(ev: Evaluation, *, max_examples=MAX_EXAMPLES):
    model = ev.model
    errors = {k for k, v in ev.values.items() if isinstance(v, XlErr)}
    reverse = defaultdict(list)
    for node, deps in ev.deps.items():
        for d in deps:
            reverse[d].append(node)
    roots = [n for n in sorted(errors) if not any(d in errors for d in ev.deps[n])]
    out = []
    for root in roots:
        seen, queue = {root}, [root]
        while queue:
            cur = queue.pop()
            for nxt in reverse.get(cur, ()):
                if nxt in errors and nxt not in seen:
                    seen.add(nxt)
                    queue.append(nxt)
        sd = model.sheets[root[0]]
        text = sd.formulas[(root[1], root[2])]
        code = ev.values[root].code
        origin = 'literal_error_in_formula' if '#REF!' in text or '#N/A' in text or '#NAME?' in text else (
            'circular_reference' if code == '#CIRC!' else 'unsupported_function' if code == '#UNSUPPORTED!' else 'evaluated_error')
        sheets = Counter(model.sheets[n[0]].name for n in seen)
        out.append({'root_cell': ev.label(root), 'error': code, 'origin': origin, 'formula': text[:140],
                    'downstream_error_cells': len(seen) - 1, 'affected_sheets': dict(sheets),
                    'in_hidden_sheet': sd.state != 'visible',
                    'examples': sorted(ev.label(n) for n in seen if n != root)[:max_examples]})
    out.sort(key=lambda x: -x['downstream_error_cells'])
    stored = sum(1 for key in ev.values for _ in [0]
                 if isinstance(model.sheets[key[0]].cached.get((key[1], key[2])), XlErr))
    return {'error_cells': len(errors), 'root_errors': out[:100], 'root_count': len(out),
            'stored_cache_error_cells': stored,
            'workbook_error_cells': sum(1 for n in errors if ev.values[n].code in WORKBOOK_ERROR_CODES),
            'method': 'roots are error cells with no error precedent; downstream counted over the dependency graph'}


# ------------------------------------------------------ (g) hardcodes in formulas
def r1c1(text, r, c):
    out = []
    for kind, val in tokenize(text[1:]):
        if kind in ('ref', 'colrange', 'rowrange'):
            sheet = ''
            body = val
            if '!' in val:
                sheet, body = val.rsplit('!', 1)
                sheet += '!'

            def one(tok):
                m = re.fullmatch(r'(\$?)([A-Za-z]{1,3})(\$?)([0-9]+)', tok)
                if not m:
                    return tok
                from openpyxl.utils.cell import column_index_from_string
                cc, rr = column_index_from_string(m.group(2).upper()), int(m.group(4))
                col = f'C{cc}' if m.group(1) else f'C[{cc - c}]'
                row = f'R{rr}' if m.group(3) else f'R[{rr - r}]'
                return row + col
            out.append(sheet + ':'.join(one(p) for p in body.split(':')))
        else:
            out.append(val)
    return '=' + ''.join(out)


def hardcode_audit(ev: Evaluation, *, allowed=(0, 1, 2, 12, 100, 365, 4, 3, 10), max_examples=MAX_EXAMPLES):
    model = ev.model
    allowed = set(float(x) for x in allowed)
    groups = defaultdict(lambda: {'cells': [], 'literals': set()})
    pure_constants = []
    for (sk, r, c), sd in _formula_keys(ev):
        node = model.ast(sd.name, r, c)
        nums = [n[1] for n in walk(node) if n[0] == 'num']
        if not nums:
            continue
        if node[0] in ('num',) or (node[0] == 'un' and node[2][0] == 'num') or all(
                n[0] in ('num', 'bin', 'un', 'paren', 'pct') for n in walk(node)):
            pure_constants.append({'cell': f'{sd.name}!{a1(r, c)}', 'formula': sd.formulas[(r, c)][:100]})
            continue
        pct = {n[1][1] for n in walk(node) if n[0] == 'pct' and n[1][0] == 'num'}
        lits = sorted({x for x in nums if x not in allowed and abs(x) not in allowed})
        if lits:
            g = groups[r1c1(sd.formulas[(r, c)], r, c)]
            g['cells'].append(f'{sd.name}!{a1(r, c)}')
            g['literals'].update(lits)
        del pct
    patterns = sorted(({'pattern': k[:160], 'literals': sorted(v['literals']), 'cell_count': len(v['cells']), 'cells': v['cells'][:max_examples]}
                       for k, v in groups.items()), key=lambda x: -x['cell_count'])
    overwritten = []
    for sd in model.sheets.values():
        by_row = defaultdict(list)
        for (r, c) in set(sd.formulas) | {k for k, v in sd.cells.items() if _is_num(v)}:
            by_row[r].append(c)
        for r, cols in by_row.items():
            fcols = sorted(c for c in cols if (r, c) in sd.formulas)
            if len(fcols) < 4:
                continue
            span = range(fcols[0], fcols[-1] + 1)
            lits = [c for c in span if (r, c) in sd.cells and _is_num(sd.cells[(r, c)])]
            if lits and len(fcols) >= 3 * len(lits):
                overwritten.append({'sheet': sd.name, 'row': r, 'label': row_label(ev, sd, r),
                                    'literal_cells': [f'{sd.name}!{a1(r, c)}' for c in lits][:8],
                                    'neighbour_formula': sd.formulas[(r, fcols[0])][:100]})
    return {'formula_patterns_with_embedded_literals': patterns[:100], 'pattern_count': len(patterns),
            'pure_constant_formulas': pure_constants[:100], 'pure_constant_formula_count': len(pure_constants),
            'literal_among_formulas_in_row': overwritten[:100], 'allowed_literals': sorted(allowed),
            'note': 'a literal may be a legitimate unit/calendar constant; these are locations to ask about, not defects'}


# ------------------------------------------------ (h) monthly = quarterly = FY
_NON_ADDITIVE = re.compile(r'%|margin|ratio|rate|per |avg|average|headcount|employees|balance|no\. of|number of|\bcount\b|growth|index|share of|opening|closing|price|arpu', re.I)


def _norm_label(text):
    t = re.sub(r'\([^)]*\)', ' ', text.lower())
    t = re.sub(r'[^a-z0-9&% ]+', ' ', t)
    words = [w[:-1] if len(w) > 3 and w.endswith('s') and not w.endswith('ss') else w for w in t.split()]
    return ' '.join(words)


def reconcile_totals(ev: Evaluation, members, *, explicit_rows=None, abs_tol=1e-6, rel_tol=1e-9, match_ratio=0.92,
                     max_examples=MAX_EXAMPLES):
    """Compare per-row period sums across granularity levels (monthly, quarterly, FY ...).

    members: [{'sheet','first_col','last_col','label_col'?,'total_col'?,'first_row'?,'last_row'?}]; columns are
    letters or 1-based indexes. Rows are auto-matched only when their normalised label is unique in every
    member (exact, then fuzzy); repeated labels are reported as ambiguous, never guessed. ``explicit_rows`` is
    [{'name', 'rows': [row_per_member...]}] (one row number per member, in member order) for caller-chosen pairs.
    Balances, rates and headcount-like labels are skipped from auto matching.
    """
    from openpyxl.utils.cell import column_index_from_string
    model = ev.model

    def col(x, default=None):
        if x is None:
            return default
        return x if isinstance(x, int) else column_index_from_string(x.upper())
    parsed = []
    for m in members:
        sd = model.sheet(m['sheet'])
        if sd is None:
            raise ValueError(f"Unknown sheet {m['sheet']}")
        parsed.append({'sd': sd, 'first': col(m['first_col']), 'last': col(m['last_col']), 'label_col': col(m.get('label_col'), 1),
                       'total': col(m.get('total_col')), 'r1': m.get('first_row', 1), 'r2': m.get('last_row', sd.max_row)})

    def row_figure(p, r):
        sd = p['sd']
        nums = [ev.cell_value(sd.name, r, c) for c in range(p['first'], p['last'] + 1)]
        err = next((x for x in nums if isinstance(x, XlErr)), None)
        total = ev.cell_value(sd.name, r, p['total']) if p['total'] else None
        label = ev.cell_value(sd.name, r, p['label_col'])
        return {'sheet': sd.name, 'row': r, 'label': label.strip() if isinstance(label, str) else '',
                'period_sum': math.fsum(x for x in nums if _is_num(x)) if err is None else None,
                'stored_total_column': total if _is_num(total) else None, 'error': err.code if err else None,
                'numeric': any(_is_num(x) for x in nums)}
    by_key = []
    for p in parsed:
        sd, items = p['sd'], defaultdict(list)
        for r in range(p['r1'], p['r2'] + 1):
            v = ev.cell_value(sd.name, r, p['label_col'])
            if isinstance(v, str) and v.strip():
                fig = row_figure(p, r)
                if fig['numeric']:
                    items[_norm_label(v)].append(fig)
        by_key.append(items)

    def judge(group):
        sums = [f['period_sum'] for f in group if f['period_sum'] is not None]
        refs = sums + [f['stored_total_column'] for f in group if f['stored_total_column'] is not None]
        spread = (max(refs) - min(refs)) if refs else 0.0
        scale = max((abs(x) for x in refs), default=0.0)
        if any(f['error'] for f in group) or spread > abs_tol + rel_tol * scale:
            return {'spread': spread, 'levels': [{k: v for k, v in f.items() if k != 'numeric'} for f in group],
                    'cells': [f"{f['sheet']}!{get_column_letter(parsed[k]['first'])}{f['row']}:{get_column_letter(parsed[k]['last'])}{f['row']}"
                              for k, f in enumerate(group)]}
        return None
    mismatches, ambiguous, unmatched, skipped, matched = [], [], [], [], 0
    base_keys = set(by_key[0])
    for key in sorted(base_keys):
        group = []
        ok = True
        for i, items in enumerate(by_key):
            hits = items.get(key)
            if hits is None and i:
                close = difflib.get_close_matches(key, list(items), n=1, cutoff=match_ratio)
                hits = items.get(close[0]) if close else None
            if hits is None:
                ok = False
                break
            group.append(hits)
        label = by_key[0][key][0]['label']
        if not ok:
            unmatched.append({'label': label, 'row': f"{parsed[0]['sd'].name}!{by_key[0][key][0]['row']}"})
            continue
        if any(len(h) > 1 for h in group):
            ambiguous.append({'label': label, 'occurrences': [len(h) for h in group]})
            continue
        if _NON_ADDITIVE.search(label):
            skipped.append(label)
            continue
        matched += 1
        bad = judge([h[0] for h in group])
        if bad:
            mismatches.append({'label': label, 'match': 'label', **bad})
    for pair in explicit_rows or []:
        if len(pair['rows']) != len(parsed):
            raise ValueError('explicit_rows needs one row per member')
        group = [row_figure(p, r) for p, r in zip(parsed, pair['rows'])]
        matched += 1
        bad = judge(group)
        if bad:
            mismatches.append({'label': pair.get('name') or group[0]['label'], 'match': 'explicit', **bad})
    mismatches.sort(key=lambda x: -x['spread'])
    return {'levels': [{'sheet': p['sd'].name, 'columns': f"{get_column_letter(p['first'])}:{get_column_letter(p['last'])}"} for p in parsed],
            'compared_rows': matched, 'mismatch_count': len(mismatches), 'mismatches': mismatches[:100],
            'ambiguous_repeated_labels': ambiguous[:50], 'unmatched_rows': unmatched[:50], 'skipped_non_additive_labels': skipped[:50],
            'method': 'sum of period columns per matched row compared across levels and with stored total columns'}
