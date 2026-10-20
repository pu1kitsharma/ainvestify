"""Number-binding checks for authored text and charts. Observation only.

``bind_claims`` finds every number in agent-authored text (memo section, slide title,
table cell) and checks that it binds, within its displayed rounding, to a workbook
cell value supplied in a table of references, with consistent unit/scale, sign and
comparative wording (up/down, %, x, CAGR). It never edits text and never decides what
a number should be: unbound or conflicting claims are reported with the nearest cells.

``chart_workbook_agreement`` compares native PPTX chart series (cached values and the
embedded chart workbook) with an evaluated workbook.
"""
from __future__ import annotations

import math
import re
from io import BytesIO

from .audits import parse_unit

_CUR = r'(?:INR|Rs\.?|₹|USD|US\$|\$|EUR|€|GBP|£)'
_NUMBER = re.compile(rf'''
 (?P<cur>{_CUR}\s*)?
 (?P<open>\()?(?P<sign>[-−–+])?
 (?P<num>\d{{1,3}}(?:,\d{{2,3}})+(?:\.\d+)?|\d+(?:\.\d+)?|\.\d+)
 (?P<close>\))?
 (?:\s*(?P<suffix>%|×|x(?![A-Za-z])|pp(?![A-Za-z])|bps(?![A-Za-z])|bp(?![A-Za-z])|(?:Mn|MM|mn|mm|million|millions|bn|Bn|billion|billions|Cr|cr|crore|crores|lakh|lakhs|lac|lacs|thousand|k|K|m|M|B)(?![A-Za-z])))?
 (?P<close2>\))?
''', re.X)
_SCALE = {'mn': 1e6, 'mm': 1e6, 'million': 1e6, 'millions': 1e6, 'm': 1e6, 'bn': 1e9, 'billion': 1e9, 'billions': 1e9, 'b': 1e9,
          'cr': 1e7, 'crore': 1e7, 'crores': 1e7, 'lakh': 1e5, 'lakhs': 1e5, 'lac': 1e5, 'lacs': 1e5, 'k': 1e3, 'thousand': 1e3}
_UP = re.compile(r'\b(up|increas\w*|grew|grow\w*|rose|rise\w*|higher|gain\w*|jump\w*|surg\w*|expand\w*|improv\w*|accelerat\w*)\b', re.I)
_DOWN = re.compile(r'\b(down|decreas\w*|declin\w*|fell|fall\w*|lower|drop\w*|contract\w*|reduc\w*|shrink\w*|shrank|slow\w*)\b', re.I)
_CAGR = re.compile(r'\bcagr\b|compound(?:ed)? annual', re.I)
_CLAUSE_BREAK = re.compile(r'[,;]|\band\b|\bbut\b|\bwhile\b|\bwhereas\b', re.I)
_STOP = {'the', 'a', 'an', 'of', 'in', 'to', 'and', 'for', 'is', 'was', 'by', 'at', 'on', 'from', 'with', 'as', 'total', 'inr'}


def _decimals(num_text):
    return len(num_text.split('.')[1]) if '.' in num_text else 0


def extract_claims(text):
    """Return numeric tokens with span, sentence, scale/percent/multiple flags."""
    out = []
    for m in _NUMBER.finditer(text):
        start, end = m.start('num'), m.end('num')
        before = text[start - 1] if start > 0 else ''
        if m.group('cur') is None and m.group('open') is None and m.group('sign') is None and (
                before.isalpha() or before in '.' or before.isdigit()):
            continue                                    # Q3, FY23, 1.2.3
        if before in '-' and start >= 2 and text[start - 2].isdigit():
            continue                                    # 2021-22 style ranges
        after = text[m.end('num'):m.end('num') + 1]
        if after.isalpha() and m.group('suffix') is None:
            continue                                    # 3D, 2nd
        raw = m.group('num')
        value = float(raw.replace(',', ''))
        suffix = (m.group('suffix') or '').strip()
        low = suffix.lower()
        negative = bool(m.group('open') and (m.group('close') or m.group('close2'))) or (m.group('sign') in ('-', '−', '–'))
        sentence_start = max(text.rfind('. ', 0, start), text.rfind('\n', 0, start)) + 1
        sentence_end = min([i for i in (text.find('. ', end), text.find('\n', end)) if i >= 0] or [len(text)])
        sentence = text[sentence_start:sentence_end].strip()
        left = max([sentence_start] + [b.end() for b in _CLAUSE_BREAK.finditer(text, sentence_start, start)])
        right = min([sentence_end] + [b.start() for b in _CLAUSE_BREAK.finditer(text, end, sentence_end)][:1])
        clause = text[left:right]
        is_year = (1900 <= value <= 2100 and '.' not in raw and ',' not in raw and not suffix and not m.group('cur'))
        out.append({'text': m.group(0).strip(), 'span': [m.start(), m.end()], 'value': -value if negative else value,
                    'abs_value': value, 'negative': negative, 'decimals': _decimals(raw), 'raw_digits': raw,
                    'percent': low == '%', 'bps': low in ('bps', 'bp'), 'points': low == 'pp',
                    'multiple': low in ('x', '×'), 'scale_word': suffix if low in _SCALE else None,
                    'scale': _SCALE.get(low), 'currency': (m.group('cur') or '').strip() or None,
                    'is_year': is_year, 'sentence': sentence, 'clause': clause})
    return out


def _tokens(text):
    return {w for w in re.findall(r'[a-z]{3,}', text.lower()) if w not in _STOP}


def _cell_scale(cell):
    unit = parse_unit(cell.get('unit') or '')
    return unit['scale'] if unit and unit['scale'] != 1.0 else (1.0 if unit else None)


def _derive(cells, limit=400):
    """Derived growth/CAGR/ratio cells from cells that share a series and carry numeric periods."""
    series = {}
    for c in cells:
        if c.get('series') is not None and isinstance(c.get('period'), (int, float)) and isinstance(c['value'], (int, float)):
            series.setdefault(c['series'], []).append(c)
    out = []
    for name, items in series.items():
        items.sort(key=lambda x: x['period'])
        for i in range(len(items)):
            for j in range(i + 1, len(items)):
                a, b = items[i], items[j]
                span = b['period'] - a['period']
                tag = f"{a['ref']}->{b['ref']}"
                base = {'label': f"{name} {a['period']}-{b['period']}", 'series': None}
                if a['value'] != 0:
                    out.append({**base, 'ref': f'derived:growth({tag})', 'value': b['value'] / a['value'] - 1, 'unit': None, 'kind': 'growth', 'derived_from': [a['ref'], b['ref']]})
                    out.append({**base, 'ref': f'derived:ratio({tag})', 'value': b['value'] / a['value'], 'unit': 'x', 'kind': 'ratio', 'derived_from': [a['ref'], b['ref']]})
                if a['value'] > 0 and b['value'] > 0 and span > 0:
                    out.append({**base, 'ref': f'derived:cagr({tag})', 'value': (b['value'] / a['value']) ** (1 / span) - 1, 'unit': None, 'kind': 'cagr', 'derived_from': [a['ref'], b['ref']]})
                out.append({**base, 'ref': f'derived:delta({tag})', 'value': b['value'] - a['value'], 'unit': a.get('unit'), 'kind': 'delta', 'derived_from': [a['ref'], b['ref']]})
                if len(out) > limit:
                    return out
    return out


def _candidates(claim, cell):
    """Yield (transform, displayed_value, note) interpretations of ``cell`` for ``claim``."""
    v = float(cell['value'])
    unit = (cell.get('unit') or '').strip().lower()
    pct_cell = unit in ('%', 'pct', 'percent', 'percentage')
    mult_cell = unit in ('x', 'times', 'multiple')
    out = []
    if claim['percent'] or claim['points']:
        if pct_cell:
            out.append(('percent_points', v, None))
        out.append(('fraction_x100', v * 100, None))
        if not pct_cell and not unit:
            out.append(('as_percent_points', v, 'cell has no unit; read as already in percent points'))
    elif claim['bps']:
        out.append(('fraction_x10000', v * 10000, None))
        out.append(('bps', v, None))
    elif claim['multiple']:
        out.append(('multiple', v, None))
    else:
        c_scale = _cell_scale(cell)
        t_scale = claim['scale']
        if t_scale and c_scale:
            out.append(('scale_converted', v * c_scale / t_scale, None))
        elif t_scale and not c_scale and not pct_cell and not mult_cell:
            out.append(('raw_unit_unverified', v, 'cell has no unit; scale word in text unverified'))
            out.append(('absolute_value', v / t_scale, 'cell read as absolute units'))
        elif not t_scale:
            out.append(('raw', v, None if not c_scale or c_scale == 1.0 else 'text states no scale; cell unit scale is not 1'))
    return out


def bind_claims(text, cells, *, ignore_numbers=(), extra_rel_tolerance=0.0, derive=True, max_nearest=3):
    """Bind each number in ``text`` to a reference cell; report unbound and conflicting claims.

    cells: [{'ref','value','unit'?,'label'?,'kind'?('level','growth','delta','ratio','cagr','share'),
             'series'?,'period'?}]. ``unit`` strings such as 'INR Mn', 'INR K', '%', 'x' set the scale expected.
    """
    if not isinstance(text, str):
        raise ValueError('text must be a string')
    cells = [dict(c) for c in cells if isinstance(c.get('value'), (int, float)) and not isinstance(c['value'], bool)
             and math.isfinite(c['value'])]
    pool = cells + (_derive(cells) if derive else [])
    ignore = {float(x) for x in ignore_numbers}
    results, counts = [], {'bound': 0, 'bound_with_warnings': 0, 'mismatch': 0, 'unbound': 0, 'ignored': 0}
    for claim in extract_claims(text):
        base = {'text': claim['text'], 'span': claim['span'], 'sentence': claim['sentence'][:240], 'value': claim['value']}
        if claim['is_year'] or claim['abs_value'] in ignore:
            results.append({**base, 'status': 'ignored', 'reason': 'year' if claim['is_year'] else 'caller_allowlist'})
            counts['ignored'] += 1
            continue
        sentence = claim['sentence']
        up, down = bool(_UP.search(claim['clause'])), bool(_DOWN.search(claim['clause']))
        cagr_word = bool(_CAGR.search(claim['clause']))
        words = _tokens(sentence)
        tol = 0.5 * 10 ** (-claim['decimals']) * 1.000001
        matches, near = [], []
        for cell in pool:
            for name, shown, note in _candidates(claim, cell):
                want = claim['value']
                direction_flexible = (down or up) and not claim['negative']
                diff = abs(shown - want)
                if direction_flexible:
                    diff = min(diff, abs(abs(shown) - claim['abs_value']))
                allowed = tol + extra_rel_tolerance * abs(want)
                overlap = len(words & _tokens(str(cell.get('label') or '') + ' ' + str(cell.get('ref') or '')))
                rec = {'ref': cell['ref'], 'cell_value': cell['value'], 'transform': name, 'shown_as': shown, 'difference': diff,
                       'label': cell.get('label'), 'kind': cell.get('kind'), 'note': note, 'overlap': overlap}
                if diff <= allowed:
                    matches.append((rec, cell))
                else:
                    near.append(rec)
        if matches:
            matches.sort(key=lambda mc: (-mc[0]['overlap'], mc[0]['difference']))
            rec, cell = matches[0]
            notes, status = [], 'bound'
            if rec['note']:
                notes.append(rec['note'])
            distinct = {m[1]['ref'] for m in matches}
            if len(distinct) > 1 and len({round(m[0]['overlap']) for m in matches[:2]}) == 1:
                notes.append(f'ambiguous: {len(distinct)} cells display the same number; best label overlap chosen')
                status = 'bound_with_warnings'
            sign_cell = cell['value']
            conflicts = []
            if claim['negative'] and sign_cell > 0 and rec['transform'] != 'absolute_value':
                conflicts.append('sign_mismatch: text is negative, cell is positive')
            if not claim['negative'] and sign_cell < 0 and not down and not up and rec['shown_as'] < 0 and (claim['percent'] or cell.get('kind') in ('growth', 'delta', 'cagr')):
                conflicts.append('sign_mismatch: text is positive, cell is negative')
            if up and not down and sign_cell < 0 and cell.get('kind') in ('growth', 'delta', 'cagr', None) and (claim['percent'] or cell.get('kind')):
                conflicts.append('direction_conflict: text says up/increase, cell is negative')
            if down and not up and sign_cell > 0 and cell.get('kind') in ('growth', 'delta', 'cagr'):
                conflicts.append('direction_conflict: text says down/decrease, signed growth cell is positive')
            if cagr_word and claim['percent'] and not (cell.get('kind') == 'cagr' or 'cagr' in str(cell.get('label') or '').lower()):
                conflicts.append('comparative_kind_mismatch: text says CAGR, bound cell is not a CAGR')
            if conflicts:
                status = 'mismatch'
                notes += conflicts
            if rec['transform'] in ('raw_unit_unverified', 'absolute_value', 'as_percent_points') and status == 'bound':
                status = 'bound_with_warnings'
            results.append({**base, 'status': status, 'binding': {k: rec[k] for k in ('ref', 'cell_value', 'transform', 'shown_as', 'label', 'kind')},
                            'other_matches': sorted(distinct - {cell['ref']})[:5], 'notes': notes})
            counts[status] += 1
            continue
        # No numeric match: check for scale-only conflicts (right digits, wrong scale) before reporting unbound.
        scale_conflict = None
        if claim['scale'] or claim['currency']:
            for cell in pool:
                if abs(float(cell['value']) - claim['value']) <= tol and (_cell_scale(cell) or 1.0) != (claim['scale'] or 1.0) and _cell_scale(cell):
                    scale_conflict = cell
                    break
        sign_only = [r for r in near if abs(abs(r['shown_as']) - claim['abs_value']) <= tol and r['shown_as'] != 0]
        if sign_only:
            r = min(sign_only, key=lambda x: -x['overlap'])
            results.append({**base, 'status': 'mismatch', 'binding': {k: r[k] for k in ('ref', 'cell_value', 'transform', 'shown_as', 'label', 'kind')},
                            'notes': [f"sign_mismatch: magnitude equals {r['ref']} but the sign differs (cell shows {r['shown_as']:g})"]})
            counts['mismatch'] += 1
            continue
        near.sort(key=lambda r: r['difference'])
        nearest = [{k: r[k] for k in ('ref', 'cell_value', 'transform', 'shown_as', 'difference', 'label')} for r in near[:max_nearest]]
        if scale_conflict is not None:
            results.append({**base, 'status': 'mismatch', 'notes': [f"scale_mismatch: digits equal {scale_conflict['ref']} ({scale_conflict.get('unit')}) but text scale is {claim['scale_word'] or claim['currency']}"],
                            'nearest': [{'ref': scale_conflict['ref'], 'cell_value': scale_conflict['value'], 'unit': scale_conflict.get('unit')}]})
            counts['mismatch'] += 1
        else:
            results.append({**base, 'status': 'unbound', 'nearest': nearest,
                            'notes': ['no cell in the supplied table displays this number within its rounding']})
            counts['unbound'] += 1
    return {'contract': 'claim-binding/v1', 'claims': results, 'counts': counts, 'cells_considered': len(pool),
            'all_bound': counts['unbound'] == 0 and counts['mismatch'] == 0,
            'note': 'binding shows the number exists in the table, not that the wording or the source data are right'}


# ---------------------------------------------------------------- documents
def document_texts(content, suffix):
    """Text units (slide title/shapes/table cells or DOCX paragraphs/table cells) with locators."""
    suffix = suffix.lower().lstrip('.')
    units = []
    if suffix == 'pptx':
        from pptx import Presentation
        deck = Presentation(BytesIO(content))
        for number, slide in enumerate(deck.slides, 1):
            for shape in slide.shapes:
                if shape.has_text_frame and shape.text_frame.text.strip():
                    units.append({'locator': f'slide {number}: {shape.name}', 'text': shape.text_frame.text})
                if getattr(shape, 'has_table', False) and shape.has_table:
                    for ri, row in enumerate(shape.table.rows, 1):
                        for ci, cell in enumerate(row.cells, 1):
                            if cell.text.strip():
                                units.append({'locator': f'slide {number}: table r{ri}c{ci}', 'text': cell.text})
    elif suffix == 'docx':
        from docx import Document
        doc = Document(BytesIO(content))
        for i, para in enumerate(doc.paragraphs, 1):
            if para.text.strip():
                units.append({'locator': f'paragraph {i}', 'text': para.text})
        for ti, table in enumerate(doc.tables, 1):
            for ri, row in enumerate(table.rows, 1):
                for ci, cell in enumerate(row.cells, 1):
                    if cell.text.strip():
                        units.append({'locator': f'table {ti} r{ri}c{ci}', 'text': cell.text})
    else:
        raise ValueError('Only .pptx and .docx documents are supported')
    return units


def bind_document(content, suffix, cells, **kwargs):
    reports, totals = [], {}
    for unit in document_texts(content, suffix):
        report = bind_claims(unit['text'], cells, **kwargs)
        for key, n in report['counts'].items():
            totals[key] = totals.get(key, 0) + n
        if any(c['status'] in ('unbound', 'mismatch', 'bound_with_warnings') for c in report['claims']):
            reports.append({'locator': unit['locator'], 'text': unit['text'][:300],
                            'claims': [c for c in report['claims'] if c['status'] in ('unbound', 'mismatch', 'bound_with_warnings')]})
    return {'contract': 'claim-binding/v1', 'counts': totals, 'all_bound': not totals.get('unbound') and not totals.get('mismatch'),
            'units_needing_attention': reports[:200]}


# ------------------------------------------------------------------ charts
_SCALES_TRIED = (1.0, 1e-3, 1e3, 1e-6, 1e6, 100.0, 0.01)


def _near(a, b, rel, absolute):
    return math.isclose(a, b, rel_tol=rel, abs_tol=absolute)


def chart_series(content):
    """Native chart series (cached values, embedded-workbook values) from PPTX bytes."""
    from pptx import Presentation
    from openpyxl import load_workbook
    import re as _re
    ns = {'c': 'http://schemas.openxmlformats.org/drawingml/2006/chart'}
    deck = Presentation(BytesIO(content))
    charts = []
    for number, slide in enumerate(deck.slides, 1):
        index = 0
        for shape in slide.shapes:
            if not getattr(shape, 'has_chart', False) or not shape.has_chart:
                continue
            index += 1
            chart = shape.chart
            embedded = None
            try:
                blob = chart.part.chart_workbook.xlsx_part.blob
                embedded = load_workbook(BytesIO(blob), data_only=True)
            except Exception:
                embedded = None
            info = {'slide': number, 'chart_index': index, 'chart_type': str(chart.chart_type), 'embedded_workbook': embedded is not None, 'series': []}
            for plot in chart.plots:
                cats = [str(c) for c in plot.categories]
                for ser in plot.series:
                    values = [None if v is None else float(v) for v in ser.values]
                    entry = {'name': ser.name, 'categories': cats, 'values': values, 'embedded_values': None, 'formula': None}
                    refs = ser._element.xpath('./c:val/c:numRef/c:f/text()')
                    if refs and embedded is not None:
                        entry['formula'] = refs[0]
                        m = _re.fullmatch(r"'?([^'!]+)'?!\$?([A-Z]+)\$?(\d+)(?::\$?([A-Z]+)\$?(\d+))?", refs[0])
                        if m and m.group(1) in embedded.sheetnames:
                            from openpyxl.utils.cell import column_index_from_string as ci
                            ws = embedded[m.group(1)]
                            c1, r1 = ci(m.group(2)), int(m.group(3))
                            c2, r2 = (ci(m.group(4)), int(m.group(5))) if m.group(4) else (c1, r1)
                            entry['embedded_values'] = [None if ws.cell(r, c).value is None else float(ws.cell(r, c).value)
                                                        for r in range(r1, r2 + 1) for c in range(c1, c2 + 1)]
                    info['series'].append(entry)
            charts.append(info)
    return charts


def _vectors(ev):
    from .formula_engine import _is_num
    for sd in ev.model.sheets.values():
        for r in range(1, sd.max_row + 1):
            row = [(c, ev.cell_value(sd.name, r, c)) for c in range(1, sd.max_col + 1)]
            yield sd.name, 'row', r, row
        for c in range(1, sd.max_col + 1):
            col = [(r, ev.cell_value(sd.name, r, c)) for r in range(1, sd.max_row + 1)]
            yield sd.name, 'col', c, col
    del _is_num


def find_series(ev, values, *, rel_tol=1e-6, abs_tol=1e-9, scales=_SCALES_TRIED, max_hits=5):
    """Locate contiguous workbook row/column runs equal to ``values`` (optionally rescaled)."""
    from .formula_engine import _is_num, a1
    n = len(values)
    hits = []
    if n == 0 or any(v is None or not math.isfinite(v) for v in values):
        return hits
    for sheet, axis, fixed, vec in _vectors(ev):
        for start in range(0, len(vec) - n + 1):
            window = vec[start:start + n]
            if not all(_is_num(v) and not isinstance(v, bool) for _, v in window):
                continue
            for scale in scales:
                if all(_near(float(v) * scale, t, rel_tol, abs_tol) for (_, v), t in zip(window, values)):
                    first, last = window[0][0], window[-1][0]
                    ref = (f'{sheet}!{a1(fixed, first)}:{a1(fixed, last)}' if axis == 'row' else f'{sheet}!{a1(first, fixed)}:{a1(last, fixed)}')
                    hits.append({'range': ref, 'workbook_to_chart_scale': scale})
                    break
            if len(hits) >= max_hits:
                return hits
    return hits


def chart_workbook_agreement(content, ev, *, bindings=None, display_decimals=None, rel_tol=1e-6, abs_tol=1e-9):
    """Compare native chart series with the evaluated workbook.

    bindings (optional): [{'slide','chart_index','series':[{'name','sheet','address'}]}] pins chart series to ranges;
    otherwise each series is searched for in the workbook (also at common unit scales, reported). With
    ``display_decimals`` the tolerance is the chart's own rounding (0.5*10^-d).
    """
    from openpyxl.utils.cell import range_boundaries
    charts = chart_series(content)
    pinned = {(b['slide'], b['chart_index']): {s['name']: s for s in b['series']} for b in (bindings or [])}
    if display_decimals is not None:
        abs_tol = max(abs_tol, 0.5 * 10 ** (-display_decimals) * 1.000001)
    report, bad = [], 0
    for chart in charts:
        item = {'slide': chart['slide'], 'chart_index': chart['chart_index'], 'chart_type': chart['chart_type'], 'series': []}
        for s in chart['series']:
            entry = {'name': s['name'], 'points': len(s['values']), 'status': None}
            if s['embedded_values'] is not None:
                same = (len(s['embedded_values']) == len(s['values']) and all(
                    (a is None and b is None) or (a is not None and b is not None and _near(a, b, rel_tol, abs_tol))
                    for a, b in zip(s['values'], s['embedded_values'])))
                entry['cache_vs_embedded_workbook'] = 'agree' if same else 'differ'
                if not same:
                    entry['cache_vs_embedded_detail'] = {'cache': s['values'][:12], 'embedded': s['embedded_values'][:12], 'formula': s['formula']}
            pin = pinned.get((chart['slide'], chart['chart_index']), {}).get(s['name'])
            if pin:
                c1, r1, c2, r2 = range_boundaries(pin['address'].replace('$', ''))
                sd = ev.model.sheet(pin['sheet'])
                got = [ev.cell_value(sd.name, r, c) for r in range(r1, r2 + 1) for c in range(c1, c2 + 1)]
                nums = [g for g in got if isinstance(g, (int, float)) and not isinstance(g, bool)]
                ok = len(nums) == len(s['values']) and all(_near(float(g), v, rel_tol, abs_tol) for g, v in zip(nums, s['values']) if v is not None)
                entry.update({'status': 'agrees_with_pinned_range' if ok else 'disagrees_with_pinned_range',
                              'pinned_range': f"{pin['sheet']}!{pin['address']}", 'workbook_values': nums[:24]})
            else:
                hits = find_series(ev, s['values'], rel_tol=rel_tol, abs_tol=abs_tol)
                if hits:
                    entry.update({'status': 'found_in_workbook', 'matches': hits})
                    if all(h['workbook_to_chart_scale'] != 1.0 for h in hits):
                        entry['status'] = 'found_only_after_rescaling'
                else:
                    entry['status'] = 'not_found_in_workbook'
            if entry['status'] in ('disagrees_with_pinned_range', 'not_found_in_workbook', 'found_only_after_rescaling') or entry.get('cache_vs_embedded_workbook') == 'differ':
                bad += 1
            item['series'].append(entry)
        report.append(item)
    return {'contract': 'chart-agreement/v1', 'charts': report, 'chart_count': len(charts), 'series_needing_attention': bad,
            'all_agree': bad == 0 and bool(charts),
            'note': 'agreement means chart numbers equal workbook numbers, not that the right series was charted'}
