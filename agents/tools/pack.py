"""Literal, agent-directed Office draft tools; no financial or prose defaults."""
from functools import wraps
from copy import deepcopy
import hashlib
import json
from pathlib import Path
import os
import tempfile

from .contracts import observation


def tool(name):
    def decorate(method):
        @wraps(method)
        def call(self, *args, **kwargs):
            arguments = {'args': args, 'kwargs': kwargs}
            before = None
            try:
                json.dumps(arguments, allow_nan=False)
                # Audit availability is required before any mutable authoring.
                with self._path('pack.trajectory.jsonl').open('a', encoding='utf-8'):
                    pass
                before = deepcopy((self.deck, self.memo, self.bindings, self.memo_bindings))
                result = observation(name, method(self, *args, **kwargs))
            except Exception as exc:
                if before is not None:
                    self.deck, self.memo, self.bindings, self.memo_bindings = before
                result = observation(name, errors=[{'code': type(exc).__name__, 'message': str(exc)}])
                try:
                    json.dumps(arguments, allow_nan=False)
                except (ValueError, TypeError):
                    arguments = {'invalid_arguments_repr': repr(arguments)}
            try:
                with self._path('pack.trajectory.jsonl').open('a', encoding='utf-8') as stream:
                    stream.write(json.dumps({'tool': name, **arguments,
                                             'observation': result}, allow_nan=False) + '\n')
            except Exception as exc:
                if before is not None:
                    self.deck, self.memo, self.bindings, self.memo_bindings = before
                result['ok'] = False
                result['errors'].append({'code': 'TrajectoryWriteFailed', 'message': str(exc)})
            return result
        return call
    return decorate


class PackTools:
    def __init__(self, sandbox_root, *, workbook=None, inline_bindings=True):
        from pptx import Presentation
        from docx import Document
        self.root = Path(sandbox_root).resolve(strict=True)
        self.deck, self.memo = Presentation(), Document()
        self.bindings, self.memo_bindings = [], []
        self.workbook_path = self._path(workbook) if workbook is not None else None
        self.inline_bindings = inline_bindings
        from docx.shared import Pt, RGBColor
        self.memo.styles['Normal'].font.name = 'Calibri'
        self.memo.styles['Normal'].font.size = Pt(11)
        self.memo.styles['Normal'].paragraph_format.space_after = Pt(6)
        for style in ('Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3'):
            self.memo.styles[style].font.color.rgb = RGBColor(0,0,0)

    def _path(self, name):
        path = (self.root / name).resolve()
        if not path.is_relative_to(self.root):
            raise ValueError('Path escapes private sandbox')
        return path

    @staticmethod
    def _text(text, maximum):
        if not isinstance(text, str) or not text.strip() or len(text) > maximum:
            raise ValueError('Literal text exceeds bounded layout')
        return text

    def _refs(self, refs):
        if not isinstance(refs, list) or not refs or len(refs) > 100:
            raise ValueError('Explicit workbook cell references required')
        for ref in refs:
            if (not isinstance(ref, dict) or set(ref) != {'workbook_sha256', 'sheet', 'cell'}
                    or len(ref['workbook_sha256']) != 64
                    or any(c not in '0123456789abcdef' for c in ref['workbook_sha256'])
                    or not ref['sheet'] or not ref['cell']):
                raise ValueError('Exact workbook hash, sheet and cell binding required')
        if self.workbook_path is not None:
            from openpyxl import load_workbook
            from openpyxl.utils.cell import coordinate_from_string
            if not 0 < self.workbook_path.stat().st_size <= 20_000_000:
                raise ValueError('Bound workbook exceeds byte budget')
            with self.workbook_path.open('rb') as stream:
                content = stream.read(20_000_001)
            if not 0 < len(content) <= 20_000_000:
                raise ValueError('Bound workbook exceeds byte budget')
            from .contracts import bounded_package
            bounded_package(content, max_bytes=20_000_000, max_cells=10000)
            digest = hashlib.sha256(content).hexdigest()
            from io import BytesIO
            book = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
            try:
                for ref in refs:
                    if ref['workbook_sha256'] != digest or ref['sheet'] not in book.sheetnames:
                        raise ValueError('Workbook identity differs from cell reference')
                    coordinate_from_string(ref['cell'])
                    if book[ref['sheet']][ref['cell']].value is None:
                        raise ValueError('Cell reference is empty or unavailable')
            finally:
                book.close()
        # Identity checks do not establish semantic or numerical agreement.
        return deepcopy(refs)

    @tool('deck.add_slide')
    def add_slide(self, layout='title'):
        if layout not in {'title', 'divider', 'table', 'chart', 'chart_callout', 'team_grid'}:
            raise ValueError('Unsupported slide layout')
        if len(self.deck.slides) >= 40:
            raise ValueError('Slide bound exceeded')
        slide = self.deck.slides.add_slide(self.deck.slide_layouts[6])
        self.bindings.append({'layout': layout, 'titles': [], 'tables': [], 'charts': []})
        return {'slide': len(self.deck.slides), 'layout': layout, 'status': 'draft'}

    @tool('deck.title')
    def title(self, slide, text, cell_refs):
        from pptx.util import Inches, Pt
        text, refs = self._text(text, 140), self._refs(cell_refs)
        if type(slide) is not int or not 1 <= slide <= len(self.deck.slides):
            raise ValueError('Invalid slide')
        shape = self.deck.slides[slide-1].shapes.add_textbox(Inches(.4), Inches(.25), Inches(9.2), Inches(.8))
        shape.text_frame.text = text
        shape.text_frame.paragraphs[0].font.size = Pt(24)
        self.bindings[slide-1]['titles'].append({'text': text, 'cell_refs': refs})
        return {'slide': slide, 'status': 'draft_binding_unverified'}

    @tool('deck.table')
    def table(self, slide, values, cell_refs):
        from pptx.util import Inches, Pt
        refs = self._refs(cell_refs)
        if (not isinstance(values, list) or not 1 <= len(values) <= 12
                or not isinstance(values[0], list) or not 1 <= len(values[0]) <= 8
                or any(not isinstance(row, list) or len(row) != len(values[0]) for row in values)
                or any(not isinstance(v, str) or len(v) > 120 for row in values for v in row)):
            raise ValueError('Bounded literal table required')
        if type(slide) is not int or not 1 <= slide <= len(self.deck.slides):
            raise ValueError('Invalid slide')
        table = self.deck.slides[slide-1].shapes.add_table(len(values), len(values[0]), Inches(.4), Inches(1.3), Inches(9.2), Inches(5)).table
        for r, row in enumerate(values):
            for c, value in enumerate(row):
                table.cell(r, c).text = value
                for paragraph in table.cell(r, c).text_frame.paragraphs:
                    paragraph.font.size = Pt(12)
        self.bindings[slide-1]['tables'].append({'values': deepcopy(values), 'cell_refs': refs})
        return {'slide': slide, 'status': 'draft_binding_unverified'}

    @tool('deck.chart')
    def chart(self, slide, categories, series, chart_type, cell_refs):
        import math
        from pptx.chart.data import CategoryChartData
        from pptx.enum.chart import XL_CHART_TYPE
        from pptx.util import Inches
        refs = self._refs(cell_refs)
        types = {'bar': XL_CHART_TYPE.COLUMN_CLUSTERED, 'line': XL_CHART_TYPE.LINE,
                 'stacked': XL_CHART_TYPE.COLUMN_STACKED, 'pie': XL_CHART_TYPE.PIE}
        if chart_type not in types:
            raise ValueError('Unsupported native chart type; waterfall is not qualified')
        if (not isinstance(categories, list) or not 1 <= len(categories) <= 24
                or any(not isinstance(v, str) or len(v) > 100 for v in categories)
                or not isinstance(series, list) or not 1 <= len(series) <= 8):
            raise ValueError('Bounded chart data required')
        data = CategoryChartData()
        data.categories = categories
        for item in series:
            name = self._text(item['name'], 100)
            values = item['values']
            if (len(values) != len(categories) or any(type(v) not in (int, float) or not math.isfinite(v) for v in values)):
                raise ValueError('Finite numeric chart series required')
            data.add_series(name, values)
        if chart_type == 'pie' and len(series) != 1:
            raise ValueError('Pie chart requires one series')
        if type(slide) is not int or not 1 <= slide <= len(self.deck.slides):
            raise ValueError('Invalid slide')
        self.deck.slides[slide-1].shapes.add_chart(types[chart_type], Inches(.4), Inches(1.3), Inches(9.2), Inches(5), data)
        self.bindings[slide-1]['charts'].append({'categories': deepcopy(categories), 'series': deepcopy(series), 'cell_refs': refs})
        return {'slide': slide, 'embedded_data': True, 'status': 'draft_binding_unverified'}

    @tool('memo.write')
    def write(self, heading, text, cell_refs):
        heading, text, refs = self._text(heading, 140), self._text(text, 10000), self._refs(cell_refs)
        if len(self.memo.paragraphs) > 200:
            raise ValueError('Memo section bound exceeded')
        self.memo.add_heading(heading, level=1)
        self.memo.add_paragraph(text)
        if self.inline_bindings:
            self.memo.add_paragraph('Cell bindings: ' + json.dumps(refs, sort_keys=True))
        self.memo_bindings.append({'heading': heading, 'text': text, 'cell_refs': refs})
        return {'status': 'draft_binding_unverified'}

    @tool('memo.write_from_cells')
    def write_from_cells(self, heading, text, values):
        """Render agent-chosen numeric placeholders from its exact workbook.

        Prose, cell choices, units and precision are supplied by the agent. The
        tool does not choose financial meaning or repair a failed conclusion.
        Values are diagnostic calculations until Office qualification passes.
        """
        import re
        import math
        from io import BytesIO
        from openpyxl import load_workbook
        from .workbook import WorkbookTools
        if self.workbook_path is None:
            raise ValueError('A bound workbook is required')
        heading, text = self._text(heading, 140), self._text(text, 10000)
        if not isinstance(values, list) or not 1 <= len(values) <= 100:
            raise ValueError('Explicit numeric cell selections required')
        if self.workbook_path.stat().st_size > 20_000_000:
            raise ValueError('Bound workbook exceeds byte budget')
        with self.workbook_path.open('rb') as stream:
            content = stream.read(20_000_001)
        if len(content) > 20_000_000:
            raise ValueError('Bound workbook exceeds byte budget')
        from .contracts import bounded_package
        bounded_package(content, max_bytes=20_000_000, max_cells=10000)
        workbook_hash = hashlib.sha256(content).hexdigest()
        book = load_workbook(BytesIO(content), data_only=False, keep_links=False)
        diagnostic = WorkbookTools.__new__(WorkbookTools)
        diagnostic.root, diagnostic.book, diagnostic.max_cells = self.root, book, 10000
        calculation = WorkbookTools.recalc.__wrapped__(diagnostic)
        computed = {(row['sheet'], row['cell']): row['value'] for row in calculation['values']}
        selected, refs = {}, []
        for value in values:
            if not isinstance(value, dict) or set(value) != {'key', 'sheet', 'cell', 'unit', 'decimals'}:
                raise ValueError('Numeric selections require key, sheet, cell, unit and decimals')
            key = value['key']
            if not isinstance(key, str) or not re.fullmatch(r'[A-Za-z][A-Za-z0-9_.-]{0,79}', key) or key in selected:
                raise ValueError('Numeric placeholder keys must be unique')
            if type(value['decimals']) is not int or not 0 <= value['decimals'] <= 6:
                raise ValueError('Numeric precision must be between zero and six')
            if not isinstance(value['unit'], str) or not value['unit'] or len(value['unit']) > 100:
                raise ValueError('Explicit numeric unit required')
            ref = {'workbook_sha256': workbook_hash, 'sheet': value['sheet'], 'cell': value['cell']}
            self._refs([ref])
            cell = book[value['sheet']][value['cell']]
            number = computed.get((value['sheet'], value['cell'])) if cell.data_type == 'f' else cell.value
            if type(number) not in (int, float) or not math.isfinite(number):
                raise ValueError('Selected cell has no supported finite numeric value')
            selected[key] = {'text': format(number, f".{value['decimals']}f"), **value}
            refs.append(ref)
        matches = list(re.finditer(r'\{\{([A-Za-z][A-Za-z0-9_.-]*)\}\}', text))
        if set(m.group(1) for m in matches) != set(selected):
            raise ValueError('Every selected numeric cell must have an exact placeholder')
        chunks, bindings, cursor, length = [], [], 0, 0
        paragraph = len(self.memo.paragraphs) + 1
        for match in matches:
            prefix = text[cursor:match.start()]
            chunks.append(prefix)
            length += len(prefix)
            item = selected[match.group(1)]
            chunks.append(item['text'])
            bindings.append({'paragraph': paragraph, 'start': length, 'end': length+len(item['text']),
                             'metric_id': item['key'], 'unit': item['unit'],
                             'sheet': item['sheet'], 'cell': item['cell'], 'decimals': item['decimals']})
            length += len(item['text'])
            cursor = match.end()
        chunks.append(text[cursor:])
        rendered = ''.join(chunks)
        result = PackTools.write.__wrapped__(self, heading, rendered, refs)
        self.memo_bindings[-1]['numeric_bindings'] = bindings
        return {**result, 'memo_bindings': bindings, 'calculation_status': 'diagnostic_only',
                'calculation_findings': calculation['findings'], 'workbook_sha256': workbook_hash}

    @tool('pack.export')
    def export(self, stem):
        return self._export_bundle(stem, include_deck=True)

    @tool('memo.export')
    def export_memo(self, stem):
        """Commit an authored memo draft without inventing a placeholder deck."""
        return self._export_bundle(stem, include_deck=False)

    def _export_bundle(self, stem, *, include_deck):
        import re
        if not isinstance(stem, str) or not re.fullmatch(r'[A-Za-z0-9][A-Za-z0-9_-]{0,79}', stem):
            raise ValueError('A simple new export stem is required')
        destination = self._path(stem + ('.pack' if include_deck else '.memo'))
        if destination.exists():
            raise FileExistsError('Exports are immutable; use a new version stem')
        if include_deck and not self.deck.slides:
            raise ValueError('A deck slide is required')
        if not include_deck and not self.memo_bindings:
            raise ValueError('An authored memo section is required')
        all_refs = [ref for slide in self.bindings for items in ('titles', 'tables', 'charts')
                    for item in slide[items] for ref in item['cell_refs']]
        all_refs += [ref for section in self.memo_bindings for ref in section['cell_refs']]
        if self.workbook_path is not None and all_refs:
            # An intervening workbook mutation invalidates this export version.
            for offset in range(0, len(all_refs), 100):
                self._refs(all_refs[offset:offset+100])
        # Stage all bytes privately; one directory rename publishes the complete
        # draft. A failed stage is retained as diagnostic evidence, never accepted.
        stage = Path(tempfile.mkdtemp(prefix='.' + stem + '-staging-', dir=self.root))
        os.chmod(stage, 0o700)
        suffixes = ('.pptx', '.docx', '.bindings.json') if include_deck else ('.docx', '.bindings.json')
        paths = [stage / (stem + suffix) for suffix in suffixes]
        try:
            if include_deck:
                self.deck.save(stage / (stem + '.pptx'))
            self.memo.save(stage / (stem + '.docx'))
            paths[-1].write_text(json.dumps({'slides': self.bindings if include_deck else [], 'memo_sections': self.memo_bindings,
                'reference_identity': 'checked' if self.workbook_path is not None else 'unverified',
                'status': 'draft_binding_unverified', 'visual_validation': 'not_run', 'release': 'blocked'}, indent=2))
            files = [{'name': path.name, 'sha256': hashlib.sha256(path.read_bytes()).hexdigest(),
                      'bytes': path.stat().st_size} for path in paths]
            manifest = {'contract': 'agentic-pack-export/v1' if include_deck else 'agentic-memo-export/v1', 'files': files,
                        'workbook_sha256': sorted(set(ref['workbook_sha256'] for ref in all_refs)),
                        'release': 'blocked', 'pdf': 'not_generated', 'visual_validation': 'not_run'}
            (stage / 'manifest.json').write_text(json.dumps(manifest, indent=2, sort_keys=True))
            for path in stage.iterdir():
                path.chmod(0o600)
            if destination.exists():
                raise FileExistsError('Export revision already exists')
            stage.rename(destination)
        except Exception as exc:
            (stage / 'failure.json').write_text(json.dumps({'status': 'failed', 'error_type': type(exc).__name__}))
            raise
        digest = hashlib.sha256((destination / 'manifest.json').read_bytes()).hexdigest()
        return {'files': [str(Path(destination.name) / p.name) for p in paths],
                'manifest': destination.name + '/manifest.json', 'manifest_sha256': digest,
                'release': 'blocked', 'pdf': 'not_generated'}

    @tool('deck.render_preview')
    def render_preview(self):
        raise NotImplementedError('Preview rendering and text-fit inspection are not qualified for this tool')
