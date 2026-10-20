"""Artifact-derived, bounded financial diagnostic audit. No inference or repairs.

Independent expectations are a caller-owned answer key, never an inferred or
model-declared approval. Passing this audit does not qualify Office recalculation,
visuals, financial assumptions, source truth or release.
"""
from __future__ import annotations

from decimal import Decimal, InvalidOperation
from hashlib import sha256
from io import BytesIO
import json
import math
from pathlib import Path
from zipfile import ZipFile
from defusedxml import ElementTree as ET
import re

from docx import Document
from openpyxl import load_workbook
from openpyxl.utils.cell import coordinate_from_string, range_boundaries

from .contracts import bounded_package
from .workbook import WorkbookTools

NUMBER_TOKEN = re.compile(r'(?<![A-Za-z0-9_.])(?:\(?[-+]?(?:\d[\d,]*(?:\.\d+)?|\.\d+)(?:[eE][-+]?\d+)?\)?%?)(?![A-Za-z0-9_])')


def audit_financial_artifacts(root, *, workbook, memo, output_spec, expectations, sources,
                              max_bytes=20_000_000, max_cells=10000):
    """Audit retained XLSX/DOCX bytes; return a versioned, non-release report.

    output_spec: metrics {id:{sheet,cell,unit}}, calculated_ranges [{sheet,range}],
      memo_bindings [{paragraph,start,end,metric_id,unit}] (zero-based spans).
    expectations: expected_outputs [{id,value,unit,tolerance,require_formula?}],
      required_memo_metrics [ids], checks [{id,kind,inputs,output,tolerance}].
      Optional expected_output.memo_units maps displayed unit to a multiplier.
    sources: nonempty relative paths; retained/hash-bound, not semantically read.
    Every literal memo number must have one exact binding. An independent check
    may use sum/difference/maximum/minimum on mapped artifact cell values.
    """
    root = Path(root).resolve(strict=True)
    failures, observations, hashes = [], {}, {}
    report = {'contract':'financial-artifact-audit/v1','status':'failed','release':'blocked',
              'scope':'bounded_objective_diagnostics','office_recalculation':'not_qualified',
              'source_semantics':'unassessed','financial_acceptance':'not_established',
              'hashes':hashes,'observations':observations,'failures':failures}
    def fail(code, **details):
        failures.append({'code':code,**details})
    def read(name):
        path = (root / name).resolve(strict=True)
        if not path.is_relative_to(root):
            raise ValueError('Artifact/source escapes sandbox')
        with path.open('rb') as stream:
            content = stream.read(max_bytes+1)
        if not 0 < len(content) <= max_bytes:
            raise ValueError('Artifact/source exceeds byte budget')
        hashes[str(name)] = sha256(content).hexdigest()
        return content
    def finite(value):
        return type(value) in (int,float) and math.isfinite(value)
    def same(actual,expected,tolerance):
        return finite(actual) and finite(expected) and finite(tolerance) and tolerance >= 0 and abs(actual-expected) <= tolerance
    try:
        # Bind exact specs too; they must be retained separately by the caller.
        for key,value in [('output_spec',output_spec),('expectations',expectations)]:
            hashes[key] = sha256(json.dumps(value,sort_keys=True,allow_nan=False,separators=(',',':')).encode()).hexdigest()
        if not isinstance(sources,list) or not sources:
            raise ValueError('At least one retained source is required')
        for source in sources:
            read(source)
        xlsx, docx = read(workbook), read(memo)
        bounded_package(xlsx,max_bytes=max_bytes,max_cells=max_cells)
        bounded_package(docx,max_bytes=max_bytes)
        book = load_workbook(BytesIO(xlsx),data_only=False,keep_links=False)
        cached = load_workbook(BytesIO(xlsx),data_only=True,keep_links=False)
        original_cells = [(ws.title,c) for ws in book for c in list(ws._cells.values())]
        formulas = [(s,c.coordinate) for s,c in original_cells if c.data_type == 'f']
        observations['formula_cells'] = [{'sheet':s,'cell':c} for s,c in formulas]
        observations['hidden_sheets'] = [ws.title for ws in book if ws.sheet_state != 'visible']
        if not formulas:
            fail('no_live_formulas')
        for s,cell in original_cells:
            cached_cell = cached[s][cell.coordinate]
            if cell.data_type == 'e' or cached_cell.data_type == 'e':
                fail('stored_cell_error',sheet=s,cell=cell.coordinate)
        reader = object.__new__(WorkbookTools)
        reader.root,reader.path,reader.book,reader.max_cells = root,root/workbook,book,max_cells
        calculation = WorkbookTools.recalc.__wrapped__(reader)
        observations['calculation'] = calculation
        for finding in calculation['findings']:
            fail('formula_evaluation_failed',finding=finding)
        evaluated = {(v['sheet'],v['cell']):v['value'] for v in calculation['values']}
        metrics = output_spec['metrics']
        if not isinstance(metrics,dict) or not metrics or len(metrics) > 1000:
            raise ValueError('Nonempty bounded metric mapping required')
        actual = {}
        coordinates = {}
        for metric,mapping in metrics.items():
            if set(mapping) != {'sheet','cell','unit'} or not isinstance(mapping['unit'],str) or not mapping['unit']:
                fail('invalid_metric_mapping',metric=metric)
                continue
            coordinate_from_string(mapping['cell'])
            sheet,cell = mapping['sheet'],mapping['cell'].replace('$','').upper()
            if sheet not in book.sheetnames:
                fail('missing_metric_sheet',metric=metric)
                continue
            target = book[sheet][cell]
            value = evaluated.get((sheet,cell)) if target.data_type == 'f' else target.value
            if not finite(value):
                fail('metric_value_unavailable',metric=metric,sheet=sheet,cell=cell)
                continue
            actual[metric], coordinates[metric] = value,(sheet,cell)
        observations['metric_values'] = actual
        calculated = output_spec['calculated_ranges']
        if not isinstance(calculated,list) or not calculated:
            fail('calculated_ranges_missing')
        else:
            total = 0
            for entry in calculated:
                if set(entry) != {'sheet','range'} or entry['sheet'] not in book.sheetnames:
                    fail('invalid_calculated_range')
                    continue
                a,b,c,d = range_boundaries(entry['range'])
                if None in (a,b,c,d) or not (1<=a<=c<=16384 and 1<=b<=d<=1048576):
                    fail('invalid_calculated_range')
                    continue
                total += (c-a+1)*(d-b+1)
                if total > max_cells:
                    raise ValueError('Calculated ranges exceed cell budget')
                for row in book[entry['sheet']].iter_rows(min_row=b,max_row=d,min_col=a,max_col=c):
                    for cell in row:
                        if cell.data_type != 'f':
                            fail('hardcode_or_missing_calculation',sheet=entry['sheet'],cell=cell.coordinate)
        expected_rows = expectations['expected_outputs']
        if not isinstance(expected_rows,list) or not expected_rows:
            raise ValueError('Independent expected outputs are required')
        expected_by_id = {}
        for expected in expected_rows:
            metric = expected['id']
            if metric in expected_by_id:
                fail('duplicate_expected_metric',metric=metric)
                continue
            expected_by_id[metric] = expected
            if metric not in actual or metrics[metric]['unit'] != expected['unit']:
                fail('expected_metric_or_unit_missing',metric=metric)
                continue
            s,c = coordinates[metric]
            comment = book[s][c].comment.text if book[s][c].comment else ''
            status = expected.get('status')
            if status not in ('historical','company_projection','illustrative'):
                fail('status_requirements_missing',metric=metric)
            elif re.search(r'(?i)status\s*:\s*'+re.escape(status)+r'\b',comment) is None:
                fail('status_metadata_missing',metric=metric)
            if re.search(r'(?i)unit\s*:\s*'+re.escape(expected['unit'])+r'(?=\s|[;,]|$)',comment) is None:
                fail('unit_metadata_missing',metric=metric)
            if not same(actual[metric],expected['value'],expected['tolerance']):
                fail('expected_value_mismatch',metric=metric,actual=actual[metric],expected=expected['value'])
            if expected.get('require_formula',False):
                s,c = coordinates[metric]
                if book[s][c].data_type != 'f':
                    fail('expected_formula_missing',metric=metric)
        forward = any(e.get('status') in ('company_projection','illustrative') for e in expected_rows)
        assumptions = output_spec.get('assumptions')
        if forward and (not isinstance(assumptions,list) or not assumptions):
            fail('assumptions_missing')
        elif assumptions is not None:
            if not isinstance(assumptions,list) or len(assumptions) > 1000:
                raise ValueError('Assumption map exceeds budget')
            for assumption in assumptions:
                if set(assumption) != {'sheet','cell'}:
                    fail('assumption_mapping_invalid')
                    continue
                cell = book[assumption['sheet']][assumption['cell']]
                comment = cell.comment.text if cell.comment else ''
                if not finite(cell.value) or cell.data_type == 'f':
                    fail('assumption_input_missing',sheet=assumption['sheet'],cell=assumption['cell'])
                if not re.search(r'(?i)(rationale|assumption)\s*:\s*\S',comment):
                    fail('assumption_rationale_missing',sheet=assumption['sheet'],cell=assumption['cell'])
                if not any(label in comment for label in ('company_projection','illustrative')):
                    fail('assumption_status_missing',sheet=assumption['sheet'],cell=assumption['cell'])
                if not re.search(r'(?i)unit\s*:\s*\S',comment):
                    fail('assumption_unit_missing',sheet=assumption['sheet'],cell=assumption['cell'])
        checks = expectations['checks']
        if not isinstance(checks,list) or not checks:
            fail('independent_arithmetic_checks_missing')
        else:
            seen = set()
            for check in checks:
                identifier = check['id']
                if identifier in seen:
                    fail('duplicate_check',check=identifier)
                    continue
                seen.add(identifier)
                inputs,output,kind = check['inputs'],check['output'],check['kind']
                if not inputs or any(key not in actual for key in inputs+[output]):
                    fail('arithmetic_check_mapping_missing',check=identifier)
                    continue
                if len({metrics[key]['unit'] for key in inputs+[output]}) != 1:
                    fail('arithmetic_check_units_mismatch',check=identifier)
                    continue
                values = [actual[key] for key in inputs]
                if kind == 'sum':
                    expected = sum(values)
                elif kind == 'difference' and len(values) == 2:
                    expected = values[0]-values[1]
                elif kind == 'maximum':
                    expected = max(values)
                elif kind == 'minimum':
                    expected = min(values)
                else:
                    fail('arithmetic_check_unsupported',check=identifier)
                    continue
                if not same(actual[output],expected,check['tolerance']):
                    fail('arithmetic_check_failed',check=identifier,actual=actual[output],expected=expected)
        document = Document(BytesIO(docx))
        # Paragraph order: body paragraphs first, then table-cell paragraphs.
        paragraphs = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    paragraphs.extend(p.text for p in cell.paragraphs)
        with ZipFile(BytesIO(docx)) as package:
            native = ET.fromstring(package.read('word/document.xml'))
            ns = {'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
            # Hyperlinks, textboxes or other content not projected by python-docx
            # must never silently escape numeric coverage.
            retained = ''.join(paragraphs)
            original = ''.join(n.text or '' for n in native.findall('.//w:t',ns))
            if retained != original:
                fail('memo_native_text_coverage_unknown')
            for name in package.namelist():
                if name.startswith(('word/header','word/footer','word/footnotes','word/endnotes')) and name.endswith('.xml'):
                    xml = ET.fromstring(package.read(name))
                    for paragraph in xml.findall('.//w:p',ns):
                        paragraphs.append(''.join(n.text or '' for n in paragraph.findall('.//w:t',ns)))
        observations['memo_paragraphs'] = paragraphs
        tokens = {}
        for index,text in enumerate(paragraphs):
            if text.startswith('Cell bindings: '):
                try:
                    metadata = json.loads(text[len('Cell bindings: '):])
                    if not isinstance(metadata,list) or not metadata:
                        raise ValueError('Missing reference metadata')
                    for ref in metadata:
                        if (set(ref) != {'workbook_sha256','sheet','cell'} or ref['workbook_sha256'] != hashes[str(workbook)]
                            or ref['sheet'] not in book.sheetnames or book[ref['sheet']][ref['cell']].value is None):
                            raise ValueError('Reference metadata mismatch')
                    continue
                except Exception:
                    fail('memo_reference_metadata_invalid',paragraph=index)
            digit_coverage = set()
            for token in NUMBER_TOKEN.finditer(text):
                tokens[(index,token.start(),token.end())] = token.group()
                digit_coverage.update(range(token.start(),token.end()))
            if any(char.isdigit() and pos not in digit_coverage for pos,char in enumerate(text)):
                fail('memo_numeric_format_unknown',paragraph=index)
        bound,covered_metrics = set(),set()
        bindings = output_spec['memo_bindings']
        if not isinstance(bindings,list) or not bindings:
            fail('memo_numeric_bindings_missing')
        else:
            for binding in bindings:
                if set(binding) != {'paragraph','start','end','metric_id','unit'}:
                    fail('invalid_memo_binding')
                    continue
                key = (binding['paragraph'],binding['start'],binding['end'])
                metric = binding['metric_id']
                if key not in tokens or key in bound or metric not in actual or metric not in expected_by_id:
                    fail('memo_binding_not_exact',metric=metric)
                    continue
                bound.add(key)
                if tokens[key].endswith('%') and binding['unit'] not in ('%','percent','percentage_points'):
                    fail('memo_percent_unit_mismatch',metric=metric)
                    continue
                literal = tokens[key].replace(',','').rstrip('%')
                if literal.startswith('(') and literal.endswith(')'):
                    literal = '-'+literal[1:-1]
                try:
                    displayed = Decimal(literal)
                except InvalidOperation:
                    fail('memo_number_unparseable',metric=metric)
                    continue
                allowed = expected_by_id[metric].get('memo_units',{expected_by_id[metric]['unit']:1})
                multiplier = allowed.get(binding['unit'])
                if not finite(multiplier):
                    fail('memo_unit_unapproved',metric=metric)
                    continue
                expected = Decimal(str(actual[metric]))*Decimal(str(multiplier))
                if abs(displayed-expected) > Decimal(str(expected_by_id[metric]['tolerance'])):
                    fail('memo_number_cell_mismatch',metric=metric,literal=tokens[key],actual=actual[metric])
                else:
                    covered_metrics.add(metric)
        for key in set(tokens)-bound:
            fail('memo_number_unbound',paragraph=key[0],start=key[1],end=key[2],literal=tokens[key])
        required_memo = expectations['required_memo_metrics']
        if not isinstance(required_memo,list) or not required_memo:
            fail('independent_memo_requirements_missing')
        else:
            for metric in required_memo:
                if metric not in covered_metrics:
                    fail('required_memo_metric_missing',metric=metric)
        book.close()
        cached.close()
    except Exception as exc:
        fail('audit_incomplete',error_type=type(exc).__name__,message=str(exc))
    report['status'] = 'objective_checks_passed' if not failures else 'failed'
    return report
