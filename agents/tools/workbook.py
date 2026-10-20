"""Sandboxed workbook authoring. No assumptions or financial content are supplied.

The bounded evaluator is diagnostic only: it supports arithmetic, direct A1
references and SUM ranges, never claims Office qualification or writes caches.
"""
from __future__ import annotations

import ast
from copy import copy
from functools import wraps
import json
import math
import operator
from pathlib import Path
import re
import tempfile
import os
from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils.cell import range_boundaries
from openpyxl.workbook.defined_name import DefinedName

from .contracts import observation, json_safe, bounded_package

REF = re.compile(r"(?<![A-Za-z0-9_.])(?:(?:'([^']+)'|([A-Za-z_][A-Za-z0-9_ ]*))!)?(\$?[A-Z]{1,3}\$?[1-9][0-9]*)(?::(\$?[A-Z]{1,3}\$?[1-9][0-9]*))?(?![A-Za-z0-9_])")


def external_formula(value):
    return isinstance(value, str) and value.startswith('=') and (
        '[' in value or '://' in value or '|' in value or
        re.search(r'(?i)\b(?:WEBSERVICE|HYPERLINK|RTD)\s*\(', value))


def tool(method):
    @wraps(method)
    def call(self, *args, **kwargs):
        original = None
        mutated = method.__name__ in {'add_sheet', 'write', 'write_block', 'name', 'style', 'comment'}
        try:
            if mutated:
                original = self._snapshot()
                # Fail before mutation if audit storage is unavailable.
                self._path(self.log_path).open('a', encoding='utf-8').close()
            result = observation('wb.' + method.__name__, method(self, *args, **kwargs))
        except Exception as exc:
            if original is not None:
                self.book = self._restore(original)
            result = observation('wb.' + method.__name__, errors=[{'code': type(exc).__name__, 'message': str(exc)}])
        try:
            self._log({'call': {'tool': 'wb.' + method.__name__, 'args': args, 'kwargs': kwargs}, 'observation': result})
        except Exception as exc:
            if original is not None and result['ok']:
                self.book = self._restore(original)
                try:
                    self._save()
                except Exception as rollback:
                    result['errors'].append({'code': 'RollbackFailed', 'message': str(rollback)})
            result['ok'] = False
            result['errors'].append({'code': 'TrajectoryWriteFailed', 'message': str(exc)})
        return result
    return call


class WorkbookTools:
    """One private sandbox per instance; source is copied, never edited.

    Existing destinations require resume=True. All paths, including log paths,
    are resolved inside sandbox_root. Callers must provision tenant ownership.
    """
    def __init__(self, sandbox_root, workbook='model.xlsx', *, source=None, resume=False, max_cells=10000, max_bytes=20_000_000):
        self.root = Path(sandbox_root).resolve(strict=True)
        self.path = self._path(workbook)
        self.log_path = self._path('workbook.trajectory.jsonl')
        self.max_cells = max_cells
        self.max_bytes = max_bytes
        if self.path.suffix.lower() != '.xlsx':
            raise ValueError('Only macro-free .xlsx workbooks are supported')
        if self.path.exists() and not resume:
            raise FileExistsError('Destination exists; explicit resume is required')
        if source is not None:
            src = self._path(source)
            if src == self.path or self.path.exists():
                raise ValueError('Source must remain distinct from a new destination')
            self.book = self._load(src)
        elif self.path.exists():
            self.book = self._load(self.path)
        else:
            self.book = Workbook()
        self._save()

    def _path(self, path):
        result = (self.root / path).resolve()
        if not result.is_relative_to(self.root):
            raise ValueError('Path escapes private workbook sandbox')
        return result

    def _load(self, path):
        if path.suffix.lower() != '.xlsx' or path.stat().st_size > self.max_bytes:
            raise ValueError('Source must be a bounded .xlsx package')
        with path.open('rb') as stream:
            content = stream.read(self.max_bytes + 1)
        if len(content) > self.max_bytes:
            raise ValueError('Source exceeds byte budget')
        bounded_package(content, max_bytes=self.max_bytes, max_cells=self.max_cells)
        return load_workbook(BytesIO(content), data_only=False, keep_links=False)

    def _snapshot(self):
        # openpyxl workbooks holding comments cannot be deep-copied; serialize instead.
        buffer = BytesIO()
        self.book.save(buffer)
        return buffer.getvalue()

    def _restore(self, content):
        return load_workbook(BytesIO(content), data_only=False, keep_links=False)

    def _save(self):
        self._path(self.path)
        self.path.parent.mkdir(parents=True, exist_ok=True)
        fd, tmp = tempfile.mkstemp(dir=self.path.parent, suffix='.xlsx')
        os.close(fd)
        try:
            self.book.save(tmp)
            if Path(tmp).stat().st_size > self.max_bytes:
                raise ValueError('Saved workbook exceeds byte budget')
            bounded_package(Path(tmp).read_bytes(), max_bytes=self.max_bytes, max_cells=self.max_cells)
            os.replace(tmp, self.path)
        finally:
            if os.path.exists(tmp):
                os.unlink(tmp)

    def _log(self, result):
        self._path(self.log_path)
        with self.log_path.open('a', encoding='utf-8') as stream:
            stream.write(json.dumps(json_safe(result), default=str, allow_nan=False) + '\n')

    def _cells(self, sheet, address):
        bounds = range_boundaries(address.replace('$', ''))
        a, b, c, d = bounds
        if None in bounds or a < 1 or b < 1 or c < a or d < b or c > 16384 or d > 1048576 or (c-a+1)*(d-b+1) > self.max_cells:
            raise ValueError('Invalid or over-budget cell range')
        return list(self.book[sheet].iter_rows(min_row=b, max_row=d, min_col=a, max_col=c))

    @tool
    def add_sheet(self, title, *, hidden=False):
        if title in self.book.sheetnames:
            raise ValueError('Sheet already exists')
        ws = self.book.create_sheet(title)
        ws.sheet_state = 'hidden' if hidden else 'visible'
        self._save()
        return {'sheet': title, 'hidden': hidden}

    @tool
    def write(self, sheet, address, values):
        return self._write_matrix(sheet,address,values)

    def _write_matrix(self, sheet, address, values, *, comment=None):
        cells = self._cells(sheet, address)
        if not isinstance(values, list) or len(values) != len(cells) or any(not isinstance(row, list) or len(row) != len(cells[0]) for row in values):
            raise ValueError('values must be a rectangular matrix matching the range')
        for row in values:
            for value in row:
                if value is not None and not isinstance(value, (str, int, float, bool)):
                    raise ValueError('Only JSON scalar values are supported')
                if isinstance(value, float) and not math.isfinite(value):
                    raise ValueError('Nonfinite numbers are unsupported')
                if external_formula(value):
                    raise ValueError('External formula references are unsupported')
        annotation = Comment(comment, 'Analyst') if comment is not None else None
        if comment is not None:
            from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
            if not isinstance(comment,str) or ILLEGAL_CHARACTERS_RE.search(comment):
                raise ValueError('A literal XML-safe source comment is required')
        for target_row, value_row in zip(cells, values):
            for cell, value in zip(target_row, value_row):
                cell.value = value
                if annotation is not None:
                    cell.comment = copy(annotation)
        self._save()
        return {'sheet': sheet, 'range': address, 'written': sum(map(len, cells))}

    @tool
    def write_block(self, sheet, start, values, *, comment=None):
        """Paste the supplied rectangle at one agent-chosen anchor.

        The matrix determines its extent, like a spreadsheet paste operation.
        It never changes numbers, formulas, row order or a previous failed call.
        """
        from openpyxl.utils.cell import coordinate_from_string, column_index_from_string, get_column_letter
        if not isinstance(start,str) or not re.fullmatch(r'\$?[A-Z]{1,3}\$?[1-9][0-9]*', start):
            raise ValueError('A single A1 start cell is required')
        if (not isinstance(values,list) or not values or not isinstance(values[0],list) or not values[0]
                or any(not isinstance(row,list) or len(row)!=len(values[0]) for row in values)):
            raise ValueError('A nonempty rectangular matrix is required')
        column, row = coordinate_from_string(start.replace('$',''))
        last_column = column_index_from_string(column) + len(values[0])-1
        last_row = row + len(values)-1
        if last_column > 16384 or last_row > 1048576 or len(values)*len(values[0]) > self.max_cells:
            raise ValueError('Pasted block exceeds worksheet or cell budget')
        address = f'{column}{row}:{get_column_letter(last_column)}{last_row}'
        result = self._write_matrix(sheet,address,values,comment=comment)
        return {**result, 'start':start, 'extent_from':'supplied_matrix'}

    @tool
    def name(self, name, sheet, address):
        self._cells(sheet, address)
        if not re.fullmatch(r'[A-Za-z_][A-Za-z0-9_]*', name) or REF.fullmatch(name.upper()):
            raise ValueError('Invalid defined name')
        escaped = sheet.replace("'", "''")
        item = DefinedName(name, attr_text=f"'{escaped}'!{address}")
        if hasattr(self.book.defined_names, 'add'):
            self.book.defined_names.add(item)
        else:
            self.book.defined_names.append(item)
        self._save()
        return {'name': name, 'sheet': sheet, 'range': address}

    @tool
    def style(self, sheet, address, *, number_format=None, font=None, fill=None, alignment=None):
        cells = self._cells(sheet, address)
        changes = {}
        if font is not None:
            changes['font'] = Font(**font)
        if fill is not None:
            changes['fill'] = PatternFill(**fill)
        if alignment is not None:
            changes['alignment'] = Alignment(**alignment)
        for row in cells:
            for cell in row:
                if number_format is not None:
                    cell.number_format = number_format
                for key, value in changes.items():
                    setattr(cell, key, copy(value))
        self._save()
        return {'sheet': sheet, 'range': address}

    @tool
    def comment(self, sheet, address, text, *, author='Analyst'):
        cells = self._cells(sheet, address)
        for row in cells:
            for cell in row:
                cell.comment = Comment(text, author)
        self._save()
        return {'sheet': sheet, 'range': address}

    @tool
    def read(self, sheet, address):
        return {'sheet': sheet, 'range': address, 'cells': [[{'cell': c.coordinate, 'value': c.value, 'type': c.data_type, 'comment': c.comment.text if c.comment else None} for c in row] for row in self._cells(sheet, address)]}

    def _dependencies(self, sheet, formula):
        return [{'sheet': m[1] or m[2] or sheet, 'range': m[3] + (':' + m[4] if m[4] else '')} for m in REF.finditer(formula)]

    @tool
    def trace(self, sheet, cell):
        rows = self._cells(sheet, cell)
        if len(rows) != 1 or len(rows[0]) != 1:
            raise ValueError('trace requires a single cell')
        value = rows[0][0].value
        return {'sheet': sheet, 'cell': cell, 'formula': value if rows[0][0].data_type == 'f' else None, 'dependencies': self._dependencies(sheet, value) if rows[0][0].data_type == 'f' else [], 'coverage': 'syntactic_A1_only; names/structured/indirect references unresolved'}

    @tool
    def recalc(self, *, backend='bounded'):
        if backend != 'bounded':
            raise ValueError('Qualified LibreOffice backend is unavailable; no external process was started')
        cache, findings = {}, []
        operations = 0
        def charge():
            nonlocal operations
            operations += 1
            if operations > 100000:
                raise ValueError('calculation_operation_budget_exceeded')
        def bounded_number(value):
            if type(value) not in (int, float, bool) or not math.isfinite(value) or abs(value) > 1e100:
                raise ValueError('numeric_magnitude_budget_exceeded')
            return value
        def calculate(sheet, coordinate, stack=()):
            charge()
            key = (sheet, coordinate.replace('$', ''))
            if key in stack:
                raise ValueError('circular_reference')
            if len(stack) >= 64:
                raise ValueError('reference_depth_budget_exceeded')
            if key in cache:
                return cache[key]
            cell = self.book[sheet][key[1]]
            value = cell.value
            if cell.data_type == 'e':
                raise ValueError('stored_error:' + value)
            if cell.data_type != 'f':
                if value is None:
                    return 0
                if isinstance(value, (int, float)):
                    return bounded_number(value)
                raise ValueError('nonnumeric_reference')
            if external_formula(value):
                raise ValueError('external_formula_unsupported')
            expr = value[1:]
            if '^' in expr or '**' in expr:
                raise ValueError('power_precedence_unqualified')
            if len(expr) > 8192 or '__cellref_' in expr:
                raise ValueError('formula_size_or_reserved_name_unsupported')
            references = {}
            def reference(match):
                identifier = '__cellref_' + str(len(references))
                references[identifier] = (match[1] or match[2] or sheet,
                    match[3] + (':' + match[4] if match[4] else ''), bool(match[4]))
                return identifier
            # Preserve references as AST leaves. In particular, unselected IF
            # branches must not traverse their references or raise their errors.
            tree = ast.parse(REF.sub(reference, expr), mode='eval')
            pending, count = [(tree,0)], 0
            while pending:
                node,depth = pending.pop()
                count += 1
                if count > 512 or depth > 64:
                    raise ValueError('formula_ast_budget_exceeded')
                pending.extend((child,depth+1) for child in ast.iter_child_nodes(node))
            ops = {ast.Add: operator.add, ast.Sub: operator.sub, ast.Mult: operator.mul, ast.Div: operator.truediv}
            def evaluate(node):
                charge()
                if isinstance(node, ast.Name) and node.id in references:
                    target_sheet,address,is_range = references[node.id]
                    rows = self._cells(target_sheet,address)
                    vals = [calculate(target_sheet,c.coordinate,stack+(key,)) for row in rows for c in row]
                    return vals if is_range else vals[0]
                if isinstance(node, ast.Constant) and type(node.value) in (int, float):
                    return bounded_number(node.value)
                if isinstance(node, ast.List):
                    return [evaluate(x) for x in node.elts]
                if isinstance(node, ast.BinOp) and type(node.op) in ops:
                    left,right = bounded_number(evaluate(node.left)),bounded_number(evaluate(node.right))
                    return bounded_number(ops[type(node.op)](left,right))
                if isinstance(node, ast.UnaryOp) and isinstance(node.op, (ast.UAdd, ast.USub)):
                    return bounded_number(evaluate(node.operand) * (-1 if isinstance(node.op, ast.USub) else 1))
                if isinstance(node, ast.Compare) and len(node.ops) == 1:
                    comparisons = {ast.Lt:operator.lt, ast.LtE:operator.le, ast.Gt:operator.gt,
                                   ast.GtE:operator.ge, ast.Eq:operator.eq, ast.NotEq:operator.ne}
                    if type(node.ops[0]) not in comparisons:
                        raise ValueError('unsupported_comparison')
                    return comparisons[type(node.ops[0])](bounded_number(evaluate(node.left)), bounded_number(evaluate(node.comparators[0])))
                if isinstance(node, ast.Call) and isinstance(node.func, ast.Name) and not node.keywords:
                    function = node.func.id.upper()
                    if function == 'IF':
                        if len(node.args) != 3:
                            raise ValueError('IF_requires_three_arguments')
                        condition = evaluate(node.args[0])
                        bounded_number(condition)
                        return evaluate(node.args[1] if condition else node.args[2])
                    args = [evaluate(x) for x in node.args]
                    flat = [item for value in args for item in (value if isinstance(value,list) else [value])]
                    if function == 'SUM':
                        result = 0
                        for item in flat:
                            charge()
                            result = bounded_number(result + item)
                        return result
                    if function in ('MIN', 'MAX') and flat:
                        return bounded_number((min if function == 'MIN' else max)(flat))
                    if function == 'ABS' and len(args) == 1:
                        return bounded_number(abs(args[0]))
                    if function == 'ROUND' and len(args) == 2 and type(args[1]) is int and abs(args[1]) <= 12:
                        from decimal import Decimal, ROUND_HALF_UP, localcontext
                        with localcontext() as context:
                            context.prec = 128
                            result = float(Decimal(str(args[0])).quantize(Decimal(1).scaleb(-args[1]), rounding=ROUND_HALF_UP))
                        return bounded_number(result)
                raise ValueError('unsupported_formula')
            computed = evaluate(tree.body)
            if isinstance(computed,list):
                raise ValueError('nonfinite_or_nonscalar_result')
            cache[key] = bounded_number(computed)
            return computed
        visited = 0
        for ws in self.book:
            for row in ws:
                for cell in row:
                    visited += 1
                    if visited > self.max_cells:
                        raise ValueError('Workbook scan exceeds cell budget')
                    if cell.data_type in ('f', 'e'):
                        try:
                            calculate(ws.title, cell.coordinate)
                        except Exception as exc:
                            findings.append({'sheet': ws.title, 'cell': cell.coordinate, 'code': str(exc)})
        return {'backend': 'bounded_arithmetic_v3', 'calculation_status': 'diagnostic_only', 'qualified': False, 'cache_written': False, 'values': [{'sheet': s, 'cell': c, 'value': v} for (s,c),v in cache.items()], 'findings': findings, 'complete': not findings, 'operations': operations, 'scope': 'bounded arithmetic, A1, SUM/MIN/MAX/ABS/ROUND and lazy IF; powers explicitly unqualified; not Office recalculation'}
