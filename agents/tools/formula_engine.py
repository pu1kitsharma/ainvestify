"""Pure-Python evaluator for the Excel formula subset used in banker models.

Observation-only: it loads a workbook (formulas + stored caches), evaluates every
formula cell without LibreOffice or Excel, and reports values, errors, circular
references and unsupported functions. It never writes caches, repairs formulas or
decides company content. Passing it does not qualify Office behaviour.

Supported: + - * / ^ & % comparisons, cross-sheet/absolute/whole-column refs, defined
names, Excel-table structured refs, array arithmetic inside SUMPRODUCT, and the
functions listed in FUNCTIONS. Volatile/external functions (OFFSET, INDIRECT, TODAY,
NOW, RAND...) are reported as unsupported and evaluate to '#UNSUPPORTED!'.
Dependencies are static, so evaluation is iterative (Tarjan SCC) and cycles are
detected exactly rather than by recursion depth.
"""
from __future__ import annotations

import calendar
import datetime as _dt
import math
import re
from decimal import Decimal, ROUND_HALF_UP, ROUND_DOWN, ROUND_UP
from io import BytesIO
from fnmatch import fnmatchcase

from openpyxl import load_workbook
from openpyxl.utils.cell import get_column_letter, column_index_from_string
from openpyxl.utils.datetime import to_excel

from .contracts import bounded_package


# --------------------------------------------------------------------- values
class XlErr:
    __slots__ = ('code',)

    def __init__(self, code):
        self.code = code

    def __eq__(self, other):
        return isinstance(other, XlErr) and other.code == self.code

    def __hash__(self):
        return hash(self.code)

    def __repr__(self):
        return self.code


DIV0, NA, NAME, NUM, REF, VALUE = (XlErr(c) for c in ('#DIV/0!', '#N/A', '#NAME?', '#NUM!', '#REF!', '#VALUE!'))
NULL, CIRC, UNSUP = XlErr('#NULL!'), XlErr('#CIRC!'), XlErr('#UNSUPPORTED!')
ERROR_BY_CODE = {e.code: e for e in (DIV0, NA, NAME, NUM, REF, VALUE, NULL, CIRC, UNSUP)}
# Codes a spreadsheet itself would display; the other two are this tool's diagnostics.
WORKBOOK_ERROR_CODES = {'#DIV/0!', '#N/A', '#NAME?', '#NULL!', '#NUM!', '#REF!', '#VALUE!'}


class _Err(Exception):
    def __init__(self, err):
        self.err = err


def a1(r, c):
    return f'{get_column_letter(c)}{r}'


class SheetData:
    def __init__(self, name, state):
        self.name, self.state = name, state
        self.cells = {}        # (r, c) -> literal value (non-formula)
        self.formulas = {}     # (r, c) -> formula text beginning '='
        self.cached = {}       # (r, c) -> stored value for formula cells
        self.max_row = self.max_col = 0
        self.hidden_rows, self.hidden_cols = set(), set()
        self.array_cells = set()


class Model:
    """Workbook contents needed for evaluation and audits (no Office calls)."""

    def __init__(self):
        self.sheets = {}       # lower name -> SheetData (insertion ordered)
        self.names = {}        # (scope_lower_or_None, name_lower) -> attr text
        self.tables = {}       # lower table name -> (sheet, r1, c1, r2, c2, [headers])
        self.hidden_names = []
        self.asts = {}
        self._name_asts = {}
        self.parse_errors = {}

    def sheet(self, name):
        return self.sheets.get(str(name).lower()) if name is not None else None

    @property
    def sheet_names(self):
        return [s.name for s in self.sheets.values()]

    def formula_count(self):
        return sum(len(s.formulas) for s in self.sheets.values())

    def literal(self, sheet, r, c):
        return self.sheets[sheet.lower()].cells.get((r, c))

    def ast(self, sheet, r, c):
        key = (sheet.lower(), r, c)
        if key not in self.asts:
            text = self.sheets[sheet.lower()].formulas[(r, c)]
            try:
                self.asts[key] = parse_formula(text)
            except ValueError as exc:
                self.parse_errors[key] = str(exc)
                self.asts[key] = ('err', NAME)
        return self.asts[key]

    def resolve_name(self, name, sheet):
        """Return (kind, payload): ('ast', node) | ('err', XlErr) | None."""
        for scope in (sheet.lower() if sheet else None, None):
            text = self.names.get((scope, name.lower()))
            if text is None:
                continue
            key = (scope, name.lower())
            if key not in self._name_asts:
                if text.strip().startswith('#'):
                    self._name_asts[key] = ('err', ERROR_BY_CODE.get(text.strip(), REF))
                else:
                    try:
                        self._name_asts[key] = parse_formula('=' + text)
                    except ValueError:
                        self._name_asts[key] = ('err', NAME)
            return self._name_asts[key]
        return None


def _to_serial(value):
    if isinstance(value, (_dt.datetime, _dt.date)):
        return float(to_excel(value))
    if isinstance(value, _dt.time):
        return (value.hour * 3600 + value.minute * 60 + value.second) / 86400
    if isinstance(value, _dt.timedelta):
        return value.total_seconds() / 86400
    return value


def _normal_value(value):
    value = _to_serial(value)
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, str) and value in ERROR_BY_CODE and value.startswith('#'):
        return ERROR_BY_CODE[value]
    return value


def load_model(content, *, max_bytes=20_000_000, max_cells=200_000):
    """Parse XLSX bytes into a Model; bounded before openpyxl touches the package."""
    if len(content) > max_bytes:
        raise ValueError('Workbook exceeds byte budget')
    bounded_package(content, max_bytes=max_bytes, max_cells=max_cells)
    import warnings
    with warnings.catch_warnings():
        warnings.simplefilter('ignore')
        formulas = load_workbook(BytesIO(content), data_only=False, keep_links=False)
        cached = load_workbook(BytesIO(content), data_only=True, keep_links=False)
    model = Model()
    for ws in formulas.worksheets:
        sd = SheetData(ws.title, ws.sheet_state)
        model.sheets[ws.title.lower()] = sd
        cws = cached[ws.title]
        for row in ws.iter_rows():
            for cell in row:
                value = cell.value
                if value is None:
                    continue
                r, c = cell.row, cell.column
                sd.max_row, sd.max_col = max(sd.max_row, r), max(sd.max_col, c)
                text = None
                if hasattr(value, 'text') and hasattr(value, 'ref'):      # ArrayFormula
                    text = value.text if str(value.text).startswith('=') else '=' + str(value.text)
                    sd.array_cells.add((r, c))
                elif type(value).__name__ == 'DataTableFormula':
                    text = '=__DATATABLE__()'
                elif isinstance(value, str) and cell.data_type == 'f' and value.startswith('='):
                    text = value
                if text is not None:
                    sd.formulas[(r, c)] = text
                    sd.cached[(r, c)] = _normal_value(cws.cell(r, c).value)
                else:
                    sd.cells[(r, c)] = _normal_value(value)
        for index, dim in ws.row_dimensions.items():
            if dim.hidden:
                sd.hidden_rows.add(index)
        for key, dim in ws.column_dimensions.items():
            if dim.hidden:
                lo = dim.min or column_index_from_string(key)
                hi = dim.max or lo
                sd.hidden_cols.update(range(lo, min(hi, 16384) + 1))
        for table in ws.tables.values():
            c1, r1, c2, r2 = _bounds(table.ref)
            headers = [str(cws.cell(r1, c).value or '') for c in range(c1, c2 + 1)]
            model.tables[table.name.lower()] = (ws.title, r1, c1, r2, c2, headers)
    names = formulas.defined_names
    for name, dn in names.items():
        model.names[(None, name.lower())] = dn.attr_text
        if getattr(dn, 'hidden', False):
            model.hidden_names.append(name)
    for ws in formulas.worksheets:
        for name, dn in ws.defined_names.items():
            model.names[(ws.title.lower(), name.lower())] = dn.attr_text
    return model


def _bounds(ref):
    from openpyxl.utils.cell import range_boundaries
    return range_boundaries(ref.replace('$', ''))


# ------------------------------------------------------------------- parsing
_SHEET = r"(?:(?:'(?:[^']|'')+'|[A-Za-z_][\w\.]*)!)"
_TOKEN = re.compile(r'''
 (?P<ws>\s+)
|(?P<str>"(?:[^"]|"")*")
|(?P<err>%(S)s?\#(?:REF!|DIV/0!|N/A|NAME\?|NULL!|NUM!|VALUE!))
|(?P<func>[A-Za-z_\\][\w\.]*(?=\())
|(?P<struct>[A-Za-z_][\w\.]*\[(?:[^\[\]]|\[[^\]]*\])*\])
|(?P<ref>%(S)s?\$?[A-Za-z]{1,3}\$?[0-9]+(?::\$?[A-Za-z]{1,3}\$?[0-9]+)?(?![\w\(!\.]))
|(?P<colrange>%(S)s?\$?[A-Za-z]{1,3}:\$?[A-Za-z]{1,3}(?![\w\(!\.]))
|(?P<rowrange>%(S)s?\$?[0-9]+:\$?[0-9]+(?![\w\(!\.]))
|(?P<num>(?:[0-9]+\.?[0-9]*|\.[0-9]+)(?:[eE][+-]?[0-9]+)?)
|(?P<name>[A-Za-z_\\][\w\.]*)
|(?P<op><=|>=|<>|[-+*/^&=<>%%(),;{}:])
''' % {'S': _SHEET}, re.X)


def tokenize(text):
    pos, out = 0, []
    while pos < len(text):
        m = _TOKEN.match(text, pos)
        if not m:
            raise ValueError(f'Unparseable formula near {text[pos:pos + 12]!r}')
        pos = m.end()
        kind = m.lastgroup
        if kind != 'ws':
            out.append((kind, m.group(kind)))
    return out


def _split_sheet(token):
    if '!' in token:
        sheet, rest = token.rsplit('!', 1)
        if sheet.startswith("'"):
            sheet = sheet[1:-1].replace("''", "'")
        return sheet, rest
    return None, token


def _cell(txt):
    m = re.fullmatch(r'\$?([A-Za-z]{1,3})\$?([0-9]+)', txt)
    col = column_index_from_string(m.group(1).upper())
    row = int(m.group(2))
    if col > 16384 or row > 1048576 or row < 1:
        raise ValueError('cell out of range')
    return row, col


class _Parser:
    def __init__(self, tokens):
        self.t, self.i = tokens, 0

    def peek(self):
        return self.t[self.i] if self.i < len(self.t) else (None, None)

    def take(self):
        tok = self.peek()
        self.i += 1
        return tok

    def op(self, *ops):
        k, v = self.peek()
        return k == 'op' and v in ops

    def expr(self):
        left = self.concat()
        while self.op('=', '<>', '<', '>', '<=', '>='):
            o = self.take()[1]
            left = ('bin', o, left, self.concat())
        return left

    def concat(self):
        left = self.additive()
        while self.op('&'):
            self.take()
            left = ('bin', '&', left, self.additive())
        return left

    def additive(self):
        left = self.term()
        while self.op('+', '-'):
            o = self.take()[1]
            left = ('bin', o, left, self.term())
        return left

    def term(self):
        left = self.power()
        while self.op('*', '/'):
            o = self.take()[1]
            left = ('bin', o, left, self.power())
        return left

    def power(self):
        left = self.unary()
        while self.op('^'):
            self.take()
            left = ('bin', '^', left, self.unary())
        return left

    def unary(self):
        if self.op('-', '+'):
            o = self.take()[1]
            operand = self.unary()
            return ('un', o, operand)
        node = self.primary()
        while self.op('%'):
            self.take()
            node = ('pct', node)
        return node

    def primary(self):
        kind, val = self.take()
        if kind == 'num':
            return ('num', float(val))
        if kind == 'str':
            return ('str', val[1:-1].replace('""', '"'))
        if kind == 'err':
            return ('err', ERROR_BY_CODE[val[val.index('#'):]])
        if kind == 'ref':
            sheet, rest = _split_sheet(val)
            parts = rest.split(':')
            r1, c1 = _cell(parts[0])
            r2, c2 = _cell(parts[1]) if len(parts) > 1 else (r1, c1)
            return ('ref', sheet, min(r1, r2), min(c1, c2), max(r1, r2), max(c1, c2))
        if kind == 'colrange':
            sheet, rest = _split_sheet(val)
            a, b = (column_index_from_string(p.replace('$', '').upper()) for p in rest.split(':'))
            return ('ref', sheet, 1, min(a, b), 0, max(a, b))
        if kind == 'rowrange':
            sheet, rest = _split_sheet(val)
            a, b = (int(p.replace('$', '')) for p in rest.split(':'))
            return ('ref', sheet, min(a, b), 1, max(a, b), 0)
        if kind == 'struct':
            tname, spec = val.split('[', 1)
            return ('struct', tname, spec[:-1])
        if kind == 'name':
            up = val.upper()
            if up == 'TRUE':
                return ('bool', True)
            if up == 'FALSE':
                return ('bool', False)
            return ('name', val)
        if kind == 'func':
            self.take()                      # '('
            args = []
            if self.op(')'):
                self.take()
            else:
                while True:
                    if self.op(',', ';'):
                        args.append(('empty',))
                    elif self.op(')'):
                        args.append(('empty',))
                    else:
                        args.append(self.expr())
                    if self.op(',', ';'):
                        self.take()
                        if self.op(')'):
                            args.append(('empty',))
                            self.take()
                            break
                        continue
                    if self.op(')'):
                        self.take()
                        break
                    raise ValueError('Expected , or ) in function call')
            name = val.upper()
            for prefix in ('_XLFN.', '_XLWS.', '_XLUDF.'):
                if name.startswith(prefix):
                    name = name[len(prefix):]
            return ('func', name, args)
        if kind == 'op' and val == '(':
            node = self.expr()
            if not self.op(')'):
                raise ValueError('Missing )')
            self.take()
            return ('paren', node)
        if kind == 'op' and val == '{':
            rows, row = [], []
            while True:
                node = self.expr()
                row.append(node)
                if self.op(','):
                    self.take()
                elif self.op(';'):
                    self.take()
                    rows.append(row)
                    row = []
                elif self.op('}'):
                    self.take()
                    rows.append(row)
                    break
                else:
                    raise ValueError('Bad array constant')
            return ('arr', rows)
        raise ValueError(f'Unexpected token {val!r}')


def parse_formula(text):
    if not text.startswith('='):
        raise ValueError('Not a formula')
    tokens = tokenize(text[1:])
    if not tokens:
        raise ValueError('Empty formula')
    parser = _Parser(tokens)
    node = parser.expr()
    if parser.i != len(tokens):
        raise ValueError('Trailing tokens in formula')
    return node


def walk(node):
    """Yield every AST node (pre-order)."""
    stack = [node]
    while stack:
        n = stack.pop()
        yield n
        t = n[0]
        if t == 'bin':
            stack += [n[3], n[2]]
        elif t in ('un',):
            stack.append(n[2])
        elif t in ('pct', 'paren'):
            stack.append(n[1])
        elif t == 'func':
            stack += list(reversed(n[2]))
        elif t == 'arr':
            for row in reversed(n[1]):
                stack += list(reversed(row))


# ---------------------------------------------------------- runtime structures
class RangeVal:
    """A reference to a rectangle in a (possibly evaluating) workbook."""
    __slots__ = ('ev', 'sheet', 'r1', 'c1', 'r2', 'c2')

    def __init__(self, ev, sheet, r1, c1, r2, c2):
        self.ev, self.sheet, self.r1, self.c1, self.r2, self.c2 = ev, sheet, r1, c1, r2, c2

    @property
    def shape(self):
        return (self.r2 - self.r1 + 1, self.c2 - self.c1 + 1)

    def get(self, i, j):
        return self.ev.cell_value(self.sheet, self.r1 + i, self.c1 + j)

    def grid(self):
        rows, cols = self.shape
        return [[self.get(i, j) for j in range(cols)] for i in range(rows)]

    def flat(self):
        sd = self.ev.model.sheets[self.sheet.lower()]
        r2, c2 = min(self.r2, max(sd.max_row, self.r1)), min(self.c2, max(sd.max_col, self.c1))
        for r in range(self.r1, r2 + 1):
            for c in range(self.c1, c2 + 1):
                yield self.ev.cell_value(self.sheet, r, c)


class Arr:
    __slots__ = ('rows',)

    def __init__(self, rows):
        self.rows = rows

    @property
    def shape(self):
        return (len(self.rows), len(self.rows[0]) if self.rows else 0)

    def grid(self):
        return self.rows

    def flat(self):
        for row in self.rows:
            yield from row


def _is_num(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def _num(v):
    if isinstance(v, XlErr):
        raise _Err(v)
    if v is None:
        return 0.0
    if isinstance(v, bool):
        return 1.0 if v else 0.0
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str):
        s = v.strip()
        try:
            if s.endswith('%'):
                return float(s[:-1].replace(',', '')) / 100
            return float(s.replace(',', ''))
        except ValueError:
            raise _Err(VALUE)
    raise _Err(VALUE)


def _bool(v):
    if isinstance(v, XlErr):
        raise _Err(v)
    if v is None:
        return False
    if isinstance(v, bool):
        return v
    if isinstance(v, (int, float)):
        return v != 0
    if isinstance(v, str):
        if v.upper() == 'TRUE':
            return True
        if v.upper() == 'FALSE':
            return False
    raise _Err(VALUE)


def _str(v):
    if isinstance(v, XlErr):
        raise _Err(v)
    if v is None:
        return ''
    if isinstance(v, bool):
        return 'TRUE' if v else 'FALSE'
    if isinstance(v, (int, float)):
        if float(v).is_integer() and abs(v) < 1e15:
            return str(int(v))
        return format(v, '.15g')
    return str(v)


def _sc(v):
    """Scalarize a reference/array (single cell -> value; larger -> top-left)."""
    if isinstance(v, RangeVal):
        return v.get(0, 0)
    if isinstance(v, Arr):
        return v.rows[0][0] if v.rows and v.rows[0] else None
    return v


def _rank(v):
    return 1 if _is_num(v) else 2 if isinstance(v, str) else 3


def _compare(op, a, b):
    if isinstance(a, XlErr):
        return a
    if isinstance(b, XlErr):
        return b
    if a is None:
        a = 0.0 if _is_num(b) or b is None else '' if isinstance(b, str) else False
    if b is None:
        b = 0.0 if _is_num(a) else '' if isinstance(a, str) else False
    ra, rb = _rank(a), _rank(b)
    if ra != rb:
        x, y = ra, rb
    elif ra == 2:
        x, y = a.lower(), b.lower()
    else:
        x, y = float(a), float(b)
        if op in ('=', '<>') and math.isclose(x, y, rel_tol=1e-15, abs_tol=0):
            x = y
    return {'=': x == y, '<>': x != y, '<': x < y, '>': x > y, '<=': x <= y, '>=': x >= y}[op]


def _round(x, digits, mode=ROUND_HALF_UP):
    if not math.isfinite(x):
        raise _Err(NUM)
    d = Decimal(repr(float(x)))
    q = Decimal(1).scaleb(-int(digits))
    return float(d.quantize(q, rounding=mode)) if digits > -300 else 0.0


# --------------------------------------------------------------------- dates
_EPOCH = _dt.date(1899, 12, 30)


def _date_from_serial(v):
    n = _num(v)
    if n < 0:
        raise _Err(NUM)
    return _EPOCH + _dt.timedelta(days=int(math.floor(n)))


def _serial(d):
    return float((d - _EPOCH).days)


def _make_date(y, m, d):
    y, m, d = int(y), int(m), int(d)
    y += (m - 1) // 12
    m = (m - 1) % 12 + 1
    if y < 0 or y > 9999:
        raise _Err(NUM)
    return _dt.date(y, m, 1) + _dt.timedelta(days=d - 1)


_MONTHS = list(calendar.month_name)[1:]


def _text_format(value, fmt):
    if isinstance(value, str):
        return value
    v = _num(value)
    bare = re.sub(r'"[^"]*"|\\.', '', fmt)
    if re.search(r'[yd]', bare, re.I) or re.search(r'mmm', bare, re.I):
        d = _date_from_serial(v)

        def sub(m):
            t = m.group(0).lower()
            return {'yyyy': f'{d.year:04d}', 'yy': f'{d.year % 100:02d}', 'mmmm': _MONTHS[d.month - 1],
                    'mmm': _MONTHS[d.month - 1][:3], 'mm': f'{d.month:02d}', 'm': str(d.month),
                    'dddd': calendar.day_name[d.weekday()], 'ddd': calendar.day_name[d.weekday()][:3],
                    'dd': f'{d.day:02d}', 'd': str(d.day)}[t]
        return re.sub(r'yyyy|yy|mmmm|mmm|mm|m|dddd|ddd|dd|d', sub, fmt, flags=re.I)
    sections = fmt.split(';')
    section = sections[0]
    if v < 0 and len(sections) > 1:
        section, v = sections[1], -v
    pct = '%' in section
    if pct:
        v *= 100
    m = re.search(r'[#0?,]*[0#?][#0?,]*(?:\.([0#?]*))?', section)
    if not m:
        return section
    decimals = len(m.group(1)) if m.group(1) is not None else 0
    grouping = ',' in m.group(0).split('.')[0]
    rounded = float(_round(abs(v), decimals)) * (-1 if v < 0 else 1)
    body = f'{rounded:,.{decimals}f}' if grouping else f'{rounded:.{decimals}f}'
    out = section[:m.start()] + body + section[m.end():]
    return re.sub(r'"([^"]*)"', r'\1', out).replace('\\', '')


# ----------------------------------------------------------------- functions
FUNCS = {}       # name -> (kind, callable)


def func(*names, kind='scalar'):
    def deco(f):
        for n in names:
            FUNCS[n] = (kind, f)
        return f
    return deco


def _collect_numbers(args):
    out = []
    for a in args:
        if isinstance(a, (RangeVal, Arr)):
            for v in a.flat():
                if isinstance(v, XlErr):
                    raise _Err(v)
                if _is_num(v):
                    out.append(float(v))
        else:
            if isinstance(a, XlErr):
                raise _Err(a)
            out.append(_num(a))
    return out


@func('SUM', kind='raw')
def f_sum(ctx, args):
    return math.fsum(_collect_numbers(args))


@func('PRODUCT', kind='raw')
def f_product(ctx, args):
    return math.prod(_collect_numbers(args))


@func('AVERAGE', kind='raw')
def f_average(ctx, args):
    n = _collect_numbers(args)
    if not n:
        raise _Err(DIV0)
    return math.fsum(n) / len(n)


@func('MIN', kind='raw')
def f_min(ctx, args):
    n = _collect_numbers(args)
    return min(n) if n else 0.0


@func('MAX', kind='raw')
def f_max(ctx, args):
    n = _collect_numbers(args)
    return max(n) if n else 0.0


@func('MEDIAN', kind='raw')
def f_median(ctx, args):
    n = sorted(_collect_numbers(args))
    if not n:
        raise _Err(NUM)
    k = len(n) // 2
    return n[k] if len(n) % 2 else (n[k - 1] + n[k]) / 2


@func('COUNT', kind='raw')
def f_count(ctx, args):
    total = 0
    for a in args:
        if isinstance(a, (RangeVal, Arr)):
            total += sum(1 for v in a.flat() if _is_num(v))
        elif _is_num(a) or (isinstance(a, str) and _try_num(a)):
            total += 1
    return float(total)


def _try_num(s):
    try:
        float(s.replace(',', ''))
        return True
    except ValueError:
        return False


@func('COUNTA', kind='raw')
def f_counta(ctx, args):
    total = 0
    for a in args:
        if isinstance(a, (RangeVal, Arr)):
            total += sum(1 for v in a.flat() if v is not None)
        elif a is not None:
            total += 1
    return float(total)


@func('SUBTOTAL', kind='raw')
def f_subtotal(ctx, args):
    code = int(_num(_sc(args[0]))) % 100
    table = {1: f_average, 2: f_count, 3: f_counta, 4: f_max, 5: f_min, 6: f_product, 9: f_sum}
    if code not in table:
        raise _Err(VALUE)
    return table[code](ctx, args[1:])


@func('ABS')
def f_abs(ctx, x):
    return abs(_num(x))


@func('SIGN')
def f_sign(ctx, x):
    n = _num(x)
    return float((n > 0) - (n < 0))


@func('INT')
def f_int(ctx, x):
    return float(math.floor(_num(x)))


@func('ROUND')
def f_round(ctx, x, d):
    return _round(_num(x), _num(d))


@func('ROUNDUP')
def f_roundup(ctx, x, d):
    return _round(_num(x), _num(d), ROUND_UP)


@func('ROUNDDOWN')
def f_rounddown(ctx, x, d):
    return _round(_num(x), _num(d), ROUND_DOWN)


@func('MOD')
def f_mod(ctx, a, b):
    a, b = _num(a), _num(b)
    if b == 0:
        raise _Err(DIV0)
    return a - b * math.floor(a / b)


@func('POWER')
def f_power(ctx, a, b):
    return _pow(_num(a), _num(b))


@func('SQRT')
def f_sqrt(ctx, a):
    a = _num(a)
    if a < 0:
        raise _Err(NUM)
    return math.sqrt(a)


@func('EXP')
def f_exp(ctx, a):
    try:
        return math.exp(_num(a))
    except OverflowError:
        raise _Err(NUM)


@func('LN')
def f_ln(ctx, a):
    a = _num(a)
    if a <= 0:
        raise _Err(NUM)
    return math.log(a)


@func('LOG10')
def f_log10(ctx, a):
    a = _num(a)
    if a <= 0:
        raise _Err(NUM)
    return math.log10(a)


@func('LOG')
def f_log(ctx, a, b=10.0):
    a, b = _num(a), _num(b)
    if a <= 0 or b <= 0 or b == 1:
        raise _Err(NUM)
    return math.log(a, b)


@func('FLOOR')
def f_floor(ctx, a, s):
    a, s = _num(a), _num(s)
    if s == 0:
        return 0.0
    return math.floor(a / s) * s


@func('CEILING')
def f_ceiling(ctx, a, s):
    a, s = _num(a), _num(s)
    if s == 0:
        return 0.0
    return math.ceil(a / s) * s


@func('NOT')
def f_not(ctx, a):
    return not _bool(a)


def _bool_args(args):
    out = []
    for a in args:
        if isinstance(a, (RangeVal, Arr)):
            for v in a.flat():
                if isinstance(v, XlErr):
                    raise _Err(v)
                if isinstance(v, bool) or _is_num(v):
                    out.append(bool(v))
        else:
            out.append(_bool(a))
    if not out:
        raise _Err(VALUE)
    return out


@func('AND', kind='raw')
def f_and(ctx, args):
    return all(_bool_args(args))


@func('OR', kind='raw')
def f_or(ctx, args):
    return any(_bool_args(args))


@func('IF', kind='lazy')
def f_if(ctx, nodes):
    if not 2 <= len(nodes) <= 3:
        return VALUE
    cond = _sc(ctx.ev.eval(nodes[0], ctx))
    try:
        truth = _bool(cond)
    except _Err as e:
        return e.err
    if truth:
        return ctx.ev.eval(nodes[1], ctx) if nodes[1][0] != 'empty' else 0.0
    if len(nodes) == 2:
        return False
    return ctx.ev.eval(nodes[2], ctx) if nodes[2][0] != 'empty' else 0.0


@func('IFERROR', kind='lazy')
def f_iferror(ctx, nodes):
    v = ctx.ev.eval(nodes[0], ctx)
    if isinstance(_sc(v), XlErr):
        return ctx.ev.eval(nodes[1], ctx)
    return v


@func('IFNA', kind='lazy')
def f_ifna(ctx, nodes):
    v = ctx.ev.eval(nodes[0], ctx)
    if _sc(v) == NA:
        return ctx.ev.eval(nodes[1], ctx)
    return v


@func('CHOOSE', kind='lazy')
def f_choose(ctx, nodes):
    try:
        k = int(_num(_sc(ctx.ev.eval(nodes[0], ctx))))
    except _Err as e:
        return e.err
    if not 1 <= k < len(nodes):
        return VALUE
    return ctx.ev.eval(nodes[k], ctx)


@func('ISNUMBER', kind='raw')
def f_isnumber(ctx, a):
    return _is_num(_sc(a[0]))


@func('ISTEXT', kind='raw')
def f_istext(ctx, a):
    return isinstance(_sc(a[0]), str)


@func('ISBLANK', kind='raw')
def f_isblank(ctx, a):
    return _sc(a[0]) is None


@func('ISLOGICAL', kind='raw')
def f_islogical(ctx, a):
    return isinstance(_sc(a[0]), bool)


@func('ISERROR', kind='raw')
def f_iserror(ctx, a):
    return isinstance(_sc(a[0]), XlErr)


@func('ISERR', kind='raw')
def f_iserr(ctx, a):
    v = _sc(a[0])
    return isinstance(v, XlErr) and v != NA


@func('ISNA', kind='raw')
def f_isna(ctx, a):
    return _sc(a[0]) == NA


@func('NA')
def f_na(ctx):
    return NA


@func('TRUE')
def f_true(ctx):
    return True


@func('FALSE')
def f_false(ctx):
    return False


@func('ROW', kind='raw')
def f_row(ctx, a):
    if not a:
        return float(ctx.r)
    if isinstance(a[0], RangeVal):
        return float(a[0].r1)
    raise _Err(VALUE)


@func('COLUMN', kind='raw')
def f_column(ctx, a):
    if not a:
        return float(ctx.c)
    if isinstance(a[0], RangeVal):
        return float(a[0].c1)
    raise _Err(VALUE)


@func('ROWS', kind='raw')
def f_rows(ctx, a):
    return float(a[0].shape[0])


@func('COLUMNS', kind='raw')
def f_columns(ctx, a):
    return float(a[0].shape[1])


def _grid(v):
    if isinstance(v, (RangeVal, Arr)):
        return v.grid()
    return [[v]]


@func('INDEX', kind='raw')
def f_index(ctx, a):
    src = a[0]
    rows, cols = src.shape if isinstance(src, (RangeVal, Arr)) else (1, 1)

    def argn(i):
        return int(_num(_sc(a[i]))) if len(a) > i and a[i] is not None else None
    r, c = argn(1), argn(2)
    if r is None:
        r = 0
    if c is None:
        if rows == 1 and cols > 1:
            r, c = 1, r
        else:
            c = 1 if cols == 1 else 0
    if r < 0 or c < 0 or r > rows or c > cols:
        raise _Err(REF)
    grid = _grid(src)
    if r == 0 and c == 0:
        return Arr(grid)
    if r == 0:
        return Arr([[row[c - 1]] for row in grid])
    if c == 0:
        return Arr([grid[r - 1]])
    return grid[r - 1][c - 1]


def _lookup_eq(a, b):
    if a is None or b is None:
        return False
    if isinstance(a, str) and isinstance(b, str):
        return a.lower() == b.lower()
    if _rank(a) != _rank(b):
        return False
    return float(a) == float(b) if _is_num(a) else a == b


def _wild(pattern, text):
    return fnmatchcase(text.lower(), pattern.lower().replace('~*', '[*]').replace('~?', '[?]'))


@func('MATCH', kind='raw')
def f_match(ctx, a):
    key = _sc(a[0])
    if isinstance(key, XlErr):
        raise _Err(key)
    mode = int(_num(_sc(a[2]))) if len(a) > 2 and a[2] is not None else 1
    vec = list(a[1].flat()) if isinstance(a[1], (RangeVal, Arr)) else [a[1]]
    if mode == 0:
        for i, v in enumerate(vec):
            if isinstance(key, str) and isinstance(v, str) and any(ch in key for ch in '*?'):
                if _wild(key, v):
                    return float(i + 1)
            elif _lookup_eq(key, v):
                return float(i + 1)
        raise _Err(NA)
    best = None
    for i, v in enumerate(vec):
        if v is None or isinstance(v, XlErr) or _rank(v) != _rank(key):
            continue
        x, y = (v.lower(), key.lower()) if isinstance(v, str) else (float(v), float(key))
        if (mode == 1 and x <= y) or (mode == -1 and x >= y):
            best = i + 1
        elif best is not None or (mode == 1 and x > y) or (mode == -1 and x < y):
            if (mode == 1 and x > y) or (mode == -1 and x < y):
                break
    if best is None:
        raise _Err(NA)
    return float(best)


@func('VLOOKUP', 'HLOOKUP', kind='raw')
def _unused(ctx, a):               # replaced below (needs the called name)
    raise _Err(NAME)


def _lookup(ctx, a, vertical):
    key = _sc(a[0])
    if isinstance(key, XlErr):
        raise _Err(key)
    grid = _grid(a[1])
    if not vertical:
        grid = [list(col) for col in zip(*grid)]
    idx = int(_num(_sc(a[2])))
    approx = _bool(_sc(a[3])) if len(a) > 3 and a[3] is not None else True
    if idx < 1 or (grid and idx > len(grid[0])):
        raise _Err(REF)
    if not approx:
        for row in grid:
            v = row[0]
            if isinstance(key, str) and isinstance(v, str) and any(ch in key for ch in '*?'):
                if _wild(key, v):
                    return row[idx - 1]
            elif _lookup_eq(key, v):
                return row[idx - 1]
        raise _Err(NA)
    found = None
    for row in grid:
        v = row[0]
        if v is None or isinstance(v, XlErr) or _rank(v) != _rank(key):
            continue
        x, y = (v.lower(), key.lower()) if isinstance(v, str) else (float(v), float(key))
        if x <= y:
            found = row
        else:
            break
    if found is None:
        raise _Err(NA)
    return found[idx - 1]


FUNCS['VLOOKUP'] = ('raw', lambda ctx, a: _lookup(ctx, a, True))
FUNCS['HLOOKUP'] = ('raw', lambda ctx, a: _lookup(ctx, a, False))


@func('SUMPRODUCT', kind='raw')
def f_sumproduct(ctx, a):
    grids = [_grid(x) for x in a]
    shape = (len(grids[0]), len(grids[0][0]) if grids[0] else 0)
    if any((len(g), len(g[0]) if g else 0) != shape for g in grids):
        raise _Err(VALUE)
    total = 0.0
    for i in range(shape[0]):
        for j in range(shape[1]):
            prod = 1.0
            for g in grids:
                v = g[i][j]
                if isinstance(v, XlErr):
                    raise _Err(v)
                prod *= float(v) if _is_num(v) and not isinstance(v, bool) else 0.0
            total += prod
    return total


def _criterion(crit):
    if isinstance(crit, XlErr):
        raise _Err(crit)
    if isinstance(crit, str):
        m = re.match(r'^(<=|>=|<>|=|<|>)?(.*)$', crit, re.S)
        op, rest = m.group(1) or '=', m.group(2)
        try:
            target = float(rest.replace(',', ''))
        except ValueError:
            target = rest
            if rest == '' and op == '=':
                return lambda v: v is None or v == ''
    else:
        op, target = '=', (float(crit) if _is_num(crit) else crit)

    def test(v):
        if v is None or isinstance(v, XlErr):
            return op == '<>' and not isinstance(v, XlErr) and target != ''
        if isinstance(target, float):
            if not _is_num(v):
                return op == '<>'
            return _compare(op, float(v), target)
        if not isinstance(v, str):
            return op == '<>'
        if op in ('=', '<>') and any(ch in target for ch in '*?'):
            hit = _wild(target, v)
            return hit if op == '=' else not hit
        return _compare(op, v, target)
    return test


@func('SUMIF', kind='raw')
def f_sumif(ctx, a):
    test = _criterion(_sc(a[1]))
    rng = list(a[0].flat())
    vals = list(a[2].flat()) if len(a) > 2 else rng
    return math.fsum(float(v) for t, v in zip(rng, vals) if test(t) and _is_num(v))


@func('COUNTIF', kind='raw')
def f_countif(ctx, a):
    test = _criterion(_sc(a[1]))
    return float(sum(1 for t in a[0].flat() if test(t)))


@func('SUMIFS', kind='raw')
def f_sumifs(ctx, a):
    tests = [(list(a[i].flat()), _criterion(_sc(a[i + 1]))) for i in range(1, len(a) - 1, 2)]
    vals = list(a[0].flat())
    return math.fsum(float(v) for k, v in enumerate(vals)
                     if _is_num(v) and all(t(r[k]) for r, t in tests))


@func('COUNTIFS', kind='raw')
def f_countifs(ctx, a):
    tests = [(list(a[i].flat()), _criterion(_sc(a[i + 1]))) for i in range(0, len(a) - 1, 2)]
    return float(sum(1 for k in range(len(tests[0][0])) if all(t(r[k]) for r, t in tests)))


@func('NPV', kind='raw')
def f_npv(ctx, a):
    rate = _num(_sc(a[0]))
    if rate == -1:
        raise _Err(DIV0)
    flows = _collect_numbers(a[1:])
    return math.fsum(v / (1 + rate) ** (i + 1) for i, v in enumerate(flows))


def _irr(flows, guess=0.1):
    if not (any(v > 0 for v in flows) and any(v < 0 for v in flows)):
        raise _Err(NUM)

    def npv(r):
        return sum(v / (1 + r) ** i for i, v in enumerate(flows))
    rate = guess
    for _ in range(100):
        d = sum(-i * v / (1 + rate) ** (i + 1) for i, v in enumerate(flows))
        if d == 0:
            break
        new = rate - npv(rate) / d
        if new <= -1:
            new = (rate - 1) / 2
        if abs(new - rate) < 1e-12:
            return new
        rate = new
    lo, hi = -0.999999, 1e6
    if npv(lo) * npv(hi) > 0:
        raise _Err(NUM)
    for _ in range(300):
        mid = (lo + hi) / 2
        if npv(lo) * npv(mid) <= 0:
            hi = mid
        else:
            lo = mid
    return (lo + hi) / 2


@func('IRR', kind='raw')
def f_irr(ctx, a):
    return _irr(_collect_numbers(a[:1]), _num(_sc(a[1])) if len(a) > 1 else 0.1)


def _dated_flows(a):
    vals = [v for v in a[0].flat()]
    dates = [v for v in a[1].flat()]
    if len(vals) != len(dates):
        raise _Err(NUM)
    for v in vals + dates:
        if isinstance(v, XlErr):
            raise _Err(v)
    return [_num(v) for v in vals], [_num(d) for d in dates]


@func('XNPV', kind='raw')
def f_xnpv(ctx, a):
    rate = _num(_sc(a[0]))
    vals, dates = _dated_flows(a[1:])
    return math.fsum(v / (1 + rate) ** ((d - dates[0]) / 365) for v, d in zip(vals, dates))


@func('XIRR', kind='raw')
def f_xirr(ctx, a):
    vals, dates = _dated_flows(a)
    if not (any(v > 0 for v in vals) and any(v < 0 for v in vals)):
        raise _Err(NUM)

    def npv(r):
        return sum(v / (1 + r) ** ((d - dates[0]) / 365) for v, d in zip(vals, dates))
    lo, hi = -0.999999, 1e6
    if npv(lo) * npv(hi) > 0:
        raise _Err(NUM)
    for _ in range(400):
        mid = (lo + hi) / 2
        if npv(lo) * npv(mid) <= 0:
            hi = mid
        else:
            lo = mid
    return (lo + hi) / 2


@func('PMT')
def f_pmt(ctx, rate, nper, pv, fv=0, typ=0):
    rate, nper, pv, fv, typ = (_num(x) for x in (rate, nper, pv, fv, typ))
    if rate == 0:
        return -(pv + fv) / nper
    f = (1 + rate) ** nper
    return -(rate * (pv * f + fv)) / ((f - 1) * (1 + rate * typ))


@func('DATE')
def f_date(ctx, y, m, d):
    y, m, d = _num(y), _num(m), _num(d)
    if 0 <= y < 1900:
        y += 1900
    return _serial(_make_date(y, m, d))


@func('EOMONTH')
def f_eomonth(ctx, start, months):
    d = _date_from_serial(start)
    total = d.year * 12 + d.month - 1 + int(_num(months))
    y, m = divmod(total, 12)
    return _serial(_dt.date(y, m + 1, calendar.monthrange(y, m + 1)[1]))


@func('EDATE')
def f_edate(ctx, start, months):
    d = _date_from_serial(start)
    total = d.year * 12 + d.month - 1 + int(_num(months))
    y, m = divmod(total, 12)
    return _serial(_dt.date(y, m + 1, min(d.day, calendar.monthrange(y, m + 1)[1])))


@func('YEAR')
def f_year(ctx, x):
    return float(_date_from_serial(x).year)


@func('MONTH')
def f_month(ctx, x):
    return float(_date_from_serial(x).month)


@func('DAY')
def f_day(ctx, x):
    return float(_date_from_serial(x).day)


@func('TEXT')
def f_text(ctx, v, fmt):
    return _text_format(v, _str(fmt))


@func('VALUE')
def f_value(ctx, v):
    return _num(v)


@func('LEN')
def f_len(ctx, s):
    return float(len(_str(s)))


@func('LEFT')
def f_left(ctx, s, n=1.0):
    return _str(s)[:max(int(_num(n)), 0)]


@func('RIGHT')
def f_right(ctx, s, n=1.0):
    n = max(int(_num(n)), 0)
    return _str(s)[len(_str(s)) - n:] if n else ''


@func('MID')
def f_mid(ctx, s, start, n):
    start, n = int(_num(start)), int(_num(n))
    if start < 1 or n < 0:
        raise _Err(VALUE)
    return _str(s)[start - 1:start - 1 + n]


@func('UPPER')
def f_upper(ctx, s):
    return _str(s).upper()


@func('LOWER')
def f_lower(ctx, s):
    return _str(s).lower()


@func('TRIM')
def f_trim(ctx, s):
    return re.sub(r' +', ' ', _str(s).strip())


@func('SUBSTITUTE')
def f_substitute(ctx, s, old, new):
    return _str(s).replace(_str(old), _str(new)) if _str(old) else _str(s)


@func('FIND')
def f_find(ctx, needle, hay, start=1.0):
    i = _str(hay).find(_str(needle), int(_num(start)) - 1)
    if i < 0:
        raise _Err(VALUE)
    return float(i + 1)


@func('SEARCH')
def f_search(ctx, needle, hay, start=1.0):
    i = _str(hay).lower().find(_str(needle).lower(), int(_num(start)) - 1)
    if i < 0:
        raise _Err(VALUE)
    return float(i + 1)


@func('CONCATENATE', 'CONCAT', kind='raw')
def f_concat(ctx, a):
    out = []
    for x in a:
        if isinstance(x, (RangeVal, Arr)):
            out += [_str(v) for v in x.flat()]
        else:
            out.append(_str(x))
    return ''.join(out)


UNSUPPORTED_FUNCS = {'OFFSET', 'INDIRECT', 'TODAY', 'NOW', 'RAND', 'RANDBETWEEN', 'CELL', 'INFO', 'HYPERLINK',
                     'WEBSERVICE', 'RTD', 'TABLE', '__DATATABLE__', 'LAMBDA', 'LET', 'FILTER', 'SORT', 'UNIQUE',
                     'XLOOKUP', 'XMATCH', 'TEXTJOIN', 'SEQUENCE', 'TRANSPOSE', 'RATE', 'IFS', 'SWITCH', 'MAXIFS', 'MINIFS',
                     'AVERAGEIF', 'AVERAGEIFS', 'LOOKUP', 'LARGE', 'SMALL', 'RANK', 'STDEV', 'DATEDIF', 'NETWORKDAYS',
                     'WORKDAY', 'YEARFRAC', 'WEEKDAY', 'DAYS', 'ISOWEEKNUM', 'COUNTBLANK', 'FV', 'PV', 'NPER', 'IPMT', 'PPMT',
                     'SLN', 'DB', 'DDB'}
FUNCTIONS = sorted(set(FUNCS) | {'SUMIFS'})


def _pow(a, b):
    if a == 0 and b == 0:
        raise _Err(NUM)
    if a == 0 and b < 0:
        raise _Err(DIV0)
    if a < 0 and not float(b).is_integer():
        raise _Err(NUM)
    try:
        r = a ** b
    except OverflowError:
        raise _Err(NUM)
    if isinstance(r, complex):
        raise _Err(NUM)
    return float(r)


# ---------------------------------------------------------------- evaluation
class Ctx:
    __slots__ = ('ev', 'sheet', 'r', 'c')

    def __init__(self, ev, sheet, r, c):
        self.ev, self.sheet, self.r, self.c = ev, sheet, r, c


class Evaluation:
    """Evaluated workbook plus structured diagnostics (see ``summary``)."""

    def __init__(self, model):
        self.model = model
        self.values = {}                 # (sheet_lower, r, c) -> value (formula cells)
        self.deps = {}                   # node -> list of formula-cell nodes
        self.circular = []               # [{'cells': [...], 'path': [...]}]
        self.unsupported = {}            # function name -> [cell labels]
        self.parse_failures = []
        self.name_problems = []
        self._current = None

    # -- public helpers
    def cell_value(self, sheet, r, c):
        key = (sheet.lower(), r, c)
        if key in self.values:
            return self.values[key]
        sd = self.model.sheets.get(sheet.lower())
        if sd is None:
            return REF
        if (r, c) in sd.formulas:
            return None                  # not yet evaluated (should not happen with static deps)
        return sd.cells.get((r, c))

    def value(self, sheet, address):
        r, c = _cell(address)
        return self.cell_value(sheet, r, c)

    def label(self, key):
        sd = self.model.sheets[key[0]]
        return f'{sd.name}!{a1(key[1], key[2])}'

    # -- AST evaluation
    def eval(self, node, ctx):
        t = node[0]
        if t == 'num' or t == 'str' or t == 'bool':
            return node[1]
        if t == 'err':
            return node[1]
        if t == 'empty':
            return None
        if t == 'paren':
            return self.eval(node[1], ctx)
        if t == 'ref':
            return self._range(node, ctx)
        if t == 'struct':
            return self._struct(node, ctx)
        if t == 'name':
            return self._name(node, ctx)
        if t == 'arr':
            return Arr([[self.eval(x, ctx) for x in row] for row in node[1]])
        if t == 'pct':
            v = _sc(self.eval(node[1], ctx))
            try:
                return _num(v) / 100
            except _Err as e:
                return e.err
        if t == 'un':
            v = self.eval(node[2], ctx)
            if isinstance(v, (RangeVal, Arr)) and v.shape != (1, 1):
                return self._elementwise(lambda x, _y: self._un(node[1], x), v, None)
            v = _sc(v)
            return self._un(node[1], v)
        if t == 'bin':
            return self._bin(node[1], self.eval(node[2], ctx), self.eval(node[3], ctx))
        if t == 'func':
            return self._call(node, ctx)
        raise ValueError(t)

    @staticmethod
    def _un(op, v):
        try:
            n = _num(v)
        except _Err as e:
            return e.err
        return -n if op == '-' else n

    def _bin(self, op, a, b):
        arr_a = isinstance(a, (RangeVal, Arr)) and a.shape != (1, 1)
        arr_b = isinstance(b, (RangeVal, Arr)) and b.shape != (1, 1)
        if arr_a or arr_b:
            return self._elementwise(lambda x, y: self._scalar_bin(op, x, y), a, b)
        return self._scalar_bin(op, _sc(a), _sc(b))

    def _elementwise(self, fn, a, b):
        ga = _grid(a) if isinstance(a, (RangeVal, Arr)) else [[a]]
        gb = _grid(b) if isinstance(b, (RangeVal, Arr)) else [[b]]
        rows = max(len(ga), len(gb))
        cols = max(len(ga[0]), len(gb[0])) if ga and gb else 0

        def pick(g, i, j):
            ri = 0 if len(g) == 1 else i
            ci = 0 if len(g[0]) == 1 else j
            if ri >= len(g) or ci >= len(g[0]):
                return NA
            return g[ri][ci]
        return Arr([[fn(pick(ga, i, j), pick(gb, i, j)) for j in range(cols)] for i in range(rows)])

    @staticmethod
    def _scalar_bin(op, a, b):
        if op in ('=', '<>', '<', '>', '<=', '>='):
            return _compare(op, a, b)
        try:
            if op == '&':
                return _str(a) + _str(b)
            x, y = _num(a), _num(b)
            if op == '+':
                return x + y
            if op == '-':
                return x - y
            if op == '*':
                return x * y
            if op == '/':
                if y == 0:
                    return DIV0
                return x / y
            if op == '^':
                return _pow(x, y)
        except _Err as e:
            return e.err
        raise ValueError(op)

    def _range(self, node, ctx):
        _, sheet, r1, c1, r2, c2 = node
        name = sheet if sheet is not None else ctx.sheet
        sd = self.model.sheet(name)
        if sd is None:
            return REF
        if r2 == 0:
            r2 = max(sd.max_row, r1)
        if c2 == 0:
            c2 = max(sd.max_col, c1)
        return RangeVal(self, sd.name, r1, c1, r2, c2)

    def _struct(self, node, ctx):
        resolved = resolve_struct(self.model, node)
        if resolved is None:
            return NAME
        return RangeVal(self, resolved[0], *resolved[1:])

    def _name(self, node, ctx):
        target = self.model.resolve_name(node[1], ctx.sheet)
        if target is None:
            return NAME
        if target[0] == 'err':
            return target[1]
        return self.eval(target, ctx)

    def _call(self, node, ctx):
        _, name, args = node
        if name in UNSUPPORTED_FUNCS or name not in FUNCS:
            self.unsupported.setdefault(name, []).append(self.label((ctx.sheet.lower(), ctx.r, ctx.c)))
            return UNSUP
        kind, fn = FUNCS[name]
        try:
            if kind == 'lazy':
                return fn(ctx, args)
            vals = [self.eval(a, ctx) for a in args]
            if kind == 'raw':
                return fn(ctx, vals)
            scal = [_sc(v) for v in vals]
            for v in scal:
                if isinstance(v, XlErr):
                    return v
            return fn(ctx, *scal)
        except _Err as e:
            return e.err
        except TypeError:
            return VALUE
        except (ValueError, OverflowError, ZeroDivisionError):
            return NUM

    def final(self, v):
        v = _sc(v)
        if v is None:
            return 0.0
        return v


def resolve_struct(model, node):
    """Excel-table structured reference -> (sheet, r1, c1, r2, c2) or None."""
    table = model.tables.get(node[1].lower())
    if table is None:
        return None
    sheet, r1, c1, r2, c2, headers = table
    spec = node[2].strip()
    items = [s.strip().strip('[]') for s in re.split(r',(?![^\[]*\])', spec)] if spec.startswith('[') else [spec]
    items = [i for i in items if i]
    cols = [i for i in items if not i.startswith('#')]
    flags = {i.lower() for i in items if i.startswith('#')}
    top, bottom = r1 + 1, r2
    if '#all' in flags:
        top = r1
    elif '#headers' in flags:
        top, bottom = r1, r1
    if not cols:
        return sheet, top, c1, bottom, c2
    wanted = []
    for col in cols:
        if col not in headers:
            return None
        wanted.append(c1 + headers.index(col))
    return sheet, top, min(wanted), bottom, max(wanted)


def ref_targets(model, node, sheet):
    """Static (sheet_name, r1, c1, r2, c2) rectangles a ref/name/struct node reads."""
    t = node[0]
    if t == 'ref':
        name = node[1] if node[1] is not None else sheet
        sd = model.sheet(name)
        if sd is None:
            return []
        r2 = node[4] or max(sd.max_row, node[2])
        c2 = node[5] or max(sd.max_col, node[3])
        return [(sd.name, node[2], node[3], r2, c2)]
    if t == 'struct':
        res = resolve_struct(model, node)
        return [res] if res else []
    if t == 'name':
        target = model.resolve_name(node[1], sheet)
        if target is None or target[0] == 'err':
            return []
        return [x for n in walk(target) for x in ref_targets(model, n, sheet) if n[0] in ('ref', 'struct')]
    return []


def formula_refs(model, sheet, r, c):
    """All rectangles read by the formula in sheet!rc (names/structured refs resolved)."""
    out = []
    for n in walk(model.ast(sheet, r, c)):
        if n[0] in ('ref', 'struct', 'name'):
            out += ref_targets(model, n, sheet)
    return out


def _formula_dependencies(model):
    deps = {}
    by_sheet = {k: set(s.formulas) for k, s in model.sheets.items()}
    for key, sd in model.sheets.items():
        for (r, c) in sd.formulas:
            node = (key, r, c)
            found = set()
            for sheet, r1, c1, r2, c2 in formula_refs(model, sd.name, r, c):
                fk = sheet.lower()
                cells = by_sheet.get(fk, ())
                area = (r2 - r1 + 1) * (c2 - c1 + 1)
                if area <= len(cells):
                    for rr in range(r1, r2 + 1):
                        for cc in range(c1, c2 + 1):
                            if (rr, cc) in cells:
                                found.add((fk, rr, cc))
                else:
                    for (rr, cc) in cells:
                        if r1 <= rr <= r2 and c1 <= cc <= c2:
                            found.add((fk, rr, cc))
            deps[node] = sorted(found)
    return deps


def _tarjan(nodes, deps):
    index, low, on, stack, out, counter = {}, {}, set(), [], [], [0]
    for root in nodes:
        if root in index:
            continue
        work = [(root, 0)]
        while work:
            v, i = work.pop()
            if i == 0:
                index[v] = low[v] = counter[0]
                counter[0] += 1
                stack.append(v)
                on.add(v)
            recurse = False
            neighbours = deps[v]
            while i < len(neighbours):
                w = neighbours[i]
                i += 1
                if w not in index:
                    work.append((v, i))
                    work.append((w, 0))
                    recurse = True
                    break
                if w in on:
                    low[v] = min(low[v], index[w])
            if recurse:
                continue
            if low[v] == index[v]:
                comp = []
                while True:
                    w = stack.pop()
                    on.discard(w)
                    comp.append(w)
                    if w == v:
                        break
                out.append(comp)
            if work:
                parent = work[-1][0]
                low[parent] = min(low[parent], low[v])
    return out


def _cycle_path(comp, deps):
    members = set(comp)
    start = min(comp)
    prev = {start: None}
    queue = [start]
    while queue:
        cur = queue.pop(0)
        for w in deps[cur]:
            if w == start:
                path = [cur]
                while prev[path[-1]] is not None:
                    path.append(prev[path[-1]])
                return list(reversed(path)) + [start]
            if w in members and w not in prev:
                prev[w] = cur
                queue.append(w)
    return list(comp)


def evaluate_model(model):
    ev = Evaluation(model)
    deps = _formula_dependencies(model)
    ev.deps = deps
    components = _tarjan(sorted(deps), deps)
    for comp in components:                         # dependencies come first
        cyclic = len(comp) > 1 or comp[0] in deps[comp[0]]
        if cyclic:
            for n in comp:
                ev.values[n] = CIRC
            ev.circular.append({'cells': sorted(ev.label(n) for n in comp)[:200], 'size': len(comp),
                                'path': [ev.label(n) for n in _cycle_path(comp, deps)][:60]})
            continue
        key = comp[0]
        sheet = model.sheets[key[0]].name
        node = model.ast(sheet, key[1], key[2])
        ctx = Ctx(ev, sheet, key[1], key[2])
        try:
            ev.values[key] = ev.final(ev.eval(node, ctx))
        except _Err as e:
            ev.values[key] = e.err
        except RecursionError:
            ev.values[key] = VALUE
    for (sk, r, c), msg in model.parse_errors.items():
        ev.parse_failures.append({'cell': f'{model.sheets[sk].name}!{a1(r, c)}', 'message': msg})
    return ev


def evaluate_workbook_bytes(content, **limits):
    return evaluate_model(load_model(content, **limits))


# ------------------------------------------------------------- reports
def values_equal(mine, cached, rel=1e-9, abs_tol=1e-9):
    if isinstance(mine, XlErr) or isinstance(cached, XlErr):
        return isinstance(mine, XlErr) and isinstance(cached, XlErr) and mine.code == cached.code
    if cached is None and mine in (0.0, '', None):
        return True
    if _is_num(mine) and _is_num(cached) and not isinstance(mine, bool) and not isinstance(cached, bool):
        return math.isclose(float(mine), float(cached), rel_tol=rel, abs_tol=abs_tol)
    return mine == cached


def compare_with_cache(ev, rel=1e-9, abs_tol=1e-9, max_listed=50):
    """Compare evaluator output with stored caches (stored caches are *unverified*)."""
    stats = {'formula_cells': 0, 'with_cache': 0, 'numeric_cache_cells': 0, 'numeric_matches': 0,
             'other_matches': 0, 'other_cells': 0, 'blocked_by_unsupported': 0}
    mismatches, blocked = [], []
    for key, mine in ev.values.items():
        sd = ev.model.sheets[key[0]]
        stats['formula_cells'] += 1
        cached = sd.cached.get((key[1], key[2]))
        if cached is None:
            continue
        stats['with_cache'] += 1
        numeric = _is_num(cached)
        stats['numeric_cache_cells' if numeric else 'other_cells'] += 1
        if mine == UNSUP:
            stats['blocked_by_unsupported'] += 1
            blocked.append(ev.label(key))
            continue
        if values_equal(mine, cached, rel, abs_tol):
            stats['numeric_matches' if numeric else 'other_matches'] += 1
        else:
            mismatches.append({'cell': ev.label(key), 'evaluated': _show(mine), 'cached': _show(cached),
                               'formula': sd.formulas[(key[1], key[2])][:160]})
    comparable = stats['numeric_cache_cells'] - sum(
        1 for b in blocked if _is_num(_cached_by_label(ev, b)))
    stats['numeric_comparable'] = comparable
    stats['numeric_match_rate'] = (stats['numeric_matches'] / comparable) if comparable else None
    stats['mismatch_count'] = len(mismatches)
    return {**stats, 'mismatches': mismatches[:max_listed], 'blocked_cells': blocked[:max_listed]}


def _cached_by_label(ev, label):
    sheet, addr = label.rsplit('!', 1)
    r, c = _cell(addr)
    return ev.model.sheets[sheet.lower()].cached.get((r, c))


def _show(v):
    return v.code if isinstance(v, XlErr) else v


def error_cells(ev):
    out = []
    for key, v in sorted(ev.values.items()):
        if isinstance(v, XlErr):
            out.append({'cell': ev.label(key), 'error': v.code,
                        'formula': ev.model.sheets[key[0]].formulas[(key[1], key[2])][:160]})
    return out


def evaluation_summary(ev, *, max_listed=100):
    errors = error_cells(ev)
    workbook_errors = [e for e in errors if e['error'] in WORKBOOK_ERROR_CODES]
    return {'sheets': [{'sheet': s.name, 'state': s.state, 'formulas': len(s.formulas)} for s in ev.model.sheets.values()],
            'formula_cells': len(ev.values),
            'workbook_error_count': len(workbook_errors),
            'workbook_errors': workbook_errors[:max_listed],
            'circular_references': ev.circular[:20],
            'unsupported_functions': {k: {'count': len(v), 'cells': v[:10]} for k, v in sorted(ev.unsupported.items())},
            'parse_failures': ev.parse_failures[:max_listed],
            'engine': 'pure_python_excel_subset', 'office_qualified': False, 'caches_written': False,
            'supported_functions': FUNCTIONS}
