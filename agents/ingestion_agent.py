"""
Ingestion Agent — architecture doc §5.2.

Deterministic parsing only. No LLM involved in this step (CLAUDE.md,
"Immediate next steps" #3) — this is what keeps DocBlock IDs a trustworthy
citation target for the Structured Extraction Agent (§5.3) instead of
something an LLM could mis-segment.

PDF: layout-aware — tables are detected and kept as table blocks (with cell
grid + bbox), not flattened into the surrounding text, which is what the
architecture doc (§1, §5.2) flags as the mistake that forces "RAG over
broken tables" downstream.

Excel: parsed row-by-row via openpyxl with sheet/row/column references kept,
not flattened to CSV first (§5.2) — CSV normalization discards the exact
cell-reference provenance §7 requires.
"""
from __future__ import annotations

from pathlib import Path
from typing import Optional

import openpyxl
import pdfplumber

from schemas import BlockType, DocBlock, Document, new_id

# Text lines separated by more than this many PDF points are treated as
# distinct paragraphs rather than merged into one text block.
PARAGRAPH_GAP_THRESHOLD = 8.0


def _bboxes_overlap(a: tuple[float, float, float, float], b: tuple[float, float, float, float]) -> bool:
    ax0, atop, ax1, abottom = a
    bx0, btop, bx1, bbottom = b
    return ax0 < bx1 and ax1 > bx0 and atop < bbottom and abottom > btop


def ingest_pdf(path: str, tenant_id: str, deal_id: str) -> Document:
    """Parse a PDF into a Document of typed, page-anchored DocBlocks."""
    pdf_path = Path(path)
    document = Document(
        id=new_id("doc"),
        tenant_id=tenant_id,
        deal_id=deal_id,
        filename=pdf_path.name,
        type="pdf",
        storage_uri=str(pdf_path.resolve()),
    )

    with pdfplumber.open(str(pdf_path)) as pdf:
        for page_number, page in enumerate(pdf.pages, start=1):
            tables = page.find_tables()
            table_bboxes = [t.bbox for t in tables]

            for table in tables:
                rows = table.extract() or []
                document.blocks.append(DocBlock(
                    id=new_id("blk"),
                    document_id=document.id,
                    page=page_number,
                    block_type=BlockType.TABLE,
                    coordinates={"bbox": list(table.bbox)},
                    content=rows,
                ))

            document.blocks.extend(
                _extract_text_blocks(page, page_number, document.id, table_bboxes)
            )

    return document


def _extract_text_blocks(page, page_number: int, document_id: str, table_bboxes) -> list[DocBlock]:
    """Group non-table text lines into paragraph-level blocks by vertical gap."""
    lines = page.extract_text_lines() or []
    blocks: list[DocBlock] = []
    paragraph: list[dict] = []

    def flush() -> Optional[DocBlock]:
        if not paragraph:
            return None
        text = "\n".join(l["text"] for l in paragraph)
        bbox = [
            min(l["x0"] for l in paragraph),
            min(l["top"] for l in paragraph),
            max(l["x1"] for l in paragraph),
            max(l["bottom"] for l in paragraph),
        ]
        paragraph.clear()
        return DocBlock(
            id=new_id("blk"),
            document_id=document_id,
            page=page_number,
            block_type=BlockType.TEXT,
            coordinates={"bbox": bbox},
            content=text,
        )

    prev_bottom = None
    for line in lines:
        bbox = (line["x0"], line["top"], line["x1"], line["bottom"])
        if any(_bboxes_overlap(bbox, tb) for tb in table_bboxes):
            block = flush()
            if block:
                blocks.append(block)
            prev_bottom = None
            continue

        if prev_bottom is not None and (line["top"] - prev_bottom) > PARAGRAPH_GAP_THRESHOLD:
            block = flush()
            if block:
                blocks.append(block)

        paragraph.append(line)
        prev_bottom = line["bottom"]

    block = flush()
    if block:
        blocks.append(block)

    return blocks


def ingest_excel(path: str, tenant_id: str, deal_id: str) -> Document:
    """Parse an Excel workbook into one DocBlock per non-empty row, keeping
    sheet name + row/column cell references for citation (§5.2, §7)."""
    xlsx_path = Path(path)
    document = Document(
        id=new_id("doc"),
        tenant_id=tenant_id,
        deal_id=deal_id,
        filename=xlsx_path.name,
        type="xlsx",
        storage_uri=str(xlsx_path.resolve()),
    )

    workbook = openpyxl.load_workbook(str(xlsx_path), data_only=True)
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            cells = {cell.coordinate: cell.value for cell in row if cell.value is not None}
            if not cells:
                continue
            row_number = row[0].row
            document.blocks.append(DocBlock(
                id=new_id("blk"),
                document_id=document.id,
                page=1,  # sheets don't paginate; sheet name carries the location instead
                block_type=BlockType.TABLE,
                coordinates={"sheet": sheet.title, "row": row_number},
                content=cells,
            ))

    return document


def ingest_document(path: str, tenant_id: str, deal_id: str) -> Document:
    """Dispatch to the right deterministic parser by file extension."""
    suffix = Path(path).suffix.lower()
    if suffix == ".pdf":
        return ingest_pdf(path, tenant_id, deal_id)
    if suffix in (".xlsx", ".xlsm"):
        return ingest_excel(path, tenant_id, deal_id)
    raise ValueError(f"Unsupported document type: {suffix} (expected .pdf or .xlsx)")
