"""New private structured intake retains raw evidence without semantic approval."""
from datetime import datetime
from io import BytesIO
import hashlib
import json
from types import SimpleNamespace

from openpyxl import Workbook

from delivery.investment_memo_stage import room_sources
from schemas import BlockType, DocBlock, Document
from scripts import private_acceptance_ingest_worker as module


def workbook_bytes():
    book = Workbook()
    sheet = book.active
    sheet.title = 'Summary'
    sheet['A1'] = 'INR million'
    sheet['B1'] = datetime(2023, 3, 31)
    sheet['A2'] = 'Company projection'
    sheet['B2'] = '=Inputs!B2*Inputs!B3'
    sheet['B2'].number_format = '#,##0.00'
    hidden = book.create_sheet('Inputs')
    hidden.sheet_state = 'hidden'
    hidden['A1'] = 'Underlying assumptions'
    hidden['A2'] = 'Transactions'
    hidden['B2'] = 10
    hidden['A3'] = 'Fee'
    hidden['B3'] = 2
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def test_financial_corpus_preserves_every_cell_formula_cache_format_date_and_hidden_sheet():
    raw = workbook_bytes()
    sha = hashlib.sha256(raw).hexdigest()
    result = module._financial_sources(raw, sha)
    cells = {(row['sheet'], row['cell']): row for row in result['inventory']['cells']}
    assert result['retained_cell_count'] == len(cells) == 9
    assert cells['Summary', 'B2']['formula'] == 'Inputs!B2*Inputs!B3'
    assert cells['Summary', 'B2']['stored'] is None
    assert cells['Summary', 'B2']['number_format'] == '#,##0.00'
    assert cells['Inputs', 'B2']['sheet_state'] == 'hidden'
    assert cells['Summary', 'B1']['value'] == datetime(2023, 3, 31)
    packed = [cell for source in result['sources']
              for cell in json.loads(source['passage'])['cells']]
    assert len(packed) == result['retained_cell_count']
    assert {row['sheet'] for row in packed} == {'Summary', 'Inputs'}
    assert result['financial_validation'] == 'not_established'
    assert result['workbook_recalculation'] == 'not_performed'
    assert result['actual_results_origin'] == 'not_established'
    assert hashlib.sha256(raw).hexdigest() == sha


def test_large_sheet_paginates_all_cells_without_first_n_loss():
    book = Workbook()
    sheet = book.active
    for row in range(1, 501):
        sheet.cell(row, 1, 'Same source row label')
        sheet.cell(row, 2, row)
    output = BytesIO()
    book.save(output)
    raw = output.getvalue()
    result = module._financial_sources(raw, hashlib.sha256(raw).hexdigest())
    assert result['retained_cell_count'] == 1000
    packed = [cell for source in result['sources']
              for cell in json.loads(source['passage'])['cells']]
    assert len(packed) == 1000
    assert len({row['cell'] for row in packed}) == 1000
    assert len(result['sources']) > 1
    assert result['semantic_coverage'] == 'unassessed'
    assert result['bounded_source_compatible'] is True


def test_room_structure_option_keeps_cell_coordinates_and_unknown_financial_origin():
    doc = Document(id='doc_a', tenant_id='tenant_a', deal_id='deal_a',
        filename='synthetic.xlsx', type='xlsx', storage_uri='private/source.xlsx',
        blocks=[DocBlock(id='block_a', document_id='doc_a', page=1,
                        block_type=BlockType.TABLE,
                        coordinates={'sheet': 'Summary', 'row': 2, 'visibility': 'visible'},
                        content={'A2': 'Revenue', 'B2': 123})])
    store = SimpleNamespace(get_documents_for_deal=lambda *args: [doc])
    lead = SimpleNamespace(tenant_id='tenant_a',
                           company_profile=SimpleNamespace(provenance={}, evidence=[]))
    room = SimpleNamespace(deal_id='deal_a')
    old, old_coverage = room_sources(store, lead, room, full_inventory=True)
    fresh, fresh_coverage = room_sources(store, lead, room, full_inventory=True,
                                       preserve_structure=True)
    assert old == []  # Legacy source filter dropped this short but useful row.
    value = json.loads(fresh[0].passage)
    assert value['coordinates'] == doc.blocks[0].coordinates
    assert value['content'] == doc.blocks[0].content
    assert value['actual_results_origin'] == 'not_established'
    assert value['financial_validation'] == 'not_established'
    assert old_coverage['source_count'] == 0
    assert fresh_coverage['source_count'] == 1


def test_v2_writes_separate_structured_sidecars_v1_output_stays_exact(tmp_path, monkeypatch):
    files = [tmp_path / 'first.pdf', tmp_path / 'second.pdf', tmp_path / 'source.xlsx']
    for path in files:
        path.write_bytes(b'synthetic private bytes')

    def ingest(path, tenant_id, deal_id):
        fmt = 'xlsx' if path.endswith('.xlsx') else 'pdf'
        return Document(id='doc', tenant_id='', deal_id='', filename=path,
            type=fmt, storage_uri=path, workbook_inventory={'findings': []},
            blocks=[DocBlock(id='block', document_id='doc', page=1,
                block_type=BlockType.TEXT,
                coordinates={'bbox': [10, 10, 100, 100]},
                content='The company reports a synthetic product offering through partners.')])

    monkeypatch.setattr(module, 'ingest_document', ingest)
    monkeypatch.setattr(module, '_financial_sources', lambda raw, sha:
                        {'sources': [], 'semantic_coverage': 'unassessed'})
    monkeypatch.setattr(module, '_pdf_visual_inventory', lambda path:
                        {1: {'embedded_images': 2, 'visual_content_status': 'unassessed'}})
    first, second = tmp_path / 'v1', tmp_path / 'v2'
    first.mkdir()
    second.mkdir()
    for root, contract in [(first, module.INGEST_V1), (second, module.INGEST_V2)]:
        (root / 'request.json').write_text(json.dumps({'contract': contract,
                                                      'files': [str(path) for path in files]}))
        monkeypatch.chdir(root)
        module.main()
    v1 = json.loads((first / 'memo_sources.json').read_text())
    v2 = json.loads((second / 'memo_sources.json').read_text())
    assert v1['sources'] == v2['sources']
    assert v1['source_sha256'] == v2['source_sha256']
    assert not (first / 'structured_source_inventory.json').exists()
    assert not (first / 'financial_sources.json').exists()
    structured = json.loads((second / 'structured_source_inventory.json').read_text())
    assert structured['documents'][0]['blocks'][0]['coordinates'] == {'bbox': [10, 10, 100, 100]}
    assert structured['source_structures']['S1']['visual_content_status'] == 'unassessed'
    assert len(json.loads((second / 'financial_sources.json').read_text())['workbooks']) == 1
    assert (second / 'financial_sources.json').stat().st_mode & 0o777 == 0o600
