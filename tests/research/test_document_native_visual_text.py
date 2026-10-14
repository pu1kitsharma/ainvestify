"""Native vector/SmartArt text retention; synthetic source, no OCR/rendering."""
from copy import deepcopy
from io import BytesIO
import hashlib
import json
from zipfile import ZipFile
from xml.etree import ElementTree as XML

import pytest
from pptx import Presentation

from agents.research import document_evidence_manifest as manifest
from scripts import private_document_manifest_worker as worker
from tests.research.test_emf_text import build,textw


IDENTITY={'tenant_id':'synthetic-private','deal_id':'synthetic-room','document_id':'synthetic-vector'}


def source_bytes():
    deck=Presentation();deck.slides.add_slide(deck.slide_layouts[6])
    buffer=BytesIO();deck.save(buffer)
    with ZipFile(BytesIO(buffer.getvalue())) as z:parts={p:z.read(p) for p in z.namelist()}
    P,A,R=manifest.NS['p'],manifest.NS['a'],manifest.REL
    slide=XML.fromstring(parts['ppt/slides/slide1.xml'])
    shapes=slide.find('{'+P+'}cSld/{'+P+'}spTree')
    pic=XML.fromstring(f'''<p:pic xmlns:p="{P}" xmlns:a="{A}" xmlns:r="{R}">
      <p:nvPicPr><p:cNvPr id="2" name="Native vector"/><p:cNvPicPr/><p:nvPr/></p:nvPicPr>
      <p:blipFill><a:blip r:embed="rIdEmf"/></p:blipFill>
      <p:spPr><a:xfrm><a:off x="100" y="200"/><a:ext cx="300" cy="400"/></a:xfrm></p:spPr></p:pic>''')
    shapes.append(pic)
    diagram='http://schemas.openxmlformats.org/drawingml/2006/diagram'
    shapes.append(XML.fromstring(f'''<p:graphicFrame xmlns:p="{P}" xmlns:a="{A}" xmlns:r="{R}" xmlns:dgm="{diagram}">
      <p:nvGraphicFramePr><p:cNvPr id="3" name="Native diagram"/><p:cNvGraphicFramePr/><p:nvPr/></p:nvGraphicFramePr>
      <p:xfrm><a:off x="500" y="600"/><a:ext cx="700" cy="800"/></p:xfrm>
      <a:graphic><a:graphicData uri="{diagram}"><dgm:relIds r:dm="rIdData"/></a:graphicData></a:graphic></p:graphicFrame>'''))
    parts['ppt/slides/slide1.xml']=XML.tostring(slide)
    package='http://schemas.openxmlformats.org/package/2006/relationships'
    relations=XML.fromstring(parts['ppt/slides/_rels/slide1.xml.rels'])
    for rid,kind,target in [('rIdEmf','image','../media/native.emf'),('rIdData','diagramData','../diagrams/data1.xml')]:
        XML.SubElement(relations,'{'+package+'}Relationship',Id=rid,Type=R+'/'+kind,Target=target)
    parts['ppt/slides/_rels/slide1.xml.rels']=XML.tostring(relations)
    parts['ppt/media/native.emf']=build([textw(' 100.35% source projection '),textw('')])
    parts['ppt/diagrams/data1.xml']=f'<dgm:dataModel xmlns:dgm="{diagram}" xmlns:a="{A}"><dgm:ptLst><dgm:pt><dgm:t><a:p><a:r><a:t>Untrusted source note.</a:t></a:r></a:p></dgm:t></dgm:pt></dgm:ptLst></dgm:dataModel>'.encode()
    parts['ppt/diagrams/drawing1.xml']=f'<dgm:drawing xmlns:dgm="{diagram}" xmlns:a="{A}"><a:p><a:r><a:t>Untrusted source note.</a:t></a:r></a:p></dgm:drawing>'.encode()
    parts['ppt/diagrams/_rels/data1.xml.rels']=f'<Relationships xmlns="{package}"><Relationship Id="rIdCache" Type="http://schemas.microsoft.com/office/2007/relationships/diagramDrawing" Target="drawing1.xml"/></Relationships>'.encode()
    output=BytesIO()
    with ZipFile(output,'w') as z:
        for name,value in parts.items():z.writestr(name,value)
    return output.getvalue()


def test_optin_keeps_native_default_exact_and_appends_raw_emf_smartart_with_separate_cache_origin():
    raw=source_bytes();base=manifest.build_document_manifest(raw,'pptx',**IDENTITY)
    before=manifest.canonical(base)
    derived=manifest.attach_document_visual_text(base,raw,expected_binding=base['binding'])
    assert manifest.canonical(base)==before
    assert derived['tables']==base['tables'] and derived['slides']==base['slides']
    assert derived['media_text_profile']=='native-pptx-visual-text-v1'
    emf=[u for u in derived['units'] if u['kind']=='emf_text']
    assert [u['content'] for u in emf]==[' 100.35% source projection ','']
    for unit in emf:
        locator=unit['source_locator'];obs=unit['coordinates']['native_emf']
        assert locator['slide']==1 and locator['shape_path']==[3]
        assert locator['media_sha256']==derived['native_emf_observations'][0]['media_sha256']
        assert obs['raw_hex']==unit['content'].encode('utf-16-le').hex()
        assert obs['reference']=={'x':10,'y':20}
    smart=[u for u in derived['units'] if u['kind'].startswith('smartart_')]
    assert len(smart)==2 and smart[0]['content']==smart[1]['content']=='Untrusted source note.'
    assert {u['origin'] for u in smart}=={'native_smartart_data_literal','stored_smartart_drawing_cache_unverified'}
    assert len({u['id'] for u in smart})==2
    assert derived['financial_validation']=='not_established' and derived['visual_coverage']=='unassessed'
    assert derived['ocr_status']=='not_performed'
    assert manifest.build_document_manifest(raw,'pptx',**IDENTITY)==base
    assert manifest.build_document_manifest(raw,'pptx',media_text_profile='native-pptx-visual-text-v1',**IDENTITY)==derived


def test_derived_reopen_reconstructs_originals_and_rejects_rehashed_changed_text():
    raw=source_bytes();base=manifest.build_document_manifest(raw,'pptx',**IDENTITY)
    derived=manifest.attach_document_visual_text(base,raw,expected_binding=base['binding'])
    assert manifest.attach_document_visual_text(derived,raw,expected_binding=base['binding'])==derived
    altered=deepcopy(derived);altered['units'][-1]['content']='manually replaced'
    altered['manifest_digest']=manifest.digest({k:v for k,v in altered.items() if k!='manifest_digest'})
    with pytest.raises(ValueError,match='Immutable'):
        manifest.attach_document_visual_text(altered,raw,expected_binding=base['binding'])
    with pytest.raises(ValueError,match='source/scope'):
        manifest.attach_document_emf(base,raw+b'changed',expected_binding=base['binding'])


def test_smartart_extension_cache_resolves_exact_slide_relationship_without_data_relationship_guessing():
    raw=source_bytes()
    with ZipFile(BytesIO(raw)) as z:parts={name:z.read(name) for name in z.namelist()}
    del parts['ppt/diagrams/_rels/data1.xml.rels']
    data=XML.fromstring(parts['ppt/diagrams/data1.xml'])
    XML.SubElement(data,'{http://schemas.microsoft.com/office/drawing/2008/diagram}dataModelExt',relId='rIdSlideCache')
    parts['ppt/diagrams/data1.xml']=XML.tostring(data)
    relations=XML.fromstring(parts['ppt/slides/_rels/slide1.xml.rels'])
    XML.SubElement(relations,'{http://schemas.openxmlformats.org/package/2006/relationships}Relationship',
        Id='rIdSlideCache',Type='http://schemas.microsoft.com/office/2007/relationships/diagramDrawing',
        Target='../diagrams/drawing1.xml')
    parts['ppt/slides/_rels/slide1.xml.rels']=XML.tostring(relations)
    output=BytesIO()
    with ZipFile(output,'w') as z:
        for name,value in parts.items():z.writestr(name,value)
    derived=manifest.build_document_manifest(output.getvalue(),'pptx',media_text_profile='native-pptx-visual-text-v1',**IDENTITY)
    cached=[u for u in derived['units'] if u['kind']=='smartart_cached_text']
    assert len(cached)==1 and cached[0]['source_locator']['part_relationship_id']=='rIdSlideCache'
    assert cached[0]['source_locator']['smartart_part']=='ppt/diagrams/drawing1.xml'
    assert derived['counts']['smartart_text_units']==2


def test_optin_worker_retains_native_manifest_and_binds_derived_parser_and_source_bytes(tmp_path):
    tmp_path.chmod(0o700);raw=source_bytes();(tmp_path/'source.pptx').write_bytes(raw)
    binding={**IDENTITY,'format':'pptx','source_sha256':hashlib.sha256(raw).hexdigest()}
    request={'source':'source.pptx','binding':binding,'media_text_profile':'native-pptx-visual-text-v1'}
    (tmp_path/'request.json').write_text(json.dumps(request))
    worker.run(tmp_path)
    base=manifest.build_document_manifest(raw,'pptx',**IDENTITY)
    assert (tmp_path/'document_evidence_manifest.json').read_bytes()==manifest.canonical(base).encode()
    derived=tmp_path/'document_evidence_manifest_visual_text.json'
    receipt=json.loads((tmp_path/'document_manifest_visual_text_receipt.json').read_text())
    assert receipt['manifest_sha256']==hashlib.sha256(derived.read_bytes()).hexdigest()
    assert receipt['native_manifest_sha256']==hashlib.sha256((tmp_path/'document_evidence_manifest.json').read_bytes()).hexdigest()
    assert receipt['parser_source_sha256'] and receipt['manifest_source_sha256']
    assert receipt['required_network_access']=='denied' and receipt['ocr']=='not_performed'
    before=derived.read_bytes();worker.run(tmp_path)
    assert derived.read_bytes()==before and (tmp_path/'source.pptx').read_bytes()==raw
    assert derived.stat().st_mode&0o777==0o600
