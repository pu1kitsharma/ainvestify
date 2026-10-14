"""Financial retrieval must preserve scope, source bytes and financial uncertainty."""
from copy import deepcopy
from datetime import datetime
import hashlib
from io import BytesIO
import json

from openpyxl import Workbook
from openpyxl.utils.datetime import CALENDAR_MAC_1904
from openpyxl.workbook.defined_name import DefinedName
import pytest

from agents.inference.model_authorship import digest
from agents.research import financial_evidence_manifest as module

PIN = {'name': 'qwen3:14b', 'digest': 'a' * 64}


def source_bytes(rows=1):
    book = Workbook()
    sheet = book.active
    sheet.title = 'Summary'
    sheet['A1'] = 'Management projection, INR million'
    sheet['B1'] = datetime(2023, 3, 31)
    sheet['A2'] = 'Projected service fees'
    sheet['B2'] = '=Inputs!B2*Inputs!B3'
    sheet['B2'].number_format = '#,##0.00'
    inputs = book.create_sheet('Inputs')
    inputs.sheet_state = 'hidden'
    inputs['A1'] = 'Underlying assumptions'
    inputs['A2'] = 'Transactions'
    inputs['B2'] = 10
    inputs['A3'] = 'Fee'
    inputs['B3'] = 2
    inputs.row_dimensions[2].hidden = True
    inputs.column_dimensions['B'].hidden = True
    book.defined_names.add(DefinedName('FeeInput', attr_text="'Inputs'!$B$3"))
    for row in range(4, rows + 4):
        inputs.cell(row, 1, 'Additional source row')
        inputs.cell(row, 2, row)
    book.create_sheet('Empty')
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def build(raw=None, **options):
    return module.build_financial_manifest(raw or source_bytes(), tenant_id='tenant_a',
        deal_id='deal_a', document_id='document_a', **options)


def cell(manifest, sheet, coordinate):
    return manifest['address_index'][sheet][coordinate]


def evidence_for(manifest):
    return module.retrieve_financial_evidence(manifest, manifest['binding'],
        ranges=[{'sheet': 'Summary', 'range': 'A1:B2'}])


def annotation_record(payload, answer):
    raw = json.dumps(answer)
    return {'id': 'response_1', 'task': module.TASK, 'input': deepcopy(payload),
            'instruction': module.INSTRUCTION,
            'schema': module.financial_semantics_schema(payload).model_json_schema(),
            'model': PIN['name'], 'routing': {'model': PIN['name']},
            'raw_response': raw, 'response_hash': digest(raw), 'answer': deepcopy(answer)}


def proposal(manifest, payload):
    return {'annotations': [{'concept': 'Service fees', 'source_status': 'projected',
        'value_cell_ids': [cell(manifest, 'Summary', 'B2')],
        'labels': [{'cell_id': cell(manifest, 'Summary', 'A2'), 'role': 'concept',
                    'quote': 'Projected service fees'},
                   {'cell_id': cell(manifest, 'Summary', 'A1'), 'role': 'source_status',
                    'quote': 'Management projection'}]}]}


def bind(manifest, payload, records):
    return module.bind_financial_annotations(manifest, manifest['binding'], payload,
        records, 'response_1', installed_model_pin=PIN)


def test_complete_inventory_and_packets_preserve_all_cells_without_memo_capacity_limit():
    raw = source_bytes(rows=400)
    manifest = build(raw)
    assert manifest['binding']['workbook_sha256'] == hashlib.sha256(raw).hexdigest()
    assert manifest['retained_cell_count'] == 809
    assert manifest['packet_count'] > 25
    packed = [item for packet in manifest['packets']
              for item in json.loads(packet['passage'])['cells']]
    assert len(packed) == len({item['id'] for item in packed}) == 809
    assert {item['id'] for item in packed} == set(manifest['cells'])
    assert all(packet['character_count'] <= 3400 for packet in manifest['packets'])
    formula = manifest['cells'][cell(manifest, 'Summary', 'B2')]
    assert formula['formula'] == 'Inputs!B2*Inputs!B3'
    assert formula['stored'] is None
    assert formula['observation_status'] == 'formula_cache_missing'
    assert formula['number_format'] == '#,##0.00'
    hidden = manifest['cells'][cell(manifest, 'Inputs', 'B2')]
    assert hidden['sheet_state'] == 'hidden' and hidden['row_hidden'] and hidden['column_hidden']
    assert manifest['cells'][cell(manifest, 'Summary', 'B1')]['value'] == '2023-03-31T00:00:00'
    assert manifest['cells'][cell(manifest, 'Summary', 'B1')]['value_kind'] == 'datetime'
    assert manifest['inventory']['defined_names'][0]['name'] == 'FeeInput'
    assert manifest['inventory']['dependency_graph']['Summary!B2'] == ['Inputs!B2', 'Inputs!B3']
    assert manifest['actual_results_origin'] == 'not_established'
    assert manifest['financial_status'] == 'unreconciled'
    assert manifest['workbook_recalculation'] == 'not_performed'
    assert manifest['semantic_coverage'] == 'unassessed'
    assert hashlib.sha256(raw).hexdigest() == manifest['binding']['workbook_sha256']


def test_exact_output_retrieval_includes_hidden_dependency_closure_and_exclusions():
    manifest = build()
    evidence = module.retrieve_financial_evidence(manifest, manifest['binding'],
        cell_ids=[cell(manifest, 'Summary', 'B2')])
    assert {(item['sheet'], item['cell']) for item in evidence['cells']} == {
        ('Summary', 'B2'), ('Inputs', 'B2'), ('Inputs', 'B3')}
    assert len(evidence['requested_cell_ids']) == 1
    assert evidence['selection_coverage'] == 'partial_selection'
    assert evidence['complete_financial_understanding'] is False
    assert evidence['excluded_packet_ids']
    assert any(item['code'] == 'missing_formula_cache' for item in evidence['findings'])
    assert evidence['financial_validation'] == 'not_established'


def test_blank_coordinates_including_empty_sheet_are_explicit_not_missing_financial_metrics():
    manifest = build()
    evidence = module.retrieve_financial_evidence(manifest, manifest['binding'],
        ranges=[{'sheet': 'Empty', 'range': 'A1:B1'}])
    assert evidence['cells'] == []
    assert evidence['blank_source_coordinates'] == [
        {'sheet': 'Empty', 'cell': 'A1', 'observation_status': 'no_nonempty_source_cell'},
        {'sheet': 'Empty', 'cell': 'B1', 'observation_status': 'no_nonempty_source_cell'}]
    assert evidence['semantic_coverage'] == 'unassessed'
    with pytest.raises(ValueError, match='retained source cells'):
        payload = module.financial_semantic_payload(manifest, manifest['binding'], evidence, model_pin=PIN)
        module.financial_semantics_schema(payload)


def test_retrieval_rejects_tampering_cross_tenant_unknown_coordinates_and_no_silent_clipping():
    manifest = build()
    other = {**manifest['binding'], 'tenant_id': 'tenant_b'}
    with pytest.raises(ValueError, match='scope'):
        module.retrieve_financial_evidence(manifest, other, ranges=[{'sheet': 'Summary', 'range': 'A1'}])
    altered = deepcopy(manifest)
    altered['cells'][cell(manifest, 'Inputs', 'B2')]['value'] = 999
    with pytest.raises(ValueError, match='digest'):
        module.retrieve_financial_evidence(altered, manifest['binding'], ranges=[{'sheet': 'Summary', 'range': 'A1'}])
    with pytest.raises(ValueError, match='Unknown financial'):
        module.retrieve_financial_evidence(manifest, manifest['binding'], cell_ids=['invented'])
    with pytest.raises(ValueError, match='bounded cell'):
        module.retrieve_financial_evidence(manifest, manifest['binding'],
            cell_ids=[cell(manifest, 'Summary', 'B2')], max_cells=2)
    with pytest.raises(ValueError, match='bounded context'):
        module.retrieve_financial_evidence(manifest, manifest['binding'],
            cell_ids=[cell(manifest, 'Summary', 'B2')], max_chars=512)
    with pytest.raises(ValueError, match='A1 coordinates'):
        module.retrieve_financial_evidence(manifest, manifest['binding'], ranges=[{'sheet': 'Summary', 'range': 'A:A'}])
    for outside in ('XFE1', 'A1048577'):
        with pytest.raises(ValueError, match='outside XLSX A1 coordinates'):
            module.retrieve_financial_evidence(manifest, manifest['binding'],
                ranges=[{'sheet': 'Summary', 'range': outside}])


def test_unsupported_dynamic_dependency_stays_a_visible_gap_without_invented_edges():
    book = Workbook()
    book.active['A1'] = '=INDIRECT("B1")'
    book.active['B1'] = 42
    output = BytesIO()
    book.save(output)
    manifest = build(output.getvalue())
    evidence = module.retrieve_financial_evidence(manifest, manifest['binding'],
        cell_ids=[cell(manifest, 'Sheet', 'A1')])
    assert evidence['dependency_coverage'] == 'partial'
    assert len(evidence['cells']) == 1
    assert any(item['code'] == 'dynamic_reference' for item in evidence['findings'])
    assert evidence['complete_financial_understanding'] is False


def test_merged_period_header_and_1904_date_system_stay_raw_source_geometry():
    book = Workbook()
    book.epoch = CALENDAR_MAC_1904
    book.active.merge_cells('A1:C1')
    book.active['A1'] = 'Source period label'
    book.active['B2'] = datetime(2024, 3, 31)
    output = BytesIO()
    book.save(output)
    manifest = build(output.getvalue())
    assert manifest['inventory']['workbook_properties']['date1904'] == '1'
    assert manifest['inventory']['date_system_observation'] == {
        'date1904': '1', 'attribute_status': 'source_attribute_present'}
    geometry = manifest['inventory']['worksheet_structures'][0]
    assert geometry['sheet'] == 'Sheet'
    assert geometry['merged_ranges'] == [{'ref': 'A1:C1'}]
    assert manifest['retained_cell_count'] == 2
    evidence = module.retrieve_financial_evidence(manifest, manifest['binding'],
        ranges=[{'sheet': 'Sheet', 'range': 'A1:C2'}])
    assert len(evidence['blank_source_coordinates']) == 4
    assert evidence['source_structure']['date_system_observation']['date1904'] == '1'
    assert evidence['source_structure']['worksheet_structures'][0]['merged_ranges'] == [{'ref': 'A1:C1'}]
    assert evidence['currency_unit_period_semantics'] == 'unassessed'
    assert manifest['cells'][cell(manifest, 'Sheet', 'B2')]['value'] == '2024-03-31T00:00:00'


def test_oversized_cell_is_retained_and_explicitly_blocks_its_packet():
    book = Workbook()
    book.active['A1'] = 'x' * 5000
    output = BytesIO()
    book.save(output)
    manifest = build(output.getvalue())
    assert manifest['retained_cell_count'] == 1
    assert len(next(iter(manifest['cells'].values()))['value']) == 5000
    assert manifest['bounded_source_compatible'] is False
    with pytest.raises(ValueError, match='bounded source'):
        module.retrieve_financial_evidence(manifest, manifest['binding'], packet_ids=['packet_000001'])


def test_exact_model_proposal_binds_local_pin_raw_input_source_spans_but_grants_no_approval():
    manifest = build()
    payload = module.financial_semantic_payload(manifest, manifest['binding'],
        evidence_for(manifest), model_pin=PIN)
    answer = proposal(manifest, payload)
    record = annotation_record(payload, answer)
    bound = bind(manifest, payload, [record])
    assert bound['annotations'] == answer['annotations']
    assert bound['annotation_status'] == 'model_proposed_pending_independent_review'
    assert bound['actual_results_origin'] == 'not_established'
    assert bound['financial_status'] == 'unreconciled'
    assert bound['release_status'] == 'diagnostic_only_no_release'
    assert bound['raw_response_sha256'] == hashlib.sha256(record['raw_response'].encode()).hexdigest()
    with pytest.raises(ValueError, match='unique recorded'):
        bind(manifest, payload, [record, record])
    for mutate in ('input', 'raw', 'model', 'quote', 'instruction'):
        changed = deepcopy(record)
        if mutate == 'input':
            changed['input']['binding']['deal_id'] = 'another_deal'
        elif mutate == 'raw':
            changed['raw_response'] += ' '
        elif mutate == 'model':
            changed['model'] = 'qwen3.5:9b'
        elif mutate == 'instruction':
            changed['instruction'] = 'Supply a predetermined company answer.'
        else:
            wrong_answer = deepcopy(answer)
            wrong_answer['annotations'][0]['labels'][0]['quote'] = 'Invented revenue'
            changed = annotation_record(payload, wrong_answer)
        with pytest.raises(ValueError):
            bind(manifest, payload, [changed])
    with pytest.raises(ValueError, match='installed local pin'):
        module.bind_financial_annotations(manifest, manifest['binding'], payload,
            [record], 'response_1', installed_model_pin={**PIN, 'digest': 'b' * 64})
    date_as_label = deepcopy(answer)
    date_as_label['annotations'][0]['labels'][1] = {
        'cell_id': cell(manifest, 'Summary', 'B1'), 'role': 'source_status',
        'quote': '2023-03-31T00:00:00'}
    with pytest.raises(ValueError, match='literal source span'):
        bind(manifest, payload, [annotation_record(payload, date_as_label)])


def test_self_rehashed_retrieval_cannot_substitute_cell_values_and_missing_is_only_model_interpretation():
    manifest = build()
    evidence = evidence_for(manifest)
    forged = deepcopy(evidence)
    forged['cells'][0]['value'] = 'Fake reported results'
    forged.pop('retrieval_digest')
    module._seal(forged, 'retrieval_digest')
    with pytest.raises(ValueError, match='exact retrieval replay'):
        module.financial_semantic_payload(manifest, manifest['binding'], forged, model_pin=PIN)
    payload = module.financial_semantic_payload(manifest, manifest['binding'], evidence, model_pin=PIN)
    answer = {'annotations': [{'concept': 'Fee evidence request', 'source_status': 'missing',
        'value_cell_ids': [], 'labels': [{'cell_id': cell(manifest, 'Summary', 'A2'),
            'role': 'concept', 'quote': 'Projected service fees'}]}]}
    bound = bind(manifest, payload, [annotation_record(payload, answer)])
    assert bound['missing_status_scope'] == 'model_interpretation_not_proven_workbook_absence'
    assert bound['semantic_coverage'] == 'unassessed'
