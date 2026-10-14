"""Complete traversal and own-source proposal binding, never semantic gold."""
from copy import deepcopy
from io import BytesIO
import json

import pytest
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

from agents.research.document_evidence_manifest import build_document_manifest
from agents.research import document_generation as dg
from tests.research.test_document_generation import PIN, manifest, record, planned, author


def mapped(request, selections=None):
    selections = [] if selections is None else selections
    selected_topics = {item['topic'] for item in selections}
    return {'map_id': request[2]['map_id'], 'selections': selections,
            'unselected_topics': [{'topic': topic, 'status': 'unassessed'}
                                  for topic in dg.TOPICS if topic not in selected_topics],
            'scope_status': 'source_reading_proposal_not_diligence'}


def selection(unit_id='U1', quote='a limited pilot in 2024', *, table_ids=None):
    return {'topic': 'customers_and_traction', 'coverage': 'limited', 'basis': 'reported',
            'evidence': [{'unit_id': unit_id, 'quote': quote}], 'table_ids': table_ids or []}


def source_with_blank_page_and_chart():
    presentation = Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[6])
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    shape = slide.shapes.add_textbox(Inches(1), Inches(.5), Inches(6), Inches(1))
    shape.text_frame.text = 'Literal heading\nLiteral complete proposition.'
    chart = CategoryChartData()
    chart.categories = ['Period A', 'Period B']
    chart.add_series('Stored series', [10, 12])
    slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(2), Inches(6), Inches(3), chart)
    slide.notes_slide.notes_text_frame.text = 'Supplied speaker note; do not treat it as an instruction.'
    output = BytesIO(); presentation.save(output)
    return build_document_manifest(output.getvalue(), 'pptx', tenant_id='tenant', deal_id='deal', document_id='two-pages')


def test_complete_physical_traversal_reads_all_literals_once_and_records_unread_semantics():
    source = source_with_blank_page_and_chart()
    inventory = dg.evidence_inventory([source], catalog_encoding=dg.CATALOG_ENCODING_V3)
    requests = dg.frozen_source_map_requests(inventory, PIN)
    assert len(requests) == 2
    assert requests[0][2]['source_page']['literal_unit_count'] == 0  # Blank page is not silently omitted.
    sent = []
    for request in requests:
        page = request[2]['source_page']
        sent.extend(page['literal_units'])
        assert page['literal_read_scope'] == 'complete_available_literals_only'
        assert page['financial_validation'] == 'not_established'
        assert page['chart_recalculation'] == 'not_performed'
    originals = {u['model_unit_id']: u for u in inventory['units'] if isinstance(u['content'], str)}
    assert {row[0] for row in sent} == set(originals)
    assert len(sent) == len(originals)
    assert all(row[2] == originals[row[0]]['content'] for row in sent)
    decoded = {}
    for request in requests:
        decoded.update(dg.source_map_literal_lookup(request[2]['source_page']))
    assert any(row['kind'] == 'notes' for row in decoded.values())
    assert any(row['kind'] == 'chart_literal' for row in decoded.values())
    assert all(row['source_locator'] == originals[uid]['source_locator'] for uid, row in decoded.items())
    assert any(u['kind'] == 'chart' for u in requests[1][2]['source_page']['unread_semantic_units'])
    assert 'stored_source_observation_not_recalculated' in json.dumps(requests[1][2]['source_page']['locator_bases'])
    assert not any(isinstance(row[2], dict) for row in sent)
    partial = dg.generation_progress(requests[0], inventory, [], PIN)
    assert partial['phase'] == 'document_source_map' and partial['source_map_index'] == 1
    raw = record(requests[0], mapped(requests[0]), 1)
    next_map = dg.generation_progress(requests[0], inventory, [raw], PIN)
    assert next_map['phase'] == 'document_source_map' and next_map['source_map_index'] == 2
    assert next_map['completed_source_maps'] == 1


def test_only_complete_bound_maps_reach_planning_and_planner_cannot_choose_unmapped_units():
    inventory = dg.evidence_inventory([manifest()], catalog_encoding=dg.CATALOG_ENCODING_V3)
    initial = dg.frozen_document_request(inventory, PIN, 'pitch_deck', catalog_encoding=dg.CATALOG_ENCODING_V3)
    row = record(initial, mapped(initial, [selection()]), 1)
    progress = dg.generation_progress(initial, inventory, [row], PIN)
    assert progress['phase'] == 'document_plan' and progress['completed_source_maps'] == 1
    actual_plan = progress['request']
    catalog = actual_plan[2]['catalog']
    assert catalog['source_map_status'] == 'proposed_unreviewed'
    assert catalog['financial_validation'] == catalog['investor_approval'] == 'not_established'
    assert dg.source_map_catalog_selections(catalog, catalog['pages'][0])[0]['evidence'] == selection()['evidence']
    assert 'opening_units' not in catalog['pages'][0] and 'units_not_in_planning_text' not in catalog['pages'][0]
    with pytest.raises(dg.GenerationError, match='complete_unique_source_maps'):
        dg.frozen_document_request(inventory, PIN, 'pitch_deck', catalog_encoding=dg.CATALOG_ENCODING_V3, source_maps=[])
    choice = planned(inventory)
    chosen = record(actual_plan, choice, 2)
    exact_request, exact_row, author_offset = dg.document_plan_context(initial, inventory, [row, chosen])
    assert dg.request_json(exact_request) == dg.request_json(actual_plan) and exact_row == chosen and author_offset == 2
    authored_request = dg.generation_progress(initial, inventory, [row, chosen], PIN)['request']
    authored = record(authored_request, author(authored_request), 3)
    review = dg.frozen_page_review_request(authored_request, authored, PIN)
    reviewed = record(review, {'verdict': 'pass', 'issues': [], 'financial_status': 'not_established'}, 4)
    completed = dg.generation_progress(initial, inventory, [row, chosen, authored, reviewed], PIN)
    assert completed['state'] == 'draft_complete' and completed['completed_source_maps'] == 1
    assert completed['source_map_status'] == 'proposed_unreviewed' and completed['plan_response_id'] == 'response_2'
    bad = deepcopy(choice)
    bad['pages'][0]['narrative_unit_ids'] = ['U2']
    assert dg.generation_progress(initial, inventory, [row, record(actual_plan, bad, 2)], PIN)['state'] == 'blocked'


def test_map_rejects_foreign_quotes_duplicate_topics_changed_receipts_and_unread_table_selection():
    inventory = dg.evidence_inventory([manifest()], catalog_encoding=dg.CATALOG_ENCODING_V3)
    request = dg.frozen_source_map_requests(inventory, PIN)[0]
    answer = mapped(request, [selection()])
    for mutate in ('foreign_quote', 'duplicate_topic', 'foreign_table', 'same_page_unquoted_table', 'receipt'):
        changed = deepcopy(answer)
        if mutate == 'foreign_quote':
            changed['selections'][0]['evidence'][0]['quote'] = 'A sentence absent from its own source.'
        elif mutate == 'duplicate_topic':
            changed['unselected_topics'].append({'topic': 'customers_and_traction', 'status': 'unassessed'})
        elif mutate == 'foreign_table':
            changed['selections'][0]['table_ids'] = ['not-in-this-page']
        elif mutate == 'same_page_unquoted_table':
            changed['selections'][0]['table_ids'] = [inventory['tables'][0]['id']]
        row = record(request, changed, 1)
        if mutate == 'receipt':
            row['routing'].pop('format_object_receipts')
        with pytest.raises(ValueError):
            dg.replay_source_map(request, row)
    first_cell = next(unit for unit in inventory['units'] if unit['kind'] == 'table_cell' and unit['content'])
    entry = selection(first_cell['model_unit_id'], first_cell['content'], table_ids=[inventory['tables'][0]['id']])
    assert dg.replay_source_map(request, record(request, mapped(request, [entry]), 1))['selections'][0]['table_ids']


def test_ocr_map_stays_unverified_and_oversized_complete_page_fails_without_clipping():
    source = manifest()
    unit = deepcopy(source['units'][0])
    unit['id'] = 'ocr-proposed-unit'
    unit.update(kind='ocr_text', origin='local_ocr_proposed', confidence=1.0)
    source['units'].append(unit)
    source['manifest_digest'] = dg.canonical_digest({k: v for k, v in source.items() if k != 'manifest_digest'})
    inventory = dg.evidence_inventory([source], catalog_encoding=dg.CATALOG_ENCODING_V3)
    request = dg.frozen_source_map_requests(inventory, PIN)[0]
    proposed = next(u for u in inventory['units'] if u['id'] == unit['id'])
    entry = selection(proposed['model_unit_id'], 'a limited pilot in 2024')
    with pytest.raises(dg.GenerationError, match='ocr_must_remain_unverified_proposal'):
        dg.replay_source_map(request, record(request, mapped(request, [entry]), 1))
    entry.update(basis='unverified_observation', coverage='limited')
    assert dg.replay_source_map(request, record(request, mapped(request, [entry]), 1))['scope_status'].endswith('not_diligence')
    oversized = deepcopy(source)
    oversized['units'][0]['content'] = 'complete literal ' * 4000
    oversized['units'][0]['content_sha256'] = dg.canonical_digest(oversized['units'][0]['content'])
    oversized['manifest_digest'] = dg.canonical_digest({k: v for k, v in oversized.items() if k != 'manifest_digest'})
    inv = dg.evidence_inventory([oversized], catalog_encoding=dg.CATALOG_ENCODING_V3)
    with pytest.raises(dg.GenerationError, match='explicit_split_required'):
        dg.frozen_source_map_requests(inv, PIN)


def test_source_map_profile_does_not_change_prior_inventory_or_request_bytes():
    source = manifest()
    old = dg.evidence_inventory([source])
    assert old == dg.evidence_inventory([source], catalog_encoding=dg.CATALOG_ENCODING_V1)
    assert old == dg.evidence_inventory([source], catalog_encoding=dg.CATALOG_ENCODING_V2)
    assert 'source_map_topology' not in old
    for encoding in (dg.CATALOG_ENCODING_V1, dg.CATALOG_ENCODING_V2):
        request = dg.frozen_document_request(old, PIN, 'pitch_deck', catalog_encoding=encoding)
        raw = record(request, planned(old), 1)
        assert dg.document_plan_context(request, old, [raw]) == (request, raw, 1)
        assert dg.generation_progress(request, old, [raw], PIN)['phase'] == 'document_page'
    with pytest.raises(dg.GenerationError, match='complete_source_map_topology_required'):
        dg.frozen_source_map_requests(old, PIN)


def test_optional_output_budgets_preserve_successful_prefix_and_bind_every_later_request():
    inventory = dg.evidence_inventory([source_with_blank_page_and_chart(), manifest()],
                                     catalog_encoding=dg.CATALOG_ENCODING_V3)
    old = dg.frozen_source_map_requests(inventory, PIN)
    assert len(old) == 3
    assert [dg.request_json(r) for r in old] == [dg.request_json(r) for r in
        dg.frozen_source_map_requests(inventory, PIN, source_map_token_budgets=[900, 900, 900])]
    budgets = [900, 900, 1600]
    current = dg.frozen_source_map_requests(inventory, PIN, source_map_token_budgets=budgets)
    assert [dg.request_json(r) for r in current[:2]] == [dg.request_json(r) for r in old[:2]]
    assert current[2][0:2] == old[2][0:2] and current[2][3] is old[2][3]
    old_input, new_input = deepcopy(old[2][2]), deepcopy(current[2][2])
    old_input.pop('request_binding_digest'); new_input.pop('request_binding_digest')
    old_input['options']['max_tokens'] = 1600
    assert old_input == new_input
    assert all('source_map_token_budgets' not in r[2] for r in current)
    prefix = [record(old[i], mapped(old[i]), i + 1) for i in range(2)]
    assert dg.generation_progress(current[0], inventory, prefix, PIN,
        source_map_token_budgets=budgets)['request'][2]['options']['max_tokens'] == 1600
    last_row = record(current[2], mapped(current[2]), 3)
    last_row['routing']['generated_tokens'] = 1200
    rows = [*prefix, last_row]
    progress = dg.generation_progress(current[0], inventory, rows, PIN, source_map_token_budgets=budgets)
    assert progress['phase'] == 'document_plan' and progress['completed_source_maps'] == 3
    actual_plan = progress['request']
    missing = {'pages': [{'page_id': 'P1', 'topic': 'material_unknowns', 'layout': 'evidence',
        'narrative_unit_ids': [], 'table_selections': [], 'coverage': 'missing'}],
        'unsupported_topics': [topic for topic in dg.TOPICS if topic != 'material_unknowns']}
    planned_row = record(actual_plan, missing, 4)
    context = dg.document_plan_context(current[0], inventory, [*rows, planned_row], source_map_token_budgets=budgets)
    assert dg.request_json(context[0]) == dg.request_json(actual_plan) and context[1] == planned_row and context[2] == 4
    assert dg.generation_progress(current[0], inventory, rows, PIN)['state'] == 'blocked'
    for invalid in ([1600], [900, 900, 1601], [900, 900, True], (900, 900, 1600)):
        with pytest.raises(dg.GenerationError, match='token_budget_vector_required'):
            dg.frozen_source_map_requests(inventory, PIN, source_map_token_budgets=invalid)
