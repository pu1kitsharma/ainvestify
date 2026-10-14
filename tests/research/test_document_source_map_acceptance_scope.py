"""Independent source-map approval boundary; one synthetic source, no gold.

A correctly bound map is a local reading proposal. It cannot grant review or
financial/investor approval, or skip the source-map prerequisite for planning.
"""
from copy import deepcopy
from io import BytesIO
import json

import pytest
from pptx import Presentation
from pptx.util import Inches

from agents.inference.model_authorship import digest
from agents.research.document_evidence_manifest import build_document_manifest
from agents.research import document_generation as generation


PIN = {'name': 'qwen3.5:9b', 'digest': 'a' * 64}


def _one_synthetic_source():
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    text = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(2))
    text.text_frame.text = 'A supplied source describes a bounded pilot.'
    output = BytesIO()
    presentation.save(output)
    manifest = build_document_manifest(output.getvalue(), 'pptx',
        tenant_id='synthetic-tenant', deal_id='synthetic-deal', document_id='synthetic-source')
    return generation.evidence_inventory([manifest], catalog_encoding=generation.CATALOG_ENCODING_V3)


def _record(request, answer):
    task, instruction, payload, schema = request
    raw = json.dumps(answer, ensure_ascii=False)
    return {'id': 'response_1', 'task': task, 'instruction': instruction,
        'input': deepcopy(payload), 'schema': schema.model_json_schema(),
        'answer': deepcopy(answer), 'raw_response': raw, 'response_hash': digest(raw),
        'model': payload['model_pin']['name'],
        'routing': {**payload['options'], 'model': payload['model_pin']['name'],
            'prompt_tokens': 1100, 'generated_tokens': 150, 'done_reason': 'stop',
            'sampling': {'presence_penalty': 0, 'repeat_penalty': 1},
            'format_object_receipts': [{'operation': 'chat',
                'model': payload['model_pin']['name'],
                'format_order_digest': payload['generation_schema_order_digest']}]}}


def test_exact_source_map_cannot_inflate_review_or_approval_before_planning():
    inventory = _one_synthetic_source()
    request = generation.frozen_source_map_requests(inventory, PIN)[0]
    literal = next(row for row in request[2]['source_page']['literal_units'] if row[2].strip())
    answer = {'map_id': request[2]['map_id'],
        'selections': [{'topic': 'customers_and_traction', 'coverage': 'limited',
            'basis': 'reported', 'evidence': [{'unit_id': literal[0], 'quote': literal[2]}],
            'table_ids': []}],
        'unselected_topics': [{'topic': topic, 'status': 'unassessed'}
            for topic in generation.TOPICS if topic != 'customers_and_traction'],
        'scope_status': 'source_reading_proposal_not_diligence'}
    row = _record(request, answer)
    # Even extra record metadata cannot appoint an independent reviewer or
    # qualify finances; only the exact typed raw map enters the plan catalog.
    row.update(independent_review='approved', financial_validation='validated',
               investor_approval='approved')
    catalog = generation._mapped_catalog(inventory, PIN, [row], kind='intro_deck', mandate='')
    assert catalog['source_map_status'] == 'proposed_unreviewed'
    assert catalog['financial_validation'] == 'not_established'
    assert catalog['investor_approval'] == 'not_established'
    assert catalog['scope_status'] == 'source_reading_proposal_not_diligence'
    assert catalog['visual_semantics'] == 'unassessed'
    page = catalog['pages'][0]
    assert page['map_id'] == answer['map_id']
    assert generation.source_map_catalog_selections(catalog, page) == answer['selections']
    decoded_unselected = [{'topic': catalog['topic_names'][index], 'status': status}
        for status, indexes in page['unselected_status_groups'].items() for index in indexes]
    assert sorted(decoded_unselected, key=lambda item: item['topic']) == sorted(
        answer['unselected_topics'], key=lambda item: item['topic'])
    assert catalog['source_map_provenance_digest'] == generation.canonical_digest([
        {'map_id': answer['map_id'], 'source_page_id': request[2]['source_page_id'],
         'response_id': row['id'], 'response_hash': row['response_hash'],
         'request_digest': generation.canonical_digest(generation.request_json(request))}])

    inflated = deepcopy(answer)
    inflated['scope_status'] = 'independently_reviewed_and_approved'
    with pytest.raises(ValueError):
        generation._mapped_catalog(inventory, PIN, [_record(request, inflated)],
            kind='intro_deck', mandate='')
    with pytest.raises(generation.GenerationError, match='complete_unique_source_maps'):
        generation._mapped_catalog(inventory, PIN, [], kind='intro_deck', mandate='')


def test_own_page_selector_binding_does_not_qualify_a_wrong_financial_topic():
    inventory = _one_synthetic_source()
    profile = {'source_map_profiles': [generation.SOURCE_MAP_UNIT_PROFILE],
               'source_map_token_budgets': [1600]}
    request = generation.frozen_source_map_requests(inventory, PIN, **profile)[0]
    literal = next(row for row in request[2]['source_page']['literal_units'] if row[2].strip())
    answer = {'map_id': request[2]['map_id'],
        **{topic: {'status': 'unassessed', 'coverage': 'limited',
            'basis': 'unverified_observation', 'unit_ids': [], 'table_ids': []}
            for topic in generation.TOPICS},
        'scope_status': 'source_reading_proposal_not_diligence'}
    # Deliberately wrong semantic classification: the source concerns a pilot,
    # not financial evidence. Own-source binding cannot approve that judgment.
    answer['financials'] = {'status': 'selected', 'coverage': 'supported',
        'basis': 'reported', 'unit_ids': [literal[0]], 'table_ids': []}
    row = _record(request, answer)
    assert generation.replay_source_map(request, row) == answer
    catalog = generation._mapped_catalog(inventory, PIN, [row],
        kind='intro_deck', mandate='', **profile)
    assert catalog['source_map_status'] == 'proposed_unreviewed'
    assert catalog['scope_status'] == 'source_reading_proposal_not_diligence'
    assert catalog['financial_validation'] == 'not_established'
    assert catalog['investor_approval'] == 'not_established'
    assert catalog['visual_semantics'] == 'unassessed'
    assert catalog['literal_unit_pool'][0]['unit_id'] == literal[0]
    assert catalog['literal_unit_pool'][0]['content'] == literal[2]
    selection = generation.source_map_catalog_selections(catalog, catalog['pages'][0])[0]
    assert selection['topic'] == 'financials'
    assert selection['evidence'] == [{'unit_id': literal[0], 'quote': literal[2]}]

    foreign = deepcopy(answer)
    foreign['financials']['unit_ids'] = ['U999999']
    with pytest.raises(generation.GenerationError, match='distinct_own_page_ids'):
        generation._mapped_catalog(inventory, PIN, [_record(request, foreign)],
            kind='intro_deck', mandate='', **profile)
    inflated = deepcopy(answer)
    inflated['scope_status'] = 'financially_qualified'
    with pytest.raises(ValueError):
        generation._mapped_catalog(inventory, PIN, [_record(request, inflated)],
            kind='intro_deck', mandate='', **profile)
