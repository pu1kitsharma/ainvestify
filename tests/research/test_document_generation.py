"""Source-value integrity and frozen authorship gates, not model-quality gold."""
from copy import deepcopy
from io import BytesIO
import json

import pytest
from pptx import Presentation
from pptx.util import Inches

from agents.inference.model_authorship import digest
from agents.research.document_evidence_manifest import build_document_manifest
from agents.research import document_generation as dg


PIN = {'name': 'qwen3.5:9b', 'digest': 'a' * 64}


def manifest():
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1))
    shape.text_frame.text = 'The source reports a limited pilot in 2024.'
    table = slide.shapes.add_table(3, 2, Inches(1), Inches(2), Inches(8), Inches(2)).table
    for row, values in enumerate([['Measure', '2024 projection'], ['Revenue', '120'], ['Cost', '']]):
        for column, value in enumerate(values):
            table.cell(row, column).text = value
    stream = BytesIO()
    presentation.save(stream)
    return build_document_manifest(stream.getvalue(), 'pptx', tenant_id='tenant', deal_id='deal', document_id='document')


def record(request, answer, number):
    task, instruction, payload, schema = request
    raw = json.dumps(answer, ensure_ascii=False)
    return {'id': f'response_{number}', 'task': task, 'instruction': instruction,
        'input': deepcopy(payload), 'schema': schema.model_json_schema(), 'answer': deepcopy(answer),
        'raw_response': raw, 'response_hash': digest(raw), 'model': payload['model_pin']['name'],
        'routing': {**payload['options'], 'model': payload['model_pin']['name'], 'prompt_tokens': 1100,
            'generated_tokens': 150, 'done_reason': 'stop',
            'sampling': {'presence_penalty': 0, 'repeat_penalty': 1},
            'format_object_receipts': [{'operation': 'chat', 'model': payload['model_pin']['name'],
                'format_order_digest': payload['generation_schema_order_digest']}]}}


def planned(inventory, *, table=False):
    page = {'page_id': 'P1', 'topic': 'financials' if table else 'customers_and_traction',
        'layout': 'table' if table else 'evidence', 'narrative_unit_ids': ['U1'],
        'table_selections': [{'table_id': inventory['tables'][0]['id'], 'column_indexes': [0, 1],
                             'row_indexes': [1, 2], 'header_row_index': 0}] if table else [],
        'coverage': 'limited'}
    return {'pages': [page], 'unsupported_topics': [t for t in dg.TOPICS if t != page['topic']]}


def author(request, *, table=False):
    answer = {'page_id': 'P1', 'title': 'Reported source evidence',
        'claims': [{'text': 'The source reports a limited pilot in 2024.', 'basis': 'reported',
                    'evidence': [{'unit_id': 'U1', 'quote': 'a limited pilot in 2024'}]}],
        'table_labels': []}
    if table:
        original = request[2]['evidence']['tables'][0]
        table_refs = []
        units = request[2]['evidence']['units']
        for cell in original['cells']:
            if cell['row'] == 1:
                uid = next(u['id'] for u in units if u['original_unit_id'] == cell['unit_id'])
                table_refs.append({'unit_id': uid, 'quote': cell['text']})
        answer['table_labels'] = [{'table_id': original['table_id'], 'column_labels':
            ['Measure', '2024 projection'], 'caption': 'Source table; financial validation pending.',
            'evidence': table_refs}]
    return answer


def prepared(*, table=False):
    inventory = dg.evidence_inventory([manifest()])
    plan_request = dg.frozen_document_request(inventory, PIN, 'pitch_deck')
    plan_row = record(plan_request, planned(inventory, table=table), 1)
    page_request = dg.frozen_page_request(plan_request, plan_row, 'P1', inventory)
    page_row = record(page_request, author(page_request, table=table), 2)
    return inventory, plan_request, plan_row, page_request, page_row


def test_complete_catalog_uses_frozen_handles_without_clipping_units():
    inventory = dg.evidence_inventory([manifest()])
    catalog = dg.planning_catalog(inventory)
    pages = catalog['pages']
    assert {u for p in pages for u in p['unit_ids']} == {u['model_unit_id'] for u in inventory['units']}
    for page in pages:
        assert set(page['unit_ids']) == set(page['units_not_in_planning_text']) | {
            u['unit_id'] for u in page['opening_units']}
    changed = deepcopy(inventory)
    changed['units'][0]['content'] += ' changed'
    with pytest.raises(dg.GenerationError, match='inventory_changed'):
        dg.planning_catalog(changed)


def test_retained_blank_paragraph_does_not_consume_planning_text_slot():
    value = manifest()
    unit = deepcopy(value['units'][0])
    unit['id'] = 'u_retained_blank'
    unit['source_locator']['paragraph'] = 999
    unit['content'] = unit['text'] = ' '
    unit['content_sha256'] = dg.canonical_digest(unit['content'])
    value['units'].insert(0, unit)
    value['manifest_digest'] = dg.canonical_digest({k: v for k, v in value.items() if k != 'manifest_digest'})
    inventory = dg.evidence_inventory([value])
    catalog = dg.planning_catalog(inventory)
    assert 'U1' in catalog['pages'][0]['units_not_in_planning_text']
    assert catalog['pages'][0]['opening_units'][0]['unit_id'] == 'U2'
    assert all(u['content'].strip() for page in catalog['pages'] for u in page['opening_units'])
    assert len(inventory['units']) == len(value['units'])


def test_explicit_compact_catalog_retains_every_unit_once_and_old_request_exact():
    inventory = dg.evidence_inventory([manifest()])
    prior = dg.planning_catalog(inventory)
    compact = dg.planning_catalog(inventory, catalog_encoding=dg.CATALOG_ENCODING_V2)
    assert dg.planning_catalog(inventory, catalog_encoding=dg.CATALOG_ENCODING_V1) == prior
    assert compact['sources'] == prior['sources']
    assert compact['inventory_revision'] == prior['inventory_revision']
    represented = []
    for old_page, page in zip(prior['pages'], compact['pages']):
        assert 'unit_ids' not in page and 'source_id' not in page
        assert page['source_ref'] == inventory['sources'][0]['citation_id']
        assert page['table_ids'] == old_page['table_ids']
        assert page['units_not_in_planning_text'] == old_page['units_not_in_planning_text']
        assert page['opening_units'] == [{'unit_id': u['unit_id'], 'content': u['content']}
                                         for u in old_page['opening_units']]
        represented.extend(page['units_not_in_planning_text'])
        represented.extend(u['unit_id'] for u in page['opening_units'])
    assert len(represented) == len(set(represented)) == len(inventory['units'])
    assert set(represented) == {u['model_unit_id'] for u in inventory['units']}
    original_request = dg.frozen_document_request(inventory, PIN, 'pitch_deck')
    assert dg.request_json(original_request) == dg.request_json(dg.frozen_document_request(
        inventory, PIN, 'pitch_deck', catalog_encoding=dg.CATALOG_ENCODING_V1))
    new_request = dg.frozen_document_request(inventory, PIN, 'pitch_deck', catalog_encoding=dg.CATALOG_ENCODING_V2)
    assert new_request[0:2] == original_request[0:2]
    assert new_request[2]['options'] == original_request[2]['options']
    assert new_request[3] is original_request[3]
    assert new_request[2]['catalog_encoding'] == dg.CATALOG_ENCODING_V2
    assert 'catalog_encoding' not in original_request[2]
    with pytest.raises(dg.GenerationError, match='unknown_document_catalog_encoding'):
        dg.planning_catalog(inventory, catalog_encoding='future-unknown')


def test_large_catalog_encoding_saves_capacity_without_erasing_complete_content():
    value = manifest()
    for i in range(800):
        unit = deepcopy(value['units'][0])
        unit['id'] = f'unit-extra-{i}'
        unit['source_locator']['paragraph'] = 1000 + i
        value['units'].append(unit)
    value['manifest_digest'] = dg.canonical_digest({k: v for k, v in value.items() if k != 'manifest_digest'})
    inventory = dg.evidence_inventory([value])
    old = dg.planning_catalog(inventory)
    compact = dg.planning_catalog(inventory, catalog_encoding=dg.CATALOG_ENCODING_V2)
    assert len(json.dumps(compact, separators=(',', ':'))) < len(json.dumps(old, separators=(',', ':'))) * .62
    assert sum(len(p['opening_units']) + len(p['units_not_in_planning_text'])
               for p in compact['pages']) == len(inventory['units'])


def test_page_encoding_compacts_only_metadata_and_rebuilds_full_replay_and_review():
    value = manifest()
    extra = deepcopy(value['units'][0])
    extra['id'] = 'u_not_selected'
    extra['source_locator']['paragraph'] = 888
    value['units'].append(extra)
    value['manifest_digest'] = dg.canonical_digest({k: v for k, v in value.items() if k != 'manifest_digest'})
    inventory = dg.evidence_inventory([value])
    old_plan = dg.frozen_document_request(inventory, PIN, 'pitch_deck')
    new_plan = dg.frozen_document_request(inventory, PIN, 'pitch_deck', catalog_encoding=dg.CATALOG_ENCODING_V2)
    choice = planned(inventory, table=True)
    old_row, new_row = record(old_plan, choice, 1), record(new_plan, choice, 1)
    old_request = dg.frozen_page_request(old_plan, old_row, 'P1', inventory)
    current = dg.frozen_page_request(new_plan, new_row, 'P1', inventory)
    old, new = old_request[2]['evidence'], current[2]['evidence']
    assert 'catalog_encoding' not in old_request[2] and 'catalog_encoding' not in old
    assert old['units_not_selected']
    assert 'units_not_selected' not in new
    assert current[2]['catalog_encoding'] == new['catalog_encoding'] == dg.CATALOG_ENCODING_V2
    assert new['units'] == old['units']  # Exact complete selected texts / hashes / qualifiers.
    selected_originals = {u['original_unit_id'] for u in new['units']}
    excluded = [u['id'] for u in inventory['units'] if u['id'] not in selected_originals]
    assert new['excluded_unit_count'] == len(excluded)
    assert new['excluded_units_digest'] == dg.canonical_digest(excluded)
    assert new['source_ledger'] == old['source_ledger']
    assert new['tables'][0]['cells'] == [{k: c[k] for k in ('unit_id', 'row', 'column', 'text')}
                                       for c in old['tables'][0]['cells']]
    old_authored, new_authored = record(old_request, author(old_request, table=True), 2), record(
        current, author(current, table=True), 2)
    assert dg.renderable_document_page(old_request, old_authored) == dg.renderable_document_page(current, new_authored)
    review = dg.frozen_page_review_request(current, new_authored, PIN)
    assert review[2]['catalog_encoding'] == review[2]['evidence']['catalog_encoding'] == dg.CATALOG_ENCODING_V2
    checked = record(review, {'verdict': 'pass', 'issues': [], 'financial_status': 'not_established'}, 3)
    assert dg.generation_progress(new_plan, inventory, [new_row, new_authored, checked], PIN)['state'] == 'draft_complete'
    mismatched = deepcopy(current)
    mismatched[2]['evidence']['excluded_units_digest'] = '0' * 64
    with pytest.raises(dg.GenerationError, match='request_binding_changed'):
        dg.replay_document_page(mismatched, new_authored)


def test_native_table_replay_copies_exact_values_and_retains_blank_status():
    inventory, plan_request, plan_row, page_request, page_row = prepared(table=True)
    section = dg.renderable_document_page(page_request, page_row)
    assert section[3] == 'table' and isinstance(section[2], str)
    table = section[4]
    assert table['evidence_revision'] == inventory['inventory_revision']
    assert table['columns'][1]['label'] == '2024 projection'
    assert [r['cells'][1]['text'] for r in table['rows']] == ['120', 'Blank in source']
    assert table['rows'][0]['cells'][1]['status'] == 'source_reported'
    assert table['rows'][1]['cells'][1]['status'] == 'missing'
    assert table['rows'][1]['cells'][1]['value_ref'] is None
    assert dg.replay_document_plan(plan_request, plan_row, inventory)['pages'][0]['table_selections'][0]['row_indexes'] == [1, 2]
    assert dg.generation_progress(plan_request, inventory, [plan_row, page_row], PIN)['phase'] == 'document_page_review'


def test_no_minimum_sentence_filler_for_table_only_page():
    _, _, _, page_request, page_row = prepared(table=True)
    answer = author(page_request, table=True)
    answer['claims'] = []
    row = record(page_request, answer, 2)
    assert dg.replay_document_page(page_request, row)['claims'] == []
    assert dg.renderable_document_page(page_request, row)[3] == 'table'


def test_confident_ocr_numeric_claim_blocks_without_native_corroboration_and_never_repairs_value():
    import hashlib
    from reportlab.pdfgen import canvas
    from agents.research.document_evidence_manifest import attach_document_ocr
    stream = BytesIO(); pdf = canvas.Canvas(stream)
    pdf.drawString(30, 700, 'Supplied source heading'); pdf.showPage(); pdf.save()
    raw = stream.getvalue()
    source = build_document_manifest(raw, 'pdf', tenant_id='tenant', deal_id='deal', document_id='ocr-source')
    proposed = [{'page': 1, 'source_sha256': hashlib.sha256(raw).hexdigest(), 'image_sha256': 'c' * 64,
        'engine': 'AppleVision-VNRecognizeTextRequest', 'coordinate_system': 'normalized-bottom-left',
        'language_correction': False, 'observations': [{'text': 'Projected amount 170',
            'confidence': 1.0, 'bbox': [.1, .2, .5, .1]}]}]
    augmented = attach_document_ocr(source, proposed, expected_binding=source['binding'])
    inventory = dg.evidence_inventory([augmented])
    unit = next(u for u in inventory['units'] if u['kind'] == 'ocr_text')
    request = dg.frozen_document_request(inventory, PIN, 'pitch_deck')
    plan = {'pages': [{'page_id': 'P1', 'topic': 'financials', 'layout': 'evidence',
        'narrative_unit_ids': [unit['model_unit_id']], 'table_selections': [], 'coverage': 'limited'}],
        'unsupported_topics': [topic for topic in dg.TOPICS if topic != 'financials']}
    plan_row = record(request, plan, 1)
    page_request = dg.frozen_page_request(request, plan_row, 'P1', inventory)
    answer = {'page_id': 'P1', 'title': 'Source observations', 'claims': [{
        'text': 'The source projects an amount of 170.', 'basis': 'projected',
        'evidence': [{'unit_id': unit['model_unit_id'], 'quote': 'Projected amount 170'}]}], 'table_labels': []}
    row = record(page_request, answer, 2)
    with pytest.raises(dg.GenerationError, match='ocr_numeric_claim_requires_independent_source_validation'):
        dg.replay_document_page(page_request, row)
    assert unit['content'] == 'Projected amount 170' and unit['confidence'] == 1.0
    assert augmented['financial_validation'] == 'not_established'
    answer['claims'][0]['text'] = 'The supplied OCR describes a projected amount; its reading is unverified.'
    row = record(page_request, answer, 2)
    section = dg.renderable_document_page(page_request, row)
    assert 'Includes unverified OCR.' in section[2]


@pytest.mark.parametrize('mutation', ['invented_number', 'wrong_source', 'lost_period', 'format_receipt', 'full_context'])
def test_unsupported_or_mutated_page_is_blocked_without_repair(mutation):
    inventory, request, plan_row, page_request, row = prepared(table=True)
    if mutation in {'invented_number', 'wrong_source', 'lost_period'}:
        answer = deepcopy(row['answer'])
        if mutation == 'invented_number':
            answer['claims'][0]['text'] = 'The pilot generated 999 of revenue.'
        elif mutation == 'wrong_source':
            answer['claims'][0]['evidence'][0]['unit_id'] = 'U9999'
        else:
            answer['table_labels'][0]['column_labels'][1] = 'Results'
        row = record(page_request, answer, 2)
    elif mutation == 'format_receipt':
        row['routing'].pop('format_object_receipts')
    else:
        row['routing']['prompt_tokens'] = row['routing']['context_tokens'] // 2 + 2
    progress = dg.generation_progress(request, inventory, [plan_row, row], PIN)
    assert progress['state'] == 'blocked'
    assert row['raw_response'] and progress['recorded_model_calls'] == 2


def test_plan_cannot_silently_omit_topics_or_select_foreign_cells():
    inventory = dg.evidence_inventory([manifest()])
    request = dg.frozen_document_request(inventory, PIN, 'pitch_deck')
    bad = planned(inventory)
    bad['unsupported_topics'].pop()
    assert dg.generation_progress(request, inventory, [record(request, bad, 1)], PIN)['state'] == 'blocked'
    with pytest.raises(dg.GenerationError, match='one_private_deal'):
        other = manifest()
        other['binding']['tenant_id'] = 'foreign'
        other['binding']['document_id'] = 'foreign-document'
        other['units'] = [{**u, 'id': u['id'] + '-foreign'} for u in other['units']]
        other['tables'] = []
        other['manifest_digest'] = dg.canonical_digest({k: v for k, v in other.items() if k != 'manifest_digest'})
        dg.evidence_inventory([manifest(), other])


def test_independent_review_must_bind_exact_phrase_and_cannot_approve_financials():
    inventory, plan_request, plan_row, page_request, page_row = prepared()
    review_request = dg.frozen_page_review_request(page_request, page_row, PIN)
    passing = record(review_request, {'verdict': 'pass', 'issues': [], 'financial_status': 'not_established'}, 3)
    complete = dg.generation_progress(plan_request, inventory, [plan_row, page_row, passing], PIN)
    assert complete['state'] == 'draft_complete'
    assert complete['investor_approval'] == complete['financial_validation'] == complete['regulatory_approval'] == 'not_established'
    bad = {'verdict': 'revise', 'issues': [{'claim_index': 0, 'exact_claim_phrase': 'nonexistent phrase',
        'unit_id': 'U1', 'exact_source_quote': 'limited pilot', 'kind': 'scope'}], 'financial_status': 'not_established'}
    assert dg.generation_progress(plan_request, inventory, [plan_row, page_row,
        record(review_request, bad, 3)], PIN)['state'] == 'blocked'
