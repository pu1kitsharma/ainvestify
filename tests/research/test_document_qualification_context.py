"""Model-only coverage correction and full original-page qualifier propagation."""
from copy import deepcopy
from io import BytesIO
import json

import pytest
from pptx import Presentation
from pptx.util import Inches

from agents.research import document_generation as dg
from agents.research import document_qualification_context as qc
from agents.research.document_evidence_manifest import build_document_manifest
from tests.research.test_document_generation import PIN, manifest, record, planned, author
from tests.research.test_document_source_unit_selector import selected


CRITIC = {'name': 'local-critic:14b', 'digest': 'b' * 64}


def review_answer(*, verdict='pass', patches=None):
    return {'verdict': verdict, 'topic_patches': patches or [],
        'review_scope': 'local_model_source_reading_only', 'financial_validation': 'not_established',
        'regulatory_approval': 'not_established', 'investor_approval': 'not_established'}


def proposal_unit(source, text, *, page=None, kind=None):
    unit = deepcopy(source['units'][0])
    unit.update(id='u_context_' + str(len(source['units'])), content=text, text=text,
                content_sha256=dg.canonical_digest(text))
    unit['source_locator']['paragraph'] = len(source['units']) + 1
    if page is not None:
        unit['source_locator']['slide'] = page
    if kind is not None:
        unit['kind'] = kind
    source['units'].append(unit)
    source['manifest_digest'] = dg.canonical_digest({key: value for key, value in source.items() if key != 'manifest_digest'})
    return unit


def prepared(*, source=None, table=False):
    # Keep the synthetic mandate common across all immutable stage requests.
    inventory = dg.evidence_inventory([source or manifest()], catalog_encoding=dg.CATALOG_ENCODING_V3)
    profiles = [dg.SOURCE_MAP_BOUND_UNIT_PROFILE]
    options = {'catalog_encoding': dg.CATALOG_ENCODING_V3, 'source_map_profiles': profiles,
               'qualification_context_profile': qc.PROFILE, 'source_coverage_reviewer_pin': CRITIC,
               'mandate': 'Private internal diligence.'}
    seed = dg.frozen_document_request(inventory, PIN, 'pitch_deck', **options)
    choice = selected(seed, topic='financials' if table else 'customers_and_traction', ids=['U1'])
    if table:
        cell = next(unit['model_unit_id'] for unit in inventory['units'] if unit['kind'] == 'table_cell' and unit['content'])
        choice['financials'].update(unit_ids=['U1', cell], table_ids=[inventory['tables'][0]['id']])
    mapped = record(seed, choice, 1)
    critic = qc.frozen_source_map_coverage_request(seed, mapped, CRITIC)
    reviewed = record(critic, review_answer(), 2)
    plan = dg.frozen_document_request(inventory, PIN, 'pitch_deck', source_maps=[mapped],
        source_map_reviews=[reviewed], **options)
    planned_row = record(plan, planned(inventory, table=table), 3)
    page = dg.frozen_page_request(plan, planned_row, 'P1', inventory)
    return inventory, profiles, seed, [mapped, reviewed, planned_row], plan, page


def test_original_page_qualifier_is_quoteable_and_recorded_basis_survives_export():
    source = manifest()
    qualifier = proposal_unit(source, 'Figures are illustrative projections for 2025.')
    inventory, profiles, seed, rows, plan, page = prepared(source=source)
    uid = next(unit['model_unit_id'] for unit in inventory['units'] if unit['id'] == qualifier['id'])
    assert uid not in page[2]['evidence']['selected_unit_ids']
    assert uid in qc._context_units(page)
    assert page[2]['evidence']['read_unit_count'] == len(inventory['units'])
    assert page[2]['evidence']['excluded_unit_count'] == 0
    assert qc._context_units(page)[uid]['content'] == qualifier['content']
    proposed = {'page_id': 'P1', 'title': 'Supplied projections', 'claims': [
        {'text': qualifier['content'], 'basis': 'projected', 'evidence': [{'unit_id': uid, 'quote': qualifier['content']}]}],
        'table_labels': []}
    authored = record(page, proposed, 4)
    assert dg.replay_document_page(page, authored) == proposed
    reviewer = dg.frozen_page_review_request(page, authored, PIN)
    assert reviewer[2]['mandate'] == page[2]['mandate'] == plan[2]['mandate']
    assert reviewer[2]['companion_source_pages'] == page[2]['companion_source_pages']
    reviewed = record(reviewer, {'verdict': 'pass', 'issues': [], 'financial_status': 'not_established'}, 5)
    section = dg.renderable_document_page(page, authored, reviewer, reviewed)
    assert section[1] == '[Projected] ' + qualifier['content'] + ' [S1]'
    assert 'unverified' in section[2] and qualifier['content'] == authored['answer']['claims'][0]['text']
    progress = dg.generation_progress(seed, inventory, [*rows, authored, reviewed], PIN,
        source_map_profiles=profiles, qualification_context_profile=qc.PROFILE, source_coverage_reviewer_pin=CRITIC)
    assert progress['state'] == 'draft_complete'
    assert progress['source_coverage_review_status'] == 'local_model_only_unqualified'
    assert progress['financial_validation'] == progress['investor_approval'] == progress['regulatory_approval'] == 'not_established'


def test_all_maps_precede_compact_corrections_and_plan_binds_both_original_model_roles():
    source = manifest()
    second = deepcopy(source['slides'][0]); second['slide'] = 2
    source['slides'].append(second)
    proposal_unit(source, 'A proposed second-page product observation.', page=2)
    inventory = dg.evidence_inventory([source], catalog_encoding=dg.CATALOG_ENCODING_V3)
    profiles = [dg.SOURCE_MAP_BOUND_UNIT_PROFILE] * 2
    maps = dg.frozen_source_map_requests(inventory, PIN, source_map_profiles=profiles)
    original = dg.frozen_document_request(inventory, PIN, 'intro_deck', catalog_encoding=dg.CATALOG_ENCODING_V3,
        source_map_profiles=profiles)
    initial = dg.frozen_document_request(inventory, PIN, 'intro_deck', catalog_encoding=dg.CATALOG_ENCODING_V3,
        source_map_profiles=profiles, qualification_context_profile=qc.PROFILE, source_coverage_reviewer_pin=CRITIC)
    assert dg.request_json(initial) == dg.request_json(original)
    kwargs = {'source_map_profiles': profiles, 'qualification_context_profile': qc.PROFILE,
              'source_coverage_reviewer_pin': CRITIC}
    rows = [record(maps[0], selected(maps[0], topic='product', ids=['U1']), 1)]
    next_call = dg.generation_progress(initial, inventory, rows, PIN, **kwargs)
    assert next_call['phase'] == 'document_source_map' and next_call['source_map_index'] == 2
    rows.append(record(maps[1], selected(maps[1]), 2))
    next_call = dg.generation_progress(initial, inventory, rows, PIN, **kwargs)
    assert next_call['phase'] == 'document_source_coverage_review' and next_call['maximum_model_calls'] == 85
    assert next_call['request'][2]['model_pin'] == CRITIC
    rows.append(record(next_call['request'], review_answer(), 3))
    next_call = dg.generation_progress(initial, inventory, rows, PIN, **kwargs)
    uid = maps[1][2]['source_page']['literal_units'][0][0]
    patch = {'topic': 'product', 'selection': {'status': 'selected', 'coverage': 'limited', 'basis': 'planned',
                                             'unit_ids': [uid], 'table_ids': []}}
    rows.append(record(next_call['request'], review_answer(verdict='corrected', patches=[patch]), 4))
    next_call = dg.generation_progress(initial, inventory, rows, PIN, **kwargs)
    assert next_call['phase'] == 'document_plan' and next_call['completed_source_coverage_reviews'] == 2
    catalog = next_call['request'][2]['catalog']
    assert {unit['unit_id'] for unit in catalog['literal_unit_pool']} == {'U1', uid}
    assert catalog['source_map_status'] == 'proposed_unreviewed'
    choice = {'pages': [{'page_id': 'P1', 'topic': 'product', 'layout': 'evidence', 'coverage': 'limited',
                         'narrative_unit_ids': [uid], 'table_selections': []}],
              'unsupported_topics': [topic for topic in dg.TOPICS if topic != 'product']}
    planned_row = record(next_call['request'], choice, 5)
    actual = dg.document_plan_context(initial, inventory, [*rows, planned_row], reviewer_pin=PIN, **kwargs)
    assert actual[1] == planned_row and actual[2] == 5
    assert rows[1]['answer']['product']['status'] == 'unassessed'  # Original mapper is preserved.
    with pytest.raises(dg.GenerationError):
        dg.document_plan_context(initial, inventory, [*rows, planned_row], source_map_profiles=profiles,
            qualification_context_profile=qc.PROFILE, reviewer_pin=PIN)  # Wrong trusted critic pin.


def test_native_table_cells_keep_original_ids_but_footer_uses_bound_model_handles():
    inventory, profiles, seed, rows, plan, page = prepared(table=True)
    proposed = author(page)
    original = page[2]['evidence']['tables'][0]
    proposed['table_labels'] = [{'table_id': original['table_id'], 'column_labels': ['Measure', '2024 projection'],
        'caption': 'Source table; financial validation pending.', 'evidence': [
            {'unit_id': cell['model_unit_id'], 'quote': cell['text']} for cell in original['cells'] if cell['row'] == 1]}]
    authored = record(page, proposed, 4)
    section = dg.renderable_document_page(page, authored)
    assert section[3] == 'table' and '[Reported]' in section[1] and '[Table semantics unassessed]' in section[1]
    assert 'slide 1' in section[2] and 'financial review pending' in section[2]
    payload = section[4]
    assert payload['columns'][1]['label'] == '2024 projection'
    assert payload['rows'][0]['cells'][1]['text'] == '120'
    assert payload['rows'][0]['cells'][1]['value_ref'] in {unit['id'] for unit in inventory['units']}
    assert payload['rows'][1]['cells'][1]['text'] == 'Blank in source'
    assert payload['rows'][1]['cells'][1]['value_ref'] is None


def test_foreign_or_blank_corrections_false_pass_duplicate_topics_and_unassessed_block():
    inventory, profiles, seed, rows, plan, page = prepared()
    review = qc.frozen_source_map_coverage_request(seed, rows[0], CRITIC)
    blank = next(unit['model_unit_id'] for unit in inventory['units'] if not unit['content'].strip())
    choice = {'status': 'selected', 'coverage': 'limited', 'basis': 'reported', 'unit_ids': ['U1'], 'table_ids': []}
    patch = {'topic': 'financials', 'selection': choice}
    for verdict, patches in [('pass', [patch]), ('corrected', [patch, patch]),
        ('corrected', [{'topic': 'financials', 'selection': {**choice, 'unit_ids': ['foreign']}}]),
        ('corrected', [{'topic': 'financials', 'selection': {**choice, 'unit_ids': [blank]}}])]:
        with pytest.raises(ValueError):
            qc.replay_source_map_coverage(review, record(review, review_answer(verdict=verdict, patches=patches), 2))
    blocked = dg.generation_progress(seed, inventory, [rows[0], record(review, review_answer(verdict='unassessed'), 2)], PIN,
        source_map_profiles=profiles, qualification_context_profile=qc.PROFILE, source_coverage_reviewer_pin=CRITIC)
    assert blocked['state'] == 'blocked' and blocked['reason'] == 'local_source_coverage_review_unassessed'


def test_complete_companion_capacity_blocks_without_mutating_or_clipping_originals():
    source = manifest()
    original = proposal_unit(source, 'Complete supplied scope. ' * 750)
    second = deepcopy(source['slides'][0]); second['slide'] = 2
    source['slides'].append(second)
    companion = proposal_unit(source, 'Another complete source scope. ' * 650, page=2)
    inventory = dg.evidence_inventory([source], catalog_encoding=dg.CATALOG_ENCODING_V3)
    profiles = [dg.SOURCE_MAP_BOUND_UNIT_PROFILE] * 2
    requests = dg.frozen_source_map_requests(inventory, PIN, kind='pitch_deck', source_map_profiles=profiles)
    request = requests[0]
    assert original['content'] in {row[2] for row in request[2]['source_page']['literal_units']}
    # Quoteable companion representation stays complete; extra planned evidence
    # may exceed the finite page budget and must never be silently shortened.
    selected_ids = [unit['model_unit_id'] for unit in inventory['units'] if unit['id'] in {original['id'], companion['id']}]
    mapped = [record(req, selected(req, topic='customers_and_traction', ids=[selected_ids[index]]), index + 1)
              for index, req in enumerate(requests)]
    reviewed = [record(qc.frozen_source_map_coverage_request(req, row, CRITIC), review_answer(), index + 3)
                for index, (req, row) in enumerate(zip(requests, mapped))]
    plan = dg.frozen_document_request(inventory, PIN, 'pitch_deck', catalog_encoding=dg.CATALOG_ENCODING_V3,
        source_map_profiles=profiles, source_maps=mapped, source_map_reviews=reviewed,
        qualification_context_profile=qc.PROFILE, source_coverage_reviewer_pin=CRITIC)
    choice = planned(inventory); choice['pages'][0]['narrative_unit_ids'] = selected_ids
    with pytest.raises(dg.GenerationError, match='complete_companion_context_requires_model_split_no_clipping'):
        dg.frozen_page_request(plan, record(plan, choice, 3), 'P1', inventory)
    assert next(unit['content'] for unit in inventory['units'] if unit['id'] == original['id']) == original['content']
    assert request[2]['source_page']['literal_units'][-1][2] == original['content']


def test_dense_native_table_uses_one_complete_companion_and_compact_exact_primary_under_bound():
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    table = slide.shapes.add_table(36, 11, Inches(1), Inches(1), Inches(10), Inches(5)).table
    for row in range(36):
        for column in range(11):
            table.cell(row, column).text = ('Column ' + str(column) if row == 0 else
                ('Measure ' + str(row) if column == 0 else str(row * column)))
    stream = BytesIO(); presentation.save(stream)
    source = build_document_manifest(stream.getvalue(), 'pptx', tenant_id='tenant', deal_id='deal', document_id='dense')
    inventory = dg.evidence_inventory([source], catalog_encoding=dg.CATALOG_ENCODING_V3)
    profiles = [dg.SOURCE_MAP_BOUND_UNIT_PROFILE]
    seed = dg.frozen_document_request(inventory, PIN, 'pitch_deck', catalog_encoding=dg.CATALOG_ENCODING_V3,
        source_map_profiles=profiles)
    original = inventory['tables'][0]
    header_id = next(unit['model_unit_id'] for unit in inventory['units'] if unit['content'] == 'Column 0')
    mapped = record(seed, selected(seed, topic='financials', ids=[header_id], tables=[original['id']]), 1)
    coverage = qc.frozen_source_map_coverage_request(seed, mapped, CRITIC)
    reviewed = record(coverage, review_answer(), 2)
    plan = dg.frozen_document_request(inventory, PIN, 'pitch_deck', catalog_encoding=dg.CATALOG_ENCODING_V3,
        source_map_profiles=profiles, source_maps=[mapped], source_map_reviews=[reviewed],
        qualification_context_profile=qc.PROFILE, source_coverage_reviewer_pin=CRITIC)
    geometry = {'narrative_unit_ids': [], 'table_selections': [{'table_id': original['id'],
        'column_indexes': list(range(7)), 'row_indexes': [1, 2, 3], 'header_row_index': 0}]}
    default_primary = dg.selected_page_evidence(inventory, geometry, catalog_encoding=dg.CATALOG_ENCODING_V3)
    companion = seed[2]['source_page']
    assert len(json.dumps(default_primary, ensure_ascii=False)) + len(json.dumps([companion], ensure_ascii=False)) > dg.MAX_PAGE_CHARACTERS
    choice = {'pages': [{'page_id': 'P1', 'topic': 'financials', 'layout': 'table', 'coverage': 'limited', **geometry}],
              'unsupported_topics': [topic for topic in dg.TOPICS if topic != 'financials']}
    planned_row = record(plan, choice, 3)
    page = dg.frozen_page_request(plan, planned_row, 'P1', inventory)
    primary = page[2]['evidence']
    assert primary['units'] == [] and len(primary['tables'][0]['cells']) == 28
    assert len(json.dumps(page[2], ensure_ascii=False)) < dg.MAX_PAGE_CHARACTERS
    assert page[2]['companion_source_pages'] == [companion]
    assert primary['read_unit_count'] == len(inventory['units']) == 396
    decoded = qc._context_units(page)
    assert all(decoded[unit['model_unit_id']]['content'] == unit['content'] for unit in inventory['units'])
    header = [cell for cell in primary['tables'][0]['cells'] if cell['row'] == 1]
    answer = {'page_id': 'P1', 'title': 'Supplied table', 'claims': [], 'table_labels': [
        {'table_id': original['id'], 'column_labels': [cell['text'] for cell in header], 'caption': 'Supplied source table.',
         'evidence': [{'unit_id': cell['model_unit_id'], 'quote': cell['text']} for cell in header]}]}
    authored = record(page, answer, 4)
    section = dg.renderable_document_page(page, authored)
    assert section[4]['rows'][0]['cells'][1]['text'] == original['rows'][1][1]
    original_cell = next(cell for cell in original['cells'] if cell['row'] == 2 and cell['column'] == 2)
    assert section[4]['rows'][0]['cells'][1]['value_ref'] == original_cell['unit_id']
    assert 'slide 1' in section[2]
    hints = plan[2]['catalog']['table_navigation'][0]['capacity_row_hints']
    assert next(hint['deck_max_body_rows'] for hint in hints if hint['column_count'] == 7) >= 3
    assert dg.selected_page_evidence(inventory, geometry, catalog_encoding=dg.CATALOG_ENCODING_V3) == default_primary


def test_planner_table_navigation_exposes_and_binds_original_header_body_and_cell_coordinates():
    inventory, profiles, seed, rows, plan, page = prepared(table=True)
    catalog = plan[2]['catalog']
    table = inventory['tables'][0]
    navigation = catalog['table_navigation'][0]
    assert navigation['index_origin'] == 0 and navigation['row_count'] == 3 and navigation['column_count'] == 2
    assert navigation['rows'][0] == ['Measure', '2024 projection']
    assert navigation['rows'][1] == ['Revenue', '120']
    assert navigation['rows'] == table['rows'] and navigation['source_locator'] == table['source_locator']
    originals = {unit['model_unit_id']: unit for unit in inventory['units']}
    cells = {(cell['row'] - 1, cell['column'] - 1): cell for cell in table['cells']}
    for row_index, unit_ids in enumerate(navigation['model_cell_unit_ids']):
        for column_index, uid in enumerate(unit_ids):
            original = originals[uid]
            assert original['id'] == cells[row_index, column_index]['unit_id']
            assert original['content'] == navigation['rows'][row_index][column_index]
    assert all(unit['source_locator'] == originals[unit['unit_id']]['source_locator']
               for unit in catalog['literal_unit_pool'])
    assert navigation['original_cells_binding_digest'] == dg.canonical_digest(table['cells'])
    assert plan[1] == qc.PLAN_CONTEXT_INSTRUCTION
    choice = planned(inventory, table=True)
    choice['pages'][0]['table_selections'][0]['row_indexes'] = [1]
    assert dg.replay_document_plan(plan, record(plan, choice, 3), inventory) == choice
    # Visible navigation cells allow original table ranges, not new narrative
    # choices outside the actual mapper/critic-selected evidence pool.
    unmapped = navigation['model_cell_unit_ids'][1][1]
    assert unmapped not in {unit['unit_id'] for unit in catalog['literal_unit_pool']}
    foreign_choice = deepcopy(choice); foreign_choice['pages'][0]['narrative_unit_ids'].append(unmapped)
    with pytest.raises(dg.GenerationError, match='unmapped_source_evidence'):
        dg.replay_document_plan(plan, record(plan, foreign_choice, 3), inventory)
    changed = deepcopy(plan[2]); changed['catalog']['table_navigation'][0]['rows'][1][1] = '999'
    changed['catalog']['table_navigation_digest'] = dg.canonical_digest(changed['catalog']['table_navigation'])
    changed.pop('request_binding_digest')
    rebound = dg._request(plan[0], plan[1], changed, plan[3])
    with pytest.raises(dg.GenerationError, match='complete_original_table_navigation_changed'):
        dg.replay_document_plan(rebound, record(rebound, choice, 3), inventory)
    legacy = dg.frozen_document_request(inventory, PIN, 'pitch_deck', mandate=plan[2]['mandate'],
        catalog_encoding=dg.CATALOG_ENCODING_V3, source_map_profiles=profiles, source_maps=[rows[0]])
    assert 'table_navigation' not in legacy[2]['catalog']
    assert all('source_locator' not in unit for unit in legacy[2]['catalog']['literal_unit_pool'])
