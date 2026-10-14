"""Fictional source/image gates; stub receipts are not live model acceptance."""
import base64
from copy import deepcopy
import hashlib
from io import BytesIO
import json
from zipfile import ZipFile

import pytest
from pydantic import ValidationError
from pptx import Presentation
from pptx.util import Inches

from agents.inference import local_document_vision as vision
from agents.inference.local_models import LocalModel, _transport_digest, generation_schema
from agents.inference.model_authorship import recorded_call, digest
from agents.preparation.preparation_budget import ACTIVE_BUDGET, PreparationBudget, preparation_budget
from agents.research import document_generation as dg
from agents.research.document_evidence_manifest import build_document_manifest
from agents.research.document_evidence_manifest import attach_document_visual_text
from tests.research.test_document_generation import PIN, manifest
from tests.research.test_document_native_visual_text import IDENTITY, source_bytes
from tests.research.test_emf_text import build, font, select, textw

IMAGE = base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+jRZkAAAAASUVORK5CYII=')


def frozen():
    inv = dg.evidence_inventory([manifest()], catalog_encoding=dg.CATALOG_ENCODING_V3)
    group = dg.frozen_source_map_requests(inv, PIN)[0][2]['source_page']
    binding = {'source_sha256': group['source_id'], 'source_page_or_slide': group['page_or_slide'],
        'rendered_pdf_sha256': 'b' * 64, 'image_sha256': hashlib.sha256(IMAGE).hexdigest(),
        'raster_engine': {'name': 'synthetic-raster', 'binary_sha256': 'c' * 64},
        'renderer_engine': {'name': 'synthetic-renderer', 'binary_sha256': 'd' * 64},
        'render_receipt_sha256': 'e' * 64}
    req = vision.frozen_table_request(IMAGE, group, PIN, raster_binding=binding)
    ids = {u['content']: uid for uid, u in dg.source_map_literal_lookup(group).items()}
    answer = {'status': 'proposed', 'tables': [{'caption': 'Reported 2024 projection',
        'caption_evidence': [{'unit_id': ids['2024 projection'], 'quote': '2024 projection'}],
        'basis': 'projected', 'header_row_index': 0,
        'rows': [[ids['Measure'], ids['2024 projection']], [ids['Revenue'], ids['120']], [ids['Cost'], None]]}]}
    return req, answer, ids


def stub_record(monkeypatch, req, answer):
    captured = []
    raw = json.dumps(answer, ensure_ascii=False)
    def transport(self, operation, **kwargs):
        ACTIVE_BUDGET.get().start_request()
        captured.append(deepcopy(kwargs))
        self._record_format_object(operation, kwargs)
        prompt_tokens = 18167 if kwargs['options']['num_ctx'] == 32768 else 1100
        self.last_call['stream_observation'] = {
            'contract': vision.STREAM_OBSERVATION_V1, 'state': 'terminal',
            'terminal_observed': True, 'client_close_completed': True,
            'model': kwargs['model'], 'returned_model': kwargs['model'], 'done_reason': 'stop',
            'prompt_tokens': prompt_tokens, 'generated_tokens': 100, 'unexpected_thinking_observed': False,
            'actual_options_digest': _transport_digest(kwargs['options']),
            'source_messages_digest': _transport_digest(kwargs['messages']),
            'evidence_digest': _transport_digest(kwargs['messages'][1]['content']),
            'format_digest': _transport_digest(kwargs['format']),
            'content_sha256': hashlib.sha256(raw.encode()).hexdigest(),
            'content_characters': len(raw), 'content_bytes': len(raw.encode()),
            'observed_chunks': 2, 'observed_content_chunks': 1,
            'first_chunk_seconds': .1, 'first_content_seconds': .1,
            'last_chunk_seconds': .2, 'elapsed_seconds': .3}
        return {'model': kwargs['model'], 'done_reason': 'stop', 'prompt_eval_count': prompt_tokens,
            'eval_count': 100, 'message': {'content': raw}}
    monkeypatch.setattr(LocalModel, '_request', transport)
    model = vision.LocalDocumentVision(PIN['name'], image_bytes=IMAGE, **req[2]['options'])
    budget = PreparationBudget(105, max_calls=1, max_requests=6)
    rows = []
    with preparation_budget(budget):
        recorded_call(model, *req, rows, lambda: None)
    assert budget.calls == 1 and budget.requests == 1 and len(captured) == 1
    return rows[0], captured[0]


def test_actual_image_transport_and_exact_native_matrix_never_qualify_financials(monkeypatch):
    req, answer, ids = frozen()
    row, transport = stub_record(monkeypatch, req, answer)
    assert transport['model'] == PIN['name'] and transport['think'] is False
    assert base64.b64decode(transport['messages'][1]['images'][0]) == IMAGE
    assert json.loads(transport['messages'][1]['content']) == req[2]
    assert transport['format'] == generation_schema(req[3], bound_strings=True)
    result = vision.replay_table(req, row, IMAGE)
    assert result['tables'][0]['cells'][1][1] == {
        'unit_id': ids['120'], 'text': '120', 'binding_status': 'exact_original_literal'}
    assert result['tables'][0]['cells'][2][1] == {
        'unit_id': None, 'text': None, 'binding_status': 'unbound'}
    assert result['semantic_review'] == 'required_not_performed'
    assert all(result[k] == 'not_established' for k in
               ('financial_validation', 'regulatory_approval', 'investor_approval'))
    assert result['release_status'] == 'draft_only'


def test_foreign_blank_ocr_ids_and_wrong_own_caption_quote_cannot_bind(monkeypatch):
    req, answer, ids = frozen()
    for uid in ('U9999', ids['']):
        bad = deepcopy(answer)
        bad['tables'][0]['rows'][1][1] = uid
        with pytest.raises(ValidationError): req[3].model_validate(bad)
    group = deepcopy(req[2]['source_page'])
    first = next(x for x in group['contexts'] if x[1] == 'text')
    first[1] = 'ocr_text'
    ocr = vision.frozen_table_request(IMAGE, group, PIN, raster_binding=req[2]['raster_binding'])
    ocr_id = next(uid for uid, u in dg.source_map_literal_lookup(group).items() if u['kind'] == 'ocr_text')
    assert ocr_id in dg.source_map_literal_lookup(ocr[2]['source_page'])  # Full context remains.
    bad = deepcopy(answer); bad['tables'][0]['rows'][1][1] = ocr_id
    with pytest.raises(ValidationError): ocr[3].model_validate(bad)
    bad = deepcopy(answer)
    bad['tables'][0]['caption_evidence'][0]['unit_id'] = ids['Revenue']
    row, _ = stub_record(monkeypatch, req, bad)
    # The quote exists elsewhere on the page, but not in its claimed own unit.
    with pytest.raises(dg.GenerationError, match='quote_not_in_own'):
        vision.replay_table(req, row, IMAGE)


def test_image_raw_model_and_full_multimodal_context_bindings_fail_closed(monkeypatch):
    req, answer, _ = frozen()
    row, _ = stub_record(monkeypatch, req, answer)
    with pytest.raises(dg.GenerationError, match='image_page_binding'):
        vision.replay_table(req, row, IMAGE + b'changed')
    changed = deepcopy(row)
    changed['routing']['vision_input_binding']['model_pin']['digest'] = 'f' * 64
    with pytest.raises(dg.GenerationError, match='actual_local_transport'):
        vision.replay_table(req, changed, IMAGE)
    for field, value in (('source_messages_digest', '0' * 64), ('terminal_observed', False)):
        changed = deepcopy(row)
        changed['routing']['stream_observation'][field] = value
        with pytest.raises(ValueError): vision.replay_table(req, changed, IMAGE)
    changed = deepcopy(row)
    changed['routing']['prompt_tokens'] = None
    with pytest.raises(ValueError, match='input token receipt'):
        vision.replay_table(req, changed, IMAGE)
    changed = deepcopy(row)
    changed['raw_response'] += ' '
    changed['response_hash'] = digest(changed['raw_response'])
    with pytest.raises(ValueError, match='exact observed bytes'):
        vision.replay_table(req, changed, IMAGE)
    changed = deepcopy(answer); changed['investor_approval'] = 'approved'
    with pytest.raises(ValidationError): req[3].model_validate(changed)


def test_image_page_engines_bounds_and_direct_unbudgeted_use_are_rejected():
    req, answer, _ = frozen()
    for key, value in (('source_page_or_slide', 2), ('image_sha256', 'f' * 64)):
        binding = deepcopy(req[2]['raster_binding']); binding[key] = value
        with pytest.raises(dg.GenerationError):
            vision.frozen_table_request(IMAGE, req[2]['source_page'], PIN, raster_binding=binding)
    binding = deepcopy(req[2]['raster_binding']); del binding['render_receipt_sha256']
    with pytest.raises(dg.GenerationError, match='render_receipt'):
        vision.frozen_table_request(IMAGE, req[2]['source_page'], PIN, raster_binding=binding)
    bad = deepcopy(answer); bad['tables'][0]['rows'][0] *= 4
    with pytest.raises(ValidationError): req[3].model_validate(bad)
    model = vision.LocalDocumentVision(PIN['name'], image_bytes=IMAGE, **req[2]['options'])
    with pytest.raises(dg.GenerationError, match='one_recorded_call'):
        model.generate(req[1], json.dumps(req[2]), req[3])


def test_explicit_32k_profile_keeps_old_bytes_and_requires_actual_full_context(monkeypatch):
    old, answer, _ = frozen()
    old_bytes = dg.request_json(old)
    larger = vision.frozen_table_request(IMAGE, old[2]['source_page'], PIN,
        raster_binding=old[2]['raster_binding'], vision_context_profile=vision.CONTEXT_PROFILE_32K)
    assert larger[2]['options']['context_tokens'] == 32768
    assert larger[2]['options']['max_tokens'] == old[2]['options']['max_tokens'] == 2200
    assert old[2]['options']['context_tokens'] == 16384
    assert 'vision_context_profile' not in old[2]
    assert dg.request_json(vision.frozen_table_request(IMAGE, old[2]['source_page'], PIN,
        raster_binding=old[2]['raster_binding'])) == old_bytes
    with pytest.raises(dg.GenerationError, match='unknown_explicit'):
        vision.frozen_table_request(IMAGE, old[2]['source_page'], PIN,
            raster_binding=old[2]['raster_binding'], vision_context_profile='arbitrary-larger-context')
    row, transport = stub_record(monkeypatch, larger, answer)
    assert transport['options']['num_ctx'] == 32768
    assert row['routing']['prompt_tokens'] == 18167
    assert vision.replay_table(larger, row, IMAGE)['status'] == 'proposed'
    for count in (16386, 31000):
        changed = deepcopy(row)
        changed['routing']['prompt_tokens'] = count
        changed['routing']['stream_observation']['prompt_tokens'] = count
        with pytest.raises(ValueError, match='context integrity'):
            vision.replay_table(larger, changed, IMAGE)
    stripped = list(deepcopy(larger))
    del stripped[2]['vision_context_profile']
    with pytest.raises(dg.GenerationError, match='frozen_request'):
        vision.replay_table(stripped, row, IMAGE)


def test_complete_matrix_profile_supports_dense_grid_single_state_and_cell_bound(monkeypatch):
    old, legacy_answer, ids = frozen()
    before = dg.request_json(old)
    # A complete synthetic dense source: every cell has a distinct original ID.
    ppt = Presentation(); slide = ppt.slides.add_slide(ppt.slide_layouts[6])
    native = slide.shapes.add_table(18, 10, Inches(1), Inches(1), Inches(8), Inches(5)).table
    for r in range(18):
        for c in range(10): native.cell(r, c).text = f'Synthetic cell {r},{c}'
    output = BytesIO(); ppt.save(output)
    source = build_document_manifest(output.getvalue(), 'pptx', tenant_id='tenant', deal_id='deal', document_id='dense')
    inventory = dg.evidence_inventory([source], catalog_encoding=dg.CATALOG_ENCODING_V3)
    group = dg._source_map_page_view(inventory, inventory['source_map_topology'][0])
    own = {u['content']: uid for uid, u in dg.source_map_literal_lookup(group).items()}
    binding = deepcopy(old[2]['raster_binding']); binding['source_sha256'] = group['source_id']
    req = vision.frozen_table_request(IMAGE, group, PIN,
        raster_binding=binding, vision_context_profile=vision.CONTEXT_PROFILE_32K,
        matrix_schema_profile=vision.MATRIX_PROFILE_V2)
    assert set(req[3].model_json_schema()['properties']) == {'tables'}
    table = {'caption': 'Synthetic cell 0,0', 'caption_evidence': [
        {'unit_id': own['Synthetic cell 0,0'], 'quote': 'Synthetic cell 0,0'}],
        'basis': 'unassessed', 'header_row_index': 0,
        'rows': [[own[f'Synthetic cell {r},{c}'] for c in range(10)] for r in range(18)]}
    answer = {'tables': [table]}
    row, _ = stub_record(monkeypatch, req, answer)
    result = vision.replay_table(req, row, IMAGE)
    assert result['status'] == 'proposed' and len(result['tables'][0]['cells']) == 18
    assert all(len(r) == 10 for r in result['tables'][0]['cells'])
    assert result['financial_validation'] == 'not_established'
    repeated = deepcopy(answer)
    repeated['tables'][0]['rows'][1][1] = repeated['tables'][0]['rows'][0][0]
    row, _ = stub_record(monkeypatch, req, repeated)
    with pytest.raises(dg.GenerationError, match='reuses_original_cell_literal'):
        vision.replay_table(req, row, IMAGE)
    oversized = deepcopy(answer)
    oversized['tables'][0]['rows'] = [[own['Synthetic cell 0,0']] * 16 for _ in range(17)]
    row, _ = stub_record(monkeypatch, req, oversized)
    with pytest.raises(dg.GenerationError, match='exceeds_cell_bound'):
        vision.replay_table(req, row, IMAGE)
    declined = {'tables': []}
    row, _ = stub_record(monkeypatch, req, declined)
    assert vision.replay_table(req, row, IMAGE)['status'] == 'unassessed'
    with pytest.raises(ValidationError): req[3].model_validate(legacy_answer)
    assert dg.request_json(vision.frozen_table_request(IMAGE, old[2]['source_page'], PIN,
        raster_binding=old[2]['raster_binding'])) == before
    with pytest.raises(ValidationError): old[3].model_validate(answer)


def test_coordinate_profile_rejects_spatial_permutations_and_binds_recorded_feedback(monkeypatch):
    # Real fictional binary positions, not a hand-corrected company matrix.
    with ZipFile(BytesIO(source_bytes())) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    parts['ppt/media/native.emf'] = build([font(1), select(1),
        textw('Header A', x=10, y=20), textw('Header B', x=100, y=20),
        textw('Stored A', x=10, y=40), textw('Stored B', x=100, y=40)])
    output = BytesIO()
    with ZipFile(output, 'w') as archive:
        for name, value in parts.items(): archive.writestr(name, value)
    raw = output.getvalue(); native = build_document_manifest(raw, 'pptx', **IDENTITY)
    derived = attach_document_visual_text(native, raw, expected_binding=native['binding'])
    inventory = dg.evidence_inventory([derived], catalog_encoding=dg.CATALOG_ENCODING_V6)
    group = dg._source_map_page_view(inventory, inventory['source_map_topology'][0])
    own = {u['content']: uid for uid, u in dg.source_map_literal_lookup(group).items()}
    old, _, _ = frozen(); binding = deepcopy(old[2]['raster_binding'])
    binding['source_sha256'] = group['source_id']
    kwargs = dict(raster_binding=binding, vision_context_profile=vision.CONTEXT_PROFILE_32K)
    previous = vision.frozen_table_request(IMAGE, group, PIN, **kwargs, matrix_schema_profile=vision.MATRIX_PROFILE_V2)
    old_bytes = dg.request_json(previous)
    req = vision.frozen_table_request(IMAGE, group, PIN, **kwargs, matrix_schema_profile=vision.MATRIX_PROFILE_V3)
    answer = {'tables': [{'caption': 'Header A', 'caption_evidence': [
        {'unit_id': own['Header A'], 'quote': 'Header A'}], 'basis': 'unassessed', 'header_row_index': 0,
        'rows': [[own['Header A'], own['Header B']], [own['Stored A'], own['Stored B']]]}]}
    row, _ = stub_record(monkeypatch, req, answer)
    assert vision.replay_table(req, row, IMAGE)['financial_validation'] == 'not_established'
    assert len(req[2]['source_positions']['frame_identity_digests']) == 1
    points = dg.source_map_literal_lookup(group)
    for uid, x, y, bounds_ref, _ in req[2]['source_positions']['rows']:
        original = points[uid]['native_emf_placement']
        assert {'x': x, 'y': y} == original['reference']
        assert req[2]['source_positions']['record_bounds'][bounds_ref] == original['record_bounds']
    swapped = deepcopy(answer)
    swapped['tables'][0]['rows'][0][0], swapped['tables'][0]['rows'][1][0] = (
        swapped['tables'][0]['rows'][1][0], swapped['tables'][0]['rows'][0][0])
    row, _ = stub_record(monkeypatch, req, swapped)
    with pytest.raises(dg.GenerationError, match='coordinate_alignment_failed'):
        vision.replay_table(req, row, IMAGE)
    previous_row, _ = stub_record(monkeypatch, previous, swapped)
    feedback = vision.proposal_feedback(previous, previous_row, IMAGE)
    assert feedback['prior_raw_response'] == previous_row['raw_response']
    assert feedback['prior_response_hash'] == previous_row['response_hash']
    assert any(e['code'] == 'row_baseline_mismatch' for e in feedback['violations'])
    repaired_request = vision.frozen_table_request(IMAGE, group, PIN, **kwargs,
        matrix_schema_profile=vision.MATRIX_PROFILE_V3, prior_proposal_feedback=feedback)
    assert repaired_request[2]['prior_proposal_feedback'] == feedback
    packed_feedback = vision.proposal_feedback(previous, previous_row, IMAGE,
        feedback_encoding=vision.FEEDBACK_ENCODING_V2)
    assert vision._decode_diagnostics(packed_feedback['violations']) == feedback['violations']
    packed_request = vision.frozen_table_request(IMAGE, group, PIN, **kwargs,
        matrix_schema_profile=vision.MATRIX_PROFILE_V3, prior_proposal_feedback=packed_feedback,
        feedback_encoding=vision.FEEDBACK_ENCODING_V2)
    assert packed_request[2]['prior_proposal_feedback']['prior_raw_response'] == feedback['prior_raw_response']
    row, _ = stub_record(monkeypatch, packed_request, answer)
    assert vision.replay_table(packed_request, row, IMAGE)['investor_approval'] == 'not_established'
    arbitrary = [{'code': 'same', 'row': 0, 'x': None, 'positions': [[0, 1], [2, 3]]},
                 {'code': 'other', 'positions': [], 'row': 0}, {'code': 'same', 'x': 0}]
    assert vision._decode_diagnostics(vision._encode_diagnostics(arbitrary)) == arbitrary
    for bad_ref in (-1, True, 999):
        altered = deepcopy(packed_feedback)
        altered['violations']['rows'][0][0] = bad_ref
        with pytest.raises(dg.GenerationError, match='diagnostic_reference'):
            vision.frozen_table_request(IMAGE, group, PIN, **kwargs,
                matrix_schema_profile=vision.MATRIX_PROFILE_V3, prior_proposal_feedback=altered,
                feedback_encoding=vision.FEEDBACK_ENCODING_V2)
    altered = deepcopy(packed_feedback)
    altered['violations']['rows'][0].append(deepcopy(altered['violations']['rows'][0][1]))
    with pytest.raises(dg.GenerationError, match='diagnostic_reference'):
        vision.frozen_table_request(IMAGE, group, PIN, **kwargs,
            matrix_schema_profile=vision.MATRIX_PROFILE_V3, prior_proposal_feedback=altered,
            feedback_encoding=vision.FEEDBACK_ENCODING_V2)
    altered = deepcopy(packed_feedback)
    assert altered['violations']['rows'][0][1][1] == 0
    altered['violations']['rows'][0][1][1] = False  # bool must never alias original integer zero.
    with pytest.raises(dg.GenerationError, match='diagnostics_changed'):
        vision.frozen_table_request(IMAGE, group, PIN, **kwargs,
            matrix_schema_profile=vision.MATRIX_PROFILE_V3, prior_proposal_feedback=altered,
            feedback_encoding=vision.FEEDBACK_ENCODING_V2)
    forged = deepcopy(feedback); forged['violations'] = []
    with pytest.raises(dg.GenerationError, match='diagnostics_changed'):
        vision.frozen_table_request(IMAGE, group, PIN, **kwargs,
            matrix_schema_profile=vision.MATRIX_PROFILE_V3, prior_proposal_feedback=forged)
    for mutation in ('unbound_column', 'x_reverse', 'unknown_frame'):
        bad = deepcopy(answer)
        if mutation == 'unbound_column':
            for r in bad['tables'][0]['rows']: r[0] = None
        elif mutation == 'x_reverse':
            for r in bad['tables'][0]['rows']: r.reverse()
        else: bad['tables'][0]['rows'][0][0] = own['Untrusted source note.']
        row, _ = stub_record(monkeypatch, req, bad)
        with pytest.raises(dg.GenerationError, match='coordinate_alignment_failed'):
            vision.replay_table(req, row, IMAGE)
    assert dg.request_json(vision.frozen_table_request(IMAGE, group, PIN, **kwargs,
        matrix_schema_profile=vision.MATRIX_PROFILE_V2)) == old_bytes
