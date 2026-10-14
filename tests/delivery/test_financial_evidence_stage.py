"""Immutable scoped workbook indexing and retained replay, using synthetic data."""
from copy import deepcopy
from datetime import datetime
import hashlib
from io import BytesIO
import json
from pathlib import Path
from types import SimpleNamespace

from openpyxl import Workbook
import pytest

from delivery import financial_evidence_stage as stage
from delivery.storage import scope_component
from schemas import Document
from scripts import private_financial_manifest_worker as parser


def setup(tmp_path, monkeypatch):
    from api.routers import deals, rooms

    uploads = tmp_path / 'uploads'
    source = uploads / scope_component('tenant_a') / scope_component('deal_a') / 'upload_a' / 'source.xlsx'
    source.parent.mkdir(parents=True)
    book = Workbook()
    book.active['A1'] = 'Management projection'
    book.active['B1'] = datetime(2023, 3, 31)
    book.active['B2'] = '=Inputs!B2*2'
    hidden = book.create_sheet('Inputs')
    hidden.sheet_state = 'hidden'
    hidden['A2'] = 'Source assumption'
    hidden['B2'] = 5
    output = BytesIO()
    book.save(output)
    source.write_bytes(output.getvalue())
    document = Document(id='document_a', tenant_id='tenant_a', deal_id='deal_a',
        filename='source.xlsx', type='xlsx', storage_uri=str(source))
    room = SimpleNamespace(id='workspace_a', tenant_id='tenant_a', deal_id='deal_a')
    store = SimpleNamespace(conn=object(),
        get_workspace=lambda tenant, workspace_id: room if (tenant, workspace_id) == ('tenant_a', 'workspace_a') else None,
        get_document=lambda tenant, document_id: document if (tenant, document_id) == ('tenant_a', document.id) else None)
    job = {'id': 'job_a', 'tenant_id': 'tenant_a', 'workspace_id': 'workspace_a',
           'actor_id': 'actor_a', 'input_revision': 'revision_a'}
    calls = []

    def worker(command, root, **options):
        assert not options.get('local_model') and not options.get('local_ipc')
        calls.append(root)
        parser.run(root)

    monkeypatch.setattr(stage, 'artifact_root', lambda: tmp_path / 'private_artifacts')
    monkeypatch.setattr(stage, 'run_private', worker)
    monkeypatch.setattr(stage, 'has_access', lambda *args: True)
    monkeypatch.setattr(deals, 'UPLOAD_ROOT', uploads)
    monkeypatch.setattr(rooms, 'revision_for', lambda *args: 'revision_a')
    return store, job, document, source, calls


def load(record, **changes):
    binding = {**record['binding'], **record['room_binding'], **changes}
    return stage.load_financial_evidence(record, **binding)


def test_private_stage_retains_exact_original_and_loader_reopens_without_parser(tmp_path, monkeypatch):
    store, job, document, source, calls = setup(tmp_path, monkeypatch)
    original = source.read_bytes()
    record = stage.run_financial_evidence_stage(store, job, document)
    assert record['state'] == 'indexed_source_only'
    assert record['binding']['workbook_sha256'] == hashlib.sha256(original).hexdigest()
    assert record['room_binding'] == {'workspace_id': 'workspace_a', 'job_id': 'job_a', 'input_revision': 'revision_a'}
    manifest = load(record)
    assert record['retained_cell_count'] == manifest['retained_cell_count'] == 5
    assert record['sheet_count'] == 2
    assert manifest['financial_status'] == 'unreconciled'
    assert manifest['financial_validation'] == 'not_established'
    assert manifest['workbook_recalculation'] == 'not_performed'
    root = Path(record['manifest_path']).parent
    assert (root / 'source.xlsx').read_bytes() == original == source.read_bytes()
    assert all((root / name).stat().st_mode & 0o777 == 0o600 for name in (
        stage.MANIFEST, stage.RECEIPT, 'source.xlsx', 'request.json', 'adapter_binding.json'))
    monkeypatch.setattr(stage, 'run_private', lambda *args, **kwargs: pytest.fail('retained load must not parse'))
    monkeypatch.setattr(parser, 'build_financial_manifest', lambda *args, **kwargs: pytest.fail('retained load must not parse'))
    assert load(record) == manifest
    assert len(calls) == 1


@pytest.mark.parametrize('name', [stage.MANIFEST, stage.RECEIPT, 'source.xlsx', 'request.json', 'adapter_binding.json'])
def test_loader_rejects_any_changed_derivative_input_or_receipt(tmp_path, monkeypatch, name):
    store, job, document, _, _ = setup(tmp_path, monkeypatch)
    record = stage.run_financial_evidence_stage(store, job, document)
    path = Path(record['manifest_path']).parent / name
    path.write_bytes(path.read_bytes() + b' ')
    with pytest.raises(ValueError, match='changed'):
        load(record)


def test_foreign_scope_stale_context_and_arbitrary_paths_fail_before_read(tmp_path, monkeypatch):
    store, job, document, _, calls = setup(tmp_path, monkeypatch)
    record = stage.run_financial_evidence_stage(store, job, document)
    for changes in ({'tenant_id': 'tenant_b'}, {'deal_id': 'deal_b'},
                    {'document_id': 'document_b'}, {'workbook_sha256': 'b' * 64},
                    {'workspace_id': 'workspace_b'}, {'job_id': 'job_b'}, {'input_revision': 'revision_b'}):
        with pytest.raises(ValueError):
            load(record, **changes)
    forged = deepcopy(record)
    forged['manifest_path'] = str(tmp_path / 'unrelated.json')
    with pytest.raises(ValueError, match='unsafe retained path'):
        load(forged)
    foreign = document.model_copy(update={'tenant_id': 'tenant_b'})
    assert stage.run_financial_evidence_stage(store, job, foreign)['reason'] == 'workbook_outside_room_scope'
    assert len(calls) == 1


def test_source_and_stage_symlinks_cannot_escape_private_scope(tmp_path, monkeypatch):
    store, job, document, source, calls = setup(tmp_path, monkeypatch)
    escaped = tmp_path / 'other.xlsx'
    escaped.write_bytes(source.read_bytes())
    source.unlink()
    source.symlink_to(escaped)
    assert stage.run_financial_evidence_stage(store, job, document)['reason'] == 'workbook_unavailable_or_unsafe'
    assert calls == []
    source.unlink()
    source.write_bytes(escaped.read_bytes())
    root = stage._stage_root('tenant_a', 'workspace_a', 'job_a', 'document_a')
    ancestor = root.parents[4]
    ancestor.parent.mkdir(parents=True)
    destination = tmp_path / 'foreign_tenant'
    destination.mkdir()
    ancestor.symlink_to(destination, target_is_directory=True)
    assert stage.run_financial_evidence_stage(store, job, document)['state'] == 'blocked'
    assert calls == [] and list(destination.iterdir()) == []


def test_original_outside_canonical_upload_tree_and_caller_stage_override_are_rejected(tmp_path, monkeypatch):
    store, job, document, source, calls = setup(tmp_path, monkeypatch)
    alternate = tmp_path / 'same_bytes.xlsx'
    alternate.write_bytes(source.read_bytes())
    document.storage_uri = str(alternate)
    assert stage.run_financial_evidence_stage(store, job, document)['reason'] == 'workbook_unavailable_or_unsafe'
    document.storage_uri = str(source)
    unsafe = tmp_path / 'arbitrary_index'
    assert stage.run_financial_evidence_stage(store, job, document, stage_root=unsafe)['reason'] == 'financial_index_directory_unsafe'
    assert calls == [] and not unsafe.exists()


def test_changed_original_or_worker_binding_is_blocked_and_failed_index_preserved(tmp_path, monkeypatch):
    store, job, document, source, calls = setup(tmp_path, monkeypatch)
    record = stage.run_financial_evidence_stage(store, job, document)
    retained = Path(record['manifest_path']).read_bytes()
    source.write_bytes(source.read_bytes() + b' ')
    assert stage.run_financial_evidence_stage(store, job, document)['state'] == 'blocked'
    assert len(calls) == 1
    assert Path(record['manifest_path']).read_bytes() == retained
    source.write_bytes((Path(record['manifest_path']).parent / 'source.xlsx').read_bytes())
    receipt_path = Path(record['manifest_path']).parent / stage.RECEIPT
    receipt = json.loads(receipt_path.read_bytes())
    receipt['binding']['tenant_id'] = 'tenant_b'
    receipt_path.write_text(json.dumps(receipt))
    assert stage.run_financial_evidence_stage(store, job, document)['state'] == 'blocked'
    assert Path(record['manifest_path']).read_bytes() == retained


def test_private_parser_rejects_linked_request_and_workbook_traversal(tmp_path):
    root = tmp_path / 'job'
    root.mkdir()
    outside = tmp_path / 'request.json'
    outside.write_text(json.dumps({'workbook': '../source.xlsx', 'binding': {}}))
    (root / 'request.json').symlink_to(outside)
    with pytest.raises(OSError):
        parser.run(root)
    (root / 'request.json').unlink()
    (root / 'request.json').write_bytes(outside.read_bytes())
    with pytest.raises(ValueError, match='Exact private workbook request'):
        parser.run(root)
