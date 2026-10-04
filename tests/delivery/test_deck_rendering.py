"""Synthetic checks for editable deck geometry and matching PDF pages."""

from io import BytesIO
import re

import pytest
from pptx import Presentation
from pypdf import PdfReader

from delivery.inspection import inspect_export_pair
from delivery.rendering import _slide_text_boxes, render_deck_pdf, render_intro


SECTIONS = [
    ('Synthetic offering', 'A cited synthetic service description [S1].',
     '[S1] Synthetic public fixture.', 'statement'),
    ('Synthetic evidence', 'The source reports a limited synthetic pilot [S1].',
     '[S1] Synthetic pilot record, dated 2025.', 'evidence'),
    ('Synthetic comparison', 'The source describes product A [S1]. '
     'The source separately describes product B [S2].',
     '[S1] Synthetic product A. [S2] Synthetic product B.', 'comparison'),
    ('Synthetic chronology', 'In 2024, the synthetic pilot began [S1]. '
     'In 2025, it remained under review [S2].',
     '[S1] Synthetic 2024 record. [S2] Synthetic 2025 record.', 'timeline'),
]


def test_deck_pdf_and_editable_slides_preserve_model_text_and_layout():
    pptx = render_intro('Synthetic company', SECTIONS)
    pdf = render_deck_pdf('Synthetic company', SECTIONS)
    deck = Presentation(BytesIO(pptx))
    pages = PdfReader(BytesIO(pdf)).pages
    assert len(deck.slides) == len(pages) == 4
    assert round(deck.slide_width / deck.slide_height, 2) == 1.78
    for slide, page, (heading, body, sources, _) in zip(deck.slides, pages, SECTIONS):
        editable_text = '\n'.join(shape.text for shape in slide.shapes if shape.has_text_frame)
        pdf_text = page.extract_text()
        for value in (heading, sources):
            assert value in editable_text
            assert value in pdf_text
        for phrase in body.split('. '):
            assert phrase in editable_text
            assert phrase in re.sub(r'\s+', ' ', pdf_text)
        assert 'DRAFT' in editable_text and 'SOURCE NOTES' in editable_text
        assert len([shape for shape in slide.shapes if shape.has_text_frame]) >= 6
    assert inspect_export_pair(pptx, 'pptx', pdf)['pair_status'] == 'pass'


def test_deck_pdf_rejects_text_that_would_not_fit_editable_box():
    # A long source list beside a short claim now fits: the notes take the room they need.
    long_sources = '[S1] ' + 'Synthetic dated evidence. ' * 40
    section = ('Synthetic claim', 'Recorded synthetic statement [S1].', long_sources, 'statement')
    assert inspect_export_pair(render_intro('Synthetic company', [section]), 'pptx',
                               render_deck_pdf('Synthetic company', [section]))['pair_status'] == 'pass'
    # A full body of wide text with a full source list cannot fit at the floor sizes:
    # the slide is refused, never clipped or shrunk below them.
    wide = ('WWWW MMMM SYNTHETIC WIDE CAPITALS [S1]. ' * 40)[:1200]
    with pytest.raises(ValueError, match='Deck PDF text exceeds'):
        render_deck_pdf('Synthetic company', [('Synthetic claim', wide, wide, 'statement')])


def test_short_evidence_sources_use_panel_height_without_changing_text():
    section = ('Synthetic evidence', 'Synthetic statement [S1].',
               '[S1] Synthetic retained record.', 'evidence')
    boxes = _slide_text_boxes(*section[:3], 'Synthetic company', section[3])
    label = next(box for box in boxes if box[4] == 'SOURCE NOTES')
    notes = next(box for box in boxes if box[4] == section[2])
    assert 2.6 < notes[1] < 4.0
    assert 0 < notes[1] - (label[1] + label[3]) < .2
    assert inspect_export_pair(render_intro('Synthetic company', [section]), 'pptx',
                               render_deck_pdf('Synthetic company', [section]))['pair_status'] == 'pass'


def _text(seed, length):
    return (seed * (length // len(seed) + 1))[:length].rstrip()


def _boxes(pptx):
    """Per slide: {text: (top, height, size, line pitch)} of the editable text boxes."""
    return [{shape.text: (shape.top, shape.height, shape.text_frame.paragraphs[0].font.size.pt,
                          shape.text_frame.paragraphs[0].line_spacing.pt)
             for shape in slide.shapes if shape.has_text_frame and shape.text}
            for slide in Presentation(BytesIO(pptx)).slides]


def test_sparse_slide_uses_larger_centred_type_and_dense_slide_keeps_its_floor():
    from delivery.rendering import _slide_text_boxes
    sparse = 'A cited synthetic service description [S1].'
    dense = _text('The synthetic source reports a limited pilot at three example sites [S1]. ', 1100)
    sections = [('Synthetic offering', sparse, '[S1] Synthetic public fixture.', layout)
                for layout in ('statement', 'evidence', 'timeline')]
    sections.append(('Synthetic detail', dense, '[S1] Synthetic public fixture.', 'statement'))
    pptx, pdf = render_intro('Synthetic company', sections), render_deck_pdf('Synthetic company', sections)
    slides = _boxes(pptx)
    for slide in slides[:3]:
        top, height, size, pitch = slide[sparse]
        # A one-sentence claim is display type, lowered into its panel.
        assert size == 40 and top > 2.6 * 914400 and round(pitch, 1) == round(40 * 1.27, 1)
        # Source notes are 12pt when there is room, not 8-9pt.
        assert slide['[S1] Synthetic public fixture.'][2] == 12
    top, height, size, _ = slides[3][dense]
    assert 12 <= size <= 14 and abs(top - 2.28 * 914400) <= 1  # a full slide starts at the top
    # Editable boxes and the PDF are drawn from one geometry, and the text is unchanged.
    for (heading, body, sources, layout), slide in zip(sections, slides):
        for x, y, w, h, value, size, _ in _slide_text_boxes(heading, body, sources,
                                                            'Synthetic company', layout):
            top, height, point = slide[value][:3]
            assert abs(top - y * 914400) <= 1 and abs(height - h * 914400) <= 1 and point == size
    assert inspect_export_pair(pptx, 'pptx', pdf)['pair_status'] == 'pass'


@pytest.mark.parametrize('layout', ['statement', 'evidence', 'comparison', 'timeline'])
def test_fitted_sizes_never_overflow_across_body_heading_and_source_lengths(layout):
    """Wide and narrow synthetic text at every length: the PDF's own overflow check,
    which uses real font metrics, accepts what the shared estimate chose."""
    from delivery.rendering import _slide_text_boxes
    seeds = ('The synthetic source reports a limited pilot at three example sites [S1]. ',
             'WWWW MMMM Synthetic WIDE Capitals AND Numbers 0000 8888 [S1]. ',
             'Internationalisation characteristically overcomplicates demonstrations [S1]. ',
             'mmmm wwww mmmmm wwwww gum mow [S1]. ',
             'Raised USD 12,500,000 in 2024; 38% growth; 1,200 sites (2025) [S12]. ')
    sizes = set()
    for seed in seeds:
        for length in range(20, 901, 40):
            section = (_text('Synthetic heading words ', 20 + length % 80), _text(seed, length),
                       _text('[S1] Synthetic dated record. ', 30 + length % 400), layout)
            render_deck_pdf('Synthetic company', [section])
            body_size = max(box[5] for box in _slide_text_boxes(*section[:3], 'T', layout)
                            if box[4] and box[4] in section[1])
            sizes.add(body_size)
            assert 12 <= body_size <= 40
    assert 40 in sizes and min(sizes) <= 16 and len(sizes) >= 5    # the whole scale is used


def test_long_heading_is_fitted_instead_of_failing_the_slide():
    heading = _text('Synthetic heading words ', 100)
    section = (heading, 'A cited synthetic service description [S1].', '[S1] Synthetic.', 'statement')
    pptx, pdf = render_intro('Synthetic company', [section]), render_deck_pdf('Synthetic company', [section])
    assert _boxes(pptx)[0][heading][2] < 27
    assert inspect_export_pair(pptx, 'pptx', pdf)['pair_status'] == 'pass'


def test_comparison_columns_share_one_size_and_baseline():
    from delivery.rendering import _slide_text_boxes
    body = ('The source describes product A in one short sentence [S1]. '
            + _text('The source separately describes product B at length [S2]. ', 300))
    boxes = _slide_text_boxes('Synthetic comparison', body, '[S1] A. [S2] B.', 'T', 'comparison')
    left, right = boxes[2], boxes[3]
    assert left[4] and right[4] and (left[4] + ' ' + right[4]) == body
    assert (left[1], left[3], left[5]) == (right[1], right[3], right[5])


def test_one_clause_comparison_uses_statement_layout_without_splitting_prose():
    from delivery.rendering import _effective_slide_layout, _slide_text_boxes
    body = ('A source states that audited financials are absent and any thesis '
            'remains conditional until the records are independently confirmed [S1].')
    layout = _effective_slide_layout('comparison', body)
    assert layout == 'statement'
    boxes = _slide_text_boxes('Risk', body, '[S1] Synthetic source.', 'Synthetic', layout)
    assert sum(value.count(body) for *_, value, size, color in boxes) == 1
    pptx = render_intro('Synthetic', [('Risk', body, '[S1] Synthetic source.', 'comparison')])
    pdf = render_deck_pdf('Synthetic', [('Risk', body, '[S1] Synthetic source.', 'comparison')])
    assert inspect_export_pair(pptx, 'pptx', pdf)['pair_status'] == 'pass'


@pytest.mark.parametrize('body', [
    'The source describes product A: a scheduling tool for clinics; it is sold with product B '
    'under one synthetic contract [S1].',
    'The source lists three synthetic items; the second is described as a pilot, the third as a '
    'plan [S1].',
    'According to the record, e.g. the 2025 synthetic filing, the pilot covers v2.5 of the tool [S1].',
])
def test_comparison_never_splits_one_continuous_sentence(body):
    from delivery.rendering import _comparison_parts, _effective_slide_layout
    assert _comparison_parts(body) == (body, '')
    assert _effective_slide_layout('comparison', body) == 'statement'
    section = ('Synthetic comparison', body, '[S1] Synthetic record.', 'comparison')
    pptx, pdf = render_intro('Synthetic', [section]), render_deck_pdf('Synthetic', [section])
    assert body in _boxes(pptx)[0]                       # one box holds the whole sentence
    assert inspect_export_pair(pptx, 'pptx', pdf)['pair_status'] == 'pass'


def test_comparison_columns_are_whole_sentences_in_the_models_order():
    from delivery.rendering import _comparison_parts
    body = ('The source describes product A: a scheduling tool [S1]. It is sold to clinics [S1]. '
            'The source separately describes product B; it is a billing tool [S2].')
    left, right = _comparison_parts(body)
    assert left + ' ' + right == body
    assert left.endswith('[S1].') and right.startswith(('It is sold', 'The source separately'))
    section = ('Synthetic comparison', body, '[S1] A. [S2] B.', 'comparison')
    pptx, pdf = render_intro('Synthetic', [section]), render_deck_pdf('Synthetic', [section])
    boxes = _boxes(pptx)[0]
    assert left in boxes and right in boxes and boxes[left] == boxes[right]
    assert inspect_export_pair(pptx, 'pptx', pdf)['pair_status'] == 'pass'


def test_source_notes_take_the_height_they_need_and_stay_legible():
    from delivery.rendering import _slide_text_boxes
    claim = 'The source reports a limited synthetic pilot at three example sites [S1].'
    short = '[S1] Synthetic pilot record, dated 2025.'
    listed = '\n'.join(f'[S{index}] Synthetic dated record number {index}, retained version.'
                       for index in range(1, 9))
    sizes = {}
    for layout in ('statement', 'timeline', 'evidence'):
        for name, sources in (('short', short), ('listed', listed)):
            section = ('Synthetic evidence', claim, sources, layout)
            boxes = _slide_text_boxes(*section[:3], 'Synthetic company', layout)
            body = next(box for box in boxes if box[4] == claim)
            notes = next(box for box in boxes if box[4] == sources)
            label = next(box for box in boxes if box[4] == 'SOURCE NOTES')
            sizes[layout, name] = (body[5], notes[5], notes[3])
            # Legible notes, a label directly above them, nothing off the slide.
            assert notes[5] >= 10 and label[5] == 9
            assert 0 <= notes[1] - (label[1] + label[3]) < .2
            assert all(0 <= box[0] and box[0] + box[2] <= 13.34 and 0 <= box[1]
                       and box[1] + box[3] <= 7.5 for box in boxes)
            if layout != 'evidence':
                assert body[1] + body[3] < label[1] and notes[1] + notes[3] <= 6.79
            pptx, pdf = (render_intro('Synthetic company', [section]),
                         render_deck_pdf('Synthetic company', [section]))
            assert inspect_export_pair(pptx, 'pptx', pdf)['pair_status'] == 'pass'
    for layout in ('statement', 'timeline'):
        # Eight notes get a taller block; the claim above them is still display type.
        assert sizes[layout, 'listed'][2] > sizes[layout, 'short'][2] * 3
        assert sizes[layout, 'short'][0] == 40 and sizes[layout, 'listed'][0] >= 26


def test_body_type_steps_down_with_length_and_every_word_reaches_both_files():
    from delivery.rendering import _slide_text_boxes
    seed = 'The synthetic source reports a limited pilot at three example sites [S1]. '
    previous = 99
    for length in (40, 120, 220, 360, 560, 820, 1150):
        body = _text(seed, length)
        section = ('Synthetic claim', body, '[S1] Synthetic record.', 'statement')
        size = next(box[5] for box in _slide_text_boxes(*section[:3], 'T', 'statement')
                    if box[4] == body)
        assert 12 <= size <= previous
        previous = size
        pptx, pdf = render_intro('Synthetic', [section]), render_deck_pdf('Synthetic', [section])
        assert body in _boxes(pptx)[0]
        printed = re.sub(r'\s+', ' ', PdfReader(BytesIO(pdf)).pages[0].extract_text())
        assert re.sub(r'\s+', ' ', body) in printed
        assert inspect_export_pair(pptx, 'pptx', pdf)['pair_status'] == 'pass'
    assert previous <= 14                                   # the longest ends near the floor


def test_decorative_shapes_follow_the_text_geometry_and_carry_no_text():
    from delivery.rendering import _slide_plan
    for section in SECTIONS:
        shapes, boxes = _slide_plan(*section[:3], 'Synthetic company', section[3])
        assert all(len(shape) == 5 and 0 <= shape[0] and shape[0] + shape[2] <= 13.34
                   and 0 <= shape[1] and shape[1] + shape[3] <= 7.5 for shape in shapes)
    slide = Presentation(BytesIO(render_intro('Synthetic company', SECTIONS))).slides[0]
    texts = [shape.text for shape in slide.shapes if shape.has_text_frame and shape.text]
    assert texts == [box[4] for box in _slide_plan(*SECTIONS[0][:3], 'Synthetic company',
                                                   SECTIONS[0][3])[1]]


# ---- editorial variety from the model's own layout choice and sentence structure

SHORT = 'The source does not report economic results or margins [S1].'
TWO = ('The source reports a scheduling tool for clinics and a pilot with one clinic [S1]. '
       'The source describes a pilot with one clinic and reports no repeat order [S1].')
THREE = ('In 2024, the synthetic pilot began at one clinic [S1]. In 2025, it remained under '
         'review [S2]. A registry lists a seed entry dated 2025-06 [S3].')
NOTES = '[S1] Example source'


def _plan(body, layout, index=0, sources=NOTES):
    from delivery.rendering import _effective_slide_layout, _slide_plan
    return _slide_plan('Synthetic heading', body, sources, 'Example Labs',
                       _effective_slide_layout(layout, body), index)


def _body_boxes(boxes, body):
    return [box for box in boxes if box[4] in body and box[4] not in ('', 'Example Labs')
            and box[4] != 'Synthetic heading']


def test_one_short_statement_is_a_display_panel_that_alternates_through_the_deck():
    dark_shapes, dark = _plan(SHORT, 'statement', 0)
    light_shapes, light = _plan(SHORT, 'statement', 1)
    (dark_box,), (light_box,) = _body_boxes(dark, SHORT), _body_boxes(light, SHORT)
    assert dark_box[4] == light_box[4] == SHORT and dark_box[5] == light_box[5] == 40
    # Same geometry; only the decoration and the text colour alternate.
    assert dark_box[:4] == light_box[:4] and (dark_box[6], light_box[6]) == ('FFFFFF', '17324D')
    panel = lambda shapes, fill: next(shape for shape in shapes if shape[4] == fill)
    for shapes, fill, box in ((dark_shapes, '17324D', dark_box), (light_shapes, 'EAF2F2', light_box)):
        x, y, w, h, _ = panel(shapes, fill)
        # The sentence sits inside its panel, vertically centred within a line.
        assert x < box[0] and box[0] + box[2] < x + w and y < box[1] and box[1] + box[3] <= y + h + .01
    # A longer or multi-sentence statement keeps the plain accent-bar slide.
    for body in (TWO, 'The source reports a pilot [S1]. ' * 8):
        shapes, boxes = _plan(body.strip(), 'statement')
        assert not any(shape[4] == '17324D' for shape in shapes)
        assert [box[4] for box in boxes].count(body.strip()) == 1


def test_evidence_sentences_become_rows_in_order_with_nothing_added():
    shapes, boxes = _plan(TWO, 'evidence')
    rows = _body_boxes(boxes, TWO)
    assert ' '.join(box[4] for box in rows) == TWO and len(rows) == 2
    assert rows[0][1] + rows[0][3] <= rows[1][1] and rows[0][5] == rows[1][5] >= 22
    assert rows[0][0] == rows[1][0] and rows[0][2] > 10            # full-width rows
    # One marker per row, level with its first line; a tinted strip holds the notes.
    markers = [shape for shape in shapes if shape[4] == 'EA5A5C' and shape[2] == shape[3] == .14]
    assert len(markers) == 2 and all(row[1] <= marker[1] <= row[1] + .5
                                     for row, marker in zip(rows, markers))
    notes = next(box for box in boxes if box[4] == NOTES)
    strip = next(shape for shape in shapes if shape[4] == 'EAF2F2')
    assert strip[1] < notes[1] and notes[1] + notes[3] <= strip[1] + strip[3] + .01
    # The only text on the slide is the model's, plus the fixed labels.
    assert [box[4] for box in boxes] == [
        'DRAFT • NOT APPROVED FOR INVESTOR DISTRIBUTION', 'Synthetic heading',
        rows[0][4], rows[1][4], 'SOURCE NOTES', NOTES,
        'Example Labs']
    # One sentence, or more than four, keeps the body-beside-notes panel.
    for body in (SHORT, ('The source reports one synthetic pilot [S1]. ' * 6).strip()):
        shapes, boxes = _plan(body, 'evidence')
        assert [box[4] for box in boxes].count(body) == 1


def test_timeline_sentences_become_steps_in_the_models_order():
    shapes, boxes = _plan(THREE, 'timeline', sources='[S1] A.\n[S2] B.\n[S3] C.')
    steps = _body_boxes(boxes, THREE)
    assert ' '.join(box[4] for box in steps) == THREE and len(steps) == 3
    assert steps[0][0] < steps[1][0] < steps[2][0]                 # left to right
    assert len({(box[1], box[2], box[3], box[5]) for box in steps}) == 1
    assert all(left[0] + left[2] < right[0] for left, right in zip(steps, steps[1:]))
    markers = [shape for shape in shapes if shape[2] == shape[3] == .19]
    assert [marker[0] for marker in markers] == [box[0] for box in steps]
    # A single long sentence stays one block.
    one = ('In 2024 the synthetic pilot began at one clinic and it was still under review at the '
           'end of 2025 with no repeat order reported [S1].')
    assert [box[4] for box in _plan(one, 'timeline')[1]].count(one) == 1


@pytest.mark.parametrize('body, expected', [
    (TWO, 2), (THREE, 3), (SHORT, 1),
    # Not split: a colon or semicolon, an abbreviation, uneven spacing, a line break,
    # a very short fragment, or more sentences than a slide can lay out.
    ('The source lists two synthetic items: a pilot; and a plan for a second site [S1].', 1),
    ('The record, e.g. the 2025 synthetic filing, covers v2.5 of the tool [S1].', 1),
    ('The source reports one pilot [S1].  The source reports no repeat order [S1].', 1),
    ('The source reports one pilot [S1].\nThe source reports no repeat order [S1].', 1),
    ('It is so [S1]. The source reports no repeat order at the clinic [S1].', 1),
    (('The source reports one synthetic pilot [S1]. ' * 5).strip(), 1),
])
def test_sentence_rows_are_exact_text_or_the_body_is_left_whole(body, expected):
    from delivery.rendering import _sentences
    parts = _sentences(body)
    assert len(parts) == expected and ' '.join(parts) == body


def test_a_deck_of_two_layout_choices_no_longer_repeats_one_look():
    """The public synthetic shape: the model chose only `statement` and `evidence`."""
    sections = [('Company Overview', TWO, NOTES, 'evidence'),
                ('Product Offering', 'The offering is a scheduling tool designed for clinics, as '
                 'reported in the source [S1].', NOTES, 'statement'),
                ('Funding Status', 'A registry lists a seed financing entry dated 2025-06 with '
                 'unknown status [S2].', '[S2] Synthetic registry', 'statement'),
                ('Pilot Risk Assessment', 'The source describes a pilot with one clinic and reports '
                 'no repeat order [S1]. This may warrant investigation, but the record does not '
                 'establish repeat demand, earned revenue, retention, or an attractive price [S1].',
                 NOTES, 'evidence'),
                ('Financial Status', SHORT, NOTES, 'statement')]
    pptx, pdf = render_intro('Example Labs', sections), render_deck_pdf('Example Labs', sections)
    assert inspect_export_pair(pptx, 'pptx', pdf)['pair_status'] == 'pass'
    deck = Presentation(BytesIO(pptx))
    looks = []
    for slide, page, (heading, body, sources, _) in zip(deck.slides, PdfReader(BytesIO(pdf)).pages,
                                                        sections):
        fills = sorted({str(shape.fill.fore_color.rgb) for shape in slide.shapes
                        if not (shape.has_text_frame and shape.text)})
        texts = [shape.text for shape in slide.shapes if shape.has_text_frame and shape.text]
        looks.append((tuple(fills), len(texts)))
        # Every model word is editable text, in order, and on the matching PDF page.
        assert texts[1] == heading and ' '.join(texts[2:-3]) == body and texts[-2] == sources
        assert texts[0].startswith('DRAFT') and texts[-3] == 'SOURCE NOTES' and texts[-1] == 'Example Labs'
        printed = re.sub(r'\s+', ' ', page.extract_text())
        assert all(re.sub(r'\s+', ' ', text) in printed for text in texts)
    # Three distinct looks across five slides, and no two neighbours alike.
    assert len(set(looks)) == 3
    assert all(first != second for first, second in zip(looks, looks[1:]))
