"""Exact synthetic table projection and editable/PDF export checks."""
from io import BytesIO

import pytest
from docx import Document
from pptx import Presentation
from pypdf import PdfReader
from pydantic import ValidationError

from delivery.contracts import EvidenceTable
from delivery.inspection import inspect_export_pair
from delivery.rendering import (render_evidence_document, render_intro,
                                render_memo, render_deck_pdf, render_memo_pdf)


def table(rows=3):
    return EvidenceTable.model_validate({
        'table_id': 'synthetic-financial-table', 'evidence_revision': 'a' * 64,
        'columns': [{'label': 'Recorded metric'}, {'label': 'Source period', 'period': '2024', 'unit': 'USD mn'},
                    {'label': 'Source plan', 'period': '2025', 'unit': 'USD mn'}],
        'rows': [{'cells': [
            {'text': f'Synthetic metric {i}', 'status': 'source_reported',
             'source_refs': ['S1'], 'value_ref': f'T1:R{i}:C1'},
            {'text': f'{i}.25', 'status': 'source_reported',
             'source_refs': ['S1'], 'value_ref': f'T1:R{i}:C2'},
            {'text': f'{i}.50', 'status': 'projection',
             'source_refs': ['S1'], 'value_ref': f'T1:R{i}:C3'},
        ]} for i in range(1, rows + 1)],
    })


def section(payload=None):
    return ('Synthetic recorded comparison', 'The supplied example table is a source report [S1].',
            '[S1] Synthetic immutable table input, financial review pending.', 'table',
            (payload or table()).model_dump(mode='json'))


def test_native_pptx_table_and_pdf_preserve_exact_text_units_periods_and_status():
    pptx = render_intro('Synthetic company', [section()])
    pdf = render_deck_pdf('Synthetic company', [section()])
    deck = Presentation(BytesIO(pptx))
    native = [shape.table for shape in deck.slides[0].shapes if shape.has_table]
    assert len(native) == 1 and len(native[0].rows) == 4
    assert native[0].cell(0, 1).text == 'Source period\n2024\nUSD mn'
    assert native[0].cell(1, 1).text == '1.25 [S1]'
    assert native[0].cell(1, 2).text == '1.50 (projection) [S1]'
    assert all(paragraph.font.size.pt >= 11 for row in native[0].rows for cell in row.cells
               for paragraph in cell.text_frame.paragraphs)
    assert inspect_export_pair(pptx, 'pptx', pdf)['pair_status'] == 'pass'
    content = PdfReader(BytesIO(pdf)).pages[0].extract_text()
    assert '1.50 (projection) [S1]' in content
    assert 'Financial review pending' in content
    # The object remains a native editable table; changing a cell changes XML
    # text, rather than editing a flattened screenshot.
    native[0].cell(1, 1).text = 'Edited synthetic value'
    rewritten = BytesIO(); deck.save(rewritten)
    assert Presentation(BytesIO(rewritten.getvalue())).slides[0].shapes[-1].table.cell(1, 1).text == 'Edited synthetic value'


def test_native_docx_table_and_memo_pdf_match_without_promoting_projection_to_actual():
    docx = render_memo('Synthetic company', [section()])
    pdf = render_memo_pdf('Synthetic company', [section()])
    doc = Document(BytesIO(docx))
    assert len(doc.tables) == 1
    assert doc.tables[0].cell(1, 2).text == '1.50 (projection) [S1]'
    assert inspect_export_pair(docx, 'docx', pdf)['pair_status'] == 'pass'
    assert 'tblHeader' in doc.tables[0].rows[0]._tr.xml
    assert all('cantSplit' in row._tr.xml for row in doc.tables[0].rows)


def test_structured_document_explicitly_allows_reference_scale_preserves_legacy_cap():
    sections = [('Synthetic source section ' + str(i), 'Recorded synthetic detail [S1].',
                 '[S1] Synthetic source.', 'statement') for i in range(37)]
    with pytest.raises(ValueError, match='sections'):
        render_intro('Synthetic', sections)
    pptx = render_evidence_document('Synthetic', sections, kind='pitch_deck', format='pptx')
    pdf = render_evidence_document('Synthetic', sections, kind='pitch_deck', format='pdf')
    assert len(Presentation(BytesIO(pptx)).slides) == len(PdfReader(BytesIO(pdf)).pages) == 37
    assert inspect_export_pair(pptx, 'pptx', pdf)['pair_status'] == 'pass'
    with pytest.raises(ValueError, match='sections'):
        render_evidence_document('Synthetic', sections + sections[:4], kind='pitch_deck', format='pdf')
    with pytest.raises(ValueError, match='kind/format'):
        render_evidence_document('Synthetic', sections, kind='financial_projection', format='xlsx')


def test_unreadable_table_blocks_both_editable_and_pdf_instead_of_clipping():
    payload = table().model_dump(mode='json')
    payload['rows'][0]['cells'][2]['text'] = 'W' * 180
    large = EvidenceTable.model_validate(payload)
    for renderer in (render_intro, render_deck_pdf, render_memo, render_memo_pdf):
        with pytest.raises(ValueError, match='unbreakable overflow'):
            renderer('Synthetic', [section(large)])
    for renderer in (render_intro, render_deck_pdf):
        with pytest.raises(ValueError, match='12 rows'):
            renderer('Synthetic', [section(table(13))])


def test_cell_provenance_shape_and_rectangularity_are_required_not_certified():
    original = table().model_dump(mode='json')
    original['rows'][0]['cells'][1]['source_refs'] = []
    with pytest.raises(ValidationError, match='immutable value and source'):
        EvidenceTable.model_validate(original)
    original = table().model_dump(mode='json')
    original['rows'][0]['cells'].pop()
    with pytest.raises(ValidationError, match='rectangular'):
        EvidenceTable.model_validate(original)
    original = table().model_dump(mode='json')
    original['rows'][0]['cells'][1]['status'] = 'missing'
    with pytest.raises(ValidationError, match='Missing table cell'):
        EvidenceTable.model_validate(original)


def test_missing_and_illustrative_status_remain_visible_and_rows_paginate_in_memo():
    payload = table(36).model_dump(mode='json')
    payload['rows'][0]['cells'][1] = {'text': 'Not supplied', 'status': 'missing', 'source_refs': []}
    payload['rows'][0]['cells'][2]['status'] = 'illustrative'
    recorded = EvidenceTable.model_validate(payload)
    docx = render_memo('Synthetic long financial appendix', [section(recorded)])
    pdf = render_memo_pdf('Synthetic long financial appendix', [section(recorded)])
    assert len(PdfReader(BytesIO(pdf)).pages) > 1
    assert inspect_export_pair(docx, 'docx', pdf)['pair_status'] == 'pass'
    all_text = '\n'.join(page.extract_text() for page in PdfReader(BytesIO(pdf)).pages)
    assert 'Not supplied (missing)' in all_text and '(illustrative)' in all_text
    assert '36.25 [S1]' in all_text and '36.50 (projection) [S1]' in all_text


def test_structured_renderer_rejects_extra_keys_unsupported_contract_and_chart_values():
    payload = table().model_dump(mode='json')
    payload['chart'] = {'series': [1, 2, 3]}
    with pytest.raises(ValidationError, match='Extra inputs'):
        render_intro('Synthetic', [section()[:4] + (payload,)])
    with pytest.raises(ValueError, match='contract'):
        render_intro('Synthetic', [section()], document_contract='unreviewed-layout')
    with pytest.raises(ValueError, match='table layout'):
        render_intro('Synthetic', [section()[:3] + ('chart', section()[4])])


def test_new_memo_contract_preserves_long_narrative_without_extending_legacy_limits():
    paragraph = ('A synthetic reported observation remains limited to the supplied source and does not establish investment suitability. ' * 7).strip()
    body = paragraph + '\n\n' + paragraph
    assert 1200 < len(body) < 2000
    sections = [('Synthetic dense diligence', body, '[S1] Synthetic original evidence.')]
    with pytest.raises(ValueError, match='exceeds'):
        render_memo('Synthetic', sections)
    docx = render_evidence_document('Synthetic', sections, kind='investment_memorandum', format='docx')
    pdf = render_evidence_document('Synthetic', sections, kind='investment_memorandum', format='pdf')
    assert inspect_export_pair(docx, 'docx', pdf)['pair_status'] == 'pass'
    assert Document(BytesIO(docx)).paragraphs[2].text == paragraph
