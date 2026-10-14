"""Finite complete-source reading precedes planning; mock responses are synthetic."""
from copy import deepcopy
from io import BytesIO
import json

import pytest
from pptx import Presentation
from pptx.util import Inches

from agents.research import document_generation as dg
from agents.research.document_evidence_manifest import build_document_manifest
from delivery import document_generation_stage as stage
from tests.research.test_document_generation import PIN, author, planned, record


def full_source(*, blank_pages=1, return_bytes=False):
    deck=Presentation()
    page=deck.slides.add_slide(deck.slide_layouts[6])
    page.shapes.add_textbox(Inches(1),Inches(1),Inches(6),Inches(1)).text='The source reports a limited pilot in 2024.'
    # A physical blank page still needs a source-reading observation.
    for _ in range(blank_pages):
        deck.slides.add_slide(deck.slide_layouts[6])
    output=BytesIO(); deck.save(output)
    value=build_document_manifest(output.getvalue(),'pptx',tenant_id='tenant',deal_id='deal',document_id='synthetic')
    return (output.getvalue(),value) if return_bytes else value


def prepare(root):
    return stage.prepare_document_generation(root,[full_source()],PIN,PIN,'pitch_deck',
                                             catalog_encoding=dg.CATALOG_ENCODING_V3)


def map_answer(request):
    payload=request[2]
    # Test fixtures describe a source proposal, never independently reviewed truth.
    rows=payload['source_page']['literal_units']
    selections=[]
    if rows:
        selections=[{'topic':'customers_and_traction','coverage':'limited','basis':'reported',
                     'evidence':[{'unit_id':rows[0][0],'quote':'a limited pilot in 2024'}],
                     'table_ids':[]}]
    return {'map_id':payload['map_id'],'selections':selections,
            'unselected_topics':[{'topic':topic,'status':'unassessed'} for topic in dg.TOPICS
                                 if topic not in {entry['topic'] for entry in selections}],
            'scope_status':'source_reading_proposal_not_diligence'}


def complete_fake_call(root,calls):
    def fake(model,task,instruction,payload,schema,attempts,save):
        request=(task,instruction,payload,schema)
        calls.append(task)
        if schema is dg.SourcePageMap:
            answer=map_answer(request)
        elif schema is dg.DocumentPlan:
            answer=planned(stage._read(root/'request.json')['inventory'])
        elif schema is dg.DocumentPage:
            answer=author(request)
        else:
            answer={'verdict':'pass','issues':[],'financial_status':'not_established'}
        attempts.append(record(request,answer,len(attempts)+1)); save()
    return fake


def advance(root):
    return stage.advance_document_generation(root,tenant_id='tenant',deal_id='deal',
        installed_digest=lambda name:PIN['digest'],model_factory=lambda *args,**kwargs:object())


def test_all_physical_maps_precede_plan_and_exact_complete_reopen_calls_nothing(tmp_path,monkeypatch):
    root=tmp_path/'source-mapped'
    prepared=prepare(root)
    body,seed,_=stage._load_run(root)
    assert prepared['source_map_count']==body['source_map_count']==2
    assert body['maximum_lifetime_model_calls']==83
    assert body['initial_request']==dg.request_json(seed)
    assert 'plan_request' not in body
    calls=[]
    monkeypatch.setattr(stage,'recorded_call',complete_fake_call(root,calls))
    assert stage.replay_document_generation(root,tenant_id='tenant',deal_id='deal')['phase']=='document_source_map'
    assert advance(root)['phase']=='document_source_map'
    assert advance(root)['phase']=='document_plan'
    assert advance(root)['phase']=='document_page'
    assert advance(root)['phase']=='document_page_review'
    result=advance(root)
    assert result['state']=='draft_complete'
    assert result['plan_response_id']=='response_3'
    assert result['financial_validation']=='not_established'
    assert len(calls)==5
    assert advance(root)==result and len(calls)==5
    assert [stage._read(root/f'call_{number}.started.json')['phase'] for number in range(1,6)]==[
        'document_source_map','document_source_map','document_plan','document_page','document_page_review']
    assert all(stage._read(root/f'budget_{n}.started.json')['max_calls']==1 for n in range(1,6))
    assert all(stage._read(root/f'budget_{n}.started.json')['max_requests']==6 for n in range(1,6))
    assert set(stage._read(root/'response_1.json')['input'])==set(seed[2])


def test_unresolved_source_map_intent_and_failed_map_never_repeat(tmp_path,monkeypatch):
    orphan=tmp_path/'orphan-map'
    prepare(orphan)
    stage._write_new(orphan/'call_1.started.json',{'unknown_transport':True})
    def never(*args,**kwargs):pytest.fail('Unresolved/failed source reading must not repeat')
    monkeypatch.setattr(stage,'recorded_call',never)
    assert advance(orphan)['reason']=='unresolved_document_call_intent'
    failed=tmp_path/'failed-map'
    prepare(failed)
    def failure(model,task,instruction,payload,schema,attempts,save):
        attempts.append({'id':'response_1','task':task,'instruction':instruction,'input':payload,
            'schema':schema.model_json_schema(),'error':'synthetic transport failure',
            'raw_response':'','response_hash':dg.digest(''),'model':PIN['name']})
        save(); raise RuntimeError('synthetic transport failure')
    monkeypatch.setattr(stage,'recorded_call',failure)
    assert advance(failed)['state']=='blocked'
    saved=(failed/'response_1.json').read_bytes()
    monkeypatch.setattr(stage,'recorded_call',never)
    assert advance(failed)['state']=='blocked'
    assert (failed/'response_1.json').read_bytes()==saved
    assert not (failed/'call_2.started.json').exists()


def test_plan_response_cannot_be_relabelled_as_source_map_prefix(tmp_path):
    root=tmp_path/'wrong-prefix'
    prepare(root)
    body,_,_=stage._load_run(root)
    old_inventory=dg.evidence_inventory([full_source()])
    plan_request=dg.frozen_document_request(old_inventory,PIN,'pitch_deck')
    row=record(plan_request,planned(old_inventory),1)
    stage._write_new(root/'call_1.started.json',{'run_binding_digest':body['run_binding_digest'],
        'request_digest':dg.canonical_digest(dg.request_json(plan_request))})
    stage._write_new(root/'response_1.json',row)
    result=stage.replay_document_generation(root,tenant_id='tenant',deal_id='deal')
    assert result['state']=='blocked'
    assert not (root/'response_2.json').exists()


@pytest.mark.parametrize('change',['count','map_request','lifetime'])
def test_resealed_input_cannot_change_map_requests_count_or_finite_lifetime(tmp_path,change):
    root=tmp_path/'mutated-input'
    prepare(root)
    body=stage._read(root/'request.json')
    if change=='count':body['source_map_count']+=1
    elif change=='map_request':body['source_map_requests'][0][2]['map_id']='M99'
    else:body['maximum_lifetime_model_calls']+=1
    body['run_binding_digest']=dg.canonical_digest({key:value for key,value in body.items() if key!='run_binding_digest'})
    (root/'request.json').write_text(json.dumps(body))
    with pytest.raises(dg.GenerationError):stage._load_run(root)


def test_v1_v2_run_bodies_preserve_historical_absence_of_source_map_fields(tmp_path):
    original=full_source()
    inventory=dg.evidence_inventory([original])
    for number,encoding in enumerate([dg.CATALOG_ENCODING_V1,dg.CATALOG_ENCODING_V2],1):
        root=tmp_path/('historical-%d'%number)
        stage.prepare_document_generation(root,[original],PIN,PIN,'pitch_deck',catalog_encoding=encoding)
        body,request,attempts=stage._load_run(root)
        expected={'contract':dg.CONTRACT,'inventory':inventory,
                  'plan_request':dg.request_json(dg.frozen_document_request(inventory,PIN,'pitch_deck',catalog_encoding=encoding)),
                  'author_pin':PIN,'reviewer_pin':PIN,'kind':'pitch_deck','mandate':'',
                  'maximum_seconds_per_advance':105,'maximum_calls_per_advance':1,
                  'maximum_requests_per_advance':6,'maximum_lifetime_model_calls':81}
        if encoding!=dg.CATALOG_ENCODING_V1:expected['catalog_encoding']=encoding
        assert body=={**expected,'run_binding_digest':dg.canonical_digest(expected)}
        assert attempts==[] and request[0]=='document_plan_v1'
        assert not {'source_map_count','source_map_requests','initial_request'} & set(body)


def adoption_fixture(tmp_path, monkeypatch):
    import hashlib
    source_root=tmp_path/'original-source'
    source_root.mkdir(mode=0o700)
    original,value=full_source(blank_pages=2,return_bytes=True)
    source=source_root/'source.pptx'; source.write_bytes(original)
    manifest_path=source_root/'document_evidence_manifest.json'
    manifest_path.write_text(json.dumps(value))
    proof={'sources':[{'manifest_path':str(manifest_path),
        'manifest_sha256':hashlib.sha256(manifest_path.read_bytes()).hexdigest(),
        'original_path':str(source),'original_sha256':hashlib.sha256(original).hexdigest()}]}
    parent=tmp_path/'parent'
    stage.prepare_document_generation(parent,[value],PIN,PIN,'pitch_deck',catalog_encoding=dg.CATALOG_ENCODING_V3)
    stage._write_new(parent/'source_provenance.json',proof)
    calls=[]
    def fake(model,task,instruction,payload,schema,attempts,save):
        calls.append(task)
        if payload['map_id']=='M3':
            attempts.append({'id':'response_3','task':task,'instruction':instruction,'input':payload,
                'schema':schema.model_json_schema(),'raw_response':'','response_hash':dg.digest(''),
                'error':'synthetic failed source map','model':PIN['name']})
            save(); raise RuntimeError('synthetic failed source map')
        request=(task,instruction,payload,schema)
        attempts.append(record(request,map_answer(request),len(attempts)+1)); save()
    monkeypatch.setattr(stage,'recorded_call',fake)
    for _ in range(3):advance(parent)
    assert len(calls)==3
    assert stage.replay_document_generation(parent,tenant_id='tenant',deal_id='deal')['state']=='blocked'
    return parent,value,proof


def new_adoption_target(tmp_path, value, proof, *, budgets=None, profiles=None, name='child'):
    root=tmp_path/name
    stage.prepare_document_generation(root,[value],PIN,PIN,'pitch_deck',catalog_encoding=dg.CATALOG_ENCODING_V3,
        source_map_token_budgets=budgets,source_map_profiles=profiles)
    stage._write_new(root/'source_provenance.json',proof)
    return root


def test_successful_raw_prefix_adoption_preserves_failed_parent_and_records_no_new_calls(tmp_path,monkeypatch):
    parent,value,proof=adoption_fixture(tmp_path,monkeypatch)
    child=new_adoption_target(tmp_path,value,proof,budgets=[900,900,1600],
                              profiles=['quoted-v1','quoted-v1','unit-ids-v2'])
    before={path.name:path.read_bytes() for path in parent.iterdir() if path.is_file()}
    count,requests,body=stage.verified_source_map_prefix(parent)
    assert count==2 and len(requests)==3 and 'source_map_profiles' not in body
    def never(*args,**kwargs):pytest.fail('Adoption or cached replay must not infer')
    monkeypatch.setattr(stage,'recorded_call',never)
    result=stage.adopt_source_map_prefix(child,parent,count=count)
    assert result['adopted_model_responses']==2
    assert result['new_model_invocations']==result['new_physical_requests']==0
    for number in (1,2):
        assert (child/f'response_{number}.json').read_bytes()==(parent/f'response_{number}.json').read_bytes()
        intent=stage._read(child/f'call_{number}.started.json')
        assert intent['execution_origin']=='adopted_recorded_source_map'
        assert not (child/f'budget_{number}.started.json').exists()
        assert not (child/f'budget_{number}.completed.json').exists()
    assert not (child/'response_3.json').exists()
    replay=stage.replay_document_generation(child,tenant_id='tenant',deal_id='deal')
    assert replay['state']=='needs_call' and replay['source_map_index']==3
    assert replay['request'][2]['options']['max_tokens']==1600
    assert replay['request'][0]!=requests[2][0]
    assert before=={path.name:path.read_bytes() for path in parent.iterdir() if path.is_file()}
    assert stage.replay_document_generation(child,tenant_id='tenant',deal_id='deal')['source_map_index']==3


def test_dynamic_bound_suffix_schema_identity_cannot_block_exact_verified_prefix_adoption(tmp_path,monkeypatch):
    parent,value,proof=adoption_fixture(tmp_path,monkeypatch)
    child=tmp_path/'bound-child'
    stage.prepare_document_generation(child,[value],PIN,PIN,'pitch_deck',
        catalog_encoding=dg.CATALOG_ENCODING_V3,
        source_map_token_budgets=[900,900,1600],
        source_map_profiles=['quoted-v1','quoted-v1','unit-ids-bound-v3'],
        qualification_context_profile='complete-page-context-v1',
        source_coverage_reviewer_pin={'name':'qwen3:14b','digest':'b'*64})
    stage._write_new(child/'source_provenance.json',proof)
    body,_,_=stage._load_run(child)
    first,second=stage._maps(body),stage._maps(body)
    # The finite selector is rebuilt on each replay. Class identity is neither
    # a request change nor evidence of any model invocation.
    assert first[2][3] is not second[2][3]
    assert [dg.request_json(request) for request in first]==[dg.request_json(request) for request in second]
    parent_before={path.name:path.read_bytes() for path in parent.iterdir() if path.is_file()}
    def never(*args,**kwargs):pytest.fail('Verified prefix adoption/reopen must not infer')
    monkeypatch.setattr(stage,'recorded_call',never)
    result=stage.adopt_source_map_prefix(child,parent,count=2)
    assert result['adopted_model_responses']==2
    assert result['new_model_invocations']==result['new_physical_requests']==0
    for number in (1,2):
        assert (child/f'response_{number}.json').read_bytes()==(parent/f'response_{number}.json').read_bytes()
        assert stage._read(child/f'call_{number}.started.json')['execution_origin']=='adopted_recorded_source_map'
    assert not list(child.glob('budget_*.json'))
    assert not (child/'response_3.json').exists()
    first_replay=stage.replay_document_generation(child,tenant_id='tenant',deal_id='deal')
    next_replay=stage.replay_document_generation(child,tenant_id='tenant',deal_id='deal')
    assert first_replay['state']==next_replay['state']=='needs_call'
    assert first_replay['source_map_index']==next_replay['source_map_index']==3
    assert first_replay['request'][0]=='document_source_map_v3'
    assert dg.request_json(first_replay['request'])==dg.request_json(next_replay['request'])
    assert parent_before=={path.name:path.read_bytes() for path in parent.iterdir() if path.is_file()}


@pytest.mark.parametrize('change',['failed_count','mismatched_budget','mismatched_profile'])
def test_unverified_or_changed_adoption_prefix_rejected_before_any_target_write(tmp_path,monkeypatch,change):
    parent,value,proof=adoption_fixture(tmp_path,monkeypatch)
    budgets=[1600,900,1600] if change=='mismatched_budget' else [900,900,1600]
    profiles=['unit-ids-v2','quoted-v1','unit-ids-v2'] if change=='mismatched_profile' else ['quoted-v1','quoted-v1','unit-ids-v2']
    child=new_adoption_target(tmp_path,value,proof,budgets=budgets,profiles=profiles)
    before={path.name:path.read_bytes() for path in child.iterdir()}
    with pytest.raises(dg.GenerationError):
        stage.adopt_source_map_prefix(child,parent,count=3 if change=='failed_count' else 2)
    assert before=={path.name:path.read_bytes() for path in child.iterdir()}


@pytest.mark.parametrize('change',['parent_raw','receipt','source_proof','fake_budget'])
def test_adopted_prefix_reopening_rechecks_parent_and_adoption_proof(tmp_path,monkeypatch,change):
    parent,value,proof=adoption_fixture(tmp_path,monkeypatch)
    child=new_adoption_target(tmp_path,value,proof,budgets=[900,900,1600],
                              profiles=['quoted-v1','quoted-v1','unit-ids-v2'])
    stage.adopt_source_map_prefix(child,parent,count=2)
    if change=='parent_raw':
        path=parent/'response_1.json'
        path.write_bytes(path.read_bytes()+b' ')
    elif change=='receipt':
        path=child/'source_map_prefix_adoption.json'
        item=stage._read(path); item['records'][0]['response_file_sha256']='b'*64
        item['adoption_digest']=dg.canonical_digest({k:v for k,v in item.items() if k!='adoption_digest'})
        path.write_text(json.dumps(item))
    elif change=='source_proof':
        path=child/'source_provenance.json'
        path.write_bytes(path.read_bytes()+b' ')
    else:
        stage._write_new(child/'budget_1.started.json',{'fabricated':True})
    with pytest.raises(dg.GenerationError):
        stage.replay_document_generation(child,tenant_id='tenant',deal_id='deal')


def test_optional_vectors_absent_defaults_unchanged_and_bad_vectors_cannot_prepare(tmp_path):
    value=full_source()
    root=tmp_path/'unchanged-default'
    stage.prepare_document_generation(root,[value],PIN,PIN,'pitch_deck',catalog_encoding=dg.CATALOG_ENCODING_V3)
    body,seed,_=stage._load_run(root)
    assert 'source_map_token_budgets' not in body and 'source_map_profiles' not in body
    assert body['source_map_requests']==[dg.request_json(request) for request in dg.frozen_source_map_requests(
        body['inventory'],PIN,kind='pitch_deck')]
    assert body['initial_request']==dg.request_json(dg.frozen_document_request(body['inventory'],PIN,'pitch_deck',catalog_encoding=dg.CATALOG_ENCODING_V3))
    with pytest.raises(dg.GenerationError):
        stage.prepare_document_generation(tmp_path/'bad-budget',[value],PIN,PIN,'pitch_deck',catalog_encoding=dg.CATALOG_ENCODING_V3,
            source_map_token_budgets=[900])
    assert not (tmp_path/'bad-budget').exists()
    with pytest.raises(dg.GenerationError):
        stage.prepare_document_generation(tmp_path/'bad-profile',[value],PIN,PIN,'pitch_deck',catalog_encoding=dg.CATALOG_ENCODING_V3,
            source_map_profiles=['quoted-v1','unknown-profile'])
    assert not (tmp_path/'bad-profile').exists()
