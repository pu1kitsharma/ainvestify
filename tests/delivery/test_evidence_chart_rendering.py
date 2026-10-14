"""Native editable chart caches copied from synthetic originals, never forecasts."""
from copy import deepcopy
from io import BytesIO
from zipfile import ZipFile

import pytest
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches
from docx import Document
from pypdf import PdfReader

from agents.research.document_evidence_manifest import build_document_manifest, digest
from delivery.rendering import (project_native_chart, render_intro, render_deck_pdf,
    render_memo, render_memo_pdf, inspect_evidence_chart_binding)
from delivery.rendering import inspect_evidence_chart_pdf_text
from delivery.inspection import inspect_bytes, inspect_export_pair


def source(chart_type=XL_CHART_TYPE.COLUMN_CLUSTERED):
    deck = Presentation(); slide = deck.slides.add_slide(deck.slide_layouts[6])
    data = CategoryChartData(); data.categories = ['FY2022', 'FY2023', 'FY2024']
    data.add_series('Recorded source', [10.25, 14.5, 12.75])
    if chart_type != XL_CHART_TYPE.PIE:
        data.add_series('Source scenario', [-2.5, 3.0, 5.25])
    slide.shapes.add_chart(chart_type, Inches(1), Inches(1), Inches(7), Inches(4), data)
    output = BytesIO(); deck.save(output)
    return build_document_manifest(output.getvalue(), 'pptx', tenant_id='synthetic-tenant',
                                   deal_id='synthetic-room', document_id='synthetic-input')


def project(manifest):
    return project_native_chart(manifest, manifest['charts'][0]['unit_id'],
        expected_binding=manifest['binding'], expected_manifest_digest=manifest['manifest_digest'],
        source_refs=['S1'])


def section(chart):
    return ('Synthetic original chart cache', 'The synthetic source cache remains unreviewed [S1].',
            '[S1] Synthetic immutable source chart.', 'chart', chart, 'evidence-chart-v1')


@pytest.mark.parametrize('chart_type', [XL_CHART_TYPE.COLUMN_CLUSTERED,
    XL_CHART_TYPE.COLUMN_STACKED, XL_CHART_TYPE.BAR_CLUSTERED, XL_CHART_TYPE.LINE_MARKERS,
    XL_CHART_TYPE.PIE])
def test_exact_source_cache_to_native_chart_editable_workbook_and_vector_pdf(chart_type):
    chart = project(source(chart_type))
    pptx = render_intro('Synthetic', [section(chart)])
    binding = inspect_evidence_chart_binding(pptx, chart)
    assert binding['native_chart_binding'] == binding['embedded_workbook_binding'] == 'pass'
    assert binding['release_status'] == 'blocked'
    deck = Presentation(BytesIO(pptx))
    assert sum(shape.has_chart for shape in deck.slides[0].shapes) == 1
    assert sum(shape.has_table for shape in deck.slides[0].shapes) == 1
    native = next(shape.chart for shape in deck.slides[0].shapes if shape.has_chart)
    if chart_type in (XL_CHART_TYPE.COLUMN_CLUSTERED, XL_CHART_TYPE.COLUMN_STACKED, XL_CHART_TYPE.BAR_CLUSTERED):
        assert all(series.invert_if_negative is False for series in native.series)
        assert native.value_axis.major_unit > 0
    pdf = render_deck_pdf('Synthetic', [section(chart)])
    assert len(PdfReader(BytesIO(pdf)).pages) == 1
    text = PdfReader(BytesIO(pdf)).pages[0].extract_text()
    assert '10.25 [S1]' in text and '12.75 [S1]' in text
    assert 'Unreviewed source chart cache' in text
    assert inspect_evidence_chart_pdf_text(pptx, pdf, chart)['source_cache_companion_text'] == 'pass'
    # Existing generic inspection remains strict about audience review of the
    # newly embedded editable workbook; this specialized binding is no waiver.
    assert any(row['code'] == 'embedded_payload_requires_audience_review'
               for row in inspect_bytes(pptx, 'pptx')['findings'])
    assert inspect_export_pair(pptx, 'pptx', pdf)['pair_status'] == 'fail'


def test_memo_chart_has_editable_exact_table_and_discloses_chart_image_limit():
    chart = project(source())
    docx, pdf = render_memo('Synthetic', [section(chart)]), render_memo_pdf('Synthetic', [section(chart)])
    doc = Document(BytesIO(docx))
    assert len(doc.inline_shapes) == 1 and len(doc.tables) == 1
    assert doc.tables[0].cell(1, 1).text == '10.25 [S1]'
    assert any('Chart image is not editable' in paragraph.text for paragraph in doc.paragraphs)
    assert inspect_export_pair(docx, 'docx', pdf)['pair_status'] == 'pass'


def test_original_binding_cache_and_generated_numeric_tampering_are_blocked():
    manifest = source(); chart = project(manifest)
    changed = deepcopy(manifest); changed['units'][0]['content']['native_xml']['text'] = 'tampered'
    with pytest.raises(ValueError, match='integrity'):
        project(changed)
    altered = deepcopy(chart); altered['series'][0]['points'][0]['numeric_literal'] = '99.0'
    with pytest.raises(ValueError):
        render_intro('Synthetic', [section(altered)])
    pptx = render_intro('Synthetic', [section(chart)])
    stream = BytesIO()
    with ZipFile(BytesIO(pptx)) as original, ZipFile(stream, 'w') as output:
        for name in original.namelist():
            content = original.read(name)
            if name == 'ppt/charts/chart1.xml':
                content = content.replace(b'<c:v>10.25</c:v>', b'<c:v>99.0</c:v>')
            output.writestr(name, content)
    assert inspect_evidence_chart_binding(stream.getvalue(), chart)['native_chart_binding'] == 'fail'


def test_external_source_chart_and_sparse_or_nonfinite_cache_are_not_guessed():
    for defect in ('external', 'sparse', 'nonfinite', 'extra_axis'):
        manifest = source()
        chart_unit = next(unit for unit in manifest['units'] if unit['kind'] == 'chart')
        if defect == 'external':
            chart_unit['content']['relationships'].append({'TargetMode': 'External', 'Target': 'https://invalid.example'})
        else:
            def nodes(node):
                yield node
                for child in node.get('children', []):
                    yield from nodes(child)
            tree = chart_unit['content']['native_xml']
            if defect == 'sparse':
                cache = next(node for node in nodes(tree) if node['tag'].endswith('}numCache'))
                cache['children'] = [node for node in cache['children'] if not
                    (node['tag'].endswith('}pt') and node['attributes'].get('idx') == '1')]
            elif defect == 'extra_axis':
                area = next(node for node in nodes(tree) if node['tag'].endswith('}plotArea'))
                axis = next(node for node in area['children'] if node['tag'].endswith('}valAx'))
                area['children'].append(deepcopy(axis))
            else:
                literal = next(unit for unit in manifest['units'] if unit['kind'] == 'chart_literal'
                               and unit['content'] == '10.25')
                literal['content'] = 'NaN'; literal['content_sha256'] = digest('NaN')
                value = next(node for node in nodes(tree) if node['tag'].endswith('}v') and node['text'] == '10.25')
                value['text'] = 'NaN'
        chart_unit['content_sha256'] = digest(chart_unit['content'])
        # Even a newly sealed manifest cannot turn an unsupported source cache
        # into valid plotted evidence. It is never repaired or interpolated.
        manifest['manifest_digest'] = digest({key: value for key, value in manifest.items() if key != 'manifest_digest'})
        with pytest.raises(ValueError):
            project(manifest)
