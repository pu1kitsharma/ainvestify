from io import BytesIO
from zipfile import ZipFile
from openpyxl import Workbook
from openpyxl.workbook.defined_name import DefinedName
from delivery.inspection import inspect_bytes, inspect_export_pair, summarize_report
from agents.core.ingestion_agent import ingest_excel


def workbook():
    w=Workbook();w.active.title='Inputs';w.active['A1']=10
    s=w.create_sheet('Hidden');s.sheet_state='hidden';s['A1']='=Inputs!A1*2';s['B1']='#REF!'
    b=BytesIO();w.save(b);return b.getvalue()


def test_hidden_formulas_errors_and_missing_caches_retained(tmp_path):
    content=workbook();path=tmp_path/'synthetic.xlsx';path.write_bytes(content)
    document=ingest_excel(str(path),'one','deal')
    report=document.workbook_inventory
    assert report['sheet_count']==2
    assert report['formula_count']==1
    assert report['formula_cells'][0]['formula']=='Inputs!A1*2'
    assert report['formula_cells'][0]['cached_value'] is None
    assert any(s.get('state')=='hidden' for s in report['sheets'])
    assert {'stored_cell_error','missing_formula_cache'} <= {f['code'] for f in report['findings']}
    assert report['calculation_status']=='not_run'
    assert path.read_bytes()==content


def test_malformed_and_traversal_packages_fail():
    assert inspect_bytes(b'not an office document','xlsx')['structural_status']=='fail'
    b=BytesIO()
    with ZipFile(b,'w') as z:z.writestr('../escape.xml','<x/>')
    assert inspect_bytes(b.getvalue(),'xlsx')['structural_status']=='fail'


def test_duplicate_parts_and_xml_entity_rejected():
    b=BytesIO()
    with ZipFile(b,'w') as z:
        z.writestr('xl/workbook.xml','<!DOCTYPE x [<!ENTITY a SYSTEM "file:///etc/passwd">]><x>&a;</x>')
        z.writestr('[Content_Types].xml','<x/>')
    assert inspect_bytes(b.getvalue(),'xlsx')['structural_status']=='fail'


def test_cross_sheet_lineage_and_cycles_remain_unqualified():
    w=Workbook();w.active.title='Inputs';w.active['A1']=10
    hidden=w.create_sheet('Hidden');hidden.sheet_state='hidden'
    hidden['B1']='=Inputs!A1*2';hidden['B2']='=B1+1'
    b=BytesIO();w.save(b)
    report=inspect_bytes(b.getvalue(),'xlsx')
    assert report['dependency_graph']['Hidden!B1']==['Inputs!A1']
    assert report['dependency_graph']['Hidden!B2']==['Hidden!B1']
    assert report['dependency_coverage']=='cell_references_only'
    assert report['calculation_status']=='not_run'
    assert 'dependency_graph' not in summarize_report(report)

    hidden['B1']='=B2+1'
    b=BytesIO();w.save(b)
    report=inspect_bytes(b.getvalue(),'xlsx')
    assert 'circular_formula_reference' in {f['code'] for f in report['findings']}
    assert report['dependency_coverage']=='partial'


def test_dynamic_and_unbounded_references_are_explicit_gaps():
    w=Workbook();w.active['A1']='=INDIRECT("B1")';w.active['A2']='=SUM(B:B)'
    b=BytesIO();w.save(b)
    report=inspect_bytes(b.getvalue(),'xlsx')
    assert {'dynamic_reference','unsupported_or_unknown_reference'} <= {f['code'] for f in report['findings']}
    assert report['dependency_coverage']=='partial'


def test_simple_defined_names_resolve_with_local_scope():
    w=Workbook();w.active.title='Inputs';w.active['A1']=10
    other=w.create_sheet('Other');other['B1']=20
    w.defined_names.add(DefinedName('Driver', attr_text="'Inputs'!$A$1"))
    other.defined_names.add(DefinedName('Driver', attr_text="'Other'!$B$1", localSheetId=1))
    w.active['C1']='=Driver*2';other['C1']='=Driver*3'
    b=BytesIO();w.save(b)
    report=inspect_bytes(b.getvalue(),'xlsx')
    assert report['dependency_graph']['Inputs!C1']==['Inputs!A1']
    assert report['dependency_graph']['Other!C1']==['Other!B1']
    assert report['dependency_coverage']=='cell_references_and_defined_names'


def test_expression_defined_name_remains_explicit_gap():
    w=Workbook();w.defined_names.add(DefinedName('Driver', attr_text='SUM(A1:A2)'))
    w.active['B1']='=Driver'
    b=BytesIO();w.save(b)
    report=inspect_bytes(b.getvalue(),'xlsx')
    assert any(f['code']=='unsupported_defined_name' for f in report['findings'])
    assert report['dependency_coverage']=='partial'


def test_exact_export_pair_catches_missing_or_extra_slide_pages():
    from reportlab.pdfgen import canvas
    from pptx import Presentation
    from pptx.util import Inches

    deck = Presentation()
    for text in ('First synthetic slide', 'Second synthetic slide'):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text = text
    editable = BytesIO()
    deck.save(editable)
    pptx = editable.getvalue()
    def pdf_with_pages(texts):
        output = BytesIO()
        document = canvas.Canvas(output)
        for text in texts:
            document.drawString(72, 720, text)
            document.showPage()
        document.save()
        return output.getvalue()
    mismatched = inspect_export_pair(pptx, 'pptx', pdf_with_pages(['First synthetic slide']))
    assert mismatched['pair_status'] == 'fail'
    assert mismatched['content_layout_status'] == 'not_run'
    assert mismatched['findings'][0]['code'] == 'export_slide_page_count_mismatch'
    matched = inspect_export_pair(pptx, 'pptx', pdf_with_pages([
        'First synthetic slide', 'Second synthetic slide']))
    assert matched['pair_status'] == matched['editable_text_status'] == 'pass'
    stale = inspect_export_pair(pptx, 'pptx', pdf_with_pages([
        'First synthetic slide', 'Stale second slide']))
    assert stale['pair_status'] == 'fail'
    assert stale['findings'][0]['code'] == 'export_editable_text_missing_from_pdf'
    added = inspect_export_pair(pptx, 'pptx', pdf_with_pages([
        'First synthetic slide. This entirely new claim is present only in the exported PDF.',
        'Second synthetic slide']))
    assert added['pair_status'] == 'fail'
    assert any(f['code'] == 'export_pdf_only_substantive_text' and f['page'] == 1
               and len(f['text_sha256']) == 64 for f in added['findings'])
    assert 'entirely new claim' not in str(added)


def test_orphan_slide_part_is_not_counted_as_a_presented_slide():
    from pptx import Presentation

    deck = Presentation()
    deck.slides.add_slide(deck.slide_layouts[6])
    output = BytesIO()
    deck.save(output)
    modified = BytesIO()
    with ZipFile(BytesIO(output.getvalue())) as source, ZipFile(modified, 'w') as target:
        for item in source.infolist():
            target.writestr(item, source.read(item.filename))
        target.writestr('ppt/slides/slide999.xml',
                        '<p:sld xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"/>')
    report = inspect_bytes(modified.getvalue(), 'pptx')
    assert report['slide_count'] == 1
    assert report['structural_status'] == 'fail'
    assert any(item['code'] == 'orphan_slide_parts' for item in report['findings'])
