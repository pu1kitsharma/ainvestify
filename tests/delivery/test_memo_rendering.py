"""The editable memo and its PDF keep the same recorded section text."""

from io import BytesIO
import re

from docx import Document
from pypdf import PdfReader

from delivery.inspection import inspect_export_pair
from delivery.rendering import render_memo, render_memo_pdf


def test_memo_pair_preserves_long_model_prose_excerpts_and_sources():
    body = (
        'The local model reports an unresolved funding status [S1]. ' * 10
        + '\n\nSource claims and exact excerpts:\n'
        + '[S1] Exact excerpt: {"status":"unknown","amount":"INR 20 crore"}.\n' * 8
    )
    sources = '[S1] https://example.invalid/record | version synthetic-v1\n' * 8
    sections = [(f'Funding reconciliation {index}', body, sources)
                for index in range(1, 4)]
    docx = render_memo('Synthetic company', sections)
    pdf = render_memo_pdf('Synthetic company', sections)
    pair = inspect_export_pair(docx, 'docx', pdf)
    assert pair['pair_status'] == pair['editable_text_status'] == 'pass'
    document = Document(BytesIO(docx))
    editable = '\n'.join(paragraph.text for paragraph in document.paragraphs)
    pages = PdfReader(BytesIO(pdf)).pages
    printed = '\n'.join(page.extract_text() for page in pages)
    for value in ('Funding reconciliation 1', 'Source claims and exact excerpts:',
                  '"status":"unknown"', 'synthetic-v1'):
        assert value in editable
        assert value in printed
    normalize = lambda value: re.sub(r'\s+', ' ', value).strip()
    for heading, model_body, model_sources in sections:
        assert heading in editable
        for line in (*model_body.splitlines(), *model_sources.splitlines()):
            if line.strip():
                assert normalize(line) in normalize(editable)
    assert len(pages) >= 2
    assert all('DRAFT' in page.extract_text() and 'Page ' in page.extract_text()
               for page in pages)


def test_memo_pair_preserves_plain_sections_without_evidence_marker():
    sections = [('Decision', 'The model asks for primary dated records [S1].',
                 '[S1] https://example.invalid/source | version v1')]
    pair = inspect_export_pair(render_memo('Synthetic company', sections), 'docx',
                               render_memo_pdf('Synthetic company', sections))
    assert pair['pair_status'] == pair['editable_text_status'] == 'pass'


def _dense_sections(count=4):
    body = ('The local model reports an unresolved funding status and asks for dated records [S1]. ' * 6
            + '\n\nSource claims and exact excerpts:\n'
            + '[S1] Exact excerpt: {"status":"unknown","amount":"INR 20 crore"}.\n' * 5)
    sources = '[S1] https://example.invalid/record | version synthetic-v1\n' * 4
    return [(f'Funding reconciliation {index}', body, sources) for index in range(1, count + 1)]


def test_editable_memo_uses_the_pdf_page_and_readable_set_off_evidence():
    from docx.oxml.ns import qn
    from delivery.rendering import _MEMO
    sections = _dense_sections()
    document = Document(BytesIO(render_memo('Synthetic company', sections)))
    pdf = PdfReader(BytesIO(render_memo_pdf('Synthetic company', sections)))
    page = document.sections[0]
    # One sheet and one set of margins for both files.
    box = pdf.pages[0].mediabox
    assert (round(page.page_width.mm), round(page.page_height.mm)) == (210, 297)
    assert (round(float(box.width) / 72 * 25.4), round(float(box.height) / 72 * 25.4)) == (210, 297)
    assert [round(value.mm) for value in (page.left_margin, page.right_margin, page.top_margin,
                                          page.bottom_margin)] == [20, 20, 25, 22]
    styles = document.styles
    size = lambda name: styles[name].font.size.pt
    assert (size('Normal'), size('Memo Evidence'), size('Memo Sources')) == (
        _MEMO['body'], _MEMO['evidence'], _MEMO['sources']) == (11, 9.5, 9)
    assert styles['Normal'].paragraph_format.widow_control is True
    # Excerpts sit behind a left rule and an indent; each section opens with a hairline.
    rule = lambda name, side: styles[name].element.pPr.find(qn('w:pBdr')).find(qn(f'w:{side}'))
    assert rule('Memo Evidence', 'left') is not None and rule('Heading 1', 'top') is not None
    assert styles['Memo Evidence'].paragraph_format.left_indent.mm > 3
    # Pagination controls: a heading or label never ends a page; an excerpt is not split.
    assert styles['Heading 1'].paragraph_format.keep_with_next
    assert styles['Memo Label'].paragraph_format.keep_with_next
    assert styles['Memo Evidence'].paragraph_format.keep_together
    by_style = {}
    for paragraph in document.paragraphs:
        by_style.setdefault(paragraph.style.name, []).append(paragraph.text)
    assert len(by_style['Memo Evidence']) == 20 and len(by_style['Memo Sources']) == 16
    # The PDF text sizes match the editable styles.
    sizes = set()
    pdf.pages[0].extract_text(visitor_text=lambda text, cm, tm, font, size_: sizes.add(
        round(size_ * tm[0], 1)) if text.strip() else None)
    assert {11, 9.5, 9, 13.5, 21} <= sizes and min(sizes) >= 8


def test_memo_pdf_never_strands_a_heading_or_label_at_the_foot_of_a_page():
    labels = ('Source claims and exact excerpts:', 'Source versions')
    for count in range(2, 9):
        sections = _dense_sections(count)
        docx, pdf = render_memo('Synthetic company', sections), render_memo_pdf('Synthetic company', sections)
        assert inspect_export_pair(docx, 'docx', pdf)['pair_status'] == 'pass'
        for page in PdfReader(BytesIO(pdf)).pages:
            lines = [line.strip() for line in page.extract_text().splitlines() if line.strip()]
            content = [line for line in lines if not line.startswith(
                ('DRAFT', 'INVESTMENT MEMORANDUM DRAFT', 'Page '))]
            assert content and content[-1] not in labels
            assert not content[-1].startswith('Funding reconciliation')


def test_reader_source_names_get_matching_editable_and_pdf_labels():
    sections = [('Synthetic section', 'Synthetic source-reported point [S1].',
                 '[S1] Synthetic official record')]
    docx = render_memo('Synthetic company', sections)
    pdf = render_memo_pdf('Synthetic company', sections)
    assert 'Sources cited' in PdfReader(BytesIO(pdf)).pages[0].extract_text()
    assert 'Sources cited' in '\n'.join(p.text for p in Document(BytesIO(docx)).paragraphs)
    assert inspect_export_pair(docx, 'docx', pdf)['pair_status'] == 'pass'


def test_paginated_memo_preserves_full_long_question_heading_in_editable_pdf_pair():
    question=('Which independently reviewed records establish the reported customer distribution, '
              'execution status, accounting policies and operating assumptions before an investment decision?')
    assert 100 < len(question) <= 250
    sections=[(question,'Request the original records and reconcile their dated scope before deciding.','')]
    docx=render_memo('Synthetic company',sections)
    pdf=render_memo_pdf('Synthetic company',sections)
    pair=inspect_export_pair(docx,'docx',pdf)
    assert pair['pair_status']==pair['editable_text_status']=='pass'
    assert question in '\n'.join(p.text for p in Document(BytesIO(docx)).paragraphs)
    normalized=lambda text:re.sub(r'\s+',' ',text).strip()
    printed='\n'.join(page.extract_text() for page in PdfReader(BytesIO(pdf)).pages)
    assert normalized(question) in normalized(printed)


def test_paginated_heading_bound_does_not_relax_slides_body_sources_or_section_cap():
    import pytest
    from delivery.rendering import render_intro,render_deck_pdf,render_research_pdf,_checked_sections
    heading=('Which dated company records should be verified before deciding on the proposed investment? '*4)[:300]
    sections=[(heading,'Recorded conditional diligence relevance.','')]
    assert _checked_sections(sections,paginated=True)==sections
    for renderer in (render_intro,render_deck_pdf):
        with pytest.raises(ValueError,match='exceeds supported layout'):renderer('Synthetic',sections)
    for renderer in (render_memo,render_memo_pdf,render_research_pdf):
        # The research PDF and both Memo formats keep the exact300-character
        # heading;300 is a layout bound, not permission to shorten model text.
        rendered=renderer('Synthetic',sections)
        assert rendered
        if renderer is not render_memo:
            printed=' '.join(page.extract_text() for page in PdfReader(BytesIO(rendered)).pages)
            assert re.sub(r'\s+',' ',heading).strip() in re.sub(r'\s+',' ',printed).strip()
        with pytest.raises(ValueError,match='exceeds supported layout'):
            renderer('Synthetic',[(heading+'x',sections[0][1],'')])
        for body,sources in [('x'*1201,''),('Recorded body','x'*1201)]:
            with pytest.raises(ValueError,match='exceeds supported layout'):
                renderer('Synthetic',[('Question',body,sources)])
        with pytest.raises(ValueError,match='sections'):
            renderer('Synthetic',[('Question','Recorded body','')]*31)


def test_source_registry_hash_uris_preserved_in_memo_pair_and_registry_style_only():
    from delivery.rendering import render_research_pdf,_source_registry_lines
    hashes=('a'*64,'b'*64,'c'*64)
    lines=[f'[S{i}] private://{value}#pages/{i} | version synthetic-{i}'
           for i,value in enumerate(hashes,start=1)]
    registry='\n'.join(lines)
    narrative='The source-reported operating model remains subject to verification before an investment decision.'
    sections=[('Recorded assessment',narrative,''),('Sources and retained versions',registry,'')]
    assert _source_registry_lines(registry)==lines
    assert _source_registry_lines(registry+'\n'+narrative)==[]
    docx=render_memo('Synthetic company',sections);pdf=render_memo_pdf('Synthetic company',sections)
    pair=inspect_export_pair(docx,'docx',pdf)
    assert pair['pair_status']==pair['editable_text_status']=='pass'
    paragraphs=Document(BytesIO(docx)).paragraphs
    for line in lines:
        paragraph=next(p for p in paragraphs if p.text==line)
        assert paragraph.style.name=='Memo Sources'
    ordinary=next(p for p in paragraphs if p.text==narrative)
    assert ordinary.style.name=='Normal'
    for rendered in (pdf,render_research_pdf('Synthetic company',sections)):
        text='\n'.join(page.extract_text() for page in PdfReader(BytesIO(rendered)).pages)
        for line in lines:
            assert 'private://'+line.split('private://',1)[1].split(' |',1)[0] in text
    # The inspector remains strict: formatting support cannot waive missing
    # or altered identifiers in an otherwise matching export.
    changed=registry.replace(hashes[0], 'd'*64)
    bad=inspect_export_pair(docx,'docx',render_memo_pdf('Synthetic company',[
        sections[0],('Sources and retained versions',changed,'')]))
    assert bad['pair_status']=='fail'


def test_multipage_source_inventories_keep_each_complete_entry_on_one_pdf_page():
    from delivery.rendering import render_research_pdf
    # Three-line entries deliberately approach page edges repeatedly. Metadata
    # stays one exact source entry, including its opaque64-character URI hash.
    entries=[f'[S{i}] private://{i:064x}#pages/{i} | '+
        (f'dated source inventory metadata for entry {i} with verification scope and original document lineage ' * 2)
        for i in range(1,25)]
    sections=[('Recorded source registry '+str(index+1),'\n'.join(entries[index:index+3]),'')
              for index in range(0,len(entries),3)]
    pdf=render_memo_pdf('Synthetic multi-page registry',sections)
    docx=render_memo('Synthetic multi-page registry',sections)
    pair=inspect_export_pair(docx,'docx',pdf)
    assert pair['pair_status']==pair['editable_text_status']=='pass'
    normalized=lambda text:re.sub(r'\s+',' ',text).strip()
    for output in (pdf,render_research_pdf('Synthetic multi-page registry',sections)):
        pages=[normalized(page.extract_text()) for page in PdfReader(BytesIO(output)).pages]
        assert len(pages)>1
        for entry in entries:
            assert sum(normalized(entry) in page for page in pages)==1
    # The separately supplied bibliography uses the same per-entry guarantee.
    source_sections=[('Recorded prose','The source-reported assessment remains unverified.',
                      '\n'.join(entries[index:index+3])) for index in range(0,len(entries),3)]
    pages=[normalized(page.extract_text()) for page in PdfReader(BytesIO(
        render_memo_pdf('Synthetic multi-page bibliography',source_sections))).pages]
    for entry in entries:assert sum(normalized(entry) in page for page in pages)==1


def _bounded_long_prose_sections():
    sections=[]
    for index in range(10):
        paragraphs=[]
        for part in range(2):
            paragraphs.append((f'Recorded narrative section {index} paragraph {part}. '+
                'The company description is a source report whose execution assumptions require independent verification before deciding on investment suitability. '*3).strip())
        sections.append((f'Conditional assessment {index}','\n\n'.join(paragraphs),''))
    return sections


def test_bounded_narrative_paragraphs_are_kept_together_in_docx_style_and_direct_pdf():
    sections=_bounded_long_prose_sections()
    docx=render_memo('Synthetic pagination canary',sections)
    doc=Document(BytesIO(docx))
    assert doc.styles['Normal'].paragraph_format.keep_together is True
    for paragraph in doc.paragraphs:
        if paragraph.text.startswith('Recorded narrative'):
            assert paragraph.style.paragraph_format.keep_together is True
    pdf=render_memo_pdf('Synthetic pagination canary',sections)
    normalized=lambda text:re.sub(r'\s+',' ',text).strip()
    pages=[normalized(page.extract_text()) for page in PdfReader(BytesIO(pdf)).pages]
    assert len(pages)>1
    for _,body,_ in sections:
        for paragraph in body.split('\n\n'):
            assert sum(normalized(paragraph) in page for page in pages)==1
    assert inspect_export_pair(docx,'docx',pdf)['pair_status']=='pass'


def test_native_libreoffice_keeps_bounded_memo_prose_paragraphs_on_one_page(tmp_path):
    import os
    import subprocess
    from pathlib import Path
    import pytest
    if os.environ.get('RUN_NATIVE_MEMO_CANARY')!='1':
        pytest.skip('Explicit native Office canary; enable RUN_NATIVE_MEMO_CANARY=1')
    executable=Path('/Applications/LibreOffice.app/Contents/MacOS/soffice')
    if not executable.is_file():pytest.skip('Pinned local LibreOffice executable unavailable')
    sections=_bounded_long_prose_sections()
    docx=render_memo('Synthetic pagination canary',sections)
    source=tmp_path/'memo.docx';source.write_bytes(docx)
    profile=(tmp_path/'isolated-office-profile').as_uri()
    result=subprocess.run([str(executable),'-env:UserInstallation='+profile,'--headless',
        '--convert-to','pdf','--outdir',str(tmp_path),str(source)],capture_output=True,text=True,timeout=60)
    assert result.returncode==0,result.stderr
    output=tmp_path/'memo.pdf';assert output.is_file()
    normalized=lambda text:re.sub(r'\s+',' ',text).strip()
    pages=[normalized(page.extract_text()) for page in PdfReader(BytesIO(output.read_bytes())).pages]
    assert len(pages)>1
    for _,body,_ in sections:
        for paragraph in body.split('\n\n'):
            assert sum(normalized(paragraph) in page for page in pages)==1
    assert inspect_export_pair(docx,'docx',output.read_bytes())['pair_status']=='pass'
