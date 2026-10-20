import json
import pytest
from agents.runtime.atomic_banker import AtomicJSONTransport,run_atomic_banker

CASE={'sources':[{'id':'H1','text':'Synthetic historical test source only'}],'requested_metric_ids':['metric']}

def envelope(content,**extra):
    return {'model':'installed-test','done':True,'done_reason':'stop','prompt_eval_count':2,'eval_count':3,
        'message':{'role':'assistant','content':json.dumps(content)},**extra}

BLOCK={'sheet':'Agent_Chosen','start':'A1','values':[['Input',5],['Cost',2],['Difference','=B1-B2']],
    'comment':'Synthetic source H1; USD; historical; model-selected layout'}
MAP={'metrics':{'metric':{'sheet':'Agent_Chosen','cell':'B3','unit':'USD'}},'calculated_ranges':[{'sheet':'Agent_Chosen','range':'B3'}]}
MEMO={'sections':[{'heading':heading,'text':'Agent-written selected value is {{metric}} USD; pending review.',
    'values':[{'key':'metric','sheet':'Agent_Chosen','cell':'B3','unit':'USD','decimals':1}]}
    for heading in ['Historical evidence','Limitations']]}

def test_scripted_three_stage_payloads_exact_and_no_native_tools(tmp_path):
    responses=iter([envelope(BLOCK),envelope(MAP),envelope(MEMO)]);requests=[]
    def transport(**kw):requests.append(kw['request']);return next(responses)
    output=tmp_path/'run'
    report=run_atomic_banker(case=CASE,output=output,model='installed-test',transport=transport)
    assert report['status']=='draft_exported' and report['calls']==3 and report['tokens']==15
    assert report['raw_capture'] is False and report['full_banker'] is False
    assert all('tools' not in request and request['think'] is False and request['format'] for request in requests)
    from openpyxl import load_workbook
    assert load_workbook(output/'model.xlsx')['Agent_Chosen']['B3'].value=='=B1-B2'
    spec=json.loads((output/'output_spec.json').read_text())
    assert spec['metrics']==MAP['metrics'] and len(spec['memo_bindings'])==2
    assert set(spec['memo_bindings'][0])=={'paragraph','start','end','metric_id','unit'}
    from docx import Document
    document=Document(output/'atomic-banker-draft.memo/atomic-banker-draft.docx')
    assert not any('workbook_sha256' in paragraph.text for paragraph in document.paragraphs)
    for binding in spec['memo_bindings']:
        assert document.paragraphs[binding['paragraph']].text[binding['start']:binding['end']]=='3.0'


def test_length_failure_preserves_partial_content_and_usage_no_retries(tmp_path):
    calls=[]
    def transport(**kw):calls.append(1);return envelope(BLOCK,done_reason='length')
    report=run_atomic_banker(case=CASE,output=tmp_path/'run',model='installed-test',transport=transport)
    assert report['status']=='failed' and report['tokens']==5 and len(calls)==1
    assert json.loads((tmp_path/'run/model-content.json').read_text())['content']==json.dumps(BLOCK)


def test_invalid_matrix_is_not_repaired(tmp_path):
    block={**BLOCK,'values':[[1,2],[3]]}
    calls=[]
    def transport(**kw):calls.append(1);return envelope(block)
    report=run_atomic_banker(case=CASE,output=tmp_path/'run',model='installed-test',transport=transport)
    assert report['status']=='failed' and len(calls)==1
    assert 'payload failed' in report['error']['message']


def test_cumulative_deadline_stops_before_write(tmp_path):
    now=[0]
    def transport(**kw):now[0]=121;return envelope(BLOCK)
    report=run_atomic_banker(case=CASE,output=tmp_path/'run',model='installed-test',transport=transport,clock=lambda:now[0])
    assert report['status']=='failed' and report['calls']==1
    from openpyxl import load_workbook
    assert 'Agent_Chosen' not in load_workbook(tmp_path/'run/model.xlsx').sheetnames


def test_memo_mapping_mismatch_fails_without_repair(tmp_path):
    memo=json.loads(json.dumps(MEMO));memo['sections'][0]['values'][0]['cell']='B1'
    responses=iter([envelope(BLOCK),envelope(MAP),envelope(memo)])
    report=run_atomic_banker(case=CASE,output=tmp_path/'run',model='installed-test',transport=lambda **kw:next(responses))
    assert report['status']=='failed' and 'disagrees' in report['error']['message']

class Response:
    status=200
    def __init__(self,body):self.body=body
    def getheader(self,name):return str(len(self.body))
    def read(self,n):part,self.body=self.body[:n],self.body[n:];return part
class Connection:
    sock=None
    def __init__(self,body):self.response=Response(body);self.requests=[]
    def request(self,*a,**kw):self.requests.append((a,kw))
    def getresponse(self):return self.response
    def close(self):pass

def test_atomic_transport_keeps_malformed_exact_bytes_before_parser_failure():
    connection=Connection(b'  malformed raw bytes\n');receipts=[]
    transport=AtomicJSONTransport(connection_factory=lambda *a,**kw:connection)
    with pytest.raises(ValueError):transport(endpoint='http://127.0.0.1:11434/v1',request={'model':'installed-test'},timeout=1,
        receipt=lambda *args:receipts.append(args))
    assert receipts[0][1]==connection.requests[0][1]['body']
    assert receipts[1][1]==b'  malformed raw bytes\n'

def test_two_stages_stops_with_actual_model_mapping_and_no_memo(tmp_path):
    calls=[];responses=iter([envelope(BLOCK),envelope(MAP)])
    def transport(**kw):calls.append(kw['request']);return next(responses)
    output=tmp_path/'two-stage'
    report=run_atomic_banker(case=CASE,output=output,model='installed-test',transport=transport,stages=2)
    assert report['status']=='model_mapped_draft' and report['calls']==2 and report['requested_stages']==2
    assert report['full_banker'] is False and report['im_claim'] is False
    assert report['stages']==['model','mapping']
    assert not (output/'memo-authored.json').exists() and not list(output.glob('*.memo'))
    spec=json.loads((output/'output_spec.json').read_text())
    assert spec['metrics']==MAP['metrics'] and spec['memo_bindings']==[] and spec['im_claim'] is False
    saved=json.loads((output/'report.json').read_text())
    assert saved['status']=='model_mapped_draft' and saved['calls']==2
