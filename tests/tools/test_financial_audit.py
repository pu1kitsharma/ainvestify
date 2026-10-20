from copy import deepcopy
from hashlib import sha256

from docx import Document
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment

from agents.tools.financial_audit import audit_financial_artifacts


def fixture(root):
    (root/'source.txt').write_text('Synthetic reported revenue 100 and cost 70; independent fixture only.')
    book=Workbook()
    ws=book.active
    ws.title='Model'
    for cell,value in {'A1':100,'A2':70,'A3':'=A1-A2','B1':120,'B2':80,'B3':'=B1-B2','C1':'=A3+B3','C2':1.2}.items():
        ws[cell]=value
        status='historical' if cell in ('A1','A2','A3') else 'company_projection'
        ws[cell].comment=Comment(f'unit: USD; status: {status}; rationale: synthetic assumption only','Agent')
    ws['C2'].comment=Comment('unit: ratio; status: company_projection; assumption: synthetic growth driver','Agent')
    book.save(root/'model.xlsx')
    doc=Document()
    doc.add_paragraph('Revenue 100. Profit 30. Cash 70.')
    doc.save(root/'memo.docx')
    spec={'metrics':{key:{'sheet':'Model','cell':cell,'unit':'USD'} for key,cell in {'revenue':'A1','cost':'A2','profit':'A3','future_profit':'B3','cash':'C1'}.items()},
          'calculated_ranges':[{'sheet':'Model','range':r} for r in ['A3','B3','C1']],
          'assumptions':[{'sheet':'Model','cell':'C2'}],
          'memo_bindings':[{'paragraph':0,'start':start,'end':end,'metric_id':metric,'unit':'USD'} for start,end,metric in [(8,11,'revenue'),(20,22,'profit'),(29,31,'cash')]]}
    # Numeric-token spans exclude sentence punctuation.
    text='Revenue 100. Profit 30. Cash 70.'
    for b,token in zip(spec['memo_bindings'],['100','30','70']):
        b['start']=text.index(token);b['end']=b['start']+len(token)
    expectations={'expected_outputs':[{'id':m,'value':v,'unit':'USD','tolerance':1e-8,'require_formula':m not in ('revenue','cost'),'status':'historical' if m in ('revenue','cost','profit') else 'company_projection'} for m,v in [('revenue',100),('cost',70),('profit',30),('future_profit',40),('cash',70)]],
                  'required_memo_metrics':['revenue','profit','cash'],
                  'checks':[{'id':'profit_tie','kind':'difference','inputs':['revenue','cost'],'output':'profit','tolerance':1e-8},
                            {'id':'cash_tie','kind':'sum','inputs':['profit','future_profit'],'output':'cash','tolerance':1e-8}]}
    return spec,expectations


def audit(root,spec,expected):
    return audit_financial_artifacts(root,workbook='model.xlsx',memo='memo.docx',output_spec=spec,expectations=expected,sources=['source.txt'])


def codes(report):
    return {r['code'] for r in report['failures']}


def test_exact_artifact_success_and_immutable_read(tmp_path):
    spec,key=fixture(tmp_path)
    before={p.name:sha256(p.read_bytes()).hexdigest() for p in tmp_path.iterdir()}
    result=audit(tmp_path,spec,key)
    assert result['status']=='objective_checks_passed',result['failures']
    assert result['release']=='blocked'
    assert result['office_recalculation']=='not_qualified'
    assert result['observations']['metric_values']['cash']==70
    assert before=={p.name:sha256(p.read_bytes()).hexdigest() for p in tmp_path.iterdir()}


def test_broken_cash_and_scenario_and_hardcoded_output(tmp_path):
    spec,key=fixture(tmp_path)
    book=load_workbook(tmp_path/'model.xlsx')
    book['Model']['B3']='=B1-B2+9'
    book['Model']['C1']=79
    book.save(tmp_path/'model.xlsx')
    result=audit(tmp_path,spec,key)
    assert {'expected_value_mismatch','hardcode_or_missing_calculation','expected_formula_missing','memo_number_cell_mismatch'} <= codes(result)


def test_hidden_error_and_unsupported_formula_fail(tmp_path):
    spec,key=fixture(tmp_path)
    book=load_workbook(tmp_path/'model.xlsx')
    ws=book.create_sheet('Hidden')
    ws.sheet_state='hidden'
    ws['A1']='=#REF!'
    ws['A2']='=NOPE(1)'
    ws['A3']='#VALUE!'
    book.save(tmp_path/'model.xlsx')
    result=audit(tmp_path,spec,key)
    assert 'Hidden' in result['observations']['hidden_sheets']
    assert {'stored_cell_error','formula_evaluation_failed'} <= codes(result)


def test_literal_unbound_and_wrong_unit_cannot_be_declared_away(tmp_path):
    spec,key=fixture(tmp_path)
    doc=Document(tmp_path/'memo.docx')
    doc.add_paragraph('Unsupported claim 999 and 1e3.')
    doc.save(tmp_path/'memo.docx')
    spec['memo_bindings'][0]['unit']='EUR'
    result=audit(tmp_path,spec,key)
    assert {'memo_unit_unapproved','memo_number_unbound','required_memo_metric_missing'} <= codes(result)
    assert any(f.get('literal')=='1e3' for f in result['failures'])


def test_missing_requirements_and_metadata_fail(tmp_path):
    spec,key=fixture(tmp_path)
    spec['memo_bindings']=[]
    spec['assumptions']=[]
    key['checks']=[]
    book=load_workbook(tmp_path/'model.xlsx')
    book['Model']['A1'].comment=None
    book.save(tmp_path/'model.xlsx')
    result=audit(tmp_path,spec,key)
    assert {'memo_numeric_bindings_missing','assumptions_missing','independent_arithmetic_checks_missing','unit_metadata_missing','status_metadata_missing'} <= codes(result)


def test_source_and_artifact_hashes_exact_and_missing_artifact_fails(tmp_path):
    spec,key=fixture(tmp_path)
    first=audit(tmp_path,spec,key)
    (tmp_path/'source.txt').write_text('changed synthetic source')
    second=audit(tmp_path,spec,key)
    assert first['hashes']['source.txt'] != second['hashes']['source.txt']
    (tmp_path/'memo.docx').unlink()
    assert 'audit_incomplete' in codes(audit(tmp_path,spec,key))


def test_headers_and_missing_required_fields_fail_closed(tmp_path):
    spec,key=fixture(tmp_path)
    doc=Document(tmp_path/'memo.docx')
    doc.sections[0].header.paragraphs[0].text='Unbound header 999'
    doc.save(tmp_path/'memo.docx')
    assert 'memo_number_unbound' in codes(audit(tmp_path,spec,key))
    del key['checks']
    assert 'audit_incomplete' in codes(audit(tmp_path,spec,key))


def test_independent_sum_check_catches_false_cash_even_if_output_key_agrees(tmp_path):
    spec,key=fixture(tmp_path)
    book=load_workbook(tmp_path/'model.xlsx')
    book['Model']['C1']='=A3+B3+5'
    book.save(tmp_path/'model.xlsx')
    for expected in key['expected_outputs']:
        if expected['id']=='cash':
            expected['value']=75
    result=audit(tmp_path,spec,key)
    assert 'arithmetic_check_failed' in codes(result)
