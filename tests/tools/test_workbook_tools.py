import hashlib
import json

from openpyxl import load_workbook
import pytest

from agents.tools import WorkbookTools


def test_anchor_paste_uses_exact_agent_matrix_and_reports_extent(tmp_path):
    tools=WorkbookTools(tmp_path)
    result=tools.write_block('Sheet','B3',[[1,2,'=B3+C3'],[3,4,'=B4+C4']],comment='Synthetic source')
    assert result['ok'] and result['result']['range']=='B3:D4'
    assert tools.book['Sheet']['D3'].value=='=B3+C3'
    before=(tmp_path/'model.xlsx').read_bytes()
    assert not tools.write_block('Sheet','B3:C3',[[9]])['ok']
    assert not tools.write_block('Sheet','A1',[[1],[2,3]])['ok']
    assert (tmp_path/'model.xlsx').read_bytes()==before


def test_session(tmp_path):
    tools = WorkbookTools(tmp_path)
    assert tools.add_sheet('Historicals')['ok']
    assert tools.write('Historicals', 'A1:C2', [[1, 2, '=SUM(A1:B1)'], [3, 4, '=C1+A2*B2']])['ok']
    assert tools.name('ReportedInputs', 'Historicals', 'A1:B2')['ok']
    assert tools.comment('Historicals', 'A1:B2', 'synthetic fixture, page 1')['ok']
    assert tools.style('Historicals', 'A1:B2', font={'color': '0000FF'}, number_format='0.00')['ok']
    result = tools.recalc()
    assert result['result']['qualified'] is False
    assert result['result']['cache_written'] is False
    assert result['result']['findings'] == []
    assert {v['cell']: v['value'] for v in result['result']['values']} == {'C1': 3, 'C2': 15}
    assert tools.trace('Historicals', 'C1')['result']['dependencies'] == [{'sheet': 'Historicals', 'range': 'A1:B1'}]
    assert tools.read('Historicals', 'A1')['result']['cells'][0][0]['comment'].startswith('synthetic')
    exported = load_workbook(tmp_path / 'model.xlsx')
    assert exported['Historicals']['C1'].value == '=SUM(A1:B1)'
    assert exported['Historicals']['A1'].font.color.rgb == '000000FF'
    records = [json.loads(line) for line in (tmp_path / 'workbook.trajectory.jsonl').read_text().splitlines()]
    assert all(set(row) == {'version', 'ok', 'tool', 'result', 'errors'} for row in [r['observation'] for r in records])
    assert all(row['version'] == 'tools/v1' for row in [r['observation'] for r in records])


def test_errors_and_hidden_sheets(tmp_path):
    tools = WorkbookTools(tmp_path)
    tools.add_sheet('Hidden inputs', hidden=True)
    tools.write('Hidden inputs', 'A1:A5', [['=A2'], ['=A1'], ['=NOPE(1)'], ['=1/0'], ['=#REF!']])
    report = tools.recalc()['result']
    assert len(report['findings']) == 5
    assert not report['complete']
    assert not tools.recalc(backend='libreoffice')['ok']
    assert not tools.write('Sheet', 'A1', [['=[external.xlsx]Sheet!A1']])['ok']
    assert not tools.write('Sheet', 'A1:B1', [[1]])['ok']
    assert tools.read('Sheet', 'A1')['result']['cells'][0][0]['value'] is None


def test_source_immutable_and_sandbox(tmp_path):
    original = WorkbookTools(tmp_path, 'original.xlsx')
    original.write('Sheet', 'A1', [[42]])
    digest = hashlib.sha256((tmp_path / 'original.xlsx').read_bytes()).hexdigest()
    derivative = WorkbookTools(tmp_path, 'derivative.xlsx', source='original.xlsx')
    derivative.write('Sheet', 'A1', [[43]])
    assert hashlib.sha256((tmp_path / 'original.xlsx').read_bytes()).hexdigest() == digest
    with pytest.raises(ValueError):
        WorkbookTools(tmp_path, '../escape.xlsx')
    with pytest.raises(FileExistsError):
        WorkbookTools(tmp_path, 'original.xlsx')
    (tmp_path / 'link').symlink_to(tmp_path.parent, target_is_directory=True)
    with pytest.raises(ValueError):
        WorkbookTools(tmp_path, 'link/escape.xlsx')


def test_cross_sheet_and_budgets(tmp_path):
    tools = WorkbookTools(tmp_path, max_cells=50)
    tools.add_sheet('Other sheet')
    tools.write('Other sheet', 'A1:B1', [[5, 6]])
    tools.write('Sheet', 'A1', [["=SUM('Other sheet'!A1:B1)"]])
    assert tools.recalc()['result']['values'][0]['value'] == 11
    assert not tools.read('Sheet', 'A1:Z100')['ok']
    assert not tools.trace('Sheet', 'A1:A2')['ok']


def test_rejected_nonfinite_and_illegal_text_are_atomic(tmp_path):
    tools = WorkbookTools(tmp_path)
    tools.write('Sheet','A1:B1',[[1,2]])
    baseline = (tmp_path / 'model.xlsx').read_bytes()
    result = tools.write('Sheet','A1:B1',[[10,float('nan')]])
    assert not result['ok']
    record = json.loads((tmp_path / 'workbook.trajectory.jsonl').read_text().splitlines()[-1])
    assert record['call']['args'][2][0][1] == {'invalid_number': 'nan'}
    assert not tools.write('Sheet','A1:B1',[[10,'\x00illegal']])['ok']
    assert tools.read('Sheet','A1:B1')['result']['cells'][0][0]['value'] == 1
    assert (tmp_path / 'model.xlsx').read_bytes() == baseline


def test_logging_failure_is_structured_and_blocks_mutation(tmp_path):
    tools = WorkbookTools(tmp_path)
    baseline = (tmp_path / 'model.xlsx').read_bytes()
    (tmp_path / 'workbook.trajectory.jsonl').mkdir()
    result = tools.write('Sheet','A1',[[10]])
    assert not result['ok']
    assert any(e['code'] == 'TrajectoryWriteFailed' for e in result['errors'])
    assert tools.book['Sheet']['A1'].value is None
    assert (tmp_path / 'model.xlsx').read_bytes() == baseline


def test_source_size_and_cell_budget_and_external_formula(tmp_path):
    tools = WorkbookTools(tmp_path)
    tools.write('Sheet','A1:B1',[[1,2]])
    with pytest.raises(ValueError, match='byte budget|bounded'):
        WorkbookTools(tmp_path,'small.xlsx',source='model.xlsx',max_bytes=10)
    with pytest.raises(ValueError, match='cell budget'):
        WorkbookTools(tmp_path,'limited.xlsx',source='model.xlsx',max_cells=1)
    # Source formulas are untrusted too, not just tool writes.
    tools.book['Sheet']['A1'] = "='[outside.xlsx]Sheet'!B1"
    findings = tools.recalc()['result']['findings']
    assert findings[0]['code'] == 'external_formula_unsupported'


def test_save_failure_restores_in_memory_state(tmp_path, monkeypatch):
    tools = WorkbookTools(tmp_path)
    def fail():
        raise OSError('synthetic storage failure')
    monkeypatch.setattr(tools,'_save',fail)
    assert not tools.add_sheet('Uncommitted')['ok']
    assert 'Uncommitted' not in tools.book.sheetnames


def test_lazy_if_reference_evaluation_preserves_unselected_errors(tmp_path):
    tools = WorkbookTools(tmp_path)
    tools.write('Sheet','A1:D1',[[1,7,'=1/0','=IF(A1>0,B1,C1)']])
    report = tools.recalc()['result']
    assert next(v['value'] for v in report['values'] if v['cell']=='D1') == 7
    assert {f['cell'] for f in report['findings']} == {'C1'}
    tools.write('Sheet','A2:D2',[[0,'=B2',8,'=IF(A2>0,B2,C2)']])
    report = tools.recalc()['result']
    assert next(v['value'] for v in report['values'] if v['cell']=='D2') == 8
    assert 'D2' not in {f['cell'] for f in report['findings']}


def test_power_rejected_and_intermediates_bounded(tmp_path):
    tools = WorkbookTools(tmp_path)
    tools.write('Sheet','A1:A5',[['=-2^2'],['=2^3^2'],['=2**16'],['=1e100*10/10'],['=SUM(1e100,1e100,-1e100)']])
    report = tools.recalc()['result']
    assert [f['code'] for f in report['findings'][:3]] == ['power_precedence_unqualified']*3
    assert all(f['code']=='numeric_magnitude_budget_exceeded' for f in report['findings'][3:])


def test_ast_budget_and_extended_arithmetic_still_supported(tmp_path):
    tools = WorkbookTools(tmp_path)
    tools.write('Sheet','A1:A5',[['=MAX(1,2,3)-MIN(1,2,3)'],['=ROUND(1.25,1)'],['=ABS(-3)'],['='+'+'.join(['1']*300)],['=-2*3']])
    report = tools.recalc()['result']
    values = {v['cell']:v['value'] for v in report['values']}
    assert (values['A1'],values['A2'],values['A3'],values['A5']) == (2,1.3,3,-6)
    assert report['findings'][0]['cell']=='A4'
    assert report['findings'][0]['code']=='formula_ast_budget_exceeded'


def test_scientific_literals_negative_refs_and_array_arithmetic_bounds(tmp_path):
    tools = WorkbookTools(tmp_path)
    tools.write('Sheet','A1:C2',[[-2,3,'=A1*A1'],[1,4,'=IF(A1<0,MAX(B1:B2),C2)']])
    tools.write('Sheet','D1:D3',[['=1E2+1'],['=A1:A2*1000000000'],['=1E100*10']])
    report = tools.recalc()['result']
    values = {v['cell']:v['value'] for v in report['values']}
    assert values['C1'] == values['C2'] == 4
    assert values['D1'] == 101
    assert {f['cell']:f['code'] for f in report['findings']} == {'D2':'numeric_magnitude_budget_exceeded','D3':'numeric_magnitude_budget_exceeded'}


def test_rollback_survives_comments_after_failed_mutation(tmp_path):
    """Regression: deepcopy rollback broke once a cell comment existed (B2 defect)."""
    tools = WorkbookTools(tmp_path)
    assert tools.add_sheet('Hist')['ok']
    assert tools.write_block('Hist', 'A1', [[1, 2], [3, 4]], comment='Synthetic source')['ok']
    ragged = tools.write_block('Hist', 'D1', [[1, 2], [3]])
    assert not ragged['ok']
    # The failed call must leave the earlier content intact and the book usable.
    assert tools.add_sheet('Next')['ok']
    assert tools.write_block('Next', 'A1', [[5]], comment='Second block')['ok']
    cells = tools.read('Hist', 'A1:B2')['result']['cells']
    assert [[c['value'] for c in row] for row in cells] == [[1, 2], [3, 4]]
