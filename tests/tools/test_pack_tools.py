from io import BytesIO
import json
import hashlib
from zipfile import ZipFile

from agents.tools.pack import PackTools
from agents.tools.export_manifest import verify_export


REF = [{'workbook_sha256': 'a'*64, 'sheet': 'Synthetic', 'cell': 'A1'}]


def test_native_draft_chart_embeds_data_and_preserves_literal_text(tmp_path):
    tools = PackTools(tmp_path)
    assert tools.add_slide('chart')['ok']
    assert tools.title(1, 'Agent supplied title', REF)['ok']
    assert tools.chart(1, ['A', 'B'], [{'name': 'Observed', 'values': [2, 4]}], 'bar', REF)['ok']
    assert tools.write('Agent supplied section', 'Literal authored text.', REF)['ok']
    result = tools.export('v1')
    assert result['ok'] and result['result']['release'] == 'blocked'
    with ZipFile(tmp_path / 'v1.pack/v1.pptx') as archive:
        embedded = [name for name in archive.namelist() if name.startswith('ppt/embeddings/')]
        assert len(embedded) == 1
        with ZipFile(BytesIO(archive.read(embedded[0]))) as workbook:
            assert 'xl/workbook.xml' in workbook.namelist()
        assert b'Agent supplied title' in archive.read('ppt/slides/slide1.xml')
        assert b'TargetMode="External"' not in archive.read('ppt/charts/_rels/chart1.xml.rels')
    assert not tools.export('v1')['ok']
    assert json.loads((tmp_path / 'v1.pack/v1.bindings.json').read_text())['visual_validation'] == 'not_run'


def test_missing_lineage_bad_paths_and_unsupported_preview_block(tmp_path):
    tools = PackTools(tmp_path)
    tools.add_slide()
    assert not tools.title(1, 'Title', [])['ok']
    assert not tools.export('../escape')['ok']
    assert not tools.render_preview()['ok']
    assert not tools.chart(1, ['A'], [{'name': 'X', 'values': [float('nan')]}], 'bar', REF)['ok']
    assert not tools.chart(1, ['A'], [{'name': 'X', 'values': [1]}], 'waterfall', REF)['ok']


def test_failed_export_never_publishes_partial_revision(tmp_path, monkeypatch):
    tools = PackTools(tmp_path)
    tools.add_slide()
    def fail(*args, **kwargs):
        raise OSError('synthetic save failure')
    monkeypatch.setattr(tools.memo, 'save', fail)
    assert not tools.export('broken')['ok']
    assert not (tmp_path / 'broken.pack').exists()
    stages = list(tmp_path.glob('.broken-staging-*'))
    assert len(stages) == 1
    assert json.loads((stages[0] / 'failure.json').read_text())['status'] == 'failed'


def test_workbook_change_blocks_export_and_manifest_detects_tamper(tmp_path):
    from agents.tools import WorkbookTools
    workbook = WorkbookTools(tmp_path)
    workbook.write('Sheet', 'A1', [[1]])
    refs = [{'workbook_sha256': hashlib.sha256((tmp_path / 'model.xlsx').read_bytes()).hexdigest(), 'sheet': 'Sheet', 'cell': 'A1'}]
    tools = PackTools(tmp_path, workbook='model.xlsx')
    tools.add_slide()
    assert tools.title(1, 'Synthetic', refs)['ok']
    exported = tools.export('good')['result']
    assert verify_export(tmp_path, exported['manifest'], expected_sha256=exported['manifest_sha256'])['release'] == 'blocked'
    (tmp_path / 'good.pack/good.pptx').write_bytes(b'tampered')
    import pytest
    with pytest.raises(ValueError):
        verify_export(tmp_path, exported['manifest'], expected_sha256=exported['manifest_sha256'])
    workbook.write('Sheet', 'A1', [[2]])
    assert not tools.export('changed')['ok']
    assert not (tmp_path / 'changed.pack').exists()


def test_missing_audit_path_prevents_mutation(tmp_path):
    (tmp_path / 'pack.trajectory.jsonl').mkdir()
    tools = PackTools(tmp_path)
    result = tools.add_slide()
    assert not result['ok']
    assert len(tools.deck.slides) == 0


def test_memo_only_export_requires_authored_sections_and_has_no_deck(tmp_path):
    tools = PackTools(tmp_path)
    assert not tools.export_memo('empty')['ok']
    assert tools.write('Synthetic evidence', 'An explicitly authored test section.', REF)['ok']
    result = tools.export_memo('authored')
    assert result['ok']
    assert {p.suffix for p in (tmp_path / 'authored.memo').iterdir()} == {'.docx', '.json'}
    receipt = result['result']
    assert verify_export(tmp_path, receipt['manifest'], expected_sha256=receipt['manifest_sha256'])['release'] == 'blocked'


def test_memo_number_insertion_uses_actual_formula_cells_and_exact_spans(tmp_path):
    from agents.tools import WorkbookTools
    from docx import Document
    workbook = WorkbookTools(tmp_path)
    workbook.write('Sheet', 'A1:B1', [[12.5, '=A1*2']])
    tools = PackTools(tmp_path, workbook='model.xlsx')
    result = tools.write_from_cells('Source discussion', 'The result is {{total}} USDk.',
        [{'key':'total','sheet':'Sheet','cell':'B1','unit':'USDk','decimals':2}])
    assert result['ok']
    assert tools.memo.paragraphs[1].text == 'The result is 25.00 USDk.'
    span = result['result']['memo_bindings'][0]
    assert tools.memo.paragraphs[span['paragraph']].text[span['start']:span['end']] == '25.00'
    assert (tmp_path/'model.xlsx').read_bytes()
    assert not tools.write_from_cells('Bad', '{{absent}}', [{'key':'total','sheet':'Sheet','cell':'B1','unit':'USDk','decimals':2}])['ok']
    assert tools.write_from_cells('Period result','{{historical.2025.total}} USDk.',
        [{'key':'historical.2025.total','sheet':'Sheet','cell':'B1','unit':'USDk','decimals':2}])['ok']
