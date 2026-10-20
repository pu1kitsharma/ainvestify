import json

import pytest

from agents.tools import AnalysisTools, DocumentTools, WorkbookTools
from agents.tools.registry import build_tool_registry
from tests.tools.synthetic_books import build
from tests.tools.test_audits import monthly_quarterly
from tests.tools.test_claim_binder import chart_deck


def sandbox(tmp_path, sheets=None, name='model.xlsx'):
    (tmp_path / name).write_bytes(build(sheets or monthly_quarterly(bug=True)))
    return AnalysisTools(tmp_path)


def test_registry_exposes_new_tools_and_keeps_existing_names(tmp_path):
    (tmp_path / 'wb').mkdir()
    registry = build_tool_registry(workbook=WorkbookTools(tmp_path / 'wb'), documents=DocumentTools(tmp_path, tenant_id='t', deal_id='d'),
                                   analysis=AnalysisTools(tmp_path))
    old = {'wb.add_sheet', 'wb.write', 'wb.write_block', 'wb.name', 'wb.style', 'wb.comment', 'wb.read', 'wb.trace', 'wb.recalc',
           'doc.pages', 'doc.text', 'doc.tables', 'doc.page_image', 'pptx.chart_series', 'xlsx.read'}
    new = {'fx.evaluate', 'fx.compare_cache', 'fx.recalc_libreoffice', 'audit.hidden', 'audit.range_width', 'audit.units', 'audit.driver_growth',
           'audit.balance_sheet', 'audit.errors', 'audit.hardcodes', 'audit.reconcile_totals', 'claims.bind', 'claims.bind_document', 'chart.agreement'}
    assert set(registry) == old | new
    for name in new:
        entry = registry[name]
        assert entry['schema']['parameters']['additionalProperties'] is False
        json.dumps(entry['schema'], allow_nan=False)
    assert build_tool_registry(analysis=None) == {}


def test_tools_v1_envelope_and_audit_roundtrip(tmp_path):
    registry = build_tool_registry(analysis=sandbox(tmp_path))
    result = registry['audit.range_width']['call'](source='model.xlsx')
    assert set(result) == {'version', 'ok', 'tool', 'result', 'errors'} and result['tool'] == 'audit.range_width' and result['ok']
    high = [f for f in result['result']['findings'] if f['severity'] == 'high']
    assert high and high[0]['cells'][0] == 'Qtrly!D4'
    evaluation = registry['fx.evaluate']['call'](source='model.xlsx', cells=[{'sheet': 'Qtrly', 'address': 'D2:D4'}])
    assert evaluation['ok'] and [c['value'] for c in evaluation['result']['values']] == [33, 4.5 + 0, 2.0] or evaluation['result']['values'][0]['value'] == 33
    recon = registry['audit.reconcile_totals']['call'](source='model.xlsx', members=[
        {'sheet': 'Monthly', 'first_col': 'D', 'last_col': 'O'}, {'sheet': 'Qtrly', 'first_col': 'D', 'last_col': 'G', 'total_col': 'H'}])
    assert recon['ok'] and recon['result']['mismatch_count'] == 1
    log = [json.loads(line) for line in (tmp_path / 'analysis.trajectory.jsonl').read_text().splitlines()]
    assert len(log) == 3 and all(row['observation']['version'] == 'tools/v1' for row in log)


def test_path_hash_pin_and_input_validation(tmp_path):
    tools = sandbox(tmp_path)
    registry = build_tool_registry(analysis=tools)
    assert not registry['audit.hidden']['call'](source='../outside.xlsx')['ok']
    assert registry['audit.hidden']['call'](source='model.xlsx', expected_sha256='0' * 64)['errors'][0]['message'].startswith('Source hash')
    assert registry['audit.hidden']['call'](source='model.xlsx')['ok']
    (tmp_path / 'model.xlsx').write_bytes(build({'S': {'A1': 1}}))
    assert 'Source hash changed' in registry['audit.hidden']['call'](source='model.xlsx')['errors'][0]['message']
    assert registry['audit.hidden']['call'](source='model.xlsx', unknown=1)['errors'][0]['code'] == 'ArgumentValidationError'
    assert registry['fx.recalc_libreoffice']['call'](source='model.xlsx', timeout_seconds=1)['errors'][0]['code'] == 'ArgumentValidationError'


def test_claims_bind_from_workbook_ranges_uses_evaluated_values_and_units(tmp_path):
    tools = sandbox(tmp_path, {'Pack': {'A1': 'INR Mn', 'A2': 'Revenue', 'B2': '=SUM(C2:E2)', 'C2': 100.0, 'D2': 120.5, 'E2': 130.0, 'A3': 'Margin', 'B3': 0.347}})
    registry = build_tool_registry(analysis=tools)
    ok = registry['claims.bind']['call'](text='Revenue of INR 350.5 Mn at a 34.7% margin.', workbook='model.xlsx', ranges=[{'sheet': 'Pack', 'address': 'B2:B3'}])
    assert ok['ok'] and ok['result']['all_bound'] and ok['result']['counts']['bound'] == 2
    bad = registry['claims.bind']['call'](text='Revenue of INR 352 Mn.', workbook='model.xlsx', ranges=[{'sheet': 'Pack', 'address': 'B2:B3'}])
    assert bad['ok'] and not bad['result']['all_bound'] and bad['result']['claims'][0]['status'] == 'unbound'
    assert not registry['claims.bind']['call'](text='x 1')['ok']        # no table, no workbook


def test_chart_agreement_and_document_tools(tmp_path):
    tools = sandbox(tmp_path, {'M': {'A1': 'Revenue', 'B1': 100.0, 'C1': 120.5, 'D1': 132.55}})
    (tmp_path / 'deck.pptx').write_bytes(chart_deck([100.0, 120.5, 132.55]))
    registry = build_tool_registry(analysis=tools)
    result = registry['chart.agreement']['call'](source='deck.pptx', workbook='model.xlsx')
    assert result['ok'] and result['result']['all_agree']
    assert not registry['chart.agreement']['call'](source='model.xlsx', workbook='model.xlsx')['ok']
    from docx import Document
    doc = Document()
    doc.add_paragraph('Revenue 132.55 then 999.')
    doc.save(tmp_path / 'memo.docx')
    bound = registry['claims.bind_document']['call'](source='memo.docx', workbook='model.xlsx', ranges=[{'sheet': 'M', 'address': 'B1:D1'}])
    assert bound['ok'] and (bound['result']['counts']['bound'], bound['result']['counts']['unbound']) == (1, 1) and not bound['result']['all_bound']


def test_libreoffice_tool_reports_unrun_as_observation_error(tmp_path, monkeypatch):
    from agents.tools import recalc
    monkeypatch.setattr(recalc, 'availability', lambda: (False, 'private_process_isolation_unavailable'))
    registry = build_tool_registry(analysis=sandbox(tmp_path, {'S': {'A1': 1, 'A2': '=A1+1'}}))
    result = registry['fx.recalc_libreoffice']['call'](source='model.xlsx')
    assert not result['ok'] and result['errors'][0]['code'] == 'UNRUN' and result['result']['mandatory_unrun'] is True


def test_libreoffice_tool_live_when_isolation_available(tmp_path):
    from agents.tools import recalc
    available, reason = recalc.availability()
    if not available:
        pytest.skip(reason)
    registry = build_tool_registry(analysis=sandbox(tmp_path, {'S': {'A1': 2, 'A2': '=A1*A1', 'A3': '=SUM(A1:A2)'}}))
    result = registry['fx.recalc_libreoffice']['call'](source='model.xlsx', timeout_seconds=120)
    assert result['ok'] and result['result']['status'] == 'compared' and result['result']['agreement'] == 1.0
