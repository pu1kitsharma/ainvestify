from hashlib import sha256

from openpyxl import Workbook
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches
from reportlab.pdfgen import canvas

from agents.tools.documents import DocumentTools


def tools(root, **kwargs):
    return DocumentTools(root, tenant_id='synthetic-tenant', deal_id='synthetic-deal', **kwargs)


def test_pptx_native_text_table_chart(tmp_path):
    p = Presentation()
    slide = p.slides.add_slide(p.slide_layouts[6])
    slide.shapes.add_textbox(0,0, Inches(5), Inches(1)).text = 'Synthetic title'
    table = slide.shapes.add_table(1,2,0,Inches(1),Inches(5),Inches(1)).table
    table.cell(0,0).text, table.cell(0,1).text = 'Metric', '7'
    data = CategoryChartData()
    data.categories = ['A','B']
    data.add_series('Series', [1,2])
    slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, 0, Inches(2), Inches(5), Inches(3), data)
    path = tmp_path / 'sample.pptx'
    p.save(path)
    t = tools(tmp_path)
    assert t.pages('sample.pptx')['result']['count'] == 1
    texts = t.text('sample.pptx',1)
    assert texts['ok']
    assert any(u['text'] == 'Synthetic title' for u in texts['result']['units'])
    tables = t.tables('sample.pptx',1)['result']['tables']
    assert tables[0]['rows'] == [['Metric', '7']]
    charts = t.chart_series('sample.pptx',1)['result']
    assert charts['units']
    assert charts['external_relationships_fetched'] is False
    digest = sha256(path.read_bytes()).hexdigest()
    assert all(u['source_id'] == digest for u in charts['units'])
    assert not t.text('sample.pptx',2)['ok']
    assert not t.page_image('sample.pptx',1)['ok']


def test_pdf(tmp_path):
    path = tmp_path / 'sample.pdf'
    c = canvas.Canvas(str(path))
    c.drawString(20,800,'Synthetic PDF 123')
    c.showPage()
    c.save()
    t = tools(tmp_path)
    assert t.pages('sample.pdf')['result']['count'] == 1
    assert any('Synthetic PDF 123' in u['text'] for u in t.text('sample.pdf',1)['result']['units'])
    assert t.tables('sample.pdf',1)['result']['tables'] == []
    assert not t.chart_series('sample.pdf',1)['ok']


def test_xlsx_hash_and_bounds(tmp_path):
    book = Workbook()
    book.active['A1'], book.active['B1'] = 7, '=A1*2'
    book.save(tmp_path / 'model.xlsx')
    t = tools(tmp_path)
    result = t.xlsx_read('model.xlsx', 'Sheet', 'A1:C1')
    assert result['ok']
    cells = result['result']['cells'][0]
    assert cells[2]['value'] is None
    assert cells[2]['cell'] == 'Sheet!C1'
    assert cells[1]['value'] == '=A1*2'
    assert cells[1]['cached_value'] is None
    assert not t.xlsx_read('model.xlsx','Sheet','A1:Z1000')['ok']
    assert not t.xlsx_read('../model.xlsx','Sheet','A1')['ok']
    assert not t.xlsx_read('model.xlsx','Sheet','A1', expected_sha256='0'*64)['ok']
    book.active['A1'] = 8
    book.save(tmp_path / 'model.xlsx')
    assert not t.xlsx_read('model.xlsx','Sheet','A1')['ok']


def test_output_and_file_budgets(tmp_path):
    path = tmp_path / 'sample.pdf'
    c = canvas.Canvas(str(path))
    c.drawString(20,800,'Synthetic PDF 123')
    c.showPage()
    c.save()
    assert not tools(tmp_path, max_bytes=10).pages('sample.pdf')['ok']
    assert not tools(tmp_path, max_output_chars=10).text('sample.pdf',1)['ok']


def test_audit_failure_and_source_cell_budget(tmp_path):
    book = Workbook()
    book.active['A1'],book.active['B1'] = 1,2
    book.save(tmp_path / 'model.xlsx')
    assert not tools(tmp_path,max_cells=1).xlsx_read('model.xlsx','Sheet','A1')['ok']
    (tmp_path / 'documents.trajectory.jsonl').unlink()
    (tmp_path / 'documents.trajectory.jsonl').mkdir()
    result = tools(tmp_path).xlsx_read('model.xlsx','Sheet','A1')
    assert not result['ok']
    assert result['errors'][-1]['code'] == 'TrajectoryWriteFailed'
