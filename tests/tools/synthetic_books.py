"""Synthetic workbook builders for the analyst tool-layer tests (no company data)."""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.worksheet.table import Table


def build(sheets, *, hidden=(), names=None, hidden_rows=None, hidden_cols=None, tables=None):
    """sheets: {title: {address: value}}; returns XLSX bytes (no stored caches)."""
    book = Workbook()
    book.remove(book.active)
    for title, cells in sheets.items():
        ws = book.create_sheet(title)
        for address, value in cells.items():
            ws[address] = value
        if title in hidden:
            ws.sheet_state = 'hidden'
        for row in (hidden_rows or {}).get(title, ()):
            ws.row_dimensions[row].hidden = True
        for col in (hidden_cols or {}).get(title, ()):
            ws.column_dimensions[col].hidden = True
    for name, ref in (names or {}).items():
        from openpyxl.workbook.defined_name import DefinedName
        book.defined_names[name] = DefinedName(name, attr_text=ref)
    for sheet, (name, ref) in (tables or {}).items():
        book[sheet].add_table(Table(displayName=name, ref=ref))
    out = BytesIO()
    book.save(out)
    return out.getvalue()
