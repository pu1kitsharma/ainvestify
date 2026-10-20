from io import BytesIO

import pytest
from docx import Document
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

from agents.tools import claim_binder as cb, formula_engine as fx
from tests.tools.synthetic_books import build

CELLS = [
    {'ref': 'M!B2', 'value': 1234.5, 'unit': 'INR Mn', 'label': 'Revenue FY24'},
    {'ref': 'M!B3', 'value': 0.347, 'unit': None, 'label': 'EBITDA margin', 'kind': 'share'},
    {'ref': 'M!B4', 'value': -0.082, 'unit': None, 'label': 'Revenue growth', 'kind': 'growth'},
    {'ref': 'M!B5', 'value': 3.2, 'unit': 'x', 'label': 'LTV to CAC'},
    {'ref': 'M!B6', 'value': 0.23, 'unit': None, 'label': 'Revenue CAGR', 'kind': 'cagr'},
    {'ref': 'M!B7', 'value': 2500000, 'unit': 'INR', 'label': 'Cash balance'},
]


def statuses(text, cells=CELLS, **kw):
    report = cb.bind_claims(text, cells, **kw)
    return {c['text']: c['status'] for c in report['claims']}, report


def test_bound_numbers_with_scale_percent_and_multiple():
    got, report = statuses('Revenue reached INR 1,234.5 Mn in FY24 with a 34.7% EBITDA margin and 3.2x LTV to CAC; cash of INR 2.5 Mn.')
    assert got == {'INR 1,234.5 Mn': 'bound', '34.7%': 'bound', '3.2x': 'bound', 'INR 2.5 Mn': 'bound'}
    assert report['all_bound']


def test_rounding_is_enforced_at_displayed_precision():
    got, _ = statuses('Margin was 35% and then 34.8% and 34.6%.')
    assert got == {'35%': 'bound', '34.8%': 'unbound', '34.6%': 'unbound'}


def test_unbound_number_reports_nearest_cells():
    _, report = statuses('Revenue was INR 1,300 Mn.')
    claim = report['claims'][0]
    assert claim['status'] == 'unbound' and claim['nearest'][0]['ref'] == 'M!B2'
    assert not report['all_bound']


def test_scale_mismatch_is_flagged_not_silently_bound():
    cells = [{'ref': 'M!C1', 'value': 1234, 'unit': 'INR K', 'label': 'Revenue'}]
    got, report = statuses('Revenue was INR 1,234 Mn.', cells)
    assert got == {'INR 1,234 Mn': 'mismatch'} and 'scale_mismatch' in report['claims'][0]['notes'][0]
    assert statuses('Revenue was INR 1.234 Mn.', cells)[0] == {'INR 1.234 Mn': 'bound'}


def test_direction_sign_and_cagr_checks():
    got, report = statuses('Revenue increased 8.2% year on year.')
    assert got == {'8.2%': 'mismatch'} and 'direction_conflict' in report['claims'][0]['notes'][0]
    assert statuses('Revenue declined 8.2% year on year.')[0] == {'8.2%': 'bound'}
    assert statuses('Revenue fell (8.2%) in FY24.')[0] == {'(8.2%)': 'bound'}
    got, report = statuses('Revenue CAGR of 23% over the period.')
    assert got == {'23%': 'bound'}
    got, report = statuses('CAGR of 34.7% was achieved.')
    assert got == {'34.7%': 'mismatch'} and 'comparative_kind_mismatch' in report['claims'][0]['notes'][0]
    got, report = statuses('Explicitly negative -34.7% margin.')
    assert got == {'-34.7%': 'mismatch'} and 'sign_mismatch' in report['claims'][0]['notes'][0]


def test_ignored_tokens_years_labels_and_allowlist():
    text = 'In FY24 and Q3 2023, per 2021-22 data and the 5G plan, 12 pillars include 3D. Revenue 1,234.5 Mn'
    report = cb.bind_claims(text, CELLS, ignore_numbers=[12])
    statuses_by_text = {c['text']: c['status'] for c in report['claims']}
    assert statuses_by_text['2023'] == 'ignored' and statuses_by_text['12'] == 'ignored'
    assert '24' not in statuses_by_text and '3' not in statuses_by_text and '22' not in statuses_by_text and '5' not in statuses_by_text
    assert statuses_by_text['1,234.5 Mn'] == 'bound'


def test_derived_growth_cagr_and_ratio_from_series():
    cells = [{'ref': 'M!B2', 'value': 100, 'series': 'rev', 'period': 2022, 'label': 'Revenue', 'unit': 'INR Mn'},
             {'ref': 'M!C2', 'value': 121, 'series': 'rev', 'period': 2024, 'label': 'Revenue', 'unit': 'INR Mn'}]
    got, report = statuses('Revenue CAGR was 10.0%, up 21% in total, a 1.21x multiple.', cells)
    assert got == {'10.0%': 'bound', '21%': 'bound', '1.21x': 'bound'}
    assert report['claims'][0]['binding']['ref'].startswith('derived:cagr')
    assert statuses('Revenue CAGR was 10.0%.', cells, derive=False)[0] == {'10.0%': 'unbound'}


def test_document_binding_docx_and_pptx():
    doc = Document()
    doc.add_paragraph('Revenue was INR 1,234.5 Mn and margin 34.7%.')
    doc.add_paragraph('Churn was 9.9%.')
    buf = BytesIO()
    doc.save(buf)
    report = cb.bind_document(buf.getvalue(), 'docx', CELLS)
    assert report['counts']['bound'] == 2 and report['counts']['unbound'] == 1 and not report['all_bound']
    assert report['units_needing_attention'][0]['locator'] == 'paragraph 2'
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = 'EBITDA margin of 34.7% on INR 1,234.5 Mn revenue'
    out = BytesIO()
    deck.save(out)
    assert cb.bind_document(out.getvalue(), 'pptx', CELLS)['all_bound']


def chart_deck(values):
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    data = CategoryChartData()
    data.categories = ['FY22', 'FY23', 'FY24']
    data.add_series('Revenue', values)
    slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1), Inches(6), Inches(4), data)
    out = BytesIO()
    deck.save(out)
    return out.getvalue()


def test_chart_workbook_agreement_found_rescaled_and_missing():
    ev = fx.evaluate_workbook_bytes(build({'M': {'A1': 'Revenue', 'B1': 100.0, 'C1': 120.5, 'D1': '=C1*1.1', 'A3': 'Rev K', 'B3': 100000, 'C3': 120500, 'D3': 132550}}))
    ok = cb.chart_workbook_agreement(chart_deck([100.0, 120.5, 132.55]), ev)
    series = ok['charts'][0]['series'][0]
    assert series['status'] == 'found_in_workbook' and series['matches'][0]['range'] == 'M!B1:D1'
    assert series['cache_vs_embedded_workbook'] == 'agree' and ok['all_agree']
    bad = cb.chart_workbook_agreement(chart_deck([100.0, 120.5, 140.0]), ev)
    assert bad['charts'][0]['series'][0]['status'] == 'not_found_in_workbook' and not bad['all_agree']
    pinned = cb.chart_workbook_agreement(chart_deck([100.0, 120.5, 140.0]), ev, bindings=[
        {'slide': 1, 'chart_index': 1, 'series': [{'name': 'Revenue', 'sheet': 'M', 'address': 'B1:D1'}]}])
    assert pinned['charts'][0]['series'][0]['status'] == 'disagrees_with_pinned_range'
    rounded = cb.chart_workbook_agreement(chart_deck([100.0, 120.5, 132.6]), ev, display_decimals=1)
    assert rounded['charts'][0]['series'][0]['status'] == 'found_in_workbook' or rounded['charts'][0]['series'][0]['matches']
