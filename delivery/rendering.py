"""Bounded draft templates. Text is supplied by recorded model projections."""
from __future__ import annotations
from io import BytesIO
import re


def _pdf_font(*, bold=False):
    """Embed a Unicode font when the qualified macOS runtime supplies one."""
    from pathlib import Path
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    name = 'RoomDejaVuSansBold' if bold else 'RoomDejaVuSans'
    filename = 'DejaVuSans-Bold.ttf' if bold else 'DejaVuSans.ttf'
    path = Path('/Applications/LibreOffice.app/Contents/Resources/fonts/truetype') / filename
    if not path.is_file():
        name = 'RoomArialUnicode'
        path = Path('/System/Library/Fonts/Supplemental/Arial Unicode.ttf')
    if path.is_file():
        if name not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont(name, str(path)))
        return name
    return 'Helvetica-Bold' if bold else 'Helvetica'


EVIDENCE_DOCUMENT_CONTRACT = 'evidence-document-v1'


def render_evidence_document(title, sections, *, kind, format):
    """Opt-in richer draft export; legacy render calls retain their 30-section cap.

    Every value/prose field has already been checked by the upstream immutable
    evidence/model replay. Exporting this view is not investor approval.
    """
    if kind in ('intro_deck', 'pitch_deck') and format in ('pptx', 'pdf'):
        renderer = render_intro if format == 'pptx' else render_deck_pdf
    elif kind == 'investment_memorandum' and format in ('docx', 'pdf'):
        renderer = render_memo if format == 'docx' else render_memo_pdf
    else:
        raise ValueError('Unsupported structured material kind/format')
    return renderer(title, sections, document_contract=EVIDENCE_DOCUMENT_CONTRACT)


def _checked_sections(sections, *, paginated=False, document_contract=None):
    if document_contract not in (None, EVIDENCE_DOCUMENT_CONTRACT):
        raise ValueError('Draft document contract is unsupported')
    limit = 40 if document_contract else 30
    if not sections or len(sections) > limit:
        raise ValueError(f"Supported draft requires 1–{limit} sections")
    for section in sections:
        if len(section) not in (3, 4, 5, 6):
            raise ValueError('Draft section shape is unsupported')
        heading, body, sources = section[:3]
        if any(not isinstance(item, str) for item in (heading, body, sources)):
            raise ValueError('Draft section text must be literal strings')
        if any(len(item) > limit for item, limit in (
                (heading, 300 if paginated else 100),
                (body, 2000 if paginated and document_contract else 1200), (sources, 1200))):
            raise ValueError("Draft section exceeds supported layout; split it before rendering")
        if len(section) == 4 and section[3] not in {
                'statement', 'evidence', 'comparison', 'timeline'}:
            raise ValueError('Draft slide layout is unsupported')
        if len(section) == 5:
            if section[3] != 'table':
                raise ValueError('Structured draft section requires table layout')
            table = _evidence_table(section[4])
            if not paginated and len(table.rows) > 12:
                raise ValueError('Draft slide table exceeds 12 rows; split it before rendering')
            if not paginated and len(body) > 400:
                raise ValueError('Draft table caption exceeds supported layout; split it before rendering')
        if len(section) == 6:
            if section[3] != 'chart' or section[5] != 'evidence-chart-v1':
                raise ValueError('Structured source chart requires its explicit chart contract')
            _evidence_chart(section[4])
            if not paginated and len(body) > 400:
                raise ValueError('Draft chart caption exceeds supported layout; split it before rendering')
    return sections


def _evidence_table(value):
    from delivery.contracts import EvidenceTable
    return value if isinstance(value, EvidenceTable) else EvidenceTable.model_validate(value)


def _evidence_chart(value):
    from delivery.contracts import EvidenceChart
    return value if isinstance(value, EvidenceChart) else EvidenceChart.model_validate(value)


def project_native_chart(manifest, chart_unit_id, *, expected_binding,
                         expected_manifest_digest, source_refs):
    """Copy a complete, mechanically bound native source chart cache.

    Read-only observation of retained OOXML and its separately indexed literal
    units. No embedded workbook executes, no formula is calculated, no data is
    obtained from a chart image, and no financial/unit/period status is inferred.
    Unsupported structures remain blocked instead of being approximated.
    """
    from agents.research.document_evidence_manifest import CONTRACT, SCOPE, digest
    from delivery.contracts import EvidenceChart
    if (manifest.get('contract') != CONTRACT or manifest.get('binding') != expected_binding or
            manifest.get('manifest_digest') != expected_manifest_digest or
            digest({key: value for key, value in manifest.items() if key != 'manifest_digest'}) !=
            expected_manifest_digest or any(manifest.get(key) != value for key, value in SCOPE.items()) or
            expected_binding.get('format') != 'pptx'):
        raise ValueError('Original chart manifest scope/binding/integrity mismatch')
    units = {unit['id']: unit for unit in manifest['units']}
    if len(units) != len(manifest['units']):
        raise ValueError('Ambiguous original chart evidence units')
    original = units.get(chart_unit_id)
    if (original is None or original.get('kind') != 'chart' or
            original.get('source_id') != expected_binding['source_sha256'] or
            original.get('content_sha256') != digest(original.get('content'))):
        raise ValueError('Original chart unit binding mismatch')
    content = original['content']
    if any(rel.get('TargetMode') == 'External' for rel in content['relationships']):
        raise ValueError('External source chart dependencies are unsupported')
    literal_units = {}
    for unit in manifest['units']:
        if unit.get('kind') != 'chart_literal' or unit.get('chart_unit_id') != chart_unit_id:
            continue
        path = unit['source_locator'].get('xml_path')
        if (path in literal_units or unit['source_id'] != expected_binding['source_sha256'] or
                unit['source_locator'].get('chart_part') != content['part'] or
                unit['content_sha256'] != digest(unit['content'])):
            raise ValueError('Original chart literal lineage mismatch')
        literal_units[path] = unit

    def tag(node):
        return node['tag'].rsplit('}', 1)[-1]

    def children(node, path):
        return [(child, path + '/' + tag(child) + '[' + str(index + 1) + ']')
                for index, child in enumerate(node.get('children', []))]

    def direct(node, path, name):
        return [(child, child_path) for child, child_path in children(node, path) if tag(child) == name]

    def one(node, path, name):
        found = direct(node, path, name)
        if len(found) != 1:
            raise ValueError('Missing or ambiguous original source chart ' + name)
        return found[0]

    def literal(node, path, *, numeric=False):
        unit = literal_units.get(path)
        if unit is None or unit['content'] != (node.get('text') or ''):
            raise ValueError('Exact source chart cache literal is not independently bound')
        return {('numeric_literal' if numeric else 'text'): unit['content'],
                'literal_unit_id': unit['id'], 'literal_sha256': unit['content_sha256']}

    def cache(parent, path, *, numeric=False):
        cache_types = ('numCache', 'numLit') if numeric else ('strCache', 'strLit', 'numCache', 'numLit')
        candidates = []
        def visit(node, where):
            if tag(node) in cache_types:
                candidates.append((node, where))
                return
            for child, child_path in children(node, where):
                visit(child, child_path)
        visit(parent, path)
        if len(candidates) != 1:
            raise ValueError('Source chart requires one complete literal cache, not a missing formula result')
        node, where = candidates[0]
        count_node, count_path = one(node, where, 'ptCount')
        count_text = count_node.get('attributes', {}).get('val') or count_node.get('text')
        if not count_text or not count_text.isdigit() or not 1 <= int(count_text) <= 12:
            raise ValueError('Source chart cache exceeds supported complete point count')
        points = direct(node, where, 'pt')
        if len(points) != int(count_text):
            raise ValueError('Sparse or missing original chart cache points are unsupported')
        indexed = {}
        for point, point_path in points:
            index = point.get('attributes', {}).get('idx')
            if index is None or not index.isdigit() or int(index) in indexed:
                raise ValueError('Ambiguous original chart cache indexes')
            value, value_path = one(point, point_path, 'v')
            indexed[int(index)] = literal(value, value_path, numeric=numeric)
        if set(indexed) != set(range(len(points))):
            raise ValueError('Original chart cache indexes must be complete and ordered')
        return [indexed[index] for index in range(len(points))]

    root = content['native_xml']
    root_path = '/' + tag(root)
    chart, chart_path = one(root, root_path, 'chart')
    area, area_path = one(chart, chart_path, 'plotArea')
    plots = [(node, path) for node, path in children(area, area_path)
             if tag(node).endswith('Chart')]
    if len(plots) != 1 or tag(plots[0][0]) not in {'barChart', 'lineChart', 'pieChart'}:
        raise ValueError('Combined, 3D or unsupported native source charts are blocked')
    plot, plot_path = plots[0]
    axes = [(node, path) for node, path in children(area, area_path) if tag(node).endswith('Ax')]
    if tag(plot) == 'pieChart':
        if axes:
            raise ValueError('Unexpected source pie axes')
        plot_type, axis_count = 'pie', 0
    else:
        if len(axes) != 2 or sorted(tag(node) for node, _ in axes) != ['catAx', 'valAx']:
            raise ValueError('Multi-axis/date-axis source charts are unsupported')
        axis_ids = [one(node, path, 'axId')[0]['attributes'].get('val') for node, path in axes]
        plotted_axis_ids = [node['attributes'].get('val') for node, _ in direct(plot, plot_path, 'axId')]
        if len(set(axis_ids)) != 2 or sorted(axis_ids) != sorted(plotted_axis_ids):
            raise ValueError('Original source chart axis binding mismatch')
        axis_count = 2
        if tag(plot) == 'lineChart':
            grouping, _ = one(plot, plot_path, 'grouping')
            if grouping['attributes'].get('val') != 'standard':
                raise ValueError('Stacked or percentage source line charts are unsupported')
            plot_type = 'line'
        else:
            grouping, _ = one(plot, plot_path, 'grouping')
            direction, _ = one(plot, plot_path, 'barDir')
            combination = (direction['attributes'].get('val'), grouping['attributes'].get('val'))
            plot_type = {('col', 'clustered'): 'column_clustered', ('col', 'stacked'): 'column_stacked',
                         ('bar', 'clustered'): 'bar_clustered'}.get(combination)
            if plot_type is None:
                raise ValueError('Percent-stacked or unsupported source bar grouping')
    series_rows = direct(plot, plot_path, 'ser')
    if not 1 <= len(series_rows) <= 3:
        raise ValueError('Source chart exceeds readable complete series count')
    categories, series = None, []
    for native, native_path in series_rows:
        category, category_path = one(native, native_path, 'cat')
        current_categories = cache(category, category_path)
        if categories is not None and [c['text'] for c in categories] != [c['text'] for c in current_categories]:
            raise ValueError('Different source series categories are unsupported')
        if categories is None:
            categories = current_categories
        values, values_path = one(native, native_path, 'val')
        points = cache(values, values_path, numeric=True)
        name_node, name_path = one(native, native_path, 'tx')
        direct_names = direct(name_node, name_path, 'v')
        if direct_names:
            if len(direct_names) != 1:
                raise ValueError('Ambiguous original series name')
            name = literal(*direct_names[0])
        else:
            names = cache(name_node, name_path)
            if len(names) != 1:
                raise ValueError('Original series name must be one literal')
            name = names[0]
        series.append({'name': name, 'points': points})
    payload = {'contract': 'evidence-chart-v1', 'chart_unit_id': chart_unit_id,
        'chart_part': content['part'], 'evidence_revision': expected_manifest_digest,
        'source_sha256': expected_binding['source_sha256'], 'plot_type': plot_type,
        'categories': categories, 'series': series, 'source_refs': list(source_refs),
        'source_status': 'source_cache_unreviewed', 'financial_validation': 'not_established',
        'recalculation': 'not_performed', 'period': 'unassessed', 'unit': 'unassessed',
        'axis_count': axis_count, 'external_dependencies': False}
    payload['cache_digest'] = digest(payload)
    return EvidenceChart.model_validate(payload).model_dump(mode='json')


def _table_cell_text(cell):
    """Preserve exact supplied text, visible source IDs and source status.

    A reported observation is not changed into an actual/reconciled number.
    Projection and illustrative labels remain visible on every affected cell.
    """
    label = {'projection': ' (projection)', 'illustrative': ' (illustrative)',
             'missing': ' (missing)', 'source_reported': ''}[cell.status]
    refs = ' '.join('[' + ref + ']' for ref in cell.source_refs)
    return cell.text + label + (' ' + refs if refs else '')


def _table_text_matrix(table):
    headers = ['\n'.join(item for item in (column.label, column.period, column.unit) if item)
               for column in table.columns]
    return [headers] + [[_table_cell_text(cell) for cell in row.cells] for row in table.rows]


_SOURCE_REGISTRY_LINE = re.compile(r'^\[S[1-9][0-9]*\]\s+[A-Za-z][A-Za-z0-9+.-]*://\S+(?:\s+.*)?$')


def _source_registry_lines(body):
    """Recognize complete reference inventories, never ordinary narrative.

    Registry text stays exact. Only its typographic role changes, allowing
    opaque version/hash URIs to fit without splitting inside their identifiers.
    """
    lines = [line for line in body.splitlines() if line.strip()]
    return lines if lines and all(_SOURCE_REGISTRY_LINE.fullmatch(line) for line in lines) else []


_SENTENCE_END = re.compile(r'(?<=[.!?])["\')\]]*\s+(?=["\'(\[]?[A-Z0-9])')


def _sentences(body):
    """The model's sentences, or the whole body when it is not cleanly two to four.

    Each returned part is exact text; joined by single spaces they are the body.
    """
    parts = _SENTENCE_END.split(body)
    if 2 <= len(parts) <= 4 and ' '.join(parts) == body and all(len(part) >= 20 for part in parts):
        return parts
    return [body]


def _comparison_parts(body):
    """Two columns only where the model wrote two sentences.

    The cut is the sentence end nearest the middle. A colon, a semicolon or a
    space is never a cut: one continuous sentence stays in one column.
    """
    midpoint = len(body) // 2
    candidates = [match.end() for match in _SENTENCE_END.finditer(body)]
    cut = min(candidates, key=lambda value: abs(value-midpoint)) if candidates else -1
    if cut < 20 or cut > len(body)-20:
        return body, ''
    return body[:cut].rstrip(), body[cut:].lstrip()


def _effective_slide_layout(layout, body):
    """A comparison needs two complete sentences; otherwise use a statement panel."""
    if layout == 'comparison' and not _comparison_parts(body)[1]:
        return 'statement'
    return layout


_LEADING = 1.27                         # line pitch, shared by the PPTX and its PDF
# Largest that fits. A one- or two-sentence claim is set as display type; a full
# slide ends at the 12pt floor.
_BODY_SIZES = (40, 34, 30, 26, 22, 19, 16, 14, 12)
_SOURCE_SIZES = (12, 11, 10, 9, 8)      # 8 only when a long list fits no other way
_HEADING_SIZES = (27, 24, 21, 18)
_TOP, _BOTTOM = 2.08, 6.78              # the content band between the two hairlines


def _estimated_lines(text, size, width):
    """A deliberately generous line count for wrapped text; no font file is read,
    so the editable slide and its PDF get the same geometry on every machine.

    Words are wrapped greedily with rough advance widths, so a few long words
    in a narrow column are counted as the lines they really take.
    """
    def ems(word):
        # Capitals, digits and m/w are wider than other lowercase; points are narrow.
        return sum(1.0 if char in 'WMmw@%' else .74 if char.isupper() else .66 if char.isdigit()
                   else .36 if char in ".,;:!'|ijl()[]" else .6 for char in word)
    room = width * 72 / (size * 1.04)               # ems per line, with a safety margin
    lines = 0
    for part in text.splitlines() or ['']:
        used, lines = 0.0, lines + 1
        for word in part.split():
            need = ems(word)
            if used and used + .34 + need > room:
                lines, used = lines + 1, 0.0
            if need > room:                          # a word longer than the line is broken
                lines += int(need // room)
                need %= room
            used += (.34 if used else 0) + need
    return lines


def _estimated_height(text, size, width):
    return _estimated_lines(text, size, width) * size * _LEADING / 72


def _fit_size(text, width, height, sizes):
    """The largest listed size whose estimated text fits the box; else the smallest."""
    for size in sizes:
        if _estimated_height(text, size, width) <= height:
            return size
    return sizes[-1]


def _centred(x, y, w, h, text, size, color):
    """Lower a short text block towards the middle of its panel.

    A sparse slide otherwise leaves its few lines pinned to the top of an
    empty panel. The box keeps its bottom edge, so nothing can overflow.
    """
    used = _estimated_height(text, size, w)
    drop = round(min((h - used) / 2, h * .3), 2) if used < h * .6 else 0
    return (x, round(y + drop, 2), w, round(h - drop, 2), text, size, color)


def _notes_strip(sources, body_need):
    """Source notes along the bottom: as tall as their text needs at a legible size,
    but never past what the body needs at its floor size.

    Returns (label box, notes box, y of the top of the strip).
    """
    cap = max(.67, min(2.4, _BOTTOM - _TOP - .85 - body_need * 1.08))
    size = _fit_size(sources, 11.5, cap, _SOURCE_SIZES)
    height = round(min(cap, max(.34, _estimated_height(sources, size, 11.5) + .06)), 2)
    y = round(_BOTTOM - height, 2)
    return ((.91,round(y - .33, 2),11.5,.25,'SOURCE NOTES',9,'3B6870'),
            (.91,y,11.5,height,sources,size,'42566A'), round(y - .47, 2))


def _hero(body, sources, index):
    """A single short sentence: display type in a full-width panel, alternately dark
    and light through the deck so consecutive statements do not look alike."""
    label, notes, top = _notes_strip(sources, _estimated_height(body, _BODY_SIZES[-1], 10.6))
    y, dark = _TOP + .12, index % 2 == 0
    h = round(top - .2 - y, 2)
    size = _fit_size(body, 10.6, h - .6, _BODY_SIZES)
    drop = round(max(.3, (h - _estimated_height(body, size, 10.6)) / 2), 2)
    shapes = [(.78, y, 11.77, h, '17324D' if dark else 'EAF2F2'),
              (.78, y, .09, h, 'EA5A5C' if dark else '4B8790')]
    return shapes, [(1.35, round(y + drop, 2), 10.6, round(h - drop, 2), body, size,
                     'FFFFFF' if dark else '17324D'), label, notes]


def _points(parts, sources):
    """Several sentences of evidence: one editable row each, with a marker and a
    hairline between rows, above a tinted source strip. None when they do not fit."""
    x, w, gap = 1.3, 10.9, .3 if len(parts) <= 2 else .22
    need = lambda size: sum(_estimated_height(part, size, w) for part in parts) + gap * (len(parts) - 1)
    label, notes, top = _notes_strip(sources, need(_BODY_SIZES[-1]))
    y, bottom = _TOP + .22, top - .25
    size = next((size for size in _BODY_SIZES if size <= 30 and need(size) <= bottom - y), None)
    if size is None:
        return None
    line = size * _LEADING / 72
    y = round(y + min((bottom - y - need(size)) / 2, (bottom - y) * .25), 2)
    shapes = [(.78, round(top - .08, 2), 11.77, round(_BOTTOM + .06 - top + .08, 2), 'EAF2F2'),
              (.78, round(top - .08, 2), .055, round(_BOTTOM + .06 - top + .08, 2), '4B8790')]
    boxes = []
    for index, part in enumerate(parts):
        height = _estimated_height(part, size, w)
        if index:
            shapes.append((x, round(y - gap / 2, 2), w, .012, 'DCE5E9'))
        shapes.append((.86, round(y + (line - .14) / 2, 2), .14, .14, 'EA5A5C'))
        boxes.append((x, y, w, round(height + .06, 2), part, size, '202E3B'))
        y = round(y + height + gap, 2)
    return shapes, boxes + [(1.0, *label[1:]), (1.0, notes[1], 11.4, *notes[3:])]


def _steps(parts, sources):
    """A model-selected timeline of several sentences: one column per sentence, in
    the model's order, hung from a rule with a marker each. None when too long."""
    gap = .4
    w = round((11.5 - gap * (len(parts) - 1)) / len(parts), 2)
    longest = max(parts, key=len)
    label, notes, top = _notes_strip(sources, _estimated_height(longest, _BODY_SIZES[-1], w))
    rule, bottom = _TOP + .42, top - .2
    y = round(rule + .4, 2)
    if _estimated_height(longest, _BODY_SIZES[-1], w) > bottom - y:
        return None
    size = min(_fit_size(longest, w, bottom - y, _BODY_SIZES), 30)
    shapes = [(.78, top, 11.77, .018, 'DCE5E9'), (.91, rule, 11.5, .03, 'B9CBD1')]
    boxes = []
    for index, part in enumerate(parts):
        x = round(.91 + index * (w + gap), 2)
        shapes.append((x, round(rule - .08, 2), .19, .19, 'EA5A5C' if index % 2 == 0 else '4B8790'))
        boxes.append((x, y, w, round(bottom - y, 2), part, size, '202E3B'))
    return shapes, boxes + [label, notes]


def _slide_plan(heading, body, sources, title, layout, index=0):
    """One geometry for the editable slide and its PDF: (decorative shapes, text boxes).

    The model chooses the layout and writes every word. Within that layout the
    arrangement follows the shape of what it wrote: one short sentence becomes a
    display panel, several sentences of evidence become rows, several sentences
    of a timeline become steps. Source notes take the height their text needs,
    the body takes the rest, and each is set in the largest type that fits.
    Text is placed, never changed; `index` only alternates decoration.
    """
    shapes = [(.0,.0,13.333,.085,'EA5A5C'),
              (.78,1.9,11.77,.025,'DCE5E9'),
              (.78,6.9,11.77,.018,'DCE5E9')]
    boxes = [(.65,.25,11.9,.3,'DRAFT • NOT APPROVED FOR INVESTOR DISTRIBUTION',9,'9B4142'),
             (.78,.93,11.85,.85,heading,_fit_size(heading, 11.85, .85, _HEADING_SIZES),'17324D')]
    footer = (.79,7.06,11.7,.22,title,8,'42566A')
    parts = _sentences(body)
    variant = None
    if layout == 'statement' and len(parts) == 1 and len(body) <= 170 and '\n' not in body:
        variant = _hero(body, sources, index)
    elif layout == 'evidence' and len(parts) > 1:
        variant = _points(parts, sources)
    elif layout == 'timeline' and len(parts) > 1:
        variant = _steps(parts, sources)
    if variant:
        return shapes + variant[0], boxes + variant[1] + [footer]
    if layout == 'evidence':
        # One sentence, or more than rows can hold: body beside a full-height notes panel.
        body_box = (.87, 2.33, 7.1, 4.05)
        notes = (8.5, 2.6, 3.7, 3.72)
        source_box = _centred(*notes, sources, _fit_size(sources, notes[2], notes[3],
                                                         _SOURCE_SIZES), '42566A')
        shapes += [(8.28,2.04,4.19,4.5,'EAF2F2'), (8.28,2.04,.055,4.5,'4B8790')]
        boxes += [_centred(*body_box, body, _fit_size(body, body_box[2], body_box[3], _BODY_SIZES),
                           '202E3B'),
                  (8.5,round(source_box[1]-.43,2),3.75,.35,'SOURCE NOTES',9,'3B6870'),
                  source_box, footer]
        return shapes, boxes
    left, right = _comparison_parts(body) if layout == 'comparison' else (body, '')
    x, w = {'comparison': (.98, 5.04), 'timeline': (1.33, 10.82)}.get(layout, (1.02, 10.93))
    y = _TOP + .2
    longest = max((left, right), key=len)
    label, notes, rule_y = _notes_strip(sources, _estimated_height(longest, _BODY_SIZES[-1], w))
    h = round(rule_y - .18 - y, 2)
    body_size = _fit_size(longest, w, h, _BODY_SIZES)
    shapes.append((.78, rule_y, 11.77, .018, 'DCE5E9'))
    if layout == 'comparison':
        # Both columns drop together so their first lines stay level.
        first = _centred(x, y, w, h, longest, body_size, '202E3B')
        panel = round(rule_y - .12 - _TOP, 2)
        shapes += [(.78,_TOP,5.55,panel,'F0F5F5'), (7.0,_TOP,5.55,panel,'F9F1F0'),
                   (.78,_TOP,5.55,.065,'4B8790'), (7.0,_TOP,5.55,.065,'EA5A5C')]
        boxes += [(x,first[1],w,first[3],left,body_size,'202E3B'),
                  (7.3,first[1],w,first[3],right,body_size,'202E3B')]
    else:
        bar = (.8, .055) if layout == 'timeline' else (.78, .065)
        shapes.append((bar[0], round(y - .1, 2), bar[1], round(h + .16, 2), 'EA5A5C'))
        boxes.append(_centred(x, y, w, h, body, body_size, '202E3B'))
    boxes += [label, notes, footer]
    return shapes, boxes


def _slide_text_boxes(heading, body, sources, title, layout, index=0):
    """Shared text geometry for editable and searchable PDF draft views."""
    return _slide_plan(heading, body, sources, title, layout, index)[1]


def _wrapped_table_lines(text, size, width, *, bold=False):
    from reportlab.lib.utils import simpleSplit
    from reportlab.pdfbase.pdfmetrics import stringWidth
    font = _pdf_font(bold=bold)
    lines = [line for part in text.splitlines() or ['']
             for line in (simpleSplit(part, font, size, width * 72) or [''])]
    if any(stringWidth(line, font, size) > width * 72 + .1 for line in lines):
        raise ValueError('Table text contains an unbreakable overflow; split it before rendering')
    return lines


def _table_slide_plan(heading, body, sources, title, table):
    """Shared actual-font geometry for native cells and their matching PDF.

    The entire table must fit at 11pt or greater. No rows/columns/qualifiers are
    dropped and no figures are parsed from prose. A dense table needs a new
    explicitly authored page selection rather than renderer truncation.
    """
    label, notes, source_top = _notes_strip(sources, 0)
    caption_size = 14 if len(body) <= 170 else 12
    caption_h = (_estimated_height(body, caption_size, 11.5) + .06) if body else 0
    top = round(_TOP + .06 + caption_h + (.15 if body else 0), 3)
    bottom = round(source_top - .48, 3)
    x, width = .91, 11.5
    first = min(3.05, width * .34)
    widths = [first] + [(width - first) / (len(table.columns) - 1)] * (len(table.columns) - 1)
    matrix = _table_text_matrix(table)
    for size in (14, 12, 11):
        try:
            heights = [max(len(_wrapped_table_lines(text, size, col_width - .16, bold=row_index == 0))
                           for text, col_width in zip(row, widths)) * size * 1.2 / 72 + .12
                       for row_index, row in enumerate(matrix)]
        except ValueError:
            if size == 11:
                raise
            continue
        if sum(heights) <= bottom - top:
            break
    else:
        raise ValueError('Draft table exceeds readable slide geometry; split it before rendering')
    shapes = [(0, 0, 13.333, .085, 'EA5A5C'), (.78, 1.9, 11.77, .025, 'DCE5E9'),
              (.78, 6.9, 11.77, .018, 'DCE5E9')]
    boxes = [(.65, .25, 11.9, .3, 'DRAFT • NOT APPROVED FOR INVESTOR DISTRIBUTION', 9, '9B4142'),
             (.78, .93, 11.85, .85, heading,
              _fit_size(heading, 11.85, .85, _HEADING_SIZES), '17324D'),
             (.91, _TOP + .06, 11.5, caption_h, body, caption_size, '202E3B'),
             (.91, round(source_top - .35, 3), 11.5, .25,
              'Figures reported by sources. Financial review pending.',
              9, '42566A'), label, notes,
             (.79, 7.06, 11.7, .22, title, 8, '42566A')]
    return shapes, boxes, {'x': x, 'y': top, 'widths': widths, 'heights': heights,
                           'matrix': matrix, 'font_size': size}


def _add_pptx_table(slide, plan):
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    shape = slide.shapes.add_table(len(plan['matrix']), len(plan['widths']),
        Inches(plan['x']), Inches(plan['y']), Inches(sum(plan['widths'])),
        Inches(sum(plan['heights'])))
    table = shape.table
    for index, width in enumerate(plan['widths']):
        table.columns[index].width = Inches(width)
    for row_index, (row, height) in enumerate(zip(plan['matrix'], plan['heights'])):
        table.rows[row_index].height = Inches(height)
        for column_index, value in enumerate(row):
            cell = table.cell(row_index, column_index)
            cell.text = value
            cell.margin_left = cell.margin_right = Inches(.08)
            cell.margin_top = cell.margin_bottom = Inches(.06)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor.from_string(
                '17324D' if row_index == 0 else 'EFF4F5' if row_index % 2 else 'FFFFFF')
            for paragraph in cell.text_frame.paragraphs:
                paragraph.font.name = 'DejaVu Sans'
                paragraph.font.size = Pt(plan['font_size'])
                paragraph.font.bold = row_index == 0
                paragraph.font.color.rgb = RGBColor.from_string('FFFFFF' if row_index == 0 else '202E3B')
                paragraph.line_spacing = Pt(plan['font_size'] * 1.2)
                paragraph.alignment = PP_ALIGN.LEFT
    return table


_CHART_COLORS = ('4B8790', 'EA5A5C', '6D79A6', 'D4A84B', '789A64', 'AD79A9',
                 '7DB6C5', 'D98248', '587492', 'A5A54E', 'A97460', 'AF95C9')
_CHART_NOTE = 'Unreviewed source chart cache. No recalculation or financial validation. Period and units unassessed.'


def _chart_source_table(chart):
    """A visible literal cache companion, never a recalculated financial table."""
    from delivery.contracts import EvidenceTable
    columns = [{'label': 'Source category'}] + [{'label': series.name.text} for series in chart.series]
    rows = []
    for index, category in enumerate(chart.categories):
        cells = [{'text': category.text, 'status': 'source_reported',
                  'source_refs': chart.source_refs, 'value_ref': category.literal_unit_id}]
        cells.extend({'text': series.points[index].numeric_literal, 'status': 'source_reported',
                      'source_refs': chart.source_refs, 'value_ref': series.points[index].literal_unit_id}
                     for series in chart.series)
        rows.append({'cells': cells})
    return EvidenceTable.model_validate({'table_id': chart.chart_unit_id,
        'evidence_revision': chart.evidence_revision, 'columns': columns, 'rows': rows})


def _chart_slide_plan(heading, body, sources, title, chart):
    label, notes, source_top = _notes_strip(sources, 0)
    caption_size = 12
    caption_h = _estimated_height(body, caption_size, 11.5) + .06 if body else 0
    top = round(_TOP + .06 + caption_h + (.15 if body else 0), 3)
    bottom = round(source_top - .48, 3)
    matrix = _table_text_matrix(_chart_source_table(chart))
    widths = [1.4] + [2.95 / len(chart.series)] * len(chart.series)
    size = 9
    heights = [max(len(_wrapped_table_lines(text, size, width - .16, bold=index == 0))
                   for text, width in zip(row, widths)) * size * 1.2 / 72 + .1
               for index, row in enumerate(matrix)]
    if sum(heights) > bottom - top or bottom - top < 2:
        raise ValueError('Complete chart cache companion exceeds readable slide; unsupported without an authored split')
    shapes = [(0, 0, 13.333, .085, 'EA5A5C'), (.78, 1.9, 11.77, .025, 'DCE5E9'),
              (.78, 6.9, 11.77, .018, 'DCE5E9')]
    boxes = [(.65, .25, 11.9, .3, 'DRAFT • NOT APPROVED FOR INVESTOR DISTRIBUTION', 9, '9B4142'),
             (.78, .93, 11.85, .85, heading, _fit_size(heading, 11.85, .85, _HEADING_SIZES), '17324D'),
             (.91, _TOP + .06, 11.5, caption_h, body, caption_size, '202E3B'),
             (.91, round(source_top - .35, 3), 11.5, .25, _CHART_NOTE, 9, '42566A'),
             label, notes, (.79, 7.06, 11.7, .22, title, 8, '42566A')]
    return shapes, boxes, {'x': 8.06, 'y': top, 'widths': widths, 'heights': heights,
                           'matrix': matrix, 'font_size': size}, (.91, top, 6.95, bottom - top)


def _chart_axis_range(chart):
    import math
    values = [[float(point.numeric_literal) for point in series.points] for series in chart.series]
    if chart.plot_type == 'column_stacked':
        observations = [sum(max(series[index], 0) for series in values) for index in range(len(chart.categories))]
        observations += [sum(min(series[index], 0) for series in values) for index in range(len(chart.categories))]
    else:
        observations = [value for series in values for value in series]
    lower, upper = min(0, min(observations)), max(0, max(observations))
    if lower == upper:
        upper = lower + 1
    span = upper - lower
    raw_step = span / 5
    magnitude = 10 ** math.floor(math.log10(raw_step))
    step = next(factor * magnitude for factor in (1, 2, 5, 10) if factor * magnitude >= raw_step)
    # Axis tick geometry is not a new observation or financial estimate. Native
    # Office and the PDF share clean bounds rather than floating-point padding.
    return (math.floor(lower / step) * step, math.ceil(upper / step) * step, step)


def _add_pptx_chart(slide, chart, geometry):
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_TICK_LABEL_POSITION
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt
    data = CategoryChartData()
    data.categories = [category.text for category in chart.categories]
    for series in chart.series:
        data.add_series(series.name.text, [float(point.numeric_literal) for point in series.points])
    native_type = {'column_clustered': XL_CHART_TYPE.COLUMN_CLUSTERED,
                   'column_stacked': XL_CHART_TYPE.COLUMN_STACKED,
                   'bar_clustered': XL_CHART_TYPE.BAR_CLUSTERED,
                   'line': XL_CHART_TYPE.LINE_MARKERS, 'pie': XL_CHART_TYPE.PIE}[chart.plot_type]
    native = slide.shapes.add_chart(native_type, *(Inches(value) for value in geometry), data).chart
    native.has_title = False
    native.has_legend = True
    native.legend.position = XL_LEGEND_POSITION.BOTTOM
    native.legend.include_in_layout = False
    native.legend.font.name, native.legend.font.size = 'DejaVu Sans', Pt(9)
    if chart.plot_type != 'pie':
        lower, upper, step = _chart_axis_range(chart)
        native.value_axis.minimum_scale, native.value_axis.maximum_scale = lower, upper
        native.value_axis.major_unit = step
        native.value_axis.tick_labels.number_format = '0.00E+00' if step < .01 or step >= 1e10 else '0.##'
        native.category_axis.tick_label_position = XL_TICK_LABEL_POSITION.LOW
        if chart.plot_type == 'bar_clustered':
            native.category_axis.reverse_order = True
            native.value_axis.tick_label_position = XL_TICK_LABEL_POSITION.NEXT_TO_AXIS
        for axis in (native.category_axis, native.value_axis):
            axis.tick_labels.font.name, axis.tick_labels.font.size = 'DejaVu Sans', Pt(9)
        for index, series in enumerate(native.series):
            if chart.plot_type == 'line':
                series.format.line.color.rgb = RGBColor.from_string(_CHART_COLORS[index])
                series.format.line.width = Pt(1.8)
                series.marker.format.fill.solid()
                series.marker.format.fill.fore_color.rgb = RGBColor.from_string(_CHART_COLORS[index])
                series.marker.format.line.color.rgb = RGBColor.from_string(_CHART_COLORS[index])
            else:
                series.invert_if_negative = False
                series.format.fill.solid()
                series.format.fill.fore_color.rgb = RGBColor.from_string(_CHART_COLORS[index])
        if chart.plot_type == 'column_stacked':
            native.plots[0].overlap = 100
    else:
        for index, point in enumerate(native.series[0].points):
            point.format.fill.solid()
            point.format.fill.fore_color.rgb = RGBColor.from_string(_CHART_COLORS[index % len(_CHART_COLORS)])
    return native


def _chart_plot_pdf(chart, width, height):
    """Standard plotting from frozen observations only; no source image copying."""
    import os
    import tempfile
    os.environ.setdefault('MPLCONFIGDIR', tempfile.gettempdir() + '/ainvestify-chart-font-cache')
    from matplotlib.figure import Figure
    figure = Figure(figsize=(width, height), dpi=120)
    axes = figure.subplots()
    labels = [category.text for category in chart.categories]
    values = [[float(point.numeric_literal) for point in series.points] for series in chart.series]
    colors = ['#' + value for value in _CHART_COLORS]
    positions = list(range(len(labels)))
    if chart.plot_type == 'pie':
        axes.pie(values[0], labels=labels, colors=[colors[i % len(colors)] for i in positions],
                 textprops={'fontsize': 9, 'fontfamily': 'DejaVu Sans'})
    elif chart.plot_type == 'line':
        for index, series in enumerate(chart.series):
            axes.plot(positions, values[index], color=colors[index], marker='o',
                      linewidth=1.5, markersize=3, label=series.name.text)
    elif chart.plot_type == 'column_stacked':
        positive, negative = [0.] * len(labels), [0.] * len(labels)
        for index, series in enumerate(chart.series):
            base = [positive[i] if value >= 0 else negative[i] for i, value in enumerate(values[index])]
            axes.bar(positions, values[index], bottom=base, color=colors[index], width=.65, label=series.name.text)
            for i, value in enumerate(values[index]):
                if value >= 0:
                    positive[i] += value
                else:
                    negative[i] += value
    else:
        bar_width = .72 / len(chart.series)
        for index, series in enumerate(chart.series):
            shifted = [position - .36 + (index + .5) * bar_width for position in positions]
            if chart.plot_type == 'bar_clustered':
                axes.barh(shifted, values[index], color=colors[index], height=bar_width, label=series.name.text)
            else:
                axes.bar(shifted, values[index], color=colors[index], width=bar_width, label=series.name.text)
    if chart.plot_type != 'pie':
        from matplotlib.ticker import MultipleLocator
        lower, upper, step = _chart_axis_range(chart)
        if chart.plot_type == 'bar_clustered':
            axes.set_yticks(positions, labels)
            axes.set_xlim(lower, upper)
            axes.xaxis.set_major_locator(MultipleLocator(step))
            axes.invert_yaxis()
            axes.grid(axis='x', color='#DCE5E9', linewidth=.4)
        else:
            long_labels = any(len(label) > 8 for label in labels)
            axes.set_xticks(positions, labels, rotation=30 if long_labels else 0,
                           ha='right' if long_labels else 'center')
            axes.set_ylim(lower, upper)
            axes.yaxis.set_major_locator(MultipleLocator(step))
            axes.grid(axis='y', color='#DCE5E9', linewidth=.4)
        axes.legend(loc='lower center', bbox_to_anchor=(.5, -.28), ncol=min(3, len(chart.series)),
                    frameon=False, fontsize=9)
        axes.set_axisbelow(True)
        axes.spines[['top', 'right']].set_visible(False)
        axes.tick_params(labelsize=9)
    figure.subplots_adjust(left=.18 if chart.plot_type == 'bar_clustered' else .13,
                           right=.96, top=.94, bottom=.25)
    from matplotlib.backends.backend_agg import FigureCanvasAgg
    canvas = FigureCanvasAgg(figure)
    canvas.draw()
    renderer = canvas.get_renderer()
    full = figure.bbox
    x_min, x_max = sorted(axes.get_xlim()); y_min, y_max = sorted(axes.get_ylim())
    x_ticks = [tick for tick in axes.get_xticklabels()
               if x_min <= tick.get_position()[0] <= x_max] if axes.axison else []
    y_ticks = [tick for tick in axes.get_yticklabels()
               if y_min <= tick.get_position()[1] <= y_max] if axes.axison else []
    legend = axes.get_legend()
    # Matplotlib creates tick artists beyond the view limits which it never
    # paints. Inspect only labels actually drawn, while retaining every source
    # category, legend and pie label in the geometry check.
    drawn_text = [*axes.texts, *x_ticks, *y_ticks, axes.title,
                  axes.xaxis.label, axes.yaxis.label,
                  axes.xaxis.offsetText, axes.yaxis.offsetText,
                  *(legend.get_texts() if legend else [])]
    for text in drawn_text:
        if not text.get_visible() or not text.get_text():
            continue
        bounds = text.get_window_extent(renderer)
        if (bounds.x0 < full.x0 - .5 or bounds.y0 < full.y0 - .5 or
                bounds.x1 > full.x1 + .5 or bounds.y1 > full.y1 + .5):
            raise ValueError('Complete original chart labels exceed readable plot geometry')
    ticks = y_ticks if chart.plot_type == 'bar_clustered' else x_ticks
    tick_bounds = [tick.get_window_extent(renderer) for tick in ticks if tick.get_text()]
    if any(first.overlaps(second) for first, second in zip(tick_bounds, tick_bounds[1:])):
        raise ValueError('Complete original chart category labels overlap; an authored split is required')
    output = BytesIO()
    figure.savefig(output, format='pdf', metadata={'Creator': 'source-cache-chart-v1'})
    return output.getvalue()


def _merge_chart_pages(pdf_content, chart_layers):
    """Overlay complete vector chart PDFs onto already-rendered matching pages."""
    from pypdf import PdfReader, PdfWriter, Transformation
    document = PdfReader(BytesIO(pdf_content)); writer = PdfWriter()
    for index, page in enumerate(document.pages):
        writer.add_page(page)
        if index in chart_layers:
            chart, (x, y, width, height) = chart_layers[index]
            layer = PdfReader(BytesIO(_chart_plot_pdf(chart, width, height))).pages[0]
            transform = Transformation().translate(x * 72, float(page.mediabox.height) - (y + height) * 72)
            writer.pages[-1].merge_transformed_page(layer, transform, over=True)
    writer.write(output := BytesIO())
    return output.getvalue()


def inspect_evidence_chart_binding(pptx_content, chart_payload, *, slide_index=0):
    """Inspect the generated native chart and its editable embedded workbook.

    This narrow mechanical check never waives the existing audience-review
    finding for embedded payloads. Generic release and exact-file professional
    reviews remain required; PDF visual parity is a separate unrun gate.
    """
    import hashlib
    from decimal import Decimal
    from openpyxl import load_workbook
    from pptx import Presentation
    from pptx.enum.chart import XL_CHART_TYPE
    from defusedxml import ElementTree as ET
    from delivery.inspection import _parts
    chart = _evidence_chart(chart_payload)
    findings = []
    expected_type = {'column_clustered': XL_CHART_TYPE.COLUMN_CLUSTERED,
                     'column_stacked': XL_CHART_TYPE.COLUMN_STACKED,
                     'bar_clustered': XL_CHART_TYPE.BAR_CLUSTERED,
                     'line': XL_CHART_TYPE.LINE_MARKERS, 'pie': XL_CHART_TYPE.PIE}[chart.plot_type]
    try:
        with _parts(pptx_content) as archive:
            native_deck = Presentation(BytesIO(pptx_content))
            native_charts = [shape.chart for shape in native_deck.slides[slide_index].shapes if shape.has_chart]
            if len(native_charts) != 1:
                raise ValueError('Expected one exact native chart on selected slide')
            native = native_charts[0]
            if native.chart_type != expected_type or len(native.plots) != 1:
                raise ValueError('Generated chart type/plot count changed')
            if [category.label for category in native.plots[0].categories] != [c.text for c in chart.categories]:
                raise ValueError('Generated source categories changed')
            if len(native.series) != len(chart.series):
                raise ValueError('Generated source series count changed')
            for generated, original in zip(native.series, chart.series):
                if generated.name != original.name.text or len(generated.values) != len(original.points):
                    raise ValueError('Generated source series identity/count changed')
                if any(Decimal(str(value)) != Decimal(point.numeric_literal)
                       for value, point in zip(generated.values, original.points)):
                    raise ValueError('Generated original cache observation changed')
            part = str(native.part.partname).lstrip('/')
            folder, filename = part.rsplit('/', 1)
            relationships = ET.fromstring(archive.read(folder + '/_rels/' + filename + '.rels'))
            targets = []
            import posixpath
            for rel in relationships:
                if rel.attrib.get('TargetMode') == 'External':
                    raise ValueError('Generated native chart has external dependency')
                if rel.attrib.get('Type', '').endswith('/package'):
                    targets.append(posixpath.normpath(posixpath.join(folder, rel.attrib['Target'])))
            if len(targets) != 1 or not targets[0].startswith('ppt/embeddings/') or not targets[0].endswith('.xlsx'):
                raise ValueError('Generated native chart must bind one internal editable workbook')
            workbook_bytes = archive.read(targets[0])
            with _parts(workbook_bytes) as nested:
                if any(any(marker in name.lower() for marker in
                           ('externallinks/', 'vbaproject', 'connections.xml', 'activex/'))
                       for name in nested.namelist()):
                    raise ValueError('Generated chart workbook has unsupported active/external parts')
                for name in nested.namelist():
                    if name.endswith('.rels') and any(rel.attrib.get('TargetMode') == 'External'
                            for rel in ET.fromstring(nested.read(name))):
                        raise ValueError('Generated chart workbook has external relationship')
            workbook = load_workbook(BytesIO(workbook_bytes), read_only=False, data_only=False)
            if len(workbook.worksheets) != 1 or workbook.defined_names:
                raise ValueError('Generated chart workbook has unexpected sheets/names')
            sheet = workbook.worksheets[0]
            if sheet.sheet_state != 'visible' or sheet.max_row != len(chart.categories) + 1 or sheet.max_column != len(chart.series) + 1:
                raise ValueError('Generated chart workbook scope changed')
            if any(cell.data_type == 'f' or cell.hyperlink or cell.comment
                   for row in sheet for cell in row):
                raise ValueError('Generated chart workbook has formulas or hidden audience payload')
            for index, series in enumerate(chart.series, 2):
                if sheet.cell(1, index).value != series.name.text:
                    raise ValueError('Generated embedded series name changed')
            for row_index, category in enumerate(chart.categories, 2):
                if sheet.cell(row_index, 1).value != category.text:
                    raise ValueError('Generated embedded category changed')
                for column_index, series in enumerate(chart.series, 2):
                    value = sheet.cell(row_index, column_index).value
                    if isinstance(value, bool) or Decimal(str(value)) != Decimal(series.points[row_index - 2].numeric_literal):
                        raise ValueError('Generated editable workbook observation changed')
            workbook.close()
    except Exception as exc:
        findings.append({'code': 'native_chart_binding_failed', 'error_type': type(exc).__name__})
    return {'contract': 'source-chart-binding-check-v1',
            'artifact_sha256': hashlib.sha256(pptx_content).hexdigest(),
            'source_cache_digest': chart.cache_digest, 'evidence_revision': chart.evidence_revision,
            'native_chart_binding': 'fail' if findings else 'pass',
            'embedded_workbook_binding': 'fail' if findings else 'pass',
            'editable_chart': 'native_pptx', 'financial_validation': 'not_established',
            'audience_review': 'not_run', 'pdf_visual_parity': 'not_run',
            'release_status': 'blocked', 'findings': findings}


def inspect_evidence_chart_pdf_text(pptx_content, pdf_content, chart_payload, *, slide_index=0):
    """Check exact source-cache companion and slide text in a concrete PDF.

    It does not reconstruct plot geometry from PDF vectors or qualify source
    numbers. Neither this check nor native cache reproduction grants release.
    """
    import hashlib
    from pptx import Presentation
    from pypdf import PdfReader
    chart = _evidence_chart(chart_payload)
    findings = []
    try:
        deck = Presentation(BytesIO(pptx_content)); pdf = PdfReader(BytesIO(pdf_content))
        if len(deck.slides) != len(pdf.pages):
            raise ValueError('Source chart slide/PDF page count mismatch')
        normalized = lambda value: re.sub(r'\s+', ' ', value).strip()
        text = normalized(pdf.pages[slide_index].extract_text() or '')
        required = []
        for shape in deck.slides[slide_index].shapes:
            if shape.has_text_frame:
                required.append(shape.text)
            if shape.has_table:
                required.extend(cell.text for row in shape.table.rows for cell in row.cells)
        required.extend(' '.join(row) for row in _table_text_matrix(_chart_source_table(chart)))
        if any(normalized(value) not in text for value in required if normalized(value)):
            raise ValueError('Source chart companion or editable slide text missing from PDF')
    except Exception as exc:
        findings.append({'code': 'source_chart_pdf_text_binding_failed', 'error_type': type(exc).__name__})
    return {'contract': 'source-chart-pdf-text-check-v1',
        'editable_sha256': hashlib.sha256(pptx_content).hexdigest(),
        'pdf_sha256': hashlib.sha256(pdf_content).hexdigest(), 'source_cache_digest': chart.cache_digest,
        'source_cache_companion_text': 'fail' if findings else 'pass',
        'plot_visual_binding': 'not_run', 'financial_validation': 'not_established',
        'audience_review': 'not_run', 'release_status': 'blocked', 'findings': findings}


def render_intro(title, sections, *, document_contract=None):
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    for index, section in enumerate(_checked_sections(sections, document_contract=document_contract)):
        heading, body, sources = section[:3]
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        table_plan = chart_plan = None
        if len(section) == 6:
            chart = _evidence_chart(section[4])
            shapes, boxes, table_plan, chart_plan = _chart_slide_plan(heading, body, sources, title, chart)
        elif len(section) == 5:
            shapes, boxes, table_plan = _table_slide_plan(heading, body, sources, title,
                                                        _evidence_table(section[4]))
        else:
            layout = _effective_slide_layout(
                section[3] if len(section) == 4 else 'statement', body)
            shapes, boxes = _slide_plan(heading, body, sources, title, layout, index)
        for x,y,w,h,fill in shapes:
            shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y),
                                           Inches(w), Inches(h))
            shape.fill.solid(); shape.fill.fore_color.rgb = RGBColor.from_string(fill)
            shape.line.fill.background()
        for x,y,w,h,text,size,color in boxes:
            box = slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
            box.text_frame.word_wrap = True
            box.text_frame.margin_left = box.text_frame.margin_right = 0
            box.text_frame.margin_top = box.text_frame.margin_bottom = 0
            paragraph = box.text_frame.paragraphs[0]
            paragraph.text = text
            paragraph.font.size = Pt(size)
            paragraph.line_spacing = Pt(round(size * _LEADING, 2))   # the PDF's line pitch
            paragraph.font.name = "DejaVu Sans"
            paragraph.font.color.rgb = RGBColor.from_string(color)
        if table_plan:
            _add_pptx_table(slide, table_plan)
        if chart_plan:
            _add_pptx_chart(slide, chart, chart_plan)
    output=BytesIO();deck.save(output);return output.getvalue()


# Point sizes shared by the editable memo and its PDF.
_MEMO = {'title': 21, 'heading': 13.5, 'body': 11, 'evidence': 9.5, 'sources': 9, 'label': 8.5}


def _memo_source_label(sources):
    """Preserve legacy URL/version labels; name reader projection lists plainly."""
    lines = [line.strip() for line in sources.splitlines() if line.strip()]
    if lines and all(re.match(r'^\[S\d+\]\s+[^|]+$', line) for line in lines):
        return 'Sources cited'
    return 'Source versions'


def _memo_table_plan(table):
    width = 170 / 25.4
    first = min(2.25, width * .34)
    widths = [first] + [(width - first) / (len(table.columns) - 1)] * (len(table.columns) - 1)
    matrix = _table_text_matrix(table)
    size = 9
    heights = [max(len(_wrapped_table_lines(text, size, col_width - .12, bold=row_index == 0))
                   for text, col_width in zip(row, widths)) * size * 1.2 / 72 + .10
               for row_index, row in enumerate(matrix)]
    if any(height > 7.5 for height in heights):
        raise ValueError('Draft memo table row cannot fit on one page; split it before rendering')
    return {'widths': widths, 'heights': heights, 'matrix': matrix, 'font_size': size}


def _add_docx_table(doc, table):
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    plan = _memo_table_plan(table)
    native = doc.add_table(rows=len(plan['matrix']), cols=len(plan['widths']))
    native.autofit = False
    native.alignment = WD_TABLE_ALIGNMENT.CENTER
    for column, width in zip(native.columns, plan['widths']):
        column.width = Inches(width)
    for row_index, (row, values) in enumerate(zip(native.rows, plan['matrix'])):
        properties = row._tr.get_or_add_trPr()
        properties.append(OxmlElement('w:cantSplit'))
        if row_index == 0:
            repeat = OxmlElement('w:tblHeader')
            repeat.set(qn('w:val'), 'true')
            properties.append(repeat)
        for cell, text, width in zip(row.cells, values, plan['widths']):
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text = text
            shade = OxmlElement('w:shd')
            shade.set(qn('w:fill'), '17324D' if row_index == 0 else 'EFF4F5' if row_index % 2 else 'FFFFFF')
            cell._tc.get_or_add_tcPr().append(shade)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.keep_together = True
                paragraph.paragraph_format.keep_with_next = False
                paragraph.paragraph_format.space_before = Pt(3)
                paragraph.paragraph_format.space_after = Pt(3)
                paragraph.paragraph_format.line_spacing = 1.2
                for run in paragraph.runs:
                    run.font.name = 'DejaVu Sans'
                    run.font.size = Pt(plan['font_size'])
                    run.font.bold = row_index == 0
                    run.font.color.rgb = RGBColor.from_string('FFFFFF' if row_index == 0 else '202E3B')
    return native


def _chart_plot_png(chart):
    import pypdfium2 as pdfium
    pdf = pdfium.PdfDocument(_chart_plot_pdf(chart, 6.6, 3.4))
    bitmap = pdf[0].render(scale=1.7).to_pil()
    output = BytesIO(); bitmap.save(output, format='PNG')
    pdf.close()
    return output.getvalue()


def render_memo(title, sections, *, document_contract=None):
    from docx import Document
    from docx.shared import Mm, Pt
    from docx.enum.style import WD_STYLE_TYPE
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import RGBColor
    doc=Document()
    # The library's blank template contains a custom bibliography XML part.
    # New drafts do not use it; omit it instead of distributing hidden payloads.
    for key,relationship in list(doc.part.rels.items()):
        if relationship.reltype.endswith('/customXml'):
            doc.part.drop_rel(key)
    page = doc.sections[0]
    # The same sheet and margins as the PDF, so the two paginate alike.
    page.page_width, page.page_height = Mm(210), Mm(297)
    page.top_margin, page.bottom_margin = Mm(25), Mm(22)
    page.left_margin = page.right_margin = Mm(20)
    normal = doc.styles['Normal']
    normal.font.name, normal.font.size = 'DejaVu Sans', Pt(_MEMO['body'])
    normal.paragraph_format.line_spacing = 1.28
    normal.paragraph_format.space_after = Pt(9)
    normal.paragraph_format.widow_control = True
    normal.paragraph_format.keep_together = True

    def edge(style, side, color, width, space):
        """A rule on one side of every paragraph in a style: layout only."""
        properties = style.element.get_or_add_pPr()
        borders = properties.find(qn('w:pBdr'))
        if borders is None:
            borders = OxmlElement('w:pBdr')
            properties.append(borders)
        rule = OxmlElement(f'w:{side}')
        for key, value in (('val', 'single'), ('sz', str(width)), ('space', str(space)),
                           ('color', color)):
            rule.set(qn(f'w:{key}'), value)
        borders.append(rule)
    title_style = doc.styles['Title']
    title_style.font.name, title_style.font.size = 'DejaVu Sans', Pt(_MEMO['title'])
    title_style.font.color.rgb = RGBColor(23, 50, 77)
    title_style.paragraph_format.space_after = Pt(17)
    heading_style = doc.styles['Heading 1']
    heading_style.font.name, heading_style.font.size = 'DejaVu Sans', Pt(_MEMO['heading'])
    heading_style.font.color.rgb = RGBColor(23, 50, 77)
    heading_style.paragraph_format.space_before = Pt(20)
    heading_style.paragraph_format.space_after = Pt(8)
    heading_style.paragraph_format.keep_with_next = True
    edge(heading_style, 'top', 'DCE5E9', 6, 8)          # a hairline opens each section
    for name, size, color, before, after in (
            ('Memo Evidence', _MEMO['evidence'], RGBColor(38, 54, 69), 0, 6),
            ('Memo Sources', _MEMO['sources'], RGBColor(66, 86, 106), 0, 5),
            ('Memo Label', _MEMO['label'], RGBColor(23, 50, 77), 12, 5)):
        style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = normal
        style.font.name, style.font.size, style.font.color.rgb = 'DejaVu Sans', Pt(size), color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_together = True
    doc.styles['Memo Label'].font.bold = True
    doc.styles['Memo Label'].paragraph_format.keep_with_next = True
    # Excerpts and source versions are set off from the prose, not shrunk below it.
    for name in ('Memo Evidence', 'Memo Sources'):
        doc.styles[name].paragraph_format.left_indent = Mm(4)
        doc.styles[name].paragraph_format.line_spacing = 1.2
    edge(doc.styles['Memo Evidence'], 'left', '4B8790', 12, 6)
    header = page.header.paragraphs[0]
    header.text = "DRAFT — NOT APPROVED FOR INVESTOR DISTRIBUTION"
    header.style = doc.styles['Memo Sources']
    doc.add_paragraph(title, style='Title')
    marker = "\n\nSource claims and exact excerpts:\n"
    for section in _checked_sections(sections, paginated=True, document_contract=document_contract):
        heading, body, sources = section[:3]
        doc.add_heading(heading, 1)
        prose, separator, evidence = body.partition(marker)
        registry = _source_registry_lines(prose)
        if registry:
            for line in registry:
                doc.add_paragraph(line, style='Memo Sources')
        else:
            for block in prose.split('\n\n'):
                if block:
                    doc.add_paragraph(block)
        if separator:
            # Keep the model's marker verbatim; only its presentation changes.
            doc.add_paragraph('Source claims and exact excerpts:', style='Memo Label')
            for line in evidence.splitlines():
                if line:
                    doc.add_paragraph(line, style='Memo Evidence')
        if len(section) == 5:
            _add_docx_table(doc, _evidence_table(section[4]))
            doc.add_paragraph('Figures reported by sources. Financial review pending.',
                              style='Memo Sources')
        if len(section) == 6:
            chart = _evidence_chart(section[4])
            doc.add_picture(BytesIO(_chart_plot_png(chart)), width=Mm(168))
            doc.add_paragraph(_CHART_NOTE, style='Memo Sources')
            doc.add_paragraph('Chart image is not editable. Exact source values remain editable below.',
                              style='Memo Sources')
            _add_docx_table(doc, _chart_source_table(chart))
        if sources:
            doc.add_paragraph(_memo_source_label(sources), style='Memo Label')
            for line in sources.splitlines():
                if line:
                    doc.add_paragraph(line, style='Memo Sources')
    footer = page.footer.paragraphs[0]
    footer.text = 'INVESTMENT MEMORANDUM DRAFT'
    footer.style = doc.styles['Memo Sources']
    output=BytesIO();doc.save(output);return output.getvalue()


def _draw_pdf_table(pdf, page_height, plan):
    from reportlab.lib import colors
    y = plan['y']
    for row_index, (row, row_height) in enumerate(zip(plan['matrix'], plan['heights'])):
        x = plan['x']
        for text, width in zip(row, plan['widths']):
            pdf.setFillColor(colors.HexColor('#17324D' if row_index == 0 else
                                            '#EFF4F5' if row_index % 2 else '#FFFFFF'))
            pdf.setStrokeColor(colors.HexColor('#DCE5E9'))
            pdf.setLineWidth(.4)
            pdf.rect(x * 72, page_height - (y + row_height) * 72,
                     width * 72, row_height * 72, fill=1, stroke=1)
            pdf.setFillColor(colors.white if row_index == 0 else colors.HexColor('#202E3B'))
            pdf.setFont(_pdf_font(bold=row_index == 0), plan['font_size'])
            lines = _wrapped_table_lines(text, plan['font_size'], width - .16, bold=row_index == 0)
            baseline = page_height - (y + .06) * 72 - plan['font_size']
            for line in lines:
                pdf.drawString((x + .08) * 72, baseline, line)
                baseline -= plan['font_size'] * 1.2
            x += width
        y += row_height


def render_deck_pdf(title, sections, *, document_contract=None):
    """A searchable PDF of the same model-authored slide text.

    This text-preserving export is a draft. The separate pair inspector checks
    exact text and page count; visual parity still needs its own review.
    """
    from reportlab.lib import colors
    from reportlab.lib.utils import simpleSplit
    from reportlab.pdfgen import canvas

    output = BytesIO()
    font = _pdf_font()
    width, height = 13.333 * 72, 7.5 * 72
    pdf = canvas.Canvas(output, pagesize=(width, height), pageCompression=0)
    pdf.setTitle(f"{title} - draft deck")
    chart_layers = {}
    for index, section in enumerate(_checked_sections(sections, document_contract=document_contract)):
        heading, body, sources = section[:3]
        pdf.setFillColor(colors.white)
        pdf.rect(0, 0, width, height, fill=1, stroke=0)
        table_plan = None
        if len(section) == 6:
            chart = _evidence_chart(section[4])
            shapes, boxes, table_plan, chart_geometry = _chart_slide_plan(heading, body, sources, title, chart)
            chart_layers[index] = (chart, chart_geometry)
        elif len(section) == 5:
            shapes, boxes, table_plan = _table_slide_plan(heading, body, sources, title,
                                                        _evidence_table(section[4]))
        else:
            layout = _effective_slide_layout(
                section[3] if len(section) == 4 else 'statement', body)
            shapes, boxes = _slide_plan(heading, body, sources, title, layout, index)
        for x,y,w,h,fill in shapes:
            pdf.setFillColor(colors.HexColor('#' + fill))
            pdf.rect(x*72, height-(y+h)*72, w*72, h*72, fill=1, stroke=0)
        for x,y,w,h,value,size,color in boxes:
            if not value:
                continue
            lines = [line for part in value.splitlines() or [""]
                     for line in (simpleSplit(part, font, size, w*72) or [""])]
            leading = size * _LEADING
            if len(lines) * leading > h*72 + 4:
                raise ValueError('Deck PDF text exceeds its matching editable box')
            pdf.setFont(font, size)
            pdf.setFillColor(colors.HexColor('#' + color))
            baseline = height - y*72 - size
            for line in lines:
                pdf.drawString(x*72, baseline, line)
                baseline -= leading
        if table_plan:
            _draw_pdf_table(pdf, height, table_plan)
        pdf.showPage()
    pdf.save()
    return _merge_chart_pages(output.getvalue(), chart_layers) if chart_layers else output.getvalue()


def render_memo_pdf(title, sections, *, document_contract=None):
    """Searchable PDF using the exact model-authored DOCX fields."""
    from html import escape
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer, KeepTogether,
                                    HRFlowable, Table, TableStyle)

    output = BytesIO()
    font = _pdf_font()
    document = SimpleDocTemplate(output, pagesize=A4, leftMargin=20*mm,
        rightMargin=20*mm, topMargin=25*mm, bottomMargin=22*mm,
        title=f"{title} - draft investment memorandum")
    base = getSampleStyleSheet()
    styles = {
        "title": ParagraphStyle("MemoTitle", parent=base["Title"],
            fontName=font, fontSize=_MEMO['title'], leading=27,
            textColor=colors.HexColor("#17324D"), alignment=0, spaceAfter=10*mm),
        "heading": ParagraphStyle("MemoHeading", parent=base["Heading2"],
            fontName=font, fontSize=_MEMO['heading'], leading=18.5,
            textColor=colors.HexColor("#17324D"), spaceBefore=2.5*mm, spaceAfter=3*mm,
            keepWithNext=1),
        "body": ParagraphStyle("MemoBody", parent=base["BodyText"],
            fontName=font, fontSize=_MEMO['body'], leading=16.5, spaceAfter=3.2*mm,
            allowWidows=0, allowOrphans=0),
        "label": ParagraphStyle("MemoLabel", parent=base["BodyText"],
            fontName=font, fontSize=_MEMO['label'], leading=12,
            textColor=colors.HexColor("#17324D"), spaceBefore=3.5*mm, spaceAfter=1.8*mm,
            keepWithNext=1),
        "evidence": ParagraphStyle("MemoEvidence", parent=base["BodyText"],
            fontName=font, fontSize=_MEMO['evidence'], leading=13.6,
            textColor=colors.HexColor("#263645"), spaceAfter=0, splitLongWords=1,
            allowWidows=0, allowOrphans=0),
        "sources": ParagraphStyle("MemoSources", parent=base["BodyText"],
            fontName=font, fontSize=_MEMO['sources'], leading=13,
            textColor=colors.HexColor("#42566A"), leftIndent=4*mm, spaceAfter=1.8*mm,
            splitLongWords=1, allowWidows=0, allowOrphans=0),
    }

    def excerpt(line):
        """One excerpt behind the same left rule the editable memo draws. It splits
        across pages only between lines, like any paragraph."""
        panel = Table([[Paragraph(escape(line), styles['evidence'])]],
                      colWidths=[document.width - 1.5*mm], hAlign='RIGHT')
        panel.setStyle(TableStyle([
            ('LINEBEFORE', (0, 0), (0, -1), 1.5, colors.HexColor('#4B8790')),
            ('LEFTPADDING', (0, 0), (-1, -1), 2.5*mm), ('RIGHTPADDING', (0, 0), (-1, -1), 0),
            ('TOPPADDING', (0, 0), (-1, -1), 0), ('BOTTOMPADDING', (0, 0), (-1, -1), 1)]))
        return [panel, Spacer(1, 2.2*mm)]

    def kept(label, flows):
        """A label never ends a page: it stays with the first item beneath it."""
        return [KeepTogether([label, *flows[0]])] + [item for flow in flows[1:] for item in flow]

    story = [Paragraph(escape(title), styles["title"])]
    marker = "\n\nSource claims and exact excerpts:\n"
    for section in _checked_sections(sections, paginated=True, document_contract=document_contract):
        heading, body, sources = section[:3]
        prose, separator, evidence = body.partition(marker)
        registry = _source_registry_lines(prose)
        blocks = ([Paragraph(escape(line), styles['sources']) for line in registry] if registry else
                  [Paragraph(escape(block).replace('\n', '<br/>'), styles['body'])
                   for block in prose.split('\n\n') if block])
        # A hairline opens each section, as in the editable memo.
        heading_flow = [Spacer(1, 4*mm),
                        HRFlowable(width='100%', thickness=.5, color=colors.HexColor('#DCE5E9'),
                                   spaceBefore=0, spaceAfter=0),
                        Paragraph(escape(heading), styles['heading'])]
        story.append(KeepTogether(heading_flow + ([blocks.pop(0)] if blocks else [])))
        story.extend(KeepTogether([block]) for block in blocks)
        lines = [line for line in evidence.splitlines() if line] if separator else []
        if separator:
            label = Paragraph('Source claims and exact excerpts:', styles['label'])
            story.extend(kept(label, [excerpt(line) for line in lines]) if lines else [label])
        if len(section) == 6:
            from reportlab.platypus import Image as PlotImage
            chart = _evidence_chart(section[4])
            story.extend([PlotImage(BytesIO(_chart_plot_png(chart)), width=168 * mm, height=86.55 * mm),
                          Paragraph(_CHART_NOTE, styles['sources']),
                          Paragraph('Chart image is not editable. Exact source values remain editable below.',
                                    styles['sources'])])
        if len(section) in (5, 6):
            table = _chart_source_table(chart) if len(section) == 6 else _evidence_table(section[4])
            plan = _memo_table_plan(table)
            cell_style = ParagraphStyle('EvidenceTableCell', fontName=font,
                fontSize=plan['font_size'], leading=plan['font_size'] * 1.2,
                textColor=colors.HexColor('#202E3B'), spaceAfter=0, spaceBefore=0)
            header_style = ParagraphStyle('EvidenceTableHeader', parent=cell_style,
                fontName=_pdf_font(bold=True), textColor=colors.white)
            cells = [[Paragraph(escape(text).replace('\n', '<br/>'),
                                header_style if row_index == 0 else cell_style)
                      for text in row] for row_index, row in enumerate(plan['matrix'])]
            native = Table(cells, colWidths=[width * 72 for width in plan['widths']],
                           repeatRows=1, splitByRow=1, hAlign='LEFT')
            native.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#17324D')),
                ('ROWBACKGROUNDS', (0, 1), (-1, -1),
                 [colors.HexColor('#EFF4F5'), colors.white]),
                ('GRID', (0, 0), (-1, -1), .35, colors.HexColor('#DCE5E9')),
                ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                ('LEFTPADDING', (0, 0), (-1, -1), .06 * 72),
                ('RIGHTPADDING', (0, 0), (-1, -1), .06 * 72),
                ('TOPPADDING', (0, 0), (-1, -1), .05 * 72),
                ('BOTTOMPADDING', (0, 0), (-1, -1), .05 * 72),
            ]))
            story.extend([Spacer(1, 2 * mm), native, Spacer(1, 2 * mm),
                Paragraph('Figures reported by sources. Financial review pending.',
                          styles['sources'])])
        versions = [line for line in sources.splitlines() if line]
        if versions:
            paragraphs = [Paragraph(escape(line), styles['sources']) for line in versions]
            story.append(KeepTogether([Paragraph(_memo_source_label(sources), styles['label']), paragraphs[0]]))
            story.extend(KeepTogether([paragraph]) for paragraph in paragraphs[1:])

    def frame(canvas, doc):
        canvas.saveState()
        width, height = A4
        canvas.setFont(font, 8)
        canvas.setFillColor(colors.HexColor('#42566A'))
        canvas.drawString(20*mm, height-15*mm, 'DRAFT — NOT APPROVED FOR INVESTOR DISTRIBUTION')
        canvas.drawString(20*mm, 12*mm, 'INVESTMENT MEMORANDUM DRAFT')
        canvas.drawRightString(width-20*mm, 12*mm, f'Page {doc.page}')
        canvas.restoreState()

    document.build(story, onFirstPage=frame, onLaterPages=frame)
    return output.getvalue()


def render_research_pdf(title, sections, *, provenance=""):
    """Render recorded model prose and exact evidence as a readable private draft."""
    from html import escape
    from datetime import date
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, KeepTogether, Table, TableStyle

    if not sections or len(sections) > 30:
        raise ValueError("Supported draft requires 1-30 recorded sections")
    output = BytesIO()
    document = SimpleDocTemplate(output, pagesize=A4, leftMargin=19*mm,
        rightMargin=19*mm, topMargin=23*mm, bottomMargin=20*mm,
        title=f"{title} - draft research", author="Local deal room")
    base = getSampleStyleSheet()
    styles = {
        "title": ParagraphStyle("ResearchTitle", parent=base["Title"], fontName="Helvetica-Bold",
            fontSize=21, leading=26, textColor=colors.HexColor("#17324D"), alignment=TA_LEFT,
            spaceAfter=5),
        "notice": ParagraphStyle("DraftNotice", parent=base["Normal"], fontName="Helvetica-Bold",
            fontSize=8, leading=11, textColor=colors.HexColor("#8B3A3A"), spaceAfter=12),
        "heading": ParagraphStyle("ResearchHeading", parent=base["Heading2"], fontName="Helvetica-Bold",
            fontSize=12, leading=16, textColor=colors.HexColor("#17324D"), spaceBefore=14,
            spaceAfter=5),
        "body": ParagraphStyle("ResearchBody", parent=base["BodyText"], fontName="Helvetica",
            fontSize=10, leading=15, spaceAfter=8),
        "evidence_label": ParagraphStyle("ResearchEvidenceLabel", parent=base["BodyText"],
            fontName="Helvetica-Bold", fontSize=8, leading=11,
            textColor=colors.HexColor("#17324D"), spaceBefore=5, spaceAfter=5),
        "evidence": ParagraphStyle("ResearchEvidence", parent=base["BodyText"],
            fontName="Helvetica", fontSize=8.5, leading=12,
            textColor=colors.HexColor("#263645"), splitLongWords=1),
        "sources": ParagraphStyle("ResearchSources", parent=base["BodyText"], fontName="Helvetica",
            fontSize=7.5, leading=11, textColor=colors.HexColor("#42566A"),
            splitLongWords=1, spaceAfter=8),
    }
    story = [Paragraph(escape(title), styles["title"]),
             Paragraph("INVESTMENT RESEARCH DRAFT  |  " + date.today().isoformat(), styles["sources"]),
             Paragraph("DRAFT - source-reported evidence; investment decision and financials unverified", styles["notice"])]
    if provenance:
        story.append(Paragraph(escape(provenance), styles["sources"]))
    marker = "\n\nSource claims and exact excerpts:\n"
    for heading, body, sources in _checked_sections(sections, paginated=True):
        section_flow = []
        prose, separator, evidence = body.partition(marker)
        registry = _source_registry_lines(prose)
        prose_blocks = registry or [part for part in prose.split("\n\n") if part.strip()]
        prose_style = styles["sources"] if registry else styles["body"]
        heading_block = Paragraph(escape(heading), styles["heading"])
        if prose_blocks:
            section_flow.extend([heading_block,
                Paragraph(escape(prose_blocks[0]).replace("\n", "<br/>"), prose_style)])
            for part in prose_blocks[1:]:
                paragraph = Paragraph(escape(part).replace("\n", "<br/>"), prose_style)
                section_flow.append(KeepTogether([paragraph]))
        else:
            section_flow.append(heading_block)
        if separator:
            section_flow.append(Paragraph("SOURCE CLAIMS AND EXACT EXCERPTS", styles["evidence_label"]))
            for line in evidence.splitlines():
                if not line.strip():
                    continue
                panel = Table([[Paragraph(escape(line), styles["evidence"])]],
                              colWidths=[document.width])
                panel.setStyle(TableStyle([
                    ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F2F6F9")),
                    ("BOX", (0, 0), (-1, -1), 0.3, colors.HexColor("#D7E1E8")),
                    ("LEFTPADDING", (0, 0), (-1, -1), 8),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 8),
                    ("TOPPADDING", (0, 0), (-1, -1), 7),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
                ]))
                section_flow.extend([panel, Spacer(1, 2*mm)])
        if sources:
            section_flow.append(Paragraph("SOURCE VERSIONS", styles["evidence_label"]))
            for line in sources.splitlines():
                if line:
                    section_flow.append(KeepTogether([Paragraph(escape(line), styles["sources"])]))
        section_flow.append(Spacer(1, 3*mm))
        first_panel = next((index for index, flowable in enumerate(section_flow)
                            if isinstance(flowable, Table)), None)
        anchor_count = first_panel + 1 if first_panel is not None else min(2, len(section_flow))
        story.append(KeepTogether(section_flow[:anchor_count]))
        story.extend(section_flow[anchor_count:])

    def page_frame(canvas, doc):
        canvas.saveState()
        width, height = A4
        canvas.setStrokeColor(colors.HexColor("#D7E1E8"))
        canvas.line(19*mm, height-17*mm, width-19*mm, height-17*mm)
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(colors.HexColor("#42566A"))
        canvas.drawString(19*mm, 12*mm, "PRIVATE RESEARCH DRAFT - NOT APPROVED FOR DISTRIBUTION")
        canvas.drawRightString(width-19*mm, 12*mm, f"Page {doc.page}")
        canvas.restoreState()

    document.build(story, onFirstPage=page_frame, onLaterPages=page_frame)
    return output.getvalue()


def qualified_sections(pack):
    """Reconstruct validated model-authored sections; never author replacements."""
    from copy import deepcopy
    from agents.preparation.authored_preparation import validate_authored_pack
    checked=validate_authored_pack(deepcopy(pack))
    facts={f['id']:f for f in checked.get('record',{}).get('facts',[])}
    result=[]
    for key,section in checked.get('sections',{}).items():
        if section.get('status')!='complete':continue
        content=section.get('content',{})
        ids=content.get('fact_ids',[])
        if not ids or any(i not in facts for i in ids):continue
        text='\n\n'.join(value for field,value in content.items() if isinstance(value,str) and field not in {'status'})
        sources='\n'.join(f"[{i}] {facts[i].get('source_url','')}" for i in ids)
        result.append((section.get('title',key),text,sources))
    return result
