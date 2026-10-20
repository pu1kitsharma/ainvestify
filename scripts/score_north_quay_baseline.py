#!/usr/bin/env python3
"""Evaluation-only scorer for a North Quay baseline run directory.

Reads the independent key (never shown to a model). Because the model chose its
own layout and metric IDs, matching is by VALUE PRESENCE: an expected figure
counts as found if any numeric workbook cell is within tolerance. That is a
generous, label-blind upper bound (coincidental matches possible), not proof of
correct wiring. Nothing is repaired or inferred.
"""
import argparse, json, sys
from decimal import Decimal
from pathlib import Path
ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
KEY = ROOT / 'evals/agentic_banker/fixtures/local_capability/answer_key/expected.json'
TOL = 0.005


def expected_values(key):
    out = {}
    for y, vals in key['historical'].items():
        for k, v in vals.items(): out[f'hist.{y}.{k}'] = float(Decimal(v))
    for sc in key['scenarios']:
        for row in sc['years']:
            for k, v in row.items():
                if k not in ('year', 'status'): out[f"{sc['name']}.{row['year']}.{k}"] = float(Decimal(v))
        for k, v in sc.items():
            if k not in ('name', 'status', 'years'):
                try: out[f"{sc['name']}.{k}"] = float(Decimal(str(v)))
                except Exception: pass
    return out


def score(run):
    run = Path(run); rep = json.loads((run / 'report.json').read_text())
    res = {'run': run.name, 'model': rep['model'], 'status': rep['status'], 'calls': rep['calls'],
           'seconds': round(rep['elapsed'], 1), 'tokens': rep['tokens'],
           'workbook_exported': False, 'sheets': [], 'numeric_cells': 0, 'formula_cells': 0,
           'recalc_findings': None, 'expected_total': 0, 'expected_found': 0, 'expected_found_by_formula': 0,
           'fy2025_balance_found': None, 'scenarios_with_funded_equity_found': [],
           'memo_exported': bool(rep.get('export', {}).get('ok')), 'error': rep.get('error')}
    exp = expected_values(json.loads(KEY.read_text()))
    res['expected_total'] = len(exp)
    wbp = run / 'model.xlsx'
    if not wbp.exists(): return res
    from openpyxl import load_workbook
    from agents.tools.workbook import WorkbookTools
    book = load_workbook(wbp, data_only=False)
    res['workbook_exported'] = True; res['sheets'] = book.sheetnames
    ev = WorkbookTools.__new__(WorkbookTools); ev.root, ev.book, ev.max_cells = run, book, 10000
    try:
        calc = WorkbookTools.recalc.__wrapped__(ev)
        vals = {(r['sheet'], r['cell']): r['value'] for r in calc['values']}
        res['recalc_findings'] = len(calc['findings'])
    except Exception as e:
        vals = {}; res['recalc_findings'] = f'recalc error {type(e).__name__}'
    cells = []
    for ws in book:
        for row in ws.iter_rows():
            for c in row:
                if c.value is None: continue
                v = vals.get((ws.title, c.coordinate)) if c.data_type == 'f' else c.value
                if c.data_type == 'f': res['formula_cells'] += 1
                if type(v) in (int, float) and v == v:
                    res['numeric_cells'] += 1; cells.append((float(v), c.data_type == 'f'))
    def find(x): return [f for v, f in cells if abs(v - x) <= TOL]
    for k, x in exp.items():
        m = find(x)
        if m: res['expected_found'] += 1; res['expected_found_by_formula'] += any(m)
    res['fy2025_balance_found'] = bool(find(452.5)) if cells else None
    for sc in ('BASE', 'UPSIDE', 'DOWNSIDE'):
        if any(find(exp[f'{sc}.{y}.funded_equity']) for y in (2026, 2027, 2028)): res['scenarios_with_funded_equity_found'].append(sc)
    return res


if __name__ == '__main__':
    p = argparse.ArgumentParser(); p.add_argument('runs', nargs='+'); a = p.parse_args()
    print(json.dumps([score(r) for r in a.runs], indent=1))
