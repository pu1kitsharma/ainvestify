"""Synthetic engine smoke/compatibility report; never a production approval."""
import hashlib
from io import BytesIO
import json
from pathlib import Path
import platform
import subprocess
import sys
import argparse
import tempfile
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))

from openpyxl import Workbook,load_workbook
from delivery.inspection import inspect_bytes,summarize_report
from delivery.libreoffice import convert,office_binary
from delivery.rendering import render_intro,render_memo


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--diagnose-unisolated',action='store_true',help='Synthetic-only functional diagnostic; never qualifies private processing')
    args=parser.parse_args()
    root=Path('runtime_qualification').resolve();root.mkdir(exist_ok=True)
    report={'synthetic_only':True,'os':platform.platform(),'checks':{},'production_qualified':False,
        'isolated':not args.diagnose_unisolated,
        'limitations':['Destination Excel/PowerPoint parity and visual acceptance remain unqualified',
            'CLI round trip is not explicit UNO calculateAll qualification']}
    try:report['libreoffice_version']=subprocess.check_output([str(office_binary()),'--version'],timeout=15,text=True).strip()
    except Exception:report['libreoffice_version']='unavailable'
    if platform.system() == 'Darwin':
        try:
            signature=subprocess.run(['codesign','--verify','--strict','--deep',str(office_binary().parents[2])],
                capture_output=True,timeout=30,check=False)
            report['libreoffice_signature'] = 'valid' if signature.returncode == 0 else 'invalid'
        except (OSError, subprocess.TimeoutExpired):
            report['libreoffice_signature'] = 'unverified'
    def conversion(content,fmt,target):
        if not args.diagnose_unisolated:return convert(content,fmt,target)
        # This path accepts only the fixed synthetic inputs generated below.
        # Runtime conversion has no option to disable the private sandbox.
        with tempfile.TemporaryDirectory(prefix='synthetic-office-diagnostic-') as directory:
            job=Path(directory);source=job/f'input.{fmt}';source.write_bytes(content)
            output=job/'output';output.mkdir()
            subprocess.run([str(office_binary()),f'-env:UserInstallation={(job/"profile").as_uri()}',
                '--headless','--convert-to',target,'--outdir',str(output),str(source)],
                timeout=90,capture_output=True,check=True)
            return (output/f'input.{target}').read_bytes()
    sections=[('Synthetic qualification','Illustrative fixture only. No company facts.','Source: synthetic test fixture')]
    files={'intro.pptx':render_intro('Synthetic fixture',sections),'memo.docx':render_memo('Synthetic fixture',sections)}
    for name,content in files.items():
        (root/name).write_bytes(content)
        try:
            pdf=conversion(content,name.split('.')[-1],'pdf')
            (root/(name+'.pdf')).write_bytes(pdf)
            inspection=inspect_bytes(pdf,'pdf')
            report['checks'][name]={'state':'pass' if inspection['structural_status']=='pass' else 'fail',
                'pdf_sha256':hashlib.sha256(pdf).hexdigest(),'inspection':summarize_report(inspection)}
        except Exception as exc:report['checks'][name]={'state':'fail','error_type':type(exc).__name__,'reason':str(exc)}
    for input_value in (10,15):
        wb=Workbook();wb.active.title='Inputs';wb.active['A1']=input_value
        calc=wb.create_sheet('Calculation');calc['A1']='=Inputs!A1*2';calc['B1']='=A1-5'
        hidden=wb.create_sheet('Hidden');hidden.sheet_state='hidden';hidden['A1']='=Calculation!B1+1'
        output=BytesIO();wb.save(output)
        name=f'synthetic-{input_value}.xlsx';(root/name).write_bytes(output.getvalue())
        try:
            recalculated=conversion(output.getvalue(),'xlsx','xlsx');(root/('recalculated-'+name)).write_bytes(recalculated)
            cached=load_workbook(BytesIO(recalculated),data_only=True)
            formulas=load_workbook(BytesIO(recalculated),data_only=False)
            actual=(cached['Calculation']['A1'].value,cached['Calculation']['B1'].value,cached['Hidden']['A1'].value)
            expected=(input_value*2,input_value*2-5,input_value*2-4)
            passed=actual==expected and formulas['Calculation']['A1'].data_type=='f' and formulas['Hidden'].sheet_state=='hidden'
            report['checks'][name]={'state':'pass' if passed else 'fail','expected':expected,'actual':actual,
                'formula_preserved':formulas['Calculation']['A1'].data_type=='f'}
            cached.close();formulas.close()
        except Exception as exc:report['checks'][name]={'state':'fail','error_type':type(exc).__name__,'reason':str(exc)}
    report_path=root/('functional-diagnostic.json' if args.diagnose_unisolated else 'report.json')
    report_path.write_text(json.dumps(report,indent=2))
    print(json.dumps(report,indent=2))
    return int(any(c['state']!='pass' for c in report['checks'].values()))


if __name__=='__main__':raise SystemExit(main())
