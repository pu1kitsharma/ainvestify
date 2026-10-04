"""Public/synthetic-only local semantic review of frozen material diagnostics."""
from __future__ import annotations

import argparse
import hashlib
import json
import re
from pathlib import Path

from pptx import Presentation

from agents.inference.local_models import LocalModel
from agents.inference.model_authorship import digest
from agents.inference.model_routing import validate_local_model_name
from scripts.evaluate_local_memo_harness import installed_models
from scripts.private_material_review_worker import review_materials


ROOT = Path(__file__).resolve().parents[1]
MATERIAL_ROOT = ROOT / 'runtime_qualification/local_material_harness'
OUTPUT_ROOT = ROOT / 'runtime_qualification/local_material_review_harness'


def _slides(path):
    presentation = Presentation(str(path))
    result = []
    for slide in presentation.slides:
        text = [shape.text.strip() for shape in slide.shapes
                if shape.has_text_frame and shape.text.strip()]
        if len(text) < 5 or text[0] != 'DRAFT • NOT APPROVED FOR INVESTOR DISTRIBUTION':
            raise ValueError('Material diagnostic editable slide structure changed')
        notes = text.index('SOURCE NOTES')
        if notes <= 2:
            raise ValueError('Material diagnostic slide body is missing')
        result.append([text[1], '\n'.join(text[2:notes]), '', 'diagnostic'])
    return result


FRESH_REVIEW_CONTRACT = 'semantic_v10'


def require_supported_contract(contract):
    """Refuse a contract the review module does not implement, instead of letting it
    fall back silently to the first, unversioned instruction."""
    from agents.research.material_review import REVIEW_INSTRUCTION, review_instruction
    if contract is not None and review_instruction({'review_contract': contract}) == REVIEW_INSTRUCTION:
        raise ValueError(f'Material review module does not implement {contract}; '
                         'refusing to fall back to the unversioned contract')
    return contract
REVIEW_OPTIONS = {'thinking': False, 'max_tokens': 1200, 'context_tokens': 16384, 'temperature': 0}


def model_slug(name: str) -> str:
    return re.sub(r'[^a-z0-9]+', '-', name.casefold()).strip('-')


def review_pin(draft_model: str, review_model, installed: dict, output: Path):
    """The frozen reviewer for a new diagnostic: the material's draft model, or an
    explicit local override.

    An override must be an installed local model, pinned by its digest; nothing
    is downloaded. Its output directory name must contain the model's slug, so a
    review by another model can never be mistaken for, or written next to, a
    default-model diagnostic.
    """
    name = draft_model if review_model is None else review_model
    validate_local_model_name(name)
    model_digest = installed.get(name)
    if not model_digest:
        raise ValueError('Review model is not installed locally with a digest; '
                         'nothing was downloaded')
    if review_model is not None and model_slug(name) not in model_slug(output.name):
        raise ValueError('An overridden review model needs its own output directory, '
                         f'named with {model_slug(name)!r}')
    pin = {'name': name, 'digest': model_digest, 'options': dict(REVIEW_OPTIONS)}
    record = {**pin, 'source': ('material_profile_draft' if review_model is None
                                else 'operator_override'),
              'material_draft_model': draft_model}
    return pin, record


def _frozen(directory: Path) -> dict:
    """Hashes of every file of an earlier diagnostic, to show this run left it alone."""
    return {str(path.relative_to(directory)): hashlib.sha256(path.read_bytes()).hexdigest()
            for path in sorted(directory.rglob('*')) if path.is_file()}


def evaluate(material_dir: Path, output: Path, *, operator_attested=False, review_model=None):
    if not operator_attested:
        raise ValueError('Public or synthetic material bytes require operator attestation')
    if not material_dir.resolve().is_relative_to(MATERIAL_ROOT.resolve()):
        raise ValueError('Material diagnostic must be under its allowlisted root')
    if not output.resolve().is_relative_to(OUTPUT_ROOT.resolve()) or output.resolve() == OUTPUT_ROOT.resolve():
        raise ValueError('Review diagnostic must be under its output root')
    report = json.loads((material_dir / 'result.json').read_text())
    if (report.get('state') != 'accepted' or
            report.get('acceptance_scope') != 'public_synthetic_diagnostic_only' or
            report.get('investor_material_accepted') is not False):
        raise ValueError('Material diagnostic lacks a historical accepted synthetic result')
    for stem in ('intro', 'pitch', 'memo'):
        for extension, key in (('pptx' if stem != 'memo' else 'docx', 'editable_sha256'),
                               ('pdf', 'pdf_sha256')):
            actual = hashlib.sha256((material_dir / f'{stem}.{extension}').read_bytes()).hexdigest()
            if actual != report['pair_reports'][stem][key]:
                raise ValueError('Material diagnostic artifact hash changed')
    old_request = json.loads((material_dir / 'material_request.json').read_text())
    if old_request['digest'] != digest({key: value for key, value in old_request.items()
                                        if key != 'digest'}):
        raise ValueError('Material diagnostic request digest changed')
    profile = json.loads((material_dir / 'model.json').read_text())
    pin, pin_record = review_pin(profile['profiles']['draft'], review_model,
                                 installed_models(), output)
    name, model_digest = pin['name'], pin['digest']
    material_before = _frozen(material_dir)
    decks = {kind: _slides(material_dir / f'{stem}.pptx') for kind, stem in
             (('intro_deck', 'intro'), ('pitch_deck', 'pitch'))}
    base = {'review_contract': require_supported_contract(FRESH_REVIEW_CONTRACT),
            'input_revision': old_request['input_revision'],
            'source_hash': old_request['source_hash'],
            'memo_digest': old_request['memo_digest'],
            'material_digest': digest({'request': old_request['digest'],
                'pair_reports': report['pair_reports']}),
            'memo_sections': old_request['sections'], 'decks': decks,
            'review_model': pin}
    output.parent.mkdir(parents=True, exist_ok=True)
    output.mkdir(parents=True, exist_ok=False, mode=0o700)
    (output / 'material_review_request.json').write_text(json.dumps({**base,
        'digest': digest(base)}, ensure_ascii=False))
    (output / 'review_model.json').write_text(json.dumps(pin_record, ensure_ascii=False))
    passes = []
    result = None
    for number in range(1, 3):
        if installed_models().get(name) != model_digest:
            result = {'state': 'blocked', 'reason': 'installed_review_model_digest_changed'}
            break
        (output / 'material_review_budget.json').write_text(json.dumps({'seconds': 105}))
        model = LocalModel(name, **REVIEW_OPTIONS)
        result = review_materials(output, model)
        passes.append({'pass': number, 'state': result['state'],
                       'reason': result.get('reason')})
        (output / 'passes.json').write_text(json.dumps(passes))
        if result['state'] != 'needs_resume':
            break
    if result.get('response_id'):
        class NoInference:
            def __init__(self, model_name):
                self.name = model_name
            def generate_for_task(self, *args, **kwargs):
                raise AssertionError('Review replay must not call inference')
        if review_materials(output, NoInference(name)) != result:
            raise ValueError('Review diagnostic differs from exact recorded model replay')
    if _frozen(material_dir) != material_before:
        raise ValueError('Review diagnostic changed its frozen material input')
    summary = {'state': result['state'], 'reason': result.get('reason'),
               'passes': len(passes), 'response_id': result.get('response_id'),
               'review': result.get('review'),
               'review_contract': FRESH_REVIEW_CONTRACT, 'review_model': pin_record,
               'acceptance_scope': 'public_synthetic_diagnostic_only',
               'investor_material_accepted': False}
    (output / 'result.json').write_text(json.dumps(summary, ensure_ascii=False, indent=2))
    return summary


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('material_dir', type=Path)
    parser.add_argument('output', type=Path)
    parser.add_argument('--operator-attested-public-or-synthetic', action='store_true')
    parser.add_argument('--review-model', default=None,
                        help='An installed local model to review with instead of the material '
                             "diagnostic's draft model. Pinned by digest; never downloaded. The "
                             'output directory name must contain the model name.')
    args = parser.parse_args()
    print(json.dumps(evaluate(args.material_dir, args.output,
        operator_attested=args.operator_attested_public_or_synthetic,
        review_model=args.review_model), indent=2))


if __name__ == '__main__':
    main()
